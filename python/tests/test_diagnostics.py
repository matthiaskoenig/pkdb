"""Actionable diagnostics retain physical provenance and legacy compatibility."""

import json

import openpyxl
import pytest

from pkdb import prepare
from pkdb.schemas.source import SourceLocation
from pkdb.schemas.validation import (
    StudyValidationError,
    ValidationIssue,
    ValidationReport,
)


def test_actual_omitted_distinguished_from_null_and_legacy_shape():
    issue = ValidationIssue(code="example", message="Example")
    assert "actual" not in issue.model_dump(mode="json")
    explicit = ValidationIssue(code="example", message="Example", actual=None)
    assert explicit.model_dump(mode="json")["actual"] is None
    assert explicit.legacy_dict() == {
        "code": "example",
        "severity": "error",
        "message": "Example",
        "source": None,
    }


def test_bounded_diagnostics_preserve_total_counts_and_order():
    report = ValidationReport(
        issues=[
            ValidationIssue(
                code="example",
                message="Example",
                actual="a" * 600,
                source=SourceLocation(file="a.xlsx", row=row),
            )
            for row in (5, 3, 4)
        ]
    ).finalize(max_issues=2)
    assert [i.source.row for i in report.issues if i.source] == [3, 4]
    assert report.error_count == 3
    assert report.returned_issue_count == 2
    assert report.omitted_issue_count == 1
    assert report.truncated and report.complete
    assert isinstance(report.issues[0].actual, str)
    assert len(report.issues[0].actual) == 512


def test_negative_value_identifies_real_cell(study_folder, vocabulary):
    if (study_folder / "Example.xlsx").exists():
        book = openpyxl.load_workbook(study_folder / "Example.xlsx")
        book["Results"]["B4"] = -2
        book.save(study_folder / "Example.xlsx")
        book.close()
        row = 4
    else:
        (study_folder / "Results.tsv").write_text("time\tmean\n0\t0\n1\t-2\n")
        row = 3
    with pytest.raises(StudyValidationError) as raised:
        prepare(study_folder, vocabulary=vocabulary)
    issue = next(i for i in raised.value.report.issues if i.code == "negative_value")
    assert issue.actual == -2
    assert issue.expected == {"minimum": 0}
    assert issue.source is not None
    assert issue.source.row == row
    assert issue.source is not None
    assert issue.source.column == "B"
    assert issue.source is not None
    assert issue.source.cell == f"B{row}"
    assert issue.source is not None
    assert issue.source.header == "mean"
    assert issue.suggestions


def test_vocabulary_constant_points_to_json_definition(study_folder, vocabulary):
    path = study_folder / "study.json"
    value = json.loads(path.read_text())
    value["outputset"]["outputs"][0]["measurement_type"] = "concentraton"
    path.write_text(json.dumps(value))
    with pytest.raises(StudyValidationError) as raised:
        prepare(study_folder, vocabulary=vocabulary)
    issue = raised.value.report.issues[0]
    assert issue.code == "unknown_measurement"
    assert issue.source is not None
    assert issue.source.file == "study.json"
    assert issue.source is not None
    assert issue.source.path == ("outputset", "outputs", 0, "measurement_type")
    assert issue.actual == "concentraton"
    assert "concentration" in issue.suggestions[0].candidates


def test_missing_column_is_honest_about_unavailable_cell(study_folder, vocabulary):
    path = study_folder / "study.json"
    value = json.loads(path.read_text())
    value["outputset"]["outputs"][0]["mean"] = "col==label"
    path.write_text(json.dumps(value))
    with pytest.raises(StudyValidationError) as raised:
        prepare(study_folder, vocabulary=vocabulary)
    report = raised.value.report
    assert not report.complete
    assert report.stopped_reason == "unknown_column"
    issue = report.issues[0]
    assert issue.actual == "label"
    assert issue.expected["available_columns"] == ["mean", "time"]
    assert issue.source is not None
    assert issue.source.cell is None


def test_physical_rows_survive_comment_and_blank_lines(tmp_path):
    from pkdb.importers.workbook import read_table

    path = tmp_path / "physical.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Results"
    sheet.append(["notes"])
    sheet.append(["time", "mean"])
    sheet.append(["# section"])
    sheet.append([None, None])
    sheet.append([1, -2])
    book.save(path)
    book.close()
    rows = read_table(path, "Results", 100)
    assert len(rows) == 1
    data, source = rows[0]
    assert data["mean"] == -2
    assert source.row == 5
    assert source.for_header("mean").cell == "B5"


@pytest.mark.parametrize("value", [float("nan"), float("inf"), float("-inf")])
def test_nonfinite_input_remains_serializable_and_explicit(value):
    issue = ValidationIssue(code="bad_number", message="Invalid number", actual=value)
    report = ValidationReport(issues=[issue]).finalize()
    serialized = json.dumps(report.model_dump(mode="json"), allow_nan=False)
    assert serialized
    assert isinstance(issue.actual, str)
    assert report.truncated


def test_all_diagnostic_surfaces_are_bounded_and_literal():
    source = SourceLocation(
        file="/private/server/file.xlsx", header="mean\x1b[31m", path=("x\x00",)
    )
    issue = ValidationIssue(
        code="example",
        message="Issue",
        source=source,
        context={"auth\x00orization": "secret", "unit\x00": "mg"},
        suggestions=[
            {
                "kind": "inspect_source",
                "message": "x" * 600,
                "candidates": ["y" * 600] * 12,
                "command": "inspect\x00",
            }
        ]
        * 12,
        related_sources=[{"label": "z" * 600, "source": source}] * 12,
        documentation_url="javascript:alert(1)",
    )
    report = ValidationReport(issues=[issue]).finalize()
    assert issue.source is not None
    assert issue.source.file == "file.xlsx"
    assert issue.source.header == "mean[31m"
    assert issue.source.path == ("x",)
    assert issue.context == {"unit": "mg"}
    assert len(issue.suggestions) == len(issue.related_sources) == 10
    assert len(issue.suggestions[0].message) == 512
    assert len(issue.suggestions[0].candidates) == 10
    assert issue.suggestions[0].command == "inspect"
    assert issue.related_sources[0].source.file == "file.xlsx"
    assert issue.documentation_url is None
    assert report.truncated
    assert "\x1b" not in json.dumps(report.model_dump())


def test_related_source_does_not_claim_truncation_when_unchanged():
    source = SourceLocation(file="study.json", path=("outputs", 0))
    report = ValidationReport(
        issues=[
            ValidationIssue(
                code="duplicate",
                message="Duplicate",
                related_sources=[{"label": "First", "source": source}],
            )
        ]
    )
    assert not report.truncated


def test_schema_failures_mark_unperformed_scientific_validation(
    study_folder, vocabulary
):
    path = study_folder / "study.json"
    value = json.loads(path.read_text())
    value["outputset"]["outputs"][0]["unsupported_field"] = "value"
    path.write_text(json.dumps(value))
    with pytest.raises(StudyValidationError) as error:
        prepare(study_folder, vocabulary=vocabulary)
    assert not error.value.report.complete
    assert error.value.report.stopped_reason == "schema_validation"
    assert error.value.report.issues[0].source is not None
    assert error.value.report.issues[0].source.path == ()


def test_missing_source_does_not_disclose_server_path(tmp_path):
    from pkdb.importers.folder import _read_json
    from pkdb.importers.workbook import read_table

    missing = tmp_path / "private-server-path" / "missing.json"
    for load in (
        lambda: _read_json(missing),
        lambda: read_table(missing.with_suffix(".tsv"), None, 10),
    ):
        with pytest.raises(StudyValidationError) as error:
            load()
        report = error.value.report.model_dump_json()
        assert str(tmp_path) not in report
        assert "private-server-path" not in report
        assert not error.value.report.complete


def test_reference_and_duplicate_reports_include_definitions(study_folder, vocabulary):
    from pkdb.domain.validation import prepare_study
    from pkdb.importers.folder import load_folder, parse_bundle

    study = parse_bundle(load_folder(study_folder))
    duplicate = study.measurements[0].model_copy(deep=True)
    duplicate.source = SourceLocation(file="Example.xlsx", sheet="Results", row=8)
    study.measurements.append(duplicate)
    study.measurements[0].group = "absent-group"
    study.measurements[0].interventions = ["absent-dose"]
    study.measurements[0].substance = "absent-substance"
    with pytest.raises(StudyValidationError) as error:
        prepare_study(study, vocabulary)
    issues = {i.code: i for i in error.value.report.issues}
    assert issues["duplicate_identifier"].actual == duplicate.key
    assert issues["duplicate_identifier"].related_sources[0].source.row == 8
    assert issues["unknown_group"].actual == "absent-group"
    assert issues["unknown_group"].expected == {"defined_identifiers": ["all"]}
    assert issues["unknown_intervention"].actual == "absent-dose"
    assert issues["unknown_substance"].expected == {"vocabulary": "substances"}


def test_stale_excel_dimensions_do_not_pad_a_million_rows(tmp_path):
    from zipfile import ZipFile

    from pkdb.importers.workbook import read_table

    path = tmp_path / "dimensions.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Results"
    sheet.append(["notes"])
    sheet.append(["time", "mean"])
    sheet.append([0, 2])
    book.save(path)
    book.close()
    with ZipFile(path) as archive:
        contents = {name: archive.read(name) for name in archive.namelist()}
    contents["xl/worksheets/sheet1.xml"] = contents["xl/worksheets/sheet1.xml"].replace(
        b'ref="A1:B3"', b'ref="A1:AMJ1048576"'
    )
    with ZipFile(path, "w") as archive:
        for name, content in contents.items():
            archive.writestr(name, content)
    rows = read_table(path, "Results", max_rows=10)
    assert len(rows) == 1
    assert rows[0][1].row == 3
    assert rows[0][0] == {"time": 0, "mean": 2}


def test_styled_empty_suffix_does_not_consume_source_row_budget(tmp_path):
    from openpyxl.styles import Font

    from pkdb.importers.workbook import read_table

    path = tmp_path / "styles.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Results"
    sheet.append(["notes"])
    sheet.append(["mean"])
    sheet.append([-2])
    sheet.cell(10000, 1).font = Font(bold=True)
    book.save(path)
    book.close()
    rows = read_table(path, "Results", max_rows=1)
    assert len(rows) == 1
    assert rows[0][1].row == 3


def test_canonical_source_preserves_legacy_serialized_fields():
    source = SourceLocation(file="study.json", path=("outputs", 0))
    assert source.model_dump(mode="json") == source.legacy_dict()
    detailed = source.model_copy(update={"cell": "B3", "header": "mean"})
    assert detailed.model_dump(mode="json")["cell"] == "B3"
    assert "cell" not in detailed.legacy_dict()


def test_unfinished_validation_is_not_success_even_without_discovered_errors():
    report = ValidationReport(complete=False, stopped_reason="time_budget")
    assert not report.valid
