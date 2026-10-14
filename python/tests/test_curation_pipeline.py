"""The curation engine syncs and formats a format 2 study before it validates it."""

from contextlib import contextmanager

import openpyxl
import pytest

from pkdb.curation import engine as module
from pkdb.curation import jobs, state, studies
from pkdb.curation import workspace as workspace_module
from pkdb.studyformat import pipeline as pipeline_module
from pkdb.studyformat.formatter import format_folder
from pkdb.studyformat.issues import make_issue
from pkdb.studyformat.sync import SyncResult, sync_study
from pkdb.studyformat.workbook.base import workbook_path
from pkdb.studyformat.workbook.view import show_place


@pytest.fixture
def workspace(tmp_path_factory, make_study, valid_files, sf_vocabulary, monkeypatch):
    folder = make_study(valid_files)
    engine = module.CurationEngine(
        folder.parent.parent,
        state_dir=tmp_path_factory.mktemp("state"),
        offline=True,
        start=False,
    )
    # The bundled vocabulary lacks the substance of the test study.
    monkeypatch.setattr(engine, "_local_vocabulary", lambda folder=None: sf_vocabulary)
    yield engine, folder
    engine.close()


def run_next(engine):
    identifier = next(iter(engine.queue))
    job = engine.queue.pop(identifier)
    engine.active = identifier
    job["status"] = "running"
    try:
        engine.run_job(job)
    finally:
        engine.active = None
    return job


def settle(engine):
    for item in engine.studies.values():
        item["_changed_at"] -= 2
    engine.schedule_changes()


def row(engine):
    return engine.snapshot()["studies"][0]


def row_list(engine):
    return engine.snapshot()["studies"]


def owned(lock):
    """Whether the current thread holds the engine lock."""
    return lock.held()


def set_mean(folder, table, value):
    """Set the mean of the second data row of a table file."""
    path = folder / table
    lines = path.read_text(encoding="utf-8").splitlines()
    header = lines[0].split("\t")
    cells = lines[2].split("\t")
    cells[header.index("mean")] = value
    lines[2] = "\t".join(cells)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8", newline="")


def set_workbook_mean(folder, sheet, value):
    """Set the mean of the second data row of a workbook sheet."""
    path = workbook_path(folder)
    book = openpyxl.load_workbook(path)
    rows = book[sheet]
    header = [cell.value for cell in rows[1]]
    rows.cell(3, header.index("mean") + 1).value = value
    book.save(path)


def workbook_mean(folder, sheet):
    rows = openpyxl.load_workbook(workbook_path(folder))[sheet]
    header = [cell.value for cell in rows[1]]
    return rows.cell(3, header.index("mean") + 1).value


def test_validation_formats_first(workspace):
    engine, folder = workspace
    table = folder / "timecourses_Fig1.tsv"
    table.write_text(
        table.read_text(encoding="utf-8") + "\n\n", encoding="utf-8", newline=""
    )  # not canonical
    settle(engine)
    job = run_next(engine)
    assert job["status"] == "succeeded", job["message"]
    assert not table.read_text(encoding="utf-8").endswith("\n\n")
    assert row(engine)["counts"] == {"errors": 0, "warnings": 0}
    assert row(engine)["sync"]["status"] == "no_workbook"
    report = engine.report(job["report_id"])
    assert "timecourses_Fig1.tsv" in report["tables_updated"]


def test_conflict_stops_before_validation(workspace, sf_vocabulary, monkeypatch):
    engine, folder = workspace
    assert sync_study(folder, sf_vocabulary).ok
    set_workbook_mean(folder, "timecourses_Fig1", 5)
    set_mean(folder, "timecourses_Fig1.tsv", "7")
    engine.scan()
    assert row(engine)["sync"] == {"status": "conflict", "changes": 0, "conflicts": 1}
    settle(engine)
    called = []
    monkeypatch.setattr(jobs, "prepare", lambda *a, **k: called.append(1))
    job = run_next(engine)
    assert job["status"] == "conflict" and not called
    assert row(engine)["status"] == "conflict"
    assert row(engine)["sync"]["status"] == "conflict"
    report = engine.report(job["report_id"])
    assert (
        report["conflicts"] and report["conflicts"][0]["file"] == "timecourses_Fig1.tsv"
    )
    assert report["conflicts"][0]["workbook_rows"][0]["row"] == 3
    assert any(issue["code"] == "sync_conflict" for issue in report["pipeline_issues"])
    # The study page lists the conflict as a problem; the overview has its count.
    problems = engine.study_detail("caffeine/Example")["problems"]
    assert "sync_conflict" in {problem["code"] for problem in problems}
    assert row(engine)["counts"]["errors"] >= 1


def test_tables_that_cannot_be_synced_make_the_study_invalid(workspace, sf_vocabulary):
    engine, folder = workspace
    assert format_folder(folder).ok
    assert sync_study(folder, sf_vocabulary).ok
    table = folder / "timecourses_Fig1.tsv"
    header, *lines = table.read_text(encoding="utf-8").splitlines()
    table.write_text(
        "\n".join([f"{header}\tcolour", *(f"{line}\tred" for line in lines)]) + "\n",
        encoding="utf-8",
        newline="",
    )
    engine.scan()
    settle(engine)
    job = run_next(engine)
    assert job["status"] == "invalid", job["message"]
    current = row(engine)
    assert current["status"] == "invalid"
    assert current["counts"]["errors"] >= 1
    assert "problems" not in current
    problems = engine.study_detail("caffeine/Example")["problems"]
    assert "unknown_column" in {problem["code"] for problem in problems}


def outside_range(folder):
    """Give the second row of timecourses_Fig1.tsv a mean outside its range: a warning."""
    table = folder / "timecourses_Fig1.tsv"
    lines = table.read_text(encoding="utf-8").splitlines()
    header = lines[0].split("\t")
    cells = lines[2].split("\t")
    cells[header.index("min")], cells[header.index("max")] = "3", "4"
    lines[2] = "\t".join(cells)
    table.write_text("\n".join(lines) + "\n", encoding="utf-8", newline="")


def test_only_warnings_of_the_validation_can_be_acknowledged(
    workspace, sf_vocabulary, monkeypatch
):
    """Acknowledge matches the warnings of the validation, never those of the sync."""
    engine, folder = workspace
    outside_range(folder)
    settle(engine)
    assert run_next(engine)["status"] == "succeeded"
    problems = engine.study_detail("caffeine/Example")["problems"]
    assert [(p["code"], p["acknowledgeable"]) for p in problems] == [
        ("outside_range", True)
    ]

    # The workbook is saved during both syncs: the job stops with the warning of the sync.
    assert sync_study(folder, sf_vocabulary).ok
    saved = make_issue(
        "workbook_changed",
        "Example.xlsx was saved during the sync; sync again",
        file="Example.xlsx",
    )
    monkeypatch.setattr(
        pipeline_module,
        "sync_study",
        lambda path, vocabulary, **_: SyncResult(
            path, workbook_path(path), workbook_action="sync_again", issues=(saved,)
        ),
    )
    engine.enqueue(["caffeine/Example"], "validate")
    assert run_next(engine)["status"] == "conflict"
    problems = engine.study_detail("caffeine/Example")["problems"]
    assert [(p["code"], p["acknowledgeable"]) for p in problems] == [
        ("workbook_changed", False)
    ]
    # The rows of the engine keep the problems as the report has them.
    assert "acknowledgeable" not in engine.studies["caffeine/Example"]["problems"][0]


def test_beyond_the_limits_the_refusal_cannot_be_acknowledged(workspace, monkeypatch):
    """The refusal of a study beyond the upload limits comes first and cannot be acknowledged;
    the warnings of the last validation still can."""
    engine, folder = workspace
    outside_range(folder)
    settle(engine)
    assert run_next(engine)["status"] == "succeeded"
    monkeypatch.setattr(studies, "MAX_ROWS", 3)
    problems = engine.study_detail("caffeine/Example")["problems"]
    assert [(p["code"], p["acknowledgeable"]) for p in problems] == [
        ("row_limit", False),
        ("outside_range", True),
    ]


def test_a_sync_stop_after_a_reference_update_is_invalid_in_the_same_job(
    workspace, sf_vocabulary, monkeypatch
):
    from pkdb import references

    engine, folder = workspace
    assert format_folder(folder).ok
    assert sync_study(folder, sf_vocabulary).ok
    table = folder / "timecourses_Fig1.tsv"
    header, *lines = table.read_text(encoding="utf-8").splitlines()
    table.write_text(
        "\n".join([f"{header}\tcolour", *(f"{line}\tred" for line in lines)]) + "\n",
        encoding="utf-8",
        newline="",
    )
    engine.scan()

    def refresh(path, resolver):
        reference = path / "reference.json"
        reference.write_text(
            reference.read_text(encoding="utf-8").replace("Example study", "Study"),
            encoding="utf-8",
            newline="",
        )
        return "Updated reference.json"

    monkeypatch.setattr(references, "sync_reference", refresh)
    settle(engine)
    job = run_next(engine)
    assert job["status"] == "invalid", job["message"]
    current = engine.studies["caffeine/Example"]
    assert current["status"] == "invalid" and current["stale"] is False
    assert current["_pending"] is False
    assert "unknown_column" in {problem["code"] for problem in current["problems"]}
    assert engine.report(job["report_id"])["reference_updated"] == (
        "Updated reference.json"
    )


def test_pipeline_runs_under_the_folder_lock(workspace, monkeypatch):
    engine, folder = workspace
    held = []

    class Lock:
        def __init__(self, path):
            held.append(("enter", path))

        def __enter__(self):
            return self

        def __exit__(self, *args):
            held.append(("exit", None))

    real = jobs.sync_and_format

    def pipeline(path, vocabulary, **kwargs):
        assert held and held[-1][0] == "enter"
        assert row(engine)["sync"]["status"] == "syncing"
        return real(path, vocabulary, **kwargs)

    monkeypatch.setattr(jobs, "folder_lock", Lock)
    monkeypatch.setattr(jobs, "sync_and_format", pipeline)
    settle(engine)
    assert run_next(engine)["status"] == "succeeded"
    assert held[0] == ("enter", folder)
    assert held[1] == ("exit", None)
    assert row(engine)["sync"]["status"] == "no_workbook"


def test_format_problems_fail_like_validation(workspace, monkeypatch):
    engine, folder = workspace
    table = folder / "timecourses_Fig1.tsv"
    table.write_text(
        table.read_text(encoding="utf-8") + "<<<<<<< HEAD\n",
        encoding="utf-8",
        newline="",
    )
    settle(engine)
    called = []
    monkeypatch.setattr(jobs, "prepare", lambda *a, **k: called.append(1))
    job = run_next(engine)
    assert job["status"] == "invalid" and not called
    assert row(engine)["status"] == "invalid"
    assert row(engine)["counts"]["errors"] >= 1
    report = engine.report(job["report_id"])
    assert any(
        issue["code"] == "merge_conflict" for issue in report["report"]["issues"]
    )


def test_sync_status_follows_the_workbook(workspace, sf_vocabulary):
    engine, folder = workspace
    assert sync_study(folder, sf_vocabulary).ok
    engine.scan()
    assert row(engine)["sync"] == {"status": "in_sync", "changes": 0, "conflicts": 0}
    set_workbook_mean(folder, "timecourses_Fig1", 5)
    engine.scan()
    assert row(engine)["sync"] == {"status": "changed", "changes": 1, "conflicts": 0}
    settle(engine)
    job = run_next(engine)
    assert job["status"] == "succeeded", job["message"]
    assert (
        "\t5\t"
        in (folder / "timecourses_Fig1.tsv").read_text(encoding="utf-8").splitlines()[2]
    )
    assert row(engine)["sync"]["status"] == "in_sync"
    (folder / f".~lock.{folder.name}.xlsx#").write_text(
        "open", encoding="utf-8", newline=""
    )
    set_mean(folder, "timecourses_Fig1.tsv", "6")
    engine.scan()
    assert row(engine)["sync"]["status"] == "workbook_open"


def test_regenerated_workbook_is_validated_by_the_same_job(workspace, sf_vocabulary):
    engine, folder = workspace
    assert format_folder(folder).ok
    assert sync_study(folder, sf_vocabulary).ok
    set_mean(folder, "timecourses_Fig1.tsv", "3")
    engine.scan()
    settle(engine)
    job = run_next(engine)
    assert job["status"] == "succeeded", job["message"]
    assert workbook_mean(folder, "timecourses_Fig1") == 3
    settle(engine)
    assert not engine.queue


def run_all(engine):
    jobs_run = []
    settle(engine)
    while engine.queue:
        jobs_run.append(run_next(engine))
        settle(engine)
    return jobs_run


@pytest.mark.parametrize(
    "moment", ["while formatting", "after the folder lock", "after the sync state"]
)
def test_a_save_during_the_job_is_never_absorbed(
    workspace, sf_vocabulary, monkeypatch, moment
):
    engine, folder = workspace
    assert format_folder(folder).ok
    assert sync_study(folder, sf_vocabulary).ok
    set_workbook_mean(folder, "timecourses_Fig1", 5)
    engine.scan()
    saved = []

    def save_again():
        if not saved:
            saved.append(moment)
            set_workbook_mean(folder, "timecourses_Fig1", 9)

    if moment == "while formatting":
        real_format = pipeline_module.format_folder

        def formatting(path, **kwargs):
            save_again()
            return real_format(path, **kwargs)

        monkeypatch.setattr(pipeline_module, "format_folder", formatting)
    elif moment == "after the folder lock":
        real_lock = jobs.folder_lock

        @contextmanager
        def lock(path):
            with real_lock(path):
                yield
            save_again()

        monkeypatch.setattr(jobs, "folder_lock", lock)
    else:
        real_describe = jobs.describe

        def describing(pipeline):
            save_again()
            return real_describe(pipeline)

        monkeypatch.setattr(jobs, "describe", describing)
    settle(engine)
    first = run_next(engine)
    assert saved
    # The job validated tables without the last save, so it must not count as current.
    assert first["status"] == "canceled", first["message"]
    assert engine.studies["caffeine/Example"]["_pending"] is True
    later = run_all(engine)
    assert [job["status"] for job in later] == ["succeeded"]
    assert (
        "\t9\t"
        in (folder / "timecourses_Fig1.tsv").read_text(encoding="utf-8").splitlines()[2]
    )
    assert row(engine)["sync"]["status"] == "in_sync"


def test_closing_the_workbook_regenerates_it(workspace, sf_vocabulary):
    engine, folder = workspace
    assert format_folder(folder).ok
    assert sync_study(folder, sf_vocabulary).ok
    engine.scan()
    assert [job["status"] for job in run_all(engine)] == ["succeeded"]
    lock = folder / f".~lock.{folder.name}.xlsx#"
    lock.write_text("open", encoding="utf-8", newline="")
    engine.scan()
    # Opening the workbook changes no source: no job.
    assert row(engine)["sync"]["status"] == "in_sync"
    settle(engine)
    assert not engine.queue
    set_mean(folder, "timecourses_Fig1.tsv", "6")
    engine.scan()
    assert row(engine)["sync"]["status"] == "workbook_open"
    assert [job["status"] for job in run_all(engine)] == ["succeeded"]
    assert row(engine)["sync"]["status"] == "workbook_open"
    assert workbook_mean(folder, "timecourses_Fig1") == 2
    lock.unlink()
    engine.scan()
    assert row(engine)["sync"] == {"status": "changed", "changes": 0, "conflicts": 0}
    assert engine.studies["caffeine/Example"]["_pending"] is True
    settle(engine)
    assert [job["action"] for job in engine.queue.values()] == ["validate"]
    assert [job["status"] for job in run_all(engine)] == ["succeeded"]
    assert workbook_mean(folder, "timecourses_Fig1") == 6
    assert row(engine)["sync"]["status"] == "in_sync"


def test_a_closed_workbook_behind_the_tables_is_changed(workspace, sf_vocabulary):
    engine, folder = workspace
    assert format_folder(folder).ok
    assert sync_study(folder, sf_vocabulary).ok
    set_mean(folder, "timecourses_Fig1.tsv", "6")
    engine.scan()
    assert row(engine)["sync"] == {"status": "changed", "changes": 0, "conflicts": 0}


def test_a_workbook_closed_during_the_job_is_regenerated(
    workspace, sf_vocabulary, monkeypatch
):
    engine, folder = workspace
    # The tables are not formatted yet, so the job writes them and scans again.
    assert sync_study(folder, sf_vocabulary).ok
    lock = folder / f".~lock.{folder.name}.xlsx#"
    lock.write_text("open", encoding="utf-8", newline="")
    set_mean(folder, "timecourses_Fig1.tsv", "6")
    engine.scan()
    real_format = pipeline_module.format_folder

    def formatting(path, **kwargs):
        lock.unlink(missing_ok=True)
        return real_format(path, **kwargs)

    monkeypatch.setattr(pipeline_module, "format_folder", formatting)
    settle(engine)
    assert run_next(engine)["status"] == "succeeded"
    assert engine.studies["caffeine/Example"]["_pending"] is True
    assert [job["status"] for job in run_all(engine)] == ["succeeded"]
    assert workbook_mean(folder, "timecourses_Fig1") == 6
    assert row(engine)["sync"]["status"] == "in_sync"


def test_a_failed_reference_lookup_still_syncs_the_workbook(workspace, sf_vocabulary):
    engine, folder = workspace
    assert sync_study(folder, sf_vocabulary).ok
    set_workbook_mean(folder, "timecourses_Fig1", 5)
    (folder / "reference.json").unlink()  # offline, nothing cached: the lookup fails
    engine.scan()
    settle(engine)
    job = run_next(engine)
    assert job["status"] == "failed" and "reference.json" in job["message"]
    assert (
        "\t5\t"
        in (folder / "timecourses_Fig1.tsv").read_text(encoding="utf-8").splitlines()[2]
    )
    assert row(engine)["sync"]["status"] == "in_sync"
    settle(engine)
    assert not engine.queue


def test_closing_the_workbook_keeps_a_queued_upload(workspace, sf_vocabulary):
    engine, folder = workspace
    assert format_folder(folder).ok
    assert sync_study(folder, sf_vocabulary).ok
    engine.scan()
    assert [job["status"] for job in run_all(engine)] == ["succeeded"]
    engine.offline = False
    engine.endpoint = "https://example.test"
    engine.api_key = "private-api-key"
    engine.account = "curator"
    engine.can_upload = True
    engine.set_mode("upload")
    lock = folder / f".~lock.{folder.name}.xlsx#"
    lock.write_text("open", encoding="utf-8", newline="")
    set_mean(folder, "timecourses_Fig1.tsv", "6")
    engine.scan()
    settle(engine)
    [upload] = engine.queue.values()
    assert upload["action"] == "upload"
    lock.unlink()
    engine.scan()
    assert row(engine)["sync"]["status"] == "changed"
    settle(engine)
    # The queued upload syncs, and so regenerates, the closed workbook itself.
    assert list(engine.queue.values()) == [upload]
    assert upload["status"] == "queued"


def test_a_workbook_closed_while_the_job_validates_is_regenerated(
    workspace, sf_vocabulary, monkeypatch
):
    engine, folder = workspace
    assert format_folder(folder).ok
    assert sync_study(folder, sf_vocabulary).ok
    lock = folder / f".~lock.{folder.name}.xlsx#"
    lock.write_text("open", encoding="utf-8", newline="")
    set_mean(folder, "timecourses_Fig1.tsv", "6")
    engine.scan()
    real_prepare = jobs.prepare

    def preparing(*args, **kwargs):
        # The watcher sees the workbook close while this job runs.
        lock.unlink()
        engine.scan()
        return real_prepare(*args, **kwargs)

    monkeypatch.setattr(jobs, "prepare", preparing)
    settle(engine)
    assert run_next(engine)["status"] == "succeeded"
    assert engine.studies["caffeine/Example"]["_pending"] is True
    monkeypatch.setattr(jobs, "prepare", real_prepare)
    assert [job["action"] for job in run_all(engine)] == ["validate"]
    assert workbook_mean(folder, "timecourses_Fig1") == 6
    assert row(engine)["sync"]["status"] == "in_sync"


def test_a_workbook_closed_during_an_upload_keeps_the_upload_pending(
    workspace, sf_vocabulary, monkeypatch
):
    engine, folder = workspace
    assert format_folder(folder).ok
    assert sync_study(folder, sf_vocabulary).ok
    engine.scan()
    assert [job["status"] for job in run_all(engine)] == ["succeeded"]
    engine.offline = False
    engine.endpoint = "https://example.test"
    engine.api_key = "private-api-key"
    engine.account = "curator"
    engine.can_upload = True
    engine.set_mode("upload")
    monkeypatch.setattr(engine, "_vocabulary", lambda client, **kwargs: sf_vocabulary)
    lock = folder / f".~lock.{folder.name}.xlsx#"
    lock.write_text("open", encoding="utf-8", newline="")
    set_mean(folder, "timecourses_Fig1.tsv", "6")
    engine.scan()
    real_prepare = jobs.prepare

    def preparing(*args, **kwargs):
        # The watcher sees the workbook close, then the curator pauses before the transfer.
        lock.unlink()
        engine.scan()
        engine.paused = True
        return real_prepare(*args, **kwargs)

    monkeypatch.setattr(jobs, "prepare", preparing)
    settle(engine)
    job = run_next(engine)
    assert job["action"] == "upload" and job["status"] == "canceled"
    engine.paused = False
    settle(engine)
    # The edit is still uploaded; that job's sync regenerates the workbook.
    assert [queued["action"] for queued in engine.queue.values()] == ["upload"]


def test_a_workbook_that_cannot_be_replaced_is_not_requeued(
    workspace, sf_vocabulary, monkeypatch
):
    from pkdb.studyformat import sync

    engine, folder = workspace
    assert format_folder(folder).ok
    assert sync_study(folder, sf_vocabulary).ok
    real_atomic_bytes = sync.atomic_bytes

    def atomic_bytes(path, data):
        if path == workbook_path(folder):
            # Windows keeps an open workbook locked, without a lock file of ours.
            raise PermissionError(13, "Permission denied", str(path))
        return real_atomic_bytes(path, data)

    monkeypatch.setattr(sync, "atomic_bytes", atomic_bytes)
    set_mean(folder, "timecourses_Fig1.tsv", "6")
    engine.scan()
    statuses = []
    for _ in range(5):
        settle(engine)
        if not engine.queue:
            break
        statuses.append(run_next(engine)["status"])
        engine.scan()
    assert statuses == ["succeeded"]
    assert row(engine)["sync"] == {
        "status": "workbook_open",
        "changes": 0,
        "conflicts": 0,
    }
    assert workbook_mean(folder, "timecourses_Fig1") == 2


def test_a_workbook_closed_during_a_failed_job_is_regenerated(
    workspace, sf_vocabulary, monkeypatch
):
    engine, folder = workspace
    assert format_folder(folder).ok
    assert sync_study(folder, sf_vocabulary).ok
    lock = folder / f".~lock.{folder.name}.xlsx#"
    lock.write_text("open", encoding="utf-8", newline="")
    set_mean(folder, "timecourses_Fig1.tsv", "6")
    engine.scan()
    assert row(engine)["sync"]["status"] == "workbook_open"
    real_pipeline = jobs.sync_and_format

    def failing(path, vocabulary, **kwargs):
        lock.unlink()
        raise OSError("disk unavailable")

    monkeypatch.setattr(jobs, "sync_and_format", failing)
    settle(engine)
    assert run_next(engine)["status"] == "failed"
    assert engine.studies["caffeine/Example"]["_pending"] is True
    monkeypatch.setattr(jobs, "sync_and_format", real_pipeline)
    assert [job["action"] for job in run_all(engine)] == ["validate"]
    assert workbook_mean(folder, "timecourses_Fig1") == 6
    assert row(engine)["sync"]["status"] == "in_sync"


def test_workbooks_are_planned_by_the_jobs_not_by_the_initial_scan(
    tmp_path_factory, make_study, valid_files, sf_vocabulary, monkeypatch
):
    folder = make_study(valid_files)
    assert format_folder(folder).ok
    assert sync_study(folder, sf_vocabulary).ok
    plans = []
    real_check = workspace_module.workbook_check

    def check(path, vocabulary):
        plans.append(path)
        return real_check(path, vocabulary)

    monkeypatch.setattr(workspace_module, "workbook_check", check)
    engine = module.CurationEngine(
        folder.parent.parent,
        state_dir=tmp_path_factory.mktemp("state"),
        offline=True,
        start=False,
    )
    try:
        monkeypatch.setattr(
            engine, "_local_vocabulary", lambda folder=None: sf_vocabulary
        )
        assert row(engine)["sync"]["status"] == "not_checked"
        settle(engine)
        assert run_next(engine)["status"] == "succeeded"
        # The job knows the state of the workbook from its sync.
        assert row(engine)["sync"] == {
            "status": "in_sync",
            "changes": 0,
            "conflicts": 0,
        }
        assert plans == []
        set_workbook_mean(folder, "timecourses_Fig1", 5)
        engine.scan()
        assert row(engine)["sync"]["status"] == "changed" and len(plans) == 1
    finally:
        engine.close()


def test_rows_that_no_job_checks_are_planned_a_few_per_tick(
    tmp_path_factory, make_study, valid_files, sf_vocabulary, monkeypatch
):
    folders = [make_study(valid_files, name=f"Example{index}") for index in range(7)]
    for folder in folders:
        assert format_folder(folder).ok
        assert sync_study(folder, sf_vocabulary).ok
    held = []
    real_check = workspace_module.workbook_check
    engine = module.CurationEngine(
        folders[0].parent.parent,
        state_dir=tmp_path_factory.mktemp("state"),
        offline=True,
        start=False,
    )

    def check(path, vocabulary):
        held.append((path.name, owned(engine.lock)))
        return real_check(path, vocabulary)

    def statuses():
        return {item["id"]: item["sync"]["status"] for item in row_list(engine)}

    try:
        monkeypatch.setattr(
            engine, "_local_vocabulary", lambda folder=None: sf_vocabulary
        )
        monkeypatch.setattr(workspace_module, "workbook_check", check)
        assert set(statuses().values()) == {"not_checked"}
        # Rows whose initial job is pending, queued or running are left to the job.
        engine._check_workbooks()
        settle(engine)
        engine._check_workbooks()
        assert held == [] and len(engine.queue) == 7
        engine.cancel_jobs([job["id"] for job in engine.queue.values()])
        engine.set_mode("off")
        engine.enqueue(["caffeine/Example6"], "validate")
        engine._check_workbooks()
        assert len(held) == 5 and not any(owned for _, owned in held)
        assert list(statuses().values()).count("in_sync") == 5
        engine._check_workbooks()
        assert len(held) == 6
        assert statuses() == {
            **{f"caffeine/Example{index}": "in_sync" for index in range(6)},
            "caffeine/Example6": "not_checked",
        }
        engine._check_workbooks()
        assert len(held) == 6
    finally:
        engine.close()


def test_a_plan_of_a_row_that_a_job_checked_meanwhile_is_dropped(
    workspace, sf_vocabulary, monkeypatch
):
    engine, folder = workspace
    engine.set_mode("off")
    real_state = engine._sync_state

    def sync_state(path):
        # A job records the state of the workbook while the watcher plans it.
        engine.studies["caffeine/Example"]["sync"] = {
            "status": "syncing",
            "changes": 0,
            "conflicts": 0,
        }
        return real_state(path)

    monkeypatch.setattr(engine, "_sync_state", sync_state)
    engine._check_workbooks()
    assert row(engine)["sync"]["status"] == "syncing"


def test_the_rescan_of_a_job_plans_no_workbook_under_the_engine_lock(
    workspace, sf_vocabulary, monkeypatch
):
    engine, folder = workspace
    assert sync_study(folder, sf_vocabulary).ok
    set_mean(folder, "timecourses_Fig1.tsv", "6")  # the job writes the workbook
    engine.scan()
    held = []
    real_check = workspace_module.workbook_check

    def check(path, vocabulary):
        held.append(owned(engine.lock))
        return real_check(path, vocabulary)

    monkeypatch.setattr(workspace_module, "workbook_check", check)
    (folder / "timecourses_Fig1.tsv").write_text(
        (folder / "timecourses_Fig1.tsv").read_text(encoding="utf-8") + "\n",
        encoding="utf-8",
        newline="",
    )  # not canonical: the job formats it and scans again
    settle(engine)
    assert run_next(engine)["status"] == "succeeded"
    assert held and not any(held)


def test_a_failing_workbook_plan_never_leaves_the_row_syncing(
    workspace, sf_vocabulary, monkeypatch
):
    engine, folder = workspace
    assert sync_study(folder, sf_vocabulary).ok

    def broken(path, vocabulary):
        raise RuntimeError("unexpected")

    def failing(path, vocabulary, **kwargs):
        raise OSError("disk unavailable")

    monkeypatch.setattr(workspace_module, "workbook_check", broken)
    set_mean(folder, "timecourses_Fig1.tsv", "6")
    engine.scan()
    assert row(engine)["sync"]["status"] == "unknown"
    monkeypatch.setattr(jobs, "sync_and_format", failing)
    settle(engine)
    job = run_next(engine)
    assert job["status"] == "failed"
    assert job["message"] == (
        f"The validation of {job['study_id']} could not read or validate the source "
        "files."
    )
    assert (job["error"]["code"], job["error"]["detail"]) == (
        "source_error",
        "OSError: disk unavailable",
    )
    assert engine.report(job["report_id"])["error"]["code"] == "source_error"
    # An expected failure keeps its traceback in the log file only.
    assert "OSError: disk unavailable" in engine.log.path.read_text(encoding="utf-8")
    assert row(engine)["sync"]["status"] == "unknown"


def test_a_waiting_row_recovers_when_its_files_return_to_the_last_scan(workspace):
    engine, folder = workspace
    assert format_folder(folder).ok
    engine.scan()
    (folder / "Example_Fig3.png").symlink_to(folder / "Example_Fig1.png")
    engine.scan()
    assert row(engine)["status"] == "waiting"
    settle(engine)
    assert not engine.queue
    # The files are as the scan before the symlink saw them.
    (folder / "Example_Fig3.png").unlink()
    engine.scan()
    assert row(engine)["status"] != "waiting"
    assert engine.studies["caffeine/Example"]["_pending"] is True
    settle(engine)
    assert list(engine.queue) == ["caffeine/Example"]


def test_a_symlink_in_the_study_waits_without_requeueing(workspace):
    engine, folder = workspace
    settle(engine)
    assert engine.queue
    # The symlink appears after the job was queued.
    (folder / "Example_Fig3.png").symlink_to(folder / "Example_Fig1.png")
    job = run_next(engine)
    assert job["status"] == "failed"
    waiting = engine.studies["caffeine/Example"]
    assert waiting["status"] == "waiting" and "Symlink" in waiting["message"]
    for _ in range(3):
        engine.scan()
        settle(engine)
        assert not engine.queue
    assert row(engine)["status"] == "waiting"
    (folder / "Example_Fig3.png").unlink()
    engine.scan()
    settle(engine)
    assert run_next(engine)["status"] == "succeeded"
    assert "message" not in engine.studies["caffeine/Example"]


@pytest.fixture
def seen(workspace, sf_vocabulary):
    """The engine and a study in step with its workbook, whose files the watcher has seen."""
    engine, folder = workspace
    assert format_folder(folder).ok
    assert sync_study(folder, sf_vocabulary).ok
    engine.scan()
    study = engine.studies["caffeine/Example"]
    study["_pending"], study["stale"] = False, False
    return engine, folder, study


def view_write(engine, folder, *, also=None):
    """Write the view of the workbook as Open tables does; `also` runs during the write."""

    def write():
        show_place(workbook_path(folder), "timecourses_Fig1", "O6")
        if also is not None:
            also()

    engine._unwatched_write(folder, workbook_path(folder).name, write)


def appended(path):
    """Save `path` with a line more, as an editor or a checkout changes it."""
    return lambda: path.write_bytes(path.read_bytes() + b"\n")


def test_a_view_write_is_no_change_for_the_watcher(seen):
    engine, folder, study = seen

    view_write(engine, folder)
    engine.scan()

    assert study["_pending"] is False and study["stale"] is False


def watcher_pass(engine):
    """One pass of the watcher without its failure handling."""
    engine.scan()
    engine.schedule_changes()
    engine._check_workbooks()


def test_a_view_write_changes_no_served_state_and_leaves_the_engine_idle(seen):
    engine, folder, study = seen
    for _ in range(2):
        watcher_pass(engine)
    served = state.state_row(study)

    view_write(engine, folder)
    # The next scan reads the workbook again, as it reads any file that was written.
    watcher_pass(engine)

    assert state.state_row(study) == served
    assert study["_pending"] is False and engine._next_job() is None
    revision = engine.lock.revision
    for _ in range(2):
        watcher_pass(engine)
    assert engine.lock.revision == revision


def test_an_edit_before_a_view_write_stays_a_change(seen):
    engine, folder, study = seen
    set_mean(folder, "timecourses_Fig1.tsv", "6")

    # No scan saw the edit before the view write.
    view_write(engine, folder)
    engine.scan()

    assert study["_pending"] is True


@pytest.mark.parametrize("name", ["timecourses_Fig1.tsv", "study.json"])
def test_an_edit_during_a_view_write_stays_a_change(seen, name):
    """An editor or a checkout saves meanwhile: only a change of the workbook is taken."""
    engine, folder, study = seen

    view_write(engine, folder, also=appended(folder / name))
    engine.scan()

    assert study["_pending"] is True


def test_a_file_added_during_a_view_write_is_a_change(seen):
    engine, folder, study = seen

    view_write(engine, folder, also=lambda: (folder / "notes.txt").write_text("new"))
    engine.scan()

    assert study["_pending"] is True


def test_a_cell_edit_after_a_view_write_is_a_change(seen):
    engine, folder, study = seen
    view_write(engine, folder)
    engine.scan()
    assert study["_pending"] is False

    set_workbook_mean(folder, "timecourses_Fig1", 7)
    engine.scan()

    assert study["_pending"] is True


def test_no_view_write_while_a_job_of_the_study_runs(seen):
    """A running job compares the source with the one it started on; a write stops it."""
    engine, folder, study = seen
    before = workbook_path(folder).read_bytes()
    engine.active = "caffeine/Example"

    with pytest.raises(state.JobRunning):
        view_write(engine, folder)
    assert workbook_path(folder).read_bytes() == before


def test_a_view_write_while_a_job_of_another_study_runs(seen):
    engine, folder, study = seen
    engine.active = "caffeine/Other"

    view_write(engine, folder)

    assert openpyxl.load_workbook(workbook_path(folder)).active.title == (
        "timecourses_Fig1"
    )


def test_opening_the_tables_at_a_place_while_a_job_runs_keeps_the_job(
    seen, monkeypatch
):
    engine, folder, study = seen
    engine.user = "curator"
    monkeypatch.setattr(studies, "open_path", lambda path, **kwargs: None)
    set_mean(folder, "timecourses_Fig1.tsv", "6")
    engine.scan()
    answers = []
    real_prepare = jobs.prepare

    def preparing(*args, **kwargs):
        # The curator opens the tables at a cell while the job validates.
        place = {"action": "open", "file": "timecourses_Fig1.tsv", "cell": "O6"}
        answers.append(engine.tables_action("caffeine/Example", place))
        return real_prepare(*args, **kwargs)

    monkeypatch.setattr(jobs, "prepare", preparing)
    settle(engine)

    job = run_next(engine)

    assert job["status"] == "succeeded", job["message"]
    [answer] = answers
    assert answer["opened"]
    assert answer["notice"] == (
        "The study is being validated or uploaded. Go to sheet timecourses_Fig1, cell O6."
    )
    # The workbook opens as it is.
    assert openpyxl.load_workbook(workbook_path(folder)).active.title == "subjects"
    assert study["_pending"] is False
