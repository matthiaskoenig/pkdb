"""Hidden TSV tables generated from a study's Excel workbook."""

from datetime import datetime, time

import openpyxl
import pytest

from pkdb.tsv import sync_tsvs


def workbook(folder, sheets):
    book = openpyxl.Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
    book.save(folder / f"{folder.name}.xlsx")


@pytest.fixture
def folder(tmp_path):
    folder = tmp_path / "Example2020"
    folder.mkdir()
    (folder / "study.json").write_text('{"name": "Example2020"}')
    workbook(
        folder,
        {
            "Tab1": [
                ["Description row is skipped"],
                ["study", "group", "count", "mean", None],
                ["Example2020", "all", 12, 1.5, None],
                ["# a comment row"],
                [None, None, None, None, None],
                ["Example2020", "young", None, 2, None],
            ],
            "Fig2": [["notes"], ["study", "time"], ["Example2020", 0.5]],
            "Notes": [["ignored"], ["a"], [1]],
            "FigEmpty": [["only a description"]],
        },
    )
    return folder


def test_sheets_are_exported_like_the_legacy_upload(folder):
    change = sync_tsvs(folder)
    assert change == "Created 2 TSV files from Example2020.xlsx"
    assert (folder / ".Example2020_Tab1.tsv").read_text() == (
        "study\tgroup\tcount\tmean\n"
        "Example2020\tall\t12.0\t1.5\n"
        "Example2020\tyoung\tNA\t2.0\n"
    )
    assert (folder / ".Example2020_Fig2.tsv").read_text() == (
        "study\ttime\nExample2020\t0.5\n"
    )
    assert sorted(p.name for p in folder.glob(".*.tsv")) == [
        ".Example2020_Fig2.tsv",
        ".Example2020_Tab1.tsv",
    ]
    before = {p.name: p.stat().st_mtime_ns for p in folder.iterdir()}
    assert sync_tsvs(folder) is None
    assert {p.name: p.stat().st_mtime_ns for p in folder.iterdir()} == before


def test_changed_and_orphaned_tables_are_updated_and_removed(folder):
    sync_tsvs(folder)
    (folder / ".Example2020_Tab1.tsv").write_text("stale\n")
    (folder / ".Example2020_TabOld.tsv").write_text("orphan\n")
    (folder / ".Example2020_TabUsed.tsv").write_text("referenced\n")
    (folder / ".Example2020_TabSheet.tsv").write_text("sheet deleted\n")
    (folder / "study.json").write_text(
        '{"name": "Example2020", "outputset": {"outputs": '
        '[{"source": ".Example2020_TabUsed.tsv"}, {"source": "TabSheet"}]}}'
    )
    assert sync_tsvs(folder) == (
        "Updated 1 and removed 1 TSV files from Example2020.xlsx"
    )
    assert (folder / ".Example2020_Tab1.tsv").read_text().startswith("study\t")
    assert not (folder / ".Example2020_TabOld.tsv").exists()
    assert (folder / ".Example2020_TabUsed.tsv").exists()
    assert (folder / ".Example2020_TabSheet.tsv").exists()


def test_folders_without_workbook_are_untouched(tmp_path):
    (tmp_path / ".Other_Tab1.tsv").write_text("kept\n")
    assert sync_tsvs(tmp_path) is None
    assert (tmp_path / ".Other_Tab1.tsv").read_text() == "kept\n"


def test_unreadable_workbook_is_reported(tmp_path):
    folder = tmp_path / "Broken2020"
    folder.mkdir()
    (folder / "Broken2020.xlsx").write_bytes(b"not a workbook")
    with pytest.raises(ValueError, match="Cannot read Broken2020.xlsx"):
        sync_tsvs(folder)


@pytest.mark.parametrize("study_folder", ["xlsx"], indirect=True)
def test_validate_creates_tsv_tables_before_validation(
    study_folder, vocabulary, tmp_path, capsys
):
    import json

    from pkdb.cli import main

    book = openpyxl.load_workbook(study_folder / "Example.xlsx")
    book.copy_worksheet(book["Results"]).title = "TabResults"
    book.save(study_folder / "Example.xlsx")
    lock = tmp_path / "vocabulary.lock.json"
    vocabulary.save(lock)
    args = ["validate", str(study_folder), "--vocabulary", str(lock), "--offline"]
    assert main([*args, "--format", "json"]) == 0
    result = json.loads(capsys.readouterr().out)
    assert result["tables_updated"] == "Created 1 TSV files from Example.xlsx"
    assert (
        (study_folder / ".Example_TabResults.tsv")
        .read_text()
        .startswith("time\tmean\n")
    )
    assert main([*args, "--format", "json"]) == 0
    assert "tables_updated" not in json.loads(capsys.readouterr().out)


N = None
# Expected tables were written by the earlier pandas implementation; the polars
# implementation must reproduce them so existing TSV files stay unchanged.
CONVERSIONS = {
    "integers": (
        [["n", "t"], [1, " 7 "], [2, "+5"], [3, "007"]],
        "n\tt\n1\t7\n2\t5\n3\t7\n",
    ),
    "floats": (
        [["a", "b", "c"], [1, "1.", 1e-05], [2, ".5", 2.5], [N, "Infinity", 1e16]],
        "a\tb\tc\n1.0\t1.0\t1e-05\n2.0\t0.5\t2.5\nNA\tinf\t1e+16\n",
    ),
    "negative zero": ([["a", "b"], [-0.0, -0.0], [1, 2.5]], "a\tb\n0\t0.0\n1\t2.5\n"),
    "missing text": (
        [["a", "b"], ["NA", "x"], ["nan", "None"], ["#N/A", "null"], [1, "N/A"]],
        "a\tb\nNA\tx\n1.0\tNA\n",
    ),
    "no-break space": ([["a"], [1], ["11.4\xa0"]], "a\n1\n11.4\xa0\n"),
    "booleans": (
        [
            ["a", "b", "c", "d"],
            [True, True, "true", True],
            [False, N, N, "true"],
            [True, False, "FALSE", False],
        ],
        "a\tb\tc\td\nTrue\t1.0\tTrue\tTrue\nFalse\tNA\tNA\ttrue\nTrue\t0.0\tFalse\tFalse\n",
    ),
    "first equal value": (
        [["a", "b"], [False, 0], [0, False], ["x", "x"], [True, 1]],
        "a\tb\nFalse\t0\nFalse\t0\nx\tx\nTrue\t1\n",
    ),
    "dates": (
        [
            ["a", "b", "c", "d"],
            [
                datetime(2020, 1, 2),
                datetime(2020, 1, 2, 3, 4, 5),
                datetime(2020, 1, 2),
                time(12, 30),
            ],
            [datetime(2020, 1, 3), datetime(2020, 1, 3), "x", time(1, 2, 3)],
        ],
        "a\tb\tc\td\n2020-01-02\t2020-01-02 03:04:05\t2020-01-02 00:00:00\t12:30:00\n"
        "2020-01-03\t2020-01-03 00:00:00\tx\t01:02:03\n",
    ),
    "comments": (
        [
            ["# skipped line"],
            ["a", "b#ignored", "dropped"],
            ["# skipped as well"],
            [1, 2, 3],
            ["x # note", 5, 6],
            ["# kept as an empty line"],
            [3, 4],
        ],
        "a\tb\n2.0\t3.0\n4.0\tNA\n",
    ),
    "quoting": (
        [["a", "b"], ["with\ttab", 'say "hi"'], ["line\nbreak", "plain"]],
        'a\tb\n"with\ttab"\t"say ""hi"""\n"line\nbreak"\tplain\n',
    ),
    "names": ([["a", N, "a", 5, "a"], [1, 2, 3, 4, 5]], "a\ta.1\t5\ta.2\n1\t3\t4\t5\n"),
}


@pytest.mark.parametrize("rows,expected", CONVERSIONS.values(), ids=CONVERSIONS)
def test_cells_are_converted_like_the_pandas_tables(tmp_path, rows, expected):
    folder = tmp_path / "Case2020"
    folder.mkdir()
    workbook(folder, {"Tab1": [["description"], *rows]})
    sync_tsvs(folder)
    assert (folder / ".Case2020_Tab1.tsv").read_text(encoding="utf-8") == expected


def test_lines_longer_than_the_header_are_reported(tmp_path):
    folder = tmp_path / "Long2020"
    folder.mkdir()
    # Comments shorten the header and the first data line, as pandas read them.
    rows = [["a", "b#c"], [1, 2, "#x"], [1, 2, 3, "#y"], [5, 6, 7, 8]]
    workbook(folder, {"Tab1": [["description"], *rows]})
    with pytest.raises(
        ValueError, match="Sheet Tab1 row 4 has 3 fields, but its header has 2"
    ):
        sync_tsvs(folder)
