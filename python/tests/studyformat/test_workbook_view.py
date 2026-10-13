"""The view of a workbook: the sheet and the cell that it opens at (#939)."""

import io
import os
import sys
import zipfile
from types import SimpleNamespace

import openpyxl
import pytest

from pkdb.studyformat.sync import sync_study
from pkdb.studyformat.workbook import view
from pkdb.studyformat.workbook.base import workbook_path
from pkdb.studyformat.workbook.view import (
    HiddenSheet,
    NoSheet,
    ViewError,
    WorkbookOpen,
    parse_cell,
    show_place,
    with_view,
)

MAIN = b"http://schemas.openxmlformats.org/spreadsheetml/2006/main"
DECLARATION = b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
WORKSHEET = b"application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"
# A workbook as another writer could save it: the main namespace with the prefix x, namespaces
# that only mc:Ignorable names, single quotes, a pane with an end tag and no selection, and a
# sheet view without content.
OTHER_WRITER = {
    "[Content_Types].xml": DECLARATION
    + b'<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    b'<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.'
    b'relationships+xml"/><Default Extension="xml" ContentType="application/xml"/>'
    b'<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-'
    b'officedocument.spreadsheetml.sheet.main+xml"/>'
    + b"".join(
        b'<Override PartName="/xl/worksheets/sheet%d.xml" ContentType="%b"/>'
        % (number, WORKSHEET)
        for number in (1, 2, 3)
    )
    + b"</Types>",
    "_rels/.rels": DECLARATION
    + b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/'
    b'relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/'
    b'officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
    b"</Relationships>",
    "xl/workbook.xml": DECLARATION
    + b'<x:workbook xmlns:x="%b" xmlns:r="http://schemas.openxmlformats.org/'
    b'officeDocument/2006/relationships" xmlns:mc="http://schemas.openxmlformats.org/'
    b'markup-compatibility/2006" mc:Ignorable="x15 xr" xmlns:x15="http://schemas.'
    b'microsoft.com/office/spreadsheetml/2010/11/main" xmlns:xr="http://schemas.'
    b'microsoft.com/office/spreadsheetml/2014/revision"><x:bookViews><x:workbookView '
    b'xWindow="0" yWindow="0" firstSheet="2" activeTab="0" xr:uid="{A}"/>'
    b'</x:bookViews><x:sheets><x:sheet name="notes" sheetId="1" r:id="rId1"/>'
    b'<x:sheet name="outputs_Fig6" sheetId="2" r:id="rId2"/><x:sheet name="raw" '
    b'sheetId="3" r:id="rId3"/></x:sheets></x:workbook>' % MAIN,
    "xl/_rels/workbook.xml.rels": DECLARATION
    + b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/'
    b'relationships">'
    + b"".join(
        b'<Relationship Id="rId%d" Type="http://schemas.openxmlformats.org/'
        b'officeDocument/2006/relationships/worksheet" Target="worksheets/sheet%d.xml"/>'
        % (number, number)
        for number in (3, 1, 2)
    )
    + b"</Relationships>",
    "xl/worksheets/sheet1.xml": DECLARATION
    + b'<x:worksheet xmlns:x="%b" xmlns:mc="http://schemas.openxmlformats.org/'
    b'markup-compatibility/2006" mc:Ignorable="x14ac" xmlns:x14ac="http://schemas.'
    b'microsoft.com/office/spreadsheetml/2009/9/ac"><x:sheetViews><x:sheetView '
    b'tabSelected=\'1\' zoomScale="120" workbookViewId="0"/></x:sheetViews>'
    b'<x:sheetFormatPr defaultRowHeight="15" x14ac:dyDescent="0.25"/><x:sheetData/>'
    b"</x:worksheet>" % MAIN,
    "xl/worksheets/sheet2.xml": DECLARATION
    + b'<x:worksheet xmlns:x="%b"><x:sheetViews><x:sheetView workbookViewId="0">'
    b'<x:pane ySplit="1" topLeftCell="A2" activePane="bottomLeft" state="frozen">'
    b"</x:pane></x:sheetView></x:sheetViews><x:sheetData/></x:worksheet>" % MAIN,
    "xl/worksheets/sheet3.xml": DECLARATION
    + b'<worksheet xmlns="%b"><sheetViews><sheetView workbookViewId="0" />'
    b"</sheetViews><sheetData/></worksheet>" % MAIN,
}


@pytest.fixture
def workbook(valid_study, sf_vocabulary):
    """The workbook of the valid study as the first sync creates it, open at subjects."""
    assert sync_study(valid_study, sf_vocabulary).ok
    return workbook_path(valid_study)


def package(parts):
    data = io.BytesIO()
    with zipfile.ZipFile(data, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in parts.items():
            archive.writestr(name, content)
    return data.getvalue()


def entries(data):
    """Name, compression, time and content of each entry of a zip, in order."""
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return [
            (info.filename, info.compress_type, info.date_time, archive.read(info))
            for info in archive.infolist()
        ]


def selected(sheet):
    """The active cell of the active pane of a sheet, as openpyxl reads it."""
    view = sheet.sheet_view
    active = view.pane.activePane if view.pane is not None else "topLeft"
    return {
        (selection.pane or "topLeft"): selection.activeCell
        for selection in view.selection
    }[active]


def test_the_workbook_opens_at_the_sheet_and_the_cell(workbook):
    assert show_place(workbook, "timecourses_Fig1", "O6")

    book = openpyxl.load_workbook(workbook)
    sheet = book.active
    assert sheet.title == "timecourses_Fig1"
    assert [page.title for page in book.worksheets if page.sheet_view.tabSelected] == [
        "timecourses_Fig1"
    ]
    assert selected(sheet) == "O6"
    assert sheet.sheet_view.selection[0].sqref == "O6"
    # The header stays frozen; the sheet shows three rows above the cell and two columns
    # before it.
    pane = sheet.sheet_view.pane
    assert (pane.ySplit, pane.state, pane.topLeftCell) == (1, "frozen", "M3")
    assert sheet.sheet_view.topLeftCell == "M1"


def test_only_the_view_changes(workbook):
    before = workbook.read_bytes()

    assert show_place(workbook, "timecourses_Fig1", "O6")

    old, new = entries(before), entries(workbook.read_bytes())
    # Every entry keeps its place, name, compression and time.
    assert [entry[:3] for entry in new] == [entry[:3] for entry in old]
    changed = {
        name: (previous, content)
        for (name, *_, content), (*_, previous) in zip(new, old, strict=True)
        if content != previous
    }
    # The workbook and the sheets that lost and got the selection: subjects and the timecourse.
    assert sorted(changed) == [
        "xl/workbook.xml",
        "xl/worksheets/sheet1.xml",
        "xl/worksheets/sheet5.xml",
    ]
    for name, (previous, content) in changed.items():
        start, end = (
            (b"<bookViews>", b"</bookViews>")
            if name == "xl/workbook.xml"
            else (b"<sheetViews>", b"</sheetViews>")
        )
        assert content.partition(start)[0] == previous.partition(start)[0]
        assert content.partition(end)[2] == previous.partition(end)[2]


def test_without_a_cell_only_the_sheet_changes(workbook):
    assert show_place(workbook, "outputs_Tab2")

    sheet = openpyxl.load_workbook(workbook).active
    assert sheet.title == "outputs_Tab2"
    assert selected(sheet) == "A1"
    assert sheet.sheet_view.pane.topLeftCell == "A2"


@pytest.mark.parametrize(
    ("cell", "pane", "view"),
    [("B4", "A2", "A1"), ("C1", "A2", "A1"), ("Z40", "X37", "X1")],
)
def test_the_sheet_scrolls_to_show_the_cell(workbook, cell, pane, view):
    assert show_place(workbook, "timecourses_Fig1", cell)

    sheet = openpyxl.load_workbook(workbook).active
    assert selected(sheet) == cell
    assert sheet.sheet_view.pane.topLeftCell == pane
    assert sheet.sheet_view.topLeftCell == view


def test_the_same_place_writes_nothing(workbook):
    assert show_place(workbook, "timecourses_Fig1", "O6")
    written = workbook.stat().st_mtime_ns

    assert not show_place(workbook, "timecourses_Fig1", "O6")
    assert workbook.stat().st_mtime_ns == written


@pytest.mark.parametrize(
    "sheet", ["outputs_Tab9", "_lists", "_base", "Timecourses_Fig1", ""]
)
def test_a_sheet_that_the_workbook_lacks_or_hides_is_no_place(workbook, sheet):
    before = workbook.read_bytes()

    with pytest.raises(NoSheet):
        show_place(workbook, sheet, "A1")
    assert workbook.read_bytes() == before


def test_a_hidden_sheet_of_a_table_is_no_place_of_its_own_kind():
    """The curator hid the sheet: it is there, but the workbook cannot open at it."""
    parts = changed(
        "xl/workbook.xml",
        b'<x:sheet name="outputs_Fig6" sheetId="2" r:id="rId2"/>',
        b'<x:sheet name="outputs_Fig6" sheetId="2" state="hidden" r:id="rId2"/>',
    )

    with pytest.raises(HiddenSheet) as error:
        with_view(package(parts), "outputs_Fig6", "Y97")
    assert isinstance(error.value, NoSheet)
    with pytest.raises(NoSheet) as error:
        with_view(package(parts), "outputs_Fig7", "Y97")
    assert not isinstance(error.value, HiddenSheet)


def test_an_open_workbook_is_left_as_it_is(workbook):
    workbook.with_name(f".~lock.{workbook.name}#").write_text(
        "curator", encoding="utf-8"
    )
    before = workbook.read_bytes()

    with pytest.raises(WorkbookOpen):
        show_place(workbook, "timecourses_Fig1", "O6")
    assert workbook.read_bytes() == before


def test_a_workbook_saved_meanwhile_is_left_as_it_is(workbook, monkeypatch):
    """A save while the app read the workbook wins: the app replaces nothing."""
    compute = view.with_view

    def saved_meanwhile(data, sheet, cell=None):
        viewed = compute(data, sheet, cell)
        book = openpyxl.load_workbook(workbook)
        book["subjects"]["A2"] = "saved meanwhile"
        book.save(workbook)
        status = workbook.stat()
        # A save within the same tick of the clock can keep the time.
        os.utime(workbook, ns=(status.st_atime_ns, status.st_mtime_ns + 10**9))
        return viewed

    monkeypatch.setattr(view, "with_view", saved_meanwhile)

    with pytest.raises(ViewError, match=r"^Example.xlsx was saved meanwhile$"):
        show_place(workbook, "timecourses_Fig1", "O6")
    book = openpyxl.load_workbook(workbook)
    assert book["subjects"]["A2"].value == "saved meanwhile"
    assert book.active.title == "subjects"


@pytest.mark.skipif(
    sys.platform == "win32" or os.geteuid() == 0,
    reason="POSIX permissions, which root ignores",
)
def test_a_workbook_in_a_folder_without_write_permission_is_not_open(workbook):
    """Only Windows refuses a replace for an open workbook; elsewhere it is no open one."""
    before = workbook.read_bytes()
    workbook.parent.chmod(0o555)
    try:
        with pytest.raises(
            ViewError, match=r"^Example.xlsx cannot be replaced: permission denied$"
        ):
            show_place(workbook, "timecourses_Fig1", "O6")
    finally:
        workbook.parent.chmod(0o755)
    assert workbook.read_bytes() == before


def test_a_workbook_that_windows_refuses_to_replace_is_open(workbook, monkeypatch):
    """Windows keeps a workbook locked while a program has it open."""

    def locked(path, data):
        raise PermissionError(13, "Permission denied")

    monkeypatch.setattr(view, "atomic_bytes", locked)
    monkeypatch.setattr(view, "os", SimpleNamespace(name="nt"))
    before = workbook.read_bytes()

    with pytest.raises(WorkbookOpen):
        show_place(workbook, "timecourses_Fig1", "O6")
    assert workbook.read_bytes() == before


@pytest.mark.parametrize(
    "cell", ["y97", "A0", "XFE1", "A1048577", "Y", "97", "A1:B2", " A1", "A01"]
)
def test_a_cell_that_is_no_workbook_address_is_refused(workbook, cell):
    with pytest.raises(ValueError):
        parse_cell(cell)
    with pytest.raises(ValueError):
        show_place(workbook, "timecourses_Fig1", cell)


def test_the_last_cell_of_a_sheet_is_a_place():
    assert parse_cell("XFD1048576") == (16384, 1048576)


def test_the_parts_of_another_writer_keep_their_prefixes_and_namespaces():
    data = with_view(package(OTHER_WRITER), "outputs_Fig6", "Y97")

    assert data is not None
    parts = {name: content for name, *_, content in entries(data)}
    assert parts["xl/workbook.xml"] == OTHER_WRITER["xl/workbook.xml"].replace(
        b'firstSheet="2" activeTab="0"', b'firstSheet="1" activeTab="1"'
    )
    assert parts["xl/worksheets/sheet1.xml"] == OTHER_WRITER[
        "xl/worksheets/sheet1.xml"
    ].replace(b"<x:sheetView tabSelected='1' zoomScale", b"<x:sheetView zoomScale")
    assert parts["xl/worksheets/sheet2.xml"] == OTHER_WRITER[
        "xl/worksheets/sheet2.xml"
    ].replace(
        b'<x:sheetView workbookViewId="0"><x:pane ySplit="1" topLeftCell="A2" '
        b'activePane="bottomLeft" state="frozen"></x:pane>',
        b'<x:sheetView workbookViewId="0" tabSelected="1" topLeftCell="W1"><x:pane '
        b'ySplit="1" topLeftCell="W94" activePane="bottomLeft" state="frozen"></x:pane>'
        b'<x:selection pane="bottomLeft" activeCell="Y97" sqref="Y97"/>',
    )
    assert parts["xl/worksheets/sheet3.xml"] == OTHER_WRITER["xl/worksheets/sheet3.xml"]
    book = openpyxl.load_workbook(io.BytesIO(data))
    assert book.active.title == "outputs_Fig6"
    assert selected(book.active) == "Y97"


def test_a_sheet_view_without_content_gets_a_selection():
    data = with_view(package(OTHER_WRITER), "raw", "C40")

    assert data is not None
    parts = {name: content for name, *_, content in entries(data)}
    assert parts["xl/worksheets/sheet3.xml"] == OTHER_WRITER[
        "xl/worksheets/sheet3.xml"
    ].replace(
        b'<sheetView workbookViewId="0" />',
        b'<sheetView workbookViewId="0" tabSelected="1" topLeftCell="A37">'
        b'<selection activeCell="C40" sqref="C40"/></sheetView>',
    )
    assert openpyxl.load_workbook(io.BytesIO(data)).active.title == "raw"


@pytest.mark.parametrize(
    ("cell", "pane"), [("Z40", "X37"), ("B1", "C2"), ("D30", "C27"), ("K2", "I2")]
)
def test_frozen_rows_and_columns_stay_and_the_pane_scrolls_to_the_cell(cell, pane):
    book = openpyxl.Workbook()
    book.active.title = "first"
    frozen = book.create_sheet("frozen")
    # Two columns and a row stay, as in the sheet of a table with a frozen label.
    frozen.freeze_panes = "C2"
    data = io.BytesIO()
    book.save(data)

    viewed = with_view(data.getvalue(), "frozen", cell)

    assert viewed is not None
    sheet = openpyxl.load_workbook(io.BytesIO(viewed)).active
    assert sheet.title == "frozen"
    assert selected(sheet) == cell
    split = sheet.sheet_view.pane
    assert (split.xSplit, split.ySplit, split.state, split.activePane) == (
        2,
        1,
        "frozen",
        "bottomRight",
    )
    assert split.topLeftCell == pane
    # The top left pane holds only frozen cells and never scrolls.
    assert sheet.sheet_view.topLeftCell == "A1"


def changed(name, old, new):
    """The parts of OTHER_WRITER with `old` replaced by `new` in the part `name`."""
    assert old in OTHER_WRITER[name]
    return {**OTHER_WRITER, name: OTHER_WRITER[name].replace(old, new)}


SHEET2 = "xl/worksheets/sheet2.xml"


@pytest.mark.parametrize("split", ["20000", "1e300"])
def test_a_frozen_pane_outside_the_sheet_is_a_view_error(split):
    parts = changed(
        SHEET2, b'<x:pane ySplit="1"', f'<x:pane xSplit="{split}" ySplit="1"'.encode()
    )

    with pytest.raises(ViewError, match="is no column"):
        with_view(package(parts), "outputs_Fig6", "Y97")


@pytest.mark.parametrize(
    ("parts", "message"),
    [
        (None, "the workbook is no xlsx file: "),
        (
            {
                name: part
                for name, part in OTHER_WRITER.items()
                if name != "xl/workbook.xml"
            },
            "the workbook cannot be read: ",
        ),
        (
            changed("xl/workbook.xml", b"</x:sheets>", b"</x:sheet>"),
            "a part of the workbook is no XML: mismatched tag",
        ),
        (
            changed("xl/_rels/workbook.xml.rels", b"</Relationships>", b""),
            "the relationships of the workbook are no XML: ",
        ),
        (
            changed(SHEET2, b"</x:pane>", b"</x:sheetView>"),
            "a part of the workbook is no XML: mismatched tag",
        ),
        (
            changed(
                "xl/workbook.xml",
                b'<x:bookViews><x:workbookView xWindow="0" yWindow="0" firstSheet="2" '
                b'activeTab="0" xr:uid="{A}"/></x:bookViews>',
                b"",
            ),
            "the workbook has no view",
        ),
        (
            changed("xl/_rels/workbook.xml.rels", b'Id="rId2"', b'Id="rId9"'),
            "the sheet outputs_Fig6 has no part",
        ),
        (
            changed(
                SHEET2,
                b'<x:sheetViews><x:sheetView workbookViewId="0"><x:pane ySplit="1" '
                b'topLeftCell="A2" activePane="bottomLeft" state="frozen"></x:pane>'
                b"</x:sheetView></x:sheetViews>",
                b"",
            ),
            "the sheet has no view",
        ),
        (
            changed("xl/workbook.xml", b'firstSheet="2"', b'firstSheet="second"'),
            "the firstSheet 'second' of the workbook is no number",
        ),
        (
            changed(SHEET2, b'ySplit="1"', b'ySplit="one"'),
            "the ySplit 'one' of the workbook is no number",
        ),
        (
            changed(SHEET2, b'ySplit="1"', b'ySplit="NaN"'),
            "the ySplit 'NaN' of the workbook is no number",
        ),
        (
            changed(SHEET2, b'<x:pane ySplit="1"', b'<x:pane xSplit="INF" ySplit="1"'),
            "the xSplit 'INF' of the workbook is no number",
        ),
    ],
)
def test_a_workbook_whose_view_cannot_be_set_is_refused(parts, message):
    """ViewError, never the ValueError of a cell that is no workbook address."""
    data = b"PK no zip" if parts is None else package(parts)

    with pytest.raises(ViewError) as error:
        with_view(data, "outputs_Fig6", "Y97")
    assert str(error.value).startswith(message)


def test_libreoffice_opens_the_workbook_at_the_place(workbook, libreoffice_resave):
    """LibreOffice reads the view: a save of LibreOffice keeps the sheet and the cell."""
    assert show_place(workbook, "timecourses_Fig1", "O6")

    sheet = openpyxl.load_workbook(libreoffice_resave(workbook)).active

    assert sheet.title == "timecourses_Fig1"
    assert selected(sheet) == "O6"


def test_a_workbook_that_libreoffice_saved_opens_at_the_place(
    workbook, libreoffice_resave
):
    """The parts as LibreOffice writes them: tabSelected="true", a selection per pane."""
    saved = libreoffice_resave(workbook)

    assert show_place(saved, "outputs_Tab2", "N3")

    book = openpyxl.load_workbook(saved)
    assert book.active.title == "outputs_Tab2"
    assert [page.title for page in book.worksheets if page.sheet_view.tabSelected] == [
        "outputs_Tab2"
    ]
    assert selected(book.active) == "N3"
