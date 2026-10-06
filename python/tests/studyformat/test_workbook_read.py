"""Tests of reading the study workbook back into canonical TSV tables.

Performance budgets, checked on CI (plan Review Focus 5):

- A 20,000-row `outputs_Tab1` sheet reads in under 15 seconds.
- A workbook whose stored `<dimension>` claims `A1:Z1048576` but which holds
  10 rows reads in under 2 seconds.
"""

import re
import time
import warnings
from datetime import UTC, datetime
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl
import pytest
from openpyxl.utils import get_column_letter

from pkdb.schemas.validation import StudyValidationError
from pkdb.studyformat.formatter import format_folder
from pkdb.studyformat.issues import REPEATED_ISSUES
from pkdb.studyformat.tables import TABLES
from pkdb.studyformat.text import render_tsv
from pkdb.studyformat.workbook import read
from pkdb.studyformat.workbook.base import BASE_SHEET
from pkdb.studyformat.workbook.read import SheetTable, read_workbook
from pkdb.studyformat.workbook.write import build_workbook

GENERATION = "0123456789abcdef0123456789abcdef"
CREATED = datetime(2026, 10, 6, 8, 30, tzinfo=UTC)
STUDY = "Example"
# Rows of the edge table of the lossless round trip (plan Review Focus 2): numbers
# with 16 and 17 significant digits, texts a spreadsheet would convert, NR in time.
EDGE_ROWS = (
    {
        "subjects": "all",
        "interventions": "D1",
        "measurement": "cmax",
        "time": "NR",
        "time_unit": "NR",
        "count": "12",
        "mean": "0.30000000000000004",
        "sd": "1234567890123456",
        "se": "123456789012345",
        "cv": "2.5e-07",
        "gmean": "1e+16",
        "median": "-3",
        "min": "abc",
        "max": "=1",
        "unit": "mg/l",
        "comment": "1-2",
    },
    {
        "subjects": "S1",
        "interventions": "D1",
        "measurement": "cmax",
        "time": "1.5",
        "mean": "1.0000000000000002",
        "gsd": "9007199254740993",
        "choice": "TRUE",
        "tissue": "007",
        "method": "=x",
        "comment": "1e3",
    },
    {
        "subjects": "S2",
        "interventions": "D1",
        "measurement": "cmax",
        "mean": "0.1",
        "choice": "FALSE",
        "unit": "50%",
        # U+0085 and U+2028 end a line for str.splitlines, but not in a TSV file.
        "comment": f"'quoted 1/2 3.0 +1 {chr(0x1F600)}{chr(0xFFFD)}{chr(0x85)}{chr(0x2028)}x",
    },
    # Spreadsheet applications read _x0041_ as the escape of A.
    {
        "subjects": "S2",
        "interventions": "D1",
        "measurement": "cmax",
        "time": "2",
        "tissue": "_x0041_",
        "method": "a_x005F_b",
        "comment": "_x005F_x0041_",
    },
    # Text that looks like an error value of a spreadsheet application.
    {
        "subjects": "S2",
        "interventions": "D1",
        "measurement": "cmax",
        "time": "3",
        "tissue": "#N/A",
        "method": "#DIV/0!",
        "comment": "#REF!",
    },
)


def tables_of(folder):
    return {
        path.name: path.read_text(encoding="utf-8")
        for path in sorted(folder.glob("*.tsv"))
    }


def build(tables, vocabulary, path, **options) -> Path:
    result = build_workbook(
        tables, vocabulary, generation=GENERATION, created=CREATED, **options
    )
    assert result.data is not None, result.issues
    path.write_bytes(result.data)
    return path


def edit(path, change) -> Path:
    workbook = openpyxl.load_workbook(path)
    change(workbook)
    workbook.save(path)
    return path


def lines(text):
    """The lines of canonical TSV text; str.splitlines would also split at U+2028."""
    return text.removesuffix("\n").split("\n")


def letter(kind, name):
    return get_column_letter(TABLES[kind].names.index(name) + 1)


def texts(content):
    return {file: table.text for file, table in content.tables.items()}


def codes(content):
    return [issue.code for issue in content.issues]


def only(content, code):
    found = [issue for issue in content.issues if issue.code == code]
    assert len(found) == 1, content.issues
    return found[0]


def at(issue):
    assert issue.source is not None
    return issue.source.sheet, issue.source.cell


def outputs(*rows, source="Tab1"):
    names = TABLES["outputs"].names
    return render_tsv(
        names,
        [
            tuple(
                {"study": STUDY, "source": source, **row}.get(name, "")
                for name in names
            )
            for row in rows
        ],
    )


def subjects(*names):
    spec = TABLES["subjects"]
    rows = [{"name": "all"}, *({"name": name, "parent": "all"} for name in names)]
    return render_tsv(
        spec.names,
        [
            tuple({"study": STUDY, **row}.get(name, "") for name in spec.names)
            for row in rows
        ],
    )


@pytest.fixture
def study_tables(valid_study):
    return tables_of(valid_study)


@pytest.fixture
def workbook(study_tables, sf_vocabulary, tmp_path):
    return build(study_tables, sf_vocabulary, tmp_path / "Example.xlsx")


def check_round_trip(content, tables, base=None):
    assert content.issues == []
    assert content.ok
    assert texts(content) == tables
    for file, table in content.tables.items():
        assert table == SheetTable(
            file, tables[file], tuple(range(1, tables[file].count("\n") + 1))
        )
    assert content.sheets == (
        "subjects",
        "interventions",
        "characteristica",
        "outputs_Tab1",
        "outputs_Tab2",
        "outputs_Tab3",
        "timecourses_Fig1",
        "scatters_Fig2",
    )
    assert content.base is not None
    assert content.base.generation == GENERATION
    assert content.base.created == CREATED
    assert dict(content.base.files) == (tables if base is None else base)


@pytest.fixture
def edge_tables(make_study, valid_files, tsv):
    """The canonical tables of the valid study plus the edge table."""
    folder = make_study({**valid_files, "outputs_Tab1.tsv": tsv("outputs", *EDGE_ROWS)})
    assert format_folder(folder).ok
    tables = tables_of(folder)
    # The edge texts survive formatting, so the round trip covers them.
    for text in (
        "0.30000000000000004",
        "1234567890123456",
        "\tNR\tNR\t",
        "\t=x\t",
        "\t_x0041_\ta_x005F_b\t",
        "\t_x005F_x0041_\n",
        "\t#N/A\t#DIV/0!\t",
        "\t#REF!\n",
    ):
        assert text in tables["outputs_Tab1.tsv"]
    return tables


@pytest.fixture
def edge_workbook(edge_tables, sf_vocabulary, tmp_path):
    return build(
        edge_tables,
        sf_vocabulary,
        tmp_path / "Example.xlsx",
        empty_sheets=["outputs_Tab3"],
    )


def libreoffice_saved(tables):
    """The edge tables as LibreOffice saves them.

    It saves the escapes of _x005F_x0041_ wrongly, which reads back as _x0041_;
    the generation warns about such text (cell_escape_text).
    """
    edge = "outputs_Tab1.tsv"
    changed = tables[edge].replace("\t_x005F_x0041_\n", "\t_x0041_\n")
    assert changed != tables[edge]
    return {**tables, edge: changed}


def test_round_trip_is_lossless(edge_workbook, edge_tables):
    check_round_trip(read_workbook(edge_workbook, STUDY), edge_tables)


def test_round_trip_is_lossless_after_a_libreoffice_save(
    edge_workbook, edge_tables, libreoffice_resave
):
    resaved = libreoffice_resave(edge_workbook)
    check_round_trip(
        read_workbook(resaved, STUDY), libreoffice_saved(edge_tables), edge_tables
    )


def test_round_trip_is_lossless_after_two_libreoffice_saves(
    edge_workbook, edge_tables, libreoffice_resave
):
    resaved = libreoffice_resave(libreoffice_resave(edge_workbook))
    check_round_trip(
        read_workbook(resaved, STUDY), libreoffice_saved(edge_tables), edge_tables
    )


def test_rows_map_canonical_lines_to_sheet_rows(sf_vocabulary, tmp_path):
    rows = (
        {"subjects": "S2", "measurement": "cmax"},
        {"subjects": "S1", "measurement": "cmax"},
    )
    tables = {"subjects.tsv": subjects("S1", "S2"), "outputs_Tab1.tsv": outputs(*rows)}
    path = build(tables, sf_vocabulary, tmp_path / "Example.xlsx")

    content = read_workbook(path, STUDY)

    assert content.ok, content.issues
    table = content.tables["outputs_Tab1.tsv"]
    assert table.text == outputs(*reversed(rows))
    assert table.rows == (1, 3, 2)


def test_empty_rows_keep_the_sheet_row_numbers(workbook, study_tables):
    def insert(book):
        book["timecourses_Fig1"].insert_rows(3, amount=2)

    content = read_workbook(edit(workbook, insert), STUDY)

    assert content.ok, content.issues
    table = content.tables["timecourses_Fig1.tsv"]
    assert table.text == study_tables["timecourses_Fig1.tsv"]
    assert table.rows == (1, 2, 5, 6)


# Column B of every table holds text, so the formula doubles the mean of row 2
# (2.5) in the sd column.
SD_CELL = f"{letter('outputs', 'sd')}2"


def double_the_mean(book):
    book["outputs_Tab2"][SD_CELL] = f"={letter('outputs', 'mean')}2*2"


def test_formula_without_value_is_an_error(workbook):
    content = read_workbook(edit(workbook, double_the_mean), STUDY)

    assert not content.ok
    issue = only(content, "formula_without_value")
    assert issue.severity == "error"
    assert at(issue) == ("outputs_Tab2", SD_CELL)
    assert "LibreOffice" in issue.suggestions[0].message


def test_formula_value_after_a_spreadsheet_save(workbook, libreoffice_resave):
    content = read_workbook(libreoffice_resave(edit(workbook, double_the_mean)), STUDY)

    assert content.ok, content.issues
    issue = only(content, "formula_value")
    assert issue.severity == "warning"
    assert at(issue) == ("outputs_Tab2", SD_CELL)
    row = lines(content.tables["outputs_Tab2.tsv"].text)[1].split("\t")
    assert row[TABLES["outputs"].names.index("sd")] == "5"


def test_formula_with_an_empty_text_value_is_an_empty_cell(
    workbook, libreoffice_resave
):
    cell = f"{letter('outputs', 'comment')}2"

    def formula(book):
        book["outputs_Tab2"][cell] = '=IF(1>2,"x","")'

    content = read_workbook(libreoffice_resave(edit(workbook, formula)), STUDY)

    assert content.ok, content.issues
    assert codes(content) == ["formula_value"]
    row = lines(content.tables["outputs_Tab2.tsv"].text)[1].split("\t")
    assert row[TABLES["outputs"].names.index("comment")] == ""


DATE_HINT = (
    "The spreadsheet converted this cell to a date; format the column as text "
    "and enter the value again."
)


@pytest.mark.parametrize(
    ("column", "value", "number_format", "code", "hint"),
    [
        ("mean", datetime(2026, 1, 2), None, "cell_date", DATE_HINT),
        ("mean", 0.2, "0%", "cell_percent", "20% is stored as 0.2; enter 20."),
        ("mean", "#DIV/0!", None, "cell_error", None),
        ("comment", "two\nlines", None, "cell_line_break", "line break"),
        ("comment", "a\tb", None, "cell_line_break", "tab"),
        # Excel stores a carriage return and a control character as escapes.
        ("comment", "a_x000D_b", None, "cell_line_break", "carriage return"),
        ("comment", "a_x0001_b", None, "illegal_character", "Remove the character"),
    ],
)
def test_converted_cells_are_errors_at_their_cell(
    workbook, column, value, number_format, code, hint
):
    cell = f"{letter('outputs', column)}2"

    def convert(book):
        target = book["outputs_Tab2"][cell]
        target.value = value
        if number_format:
            target.number_format = number_format

    content = read_workbook(edit(workbook, convert), STUDY)

    assert not content.ok
    issue = only(content, code)
    assert issue.severity == "error"
    assert at(issue) == ("outputs_Tab2", cell)
    if hint:
        assert hint in issue.suggestions[0].message
    # The table text stays one line per row.
    assert content.tables["outputs_Tab2.tsv"].text.count("\n") == 2


@pytest.mark.parametrize("text", ["#N/A", "#DIV/0!", " #REF!"])
def test_text_that_looks_like_an_error_value_is_text(workbook, study_tables, text):
    def convert(book):
        target = book["outputs_Tab2"][f"{letter('outputs', 'comment')}2"]
        target.value = text
        target.data_type = "s"

    content = read_workbook(edit(workbook, convert), STUDY)

    assert content.issues == []
    [_, row] = lines(content.tables["outputs_Tab2.tsv"].text)
    # The canonical text has no surrounding spaces.
    assert row.split("\t")[TABLES["outputs"].names.index("comment")] == text.strip()


def test_a_formula_error_in_a_date_formatted_cell_is_an_error(
    workbook, libreoffice_resave
):
    cell = f"{letter('outputs', 'mean')}2"

    def convert(book):
        target = book["outputs_Tab2"][cell]
        target.value = '="a"+1'
        target.number_format = "yyyy-mm-dd"

    content = read_workbook(libreoffice_resave(edit(workbook, convert)), STUDY)

    issue = only(content, "cell_error")
    assert "#VALUE!" in issue.message
    assert at(issue) == ("outputs_Tab2", cell)
    assert "cell_date" not in codes(content)


@pytest.mark.parametrize(
    ("value", "code"),
    [
        # openpyxl turns a date serial beyond the calendar into #VALUE! and warns.
        (1e10, "cell_date"),
        ("#DIV/0!", "cell_error"),
    ],
)
def test_a_date_formatted_cell_beyond_the_calendar_is_a_date(workbook, value, code):
    cell = f"{letter('outputs', 'mean')}2"

    def convert(book):
        target = book["outputs_Tab2"][cell]
        target.value = value
        target.number_format = "yyyy-mm-dd"

    path = edit(workbook, convert)
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        content = read_workbook(path, STUDY)

    assert caught == []
    assert codes(content) == [code]
    assert at(content.issues[0]) == ("outputs_Tab2", cell)
    if code == "cell_date":
        assert content.issues[0].suggestions[0].message == DATE_HINT


def test_booleans_and_numbers_become_canonical_text(workbook):
    def convert(book):
        sheet = book["outputs_Tab2"]
        sheet[f"{letter('outputs', 'comment')}2"] = True
        sheet[f"{letter('outputs', 'mean')}2"] = 2.50
        sheet[f"{letter('outputs', 'sd')}2"] = 3.0
        sheet[f"{letter('outputs', 'choice')}2"] = 7

    content = read_workbook(edit(workbook, convert), STUDY)

    assert content.ok, content.issues
    row = lines(content.tables["outputs_Tab2.tsv"].text)[1].split("\t")
    names = TABLES["outputs"].names
    assert row[names.index("comment")] == "TRUE"
    assert row[names.index("mean")] == "2.5"
    assert row[names.index("sd")] == "3"
    assert row[names.index("choice")] == "7"


def test_repeated_cell_issues_are_capped_per_sheet_and_code(sf_vocabulary, tmp_path):
    rows = [{"subjects": "all", "measurement": "cmax"}] * 12
    tables = {"subjects.tsv": subjects(), "outputs_Tab1.tsv": outputs(*rows)}
    path = build(tables, sf_vocabulary, tmp_path / "Example.xlsx")

    def convert(book):
        for row in range(2, 14):
            book["outputs_Tab1"][f"{letter('outputs', 'comment')}{row}"] = "a\nb"
            book["outputs_Tab1"][f"{letter('outputs', 'mean')}{row}"] = "#REF!"

    content = read_workbook(edit(path, convert), STUDY)

    for code in ("cell_line_break", "cell_error"):
        found = [issue for issue in content.issues if issue.code == code]
        assert len(found) == REPEATED_ISSUES + 1
        assert "12" in found[-1].message
        assert at(found[-1]) == ("outputs_Tab1", None)


def test_unknown_sheet_is_an_error(workbook):
    def copy(book):
        book.copy_worksheet(book["outputs_Tab2"]).title = "outputs_Tab2 (2)"

    content = read_workbook(edit(workbook, copy), STUDY)

    assert not content.ok
    issue = only(content, "unknown_sheet")
    assert at(issue) == ("outputs_Tab2 (2)", None)
    hint = issue.suggestions[0].message
    assert "<kind>_<source>" in hint
    assert "scratch" in hint
    assert "outputs_Tab2 (2)" not in content.sheets
    assert "outputs_Tab2.tsv" in content.tables


def test_scratch_sheets_are_ignored(workbook, study_tables):
    def scratch(book):
        notes = book.create_sheet("_notes")
        notes["A1"] = "anything"
        notes["B2"] = "=1+1"

    content = read_workbook(edit(workbook, scratch), STUDY)

    assert content.issues == []
    assert "_notes" not in content.sheets
    assert texts(content) == study_tables


def test_a_missing_subjects_sheet_is_an_error(workbook):
    def remove(book):
        book.remove(book["subjects"])

    content = read_workbook(edit(workbook, remove), STUDY)

    assert not content.ok
    issue = only(content, "missing_sheet")
    assert at(issue) == ("subjects", None)
    assert "subjects.tsv" not in content.tables
    assert "outputs_Tab2.tsv" in content.tables


def test_a_value_outside_the_header_columns_is_an_error(workbook):
    def outside(book):
        book["subjects"]["Z3"] = "x"

    content = read_workbook(edit(workbook, outside), STUDY)

    assert not content.ok
    assert codes(content) == ["value_outside_table"]
    assert at(content.issues[0]) == ("subjects", "Z3")


def test_an_unknown_header_column_is_reported_at_its_cell(workbook):
    def rename(book):
        book["subjects"]["G1"] = "note"

    content = read_workbook(edit(workbook, rename), STUDY)

    assert not content.ok
    issue = only(content, "unknown_column")
    assert at(issue) == ("subjects", "G1")
    assert "subjects.tsv" not in content.tables
    assert "subjects" in content.sheets


def test_a_missing_base_is_a_warning(workbook, study_tables):
    def remove(book):
        book.remove(book[BASE_SHEET])

    content = read_workbook(edit(workbook, remove), STUDY)

    assert content.ok
    assert content.base is None
    issue = only(content, "workbook_base_missing")
    assert issue.severity == "warning"
    assert texts(content) == study_tables


def test_an_invalid_base_is_a_warning(workbook):
    def damage(book):
        book[BASE_SHEET]["A1"] = "something else"

    content = read_workbook(edit(workbook, damage), STUDY)

    assert content.ok
    assert content.base is None
    assert only(content, "workbook_base_invalid").severity == "warning"


def test_a_newer_workbook_format_is_an_error(workbook):
    def newer(book):
        book[BASE_SHEET]["B1"] = 99

    content = read_workbook(edit(workbook, newer), STUDY)

    assert not content.ok
    assert content.base is None
    assert only(content, "workbook_newer").severity == "error"


def test_a_file_that_is_no_workbook_is_unreadable(tmp_path):
    path = tmp_path / "Example.xlsx"
    path.write_bytes(b"not a workbook")

    content = read_workbook(path, STUDY)

    assert codes(content) == ["workbook_unreadable"]
    assert not content.ok
    assert content.tables == {}
    assert content.base is None


def test_a_damaged_sheet_is_unreadable(workbook, tmp_path):
    damaged = tmp_path / "damaged.xlsx"
    with ZipFile(workbook) as source, ZipFile(damaged, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data[: len(data) // 2]
            target.writestr(item, data)

    content = read_workbook(damaged, STUDY)

    assert codes(content) == ["workbook_unreadable"]


def test_the_row_limit_counts_the_rows_of_every_sheet(workbook):
    assert read_workbook(workbook, STUDY, max_rows=15).ok
    with pytest.raises(StudyValidationError) as error:
        read_workbook(workbook, STUDY, max_rows=14)
    assert error.value.report.issues[0].code == "row_limit"


def test_reading_a_large_sheet_is_fast(sf_vocabulary, tmp_path):
    rows = [
        {
            "subjects": "all",
            "interventions": "D1",
            "measurement": "cmax",
            "substance": "drug",
            "tissue": "plasma",
            "time": str(index % 24),
            "time_unit": "h",
            "mean": str(index * 0.25),
            "sd": "0.5",
            "unit": "mg/l",
            "comment": f"row {index}",
        }
        for index in range(20_000)
    ]
    tables = {"subjects.tsv": subjects(), "outputs_Tab1.tsv": outputs(*rows)}
    path = build(tables, sf_vocabulary, tmp_path / "Example.xlsx")

    start = time.perf_counter()
    content = read_workbook(path, STUDY)
    elapsed = time.perf_counter() - start

    assert content.ok, content.issues[:3]
    assert len(content.tables["outputs_Tab1.tsv"].rows) == 20_001
    assert elapsed < 15, f"reading 20,000 rows took {elapsed:.1f} s"


def with_dimension(path, dimension) -> Path:
    """A copy of a workbook whose sheets all claim the stored dimension."""
    changed = path.with_name(f"dimension-{path.name}")
    with ZipFile(path) as source, ZipFile(changed, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item)
            if item.filename.startswith("xl/worksheets/sheet"):
                data, count = re.subn(
                    rb'<dimension ref="[^"]*"\s*/>',
                    f'<dimension ref="{dimension}"/>'.encode(),
                    data,
                )
                assert count == 1
            target.writestr(item, data)
    return changed


@pytest.mark.parametrize("dimension", ["A1:Z1048576", "A1"])
def test_the_stored_dimension_is_ignored(sf_vocabulary, tmp_path, dimension):
    # The comment column AA lies beyond both dimensions.
    rows = [
        {"subjects": "all", "measurement": "cmax", "mean": str(n), "comment": "c"}
        for n in range(9)
    ]
    tables = {"subjects.tsv": subjects(), "outputs_Tab1.tsv": outputs(*rows)}
    path = with_dimension(
        build(tables, sf_vocabulary, tmp_path / "Example.xlsx"), dimension
    )

    start = time.perf_counter()
    content = read_workbook(path, STUDY)
    elapsed = time.perf_counter() - start

    assert content.ok, content.issues
    assert texts(content) == tables
    assert content.base is not None
    assert dict(content.base.files) == tables
    assert elapsed < 2, f"reading took {elapsed:.1f} s"


def test_the_workbook_is_read_from_one_snapshot(workbook, study_tables, monkeypatch):
    # Both passes read the same bytes, even if the file is replaced in between.
    calls = []

    class Reader(read._Reader):
        def __init__(self, source, **options):
            calls.append(source)
            if len(calls) == 1:
                workbook.write_bytes(b"replaced while reading")
            super().__init__(source, **options)

    monkeypatch.setattr(read, "_Reader", Reader)
    content = read_workbook(workbook, STUDY)

    assert len(calls) == 2
    assert content.ok, content.issues
    assert texts(content) == study_tables
