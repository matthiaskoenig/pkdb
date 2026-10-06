import openpyxl
from openpyxl.utils import get_column_letter

from pkdb.studyformat.pipeline import sync_and_format
from pkdb.studyformat.sync import sync_study
from pkdb.studyformat.tables import parse_table_file
from pkdb.studyformat.workbook.base import workbook_path


def test_without_workbook_only_formats(valid_study, sf_vocabulary):
    table = valid_study / "timecourses_Fig1.tsv"
    table.write_text(table.read_text() + "\n")
    stages = []
    result = sync_and_format(valid_study, sf_vocabulary, on_stage=stages.append)
    assert result.ok and result.syncs == ()
    assert result.formatted is not None
    assert [change.file for change in result.formatted.changes] == [
        "timecourses_Fig1.tsv"
    ]
    assert stages == ["sync", "format"]
    assert not workbook_path(valid_study).exists()


def test_conflict_stops_before_format(valid_study, sf_vocabulary):
    assert sync_study(valid_study, sf_vocabulary).ok
    table = valid_study / "timecourses_Fig1.tsv"
    parsed = parse_table_file(table.name)
    assert parsed is not None
    mean = get_column_letter(parsed[0].names.index("mean") + 1)
    path = workbook_path(valid_study)
    book = openpyxl.load_workbook(path)
    book["timecourses_Fig1"][f"{mean}3"] = 5  # the second data row
    book.save(path)
    lines = table.read_text().splitlines()
    lines[2] = lines[2].replace("\t2\t", "\t7\t", 1)
    table.write_text("\n".join(lines) + "\n")
    stages = []
    result = sync_and_format(valid_study, sf_vocabulary, on_stage=stages.append)
    assert result.stopped == "sync" and result.formatted is None
    assert stages == ["sync"]
    assert any(issue.code == "sync_conflict" for issue in result.issues)
