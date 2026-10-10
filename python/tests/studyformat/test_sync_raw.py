"""Raw tables in the workbook: text sheets that the sync merges without sorting."""

import openpyxl

from pkdb.studyformat.sync import add_table, sync_study
from pkdb.studyformat.workbook.base import workbook_path

RAW = "Group\tAge (yr)\tWeight\nmen\t1.50\t007\nwomen\t12 ± 3\t=A1\nall\t1-2\tTRUE\nnote\t1e3\t\n"


def _sheet_rows(folder, sheet):
    book = openpyxl.load_workbook(workbook_path(folder))
    return [
        tuple("" if cell.value is None else str(cell.value) for cell in row)
        for row in book[sheet].iter_rows()
    ], book[sheet]


def test_raw_table_round_trips_as_text_cells(valid_study, sf_vocabulary):
    (valid_study / "Example_Tab2.tsv").write_text(RAW, encoding="utf-8", newline="")
    assert sync_study(valid_study, sf_vocabulary).ok
    rows, sheet = _sheet_rows(valid_study, "Example_Tab2")
    assert rows[1] == ("men", "1.50", "007")
    assert all(
        cell.data_type == "s"
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    )
    assert sheet.column_dimensions["A"].number_format == "@"
    book = openpyxl.load_workbook(workbook_path(valid_study))
    assert book.sheetnames.index("Example_Tab2") > book.sheetnames.index(
        "scatters_Fig2"
    )
    # A second sync changes nothing.
    again = sync_study(valid_study, sf_vocabulary)
    assert again.ok and not again.changes
    assert (valid_study / "Example_Tab2.tsv").read_text(encoding="utf-8") == RAW


def test_raw_row_inserted_in_workbook_merges_with_table_edit(
    valid_study, sf_vocabulary
):
    (valid_study / "Example_Tab2.tsv").write_text(RAW, encoding="utf-8", newline="")
    assert sync_study(valid_study, sf_vocabulary).ok
    path = workbook_path(valid_study)
    book = openpyxl.load_workbook(path)
    sheet = book["Example_Tab2"]
    sheet.insert_rows(2)
    sheet["A2"], sheet["B2"] = "children", "8"
    book.save(path)
    text = (valid_study / "Example_Tab2.tsv").read_text(encoding="utf-8")
    (valid_study / "Example_Tab2.tsv").write_text(
        text.replace("note\t1e3", "note\t2e3"), encoding="utf-8", newline=""
    )
    result = sync_study(valid_study, sf_vocabulary)
    assert result.ok, result.issues
    assert (valid_study / "Example_Tab2.tsv").read_text(
        encoding="utf-8"
    ).splitlines() == [
        "Group\tAge (yr)\tWeight",
        "children\t8\t",
        "men\t1.50\t007",
        "women\t12 ± 3\t=A1",
        "all\t1-2\tTRUE",
        "note\t2e3\t",
    ]


def test_add_raw_sheet_stays_until_it_has_cells(valid_study, sf_vocabulary):
    assert sync_study(valid_study, sf_vocabulary).ok
    added = add_table(valid_study, sf_vocabulary, "Example_Tab2")
    assert added.ok, added.issues
    assert sync_study(valid_study, sf_vocabulary).ok
    assert (
        "Example_Tab2" in openpyxl.load_workbook(workbook_path(valid_study)).sheetnames
    )
    assert not (valid_study / "Example_Tab2.tsv").exists()


def test_raw_sheet_of_libreoffice_save_is_unchanged(
    valid_study, sf_vocabulary, libreoffice_resave
):
    (valid_study / "Example_Tab2.tsv").write_text(RAW, encoding="utf-8", newline="")
    assert sync_study(valid_study, sf_vocabulary).ok
    path = workbook_path(valid_study)
    path.write_bytes(libreoffice_resave(path).read_bytes())
    assert "Example_Tab2" in openpyxl.load_workbook(path).sheetnames
    result = sync_study(valid_study, sf_vocabulary)
    assert result.ok and not result.changes
    assert (valid_study / "Example_Tab2.tsv").read_text(encoding="utf-8") == RAW


def _set_cell(folder, sheet, cell, value):
    path = workbook_path(folder)
    book = openpyxl.load_workbook(path)
    book[sheet][cell] = value
    book.save(path)


def test_raw_cell_reads_from_the_workbook_as_the_table_file_would(
    valid_study, sf_vocabulary
):
    (valid_study / "Example_Tab2.tsv").write_text(RAW, encoding="utf-8", newline="")
    assert sync_study(valid_study, sf_vocabulary).ok
    # Spreadsheet export quoting is removed, as when the TSV file is read.
    _set_cell(valid_study, "Example_Tab2", "D5", ' "spare" ')
    assert sync_study(valid_study, sf_vocabulary).ok
    lines = (valid_study / "Example_Tab2.tsv").read_text(encoding="utf-8").splitlines()
    assert lines[0] == "Group\tAge (yr)\tWeight\t"
    assert lines[4] == "note\t1e3\t\tspare"
    again = sync_study(valid_study, sf_vocabulary)
    assert again.ok and not again.changes


def test_raw_line_that_a_table_file_cannot_hold_is_an_issue(valid_study, sf_vocabulary):
    (valid_study / "Example_Tab2.tsv").write_text(RAW, encoding="utf-8", newline="")
    assert sync_study(valid_study, sf_vocabulary).ok
    _set_cell(valid_study, "Example_Tab2", "A3", ">>>>>>> women")
    result = sync_study(valid_study, sf_vocabulary)
    assert not result.ok and not result.changes
    [issue] = [issue for issue in result.issues if issue.severity == "error"]
    assert issue.code == "merge_conflict"
    assert issue.source is not None
    assert (issue.source.file, issue.source.sheet, issue.source.row) == (
        "Example.xlsx",
        "Example_Tab2",
        3,
    )
    assert (valid_study / "Example_Tab2.tsv").read_text(encoding="utf-8") == RAW


def _conflict_message(folder, vocabulary):
    result = sync_study(folder, vocabulary)
    [issue] = [issue for issue in result.issues if issue.code == "sync_conflict"]
    assert issue.source is not None and issue.source.row == 1
    return issue.message


def test_raw_conflict_at_the_first_line_is_before_it(valid_study, sf_vocabulary):
    raw = valid_study / "Example_Tab2.tsv"
    raw.write_text(RAW, encoding="utf-8", newline="")
    assert sync_study(valid_study, sf_vocabulary).ok
    path = workbook_path(valid_study)
    book = openpyxl.load_workbook(path)
    book["Example_Tab2"].delete_rows(1)
    book.save(path)
    raw.write_text(RAW.replace("Age (yr)", "Age (y)"), encoding="utf-8", newline="")
    assert _conflict_message(valid_study, sf_vocabulary).endswith(
        "rows removed before row 1 of the sheet, line 1 of Example_Tab2.tsv"
    )


def test_raw_conflict_with_the_first_line_removed_from_the_table(
    valid_study, sf_vocabulary
):
    raw = valid_study / "Example_Tab2.tsv"
    raw.write_text(RAW, encoding="utf-8", newline="")
    assert sync_study(valid_study, sf_vocabulary).ok
    _set_cell(valid_study, "Example_Tab2", "B1", "Age (y)")
    raw.write_text(RAW.split("\n", 1)[1], encoding="utf-8", newline="")
    assert _conflict_message(valid_study, sf_vocabulary).endswith(
        "row 1 of the sheet, lines removed before line 1 of Example_Tab2.tsv"
    )
