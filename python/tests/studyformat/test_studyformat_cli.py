import json

import pytest

from pkdb.cli import main


def lines(capsys):
    return [
        json.loads(line)
        for line in capsys.readouterr().out.splitlines()
        if line.startswith("{")
    ]


def test_format_command(make_study, valid_files, capsys):
    folder = make_study(valid_files)
    assert main(["format", str(folder), "--check", "--format", "json"]) == 1
    entry = lines(capsys)[0]
    assert entry["ok"] and "subjects.tsv" in {
        change["file"] for change in entry["changes"]
    }
    assert main(["format", str(folder), "--format", "json"]) == 0
    capsys.readouterr()
    assert main(["format", str(folder), "--check", "--format", "json"]) == 0
    assert lines(capsys)[0]["changes"] == []


def test_format_skips_format_1(tmp_path, capsys):
    folder = tmp_path / "Old"
    folder.mkdir()
    (folder / "study.json").write_text('{"sid": "X"}')
    assert main(["format", str(folder), "--format", "json"]) == 0
    assert "format 1" in lines(capsys)[0]["skipped"]


def test_format_reports_problems_for_people(make_study, valid_files, capsys):
    folder = make_study({**valid_files, "outputs_Tab2.tsv": "group\nall\n"})
    assert main(["format", str(folder), "--format", "human"]) == 1
    out = capsys.readouterr().out
    assert "caffeine/Example" in out
    assert "unknown_column" in out and "subjects" in out


def test_validate_format_2(valid_study, sf_vocabulary, tmp_path, capsys):
    lock = tmp_path / "vocabulary.json"
    sf_vocabulary.save(lock)
    args = [
        "validate",
        str(valid_study),
        "--offline",
        "--vocabulary",
        str(lock),
        "--format",
        "json",
    ]
    assert main(args) == 0
    result = lines(capsys)[-1]
    assert result["ok"] and result["study_format"] == 2
    assert result["sid"] == "caffeine/Example"
    assert result["report"]["issues"] == []
    assert "workbook" not in result


def test_validate_format_2_failure(valid_study, sf_vocabulary, tmp_path, capsys):
    (valid_study / "notes.csv").write_text("x\n")
    lock = tmp_path / "vocabulary.json"
    sf_vocabulary.save(lock)
    args = [
        "validate",
        str(valid_study),
        "--offline",
        "--vocabulary",
        str(lock),
        "--format",
        "json",
    ]
    assert main(args) == 1
    result = lines(capsys)[-1]
    assert not result["ok"]
    assert [issue["code"] for issue in result["report"]["issues"]] == ["unknown_file"]


def test_prepare_accepts_format_2(valid_study, sf_vocabulary, tmp_path, capsys):
    before = snapshot(valid_study)
    lock = tmp_path / "vocabulary.json"
    sf_vocabulary.save(lock)
    args = ["prepare", str(valid_study), "--offline", "--vocabulary", str(lock)]
    assert main([*args, "--format", "json"]) == 0
    result = lines(capsys)[-1]
    assert result["ok"] and result["study_format"] == 2
    assert result["sid"] == result["study"]["sid"] == "caffeine/Example"
    assert result["report"]["issues"] == []
    assert result["processing_version"] == "9"
    assert set(result["file_hashes"]) == set(before)
    assert snapshot(valid_study) == before


def test_prepare_reports_format_2_problems(
    valid_study, sf_vocabulary, tmp_path, capsys
):
    (valid_study / "notes.csv").write_text("x\n")
    lock = tmp_path / "vocabulary.json"
    sf_vocabulary.save(lock)
    args = ["prepare", str(valid_study), "--offline", "--vocabulary", str(lock)]
    assert main([*args, "--format", "json"]) == 1
    result = lines(capsys)[-1]
    assert not result["ok"] and result["study_format"] == 2
    assert [issue["code"] for issue in result["report"]["issues"]] == ["unknown_file"]


def test_schema_commands(tmp_path, capsys):
    assert main(["schema", "export", "--output", str(tmp_path / "schemas")]) == 0
    assert (tmp_path / "schemas" / "outputs.row.schema.json").is_file()
    assert main(["schema", "docs", "--output", str(tmp_path / "reference.md")]) == 0
    assert "# Study format" in (tmp_path / "reference.md").read_text(encoding="utf-8")


def snapshot(folder):
    return {path.name: path.read_bytes() for path in folder.iterdir()}


def test_format_walks_a_parent_directory(make_study, valid_files, capsys):
    folder = make_study(valid_files)
    old = folder.parent / "Old"
    old.mkdir()
    (old / "study.json").write_text('{"sid": "X"}')
    assert main(["format", str(folder.parent), "--format", "json"]) == 0
    entries = {entry["path"]: entry for entry in lines(capsys)}
    assert "format 1" in entries[str(old)]["skipped"]
    assert entries[str(folder)]["changes"]


def test_format_human_output(make_study, valid_files, capsys):
    folder = make_study(valid_files)
    assert main(["format", str(folder), "--check", "--format", "human"]) == 1
    out = capsys.readouterr().out
    assert "caffeine/Example: would rewrite subjects.tsv" in out
    assert "already formatted" not in out
    assert main(["format", str(folder), "--format", "human"]) == 0
    assert "caffeine/Example: rewrote subjects.tsv" in capsys.readouterr().out
    assert main(["format", str(folder), "--format", "human"]) == 0
    assert "caffeine/Example: already formatted" in capsys.readouterr().out


def test_format_without_study_folder(tmp_path, capsys):
    assert main(["format", str(tmp_path), "--format", "json"]) == 1
    assert "No study.json files found" in capsys.readouterr().err


def test_format_reports_unwritable_folder(make_study, valid_files, capsys, monkeypatch):
    def fail(*args, **kwargs):
        raise PermissionError(13, "Permission denied")

    monkeypatch.setattr("pkdb.studyformat.formatter.atomic_text", fail)
    folder = make_study(valid_files)
    assert main(["format", str(folder), "--format", "json"]) == 1
    entry = lines(capsys)[0]
    assert not entry["ok"] and "Permission denied" in entry["error"]


def test_schema_export_reports_unwritable_output(tmp_path, capsys):
    target = tmp_path / "file"
    target.write_text("x")
    assert main(["schema", "export", "--output", str(target)]) == 1
    assert "Cannot write" in capsys.readouterr().err


def test_format_human_marks_unformatted_folder_with_problems(valid_study, capsys):
    (valid_study / "outputs_Tab2.tsv").write_text("group\nall\n")
    assert main(["format", str(valid_study), "--format", "human"]) == 1
    out = capsys.readouterr().out
    assert "caffeine/Example: not formatted, fix the problems below" in out
    assert "already formatted" not in out


def test_format_human_issue_lines_name_the_file_once(make_study, valid_files, capsys):
    folder = make_study({**valid_files, "outputs_Tab2.tsv": "group\nall\n"})
    assert main(["format", str(folder), "--format", "human"]) == 1
    issue = next(
        line
        for line in capsys.readouterr().out.splitlines()
        if "[unknown_column]" in line
    )
    assert issue.startswith("  outputs_Tab2.tsv ")
    assert issue.count("outputs_Tab2.tsv") == 1


def test_format_human_locations(capsys):
    from pkdb.schemas.source import SourceLocation
    from pkdb.studyformat_cli import _location

    assert _location(SourceLocation(file="a.tsv", cell="A1", row=1)) == "a.tsv A1"
    assert _location(SourceLocation(file="a.tsv", row=3)) == "a.tsv line 3"
    assert _location(SourceLocation(file="a.tsv")) == "a.tsv"
    assert _location(None) == ""


def test_format_human_strips_terminal_controls(make_study, valid_files, capsys):
    header = "gr" + chr(0x1B) + "[31moup"
    folder = make_study({**valid_files, "outputs_Tab2.tsv": f"{header}\nall\n"})
    assert main(["format", str(folder), "--format", "human"]) == 1
    out = capsys.readouterr().out
    assert chr(0x1B) not in out and "unknown_column" in out


def test_format_and_validate_name_the_current_folder(
    valid_study, sf_vocabulary, tmp_path, capsys, monkeypatch
):
    monkeypatch.chdir(valid_study)
    assert main(["format", ".", "--format", "human"]) == 0
    assert capsys.readouterr().out.startswith("caffeine/Example: already formatted")
    lock = tmp_path / "vocabulary.json"
    sf_vocabulary.save(lock)
    args = ["validate", ".", "--offline", "--vocabulary", str(lock), "--format", "json"]
    assert main(args) == 0
    result = lines(capsys)[-1]
    assert result["sid"] == "caffeine/Example" and result["name"] == "Example"


def test_validate_reports_format_1_and_format_2_studies_together(
    study_folder, valid_study, sf_vocabulary, tmp_path, capsys
):
    lock = tmp_path / "vocabulary.json"
    sf_vocabulary.save(lock)
    args = [
        "validate",
        str(tmp_path),
        "--offline",
        "--vocabulary",
        str(lock),
        "--format",
        "json",
    ]
    assert main(args) == 0
    old, new = lines(capsys)[:2]
    assert old["sid"] == "TEST1" and "study_format" not in old
    assert old["ok"] and old["relative_path"] == "Example"
    assert old["report"]["issues"] == [] and old["vocabulary_version"]
    assert new["sid"] == "caffeine/Example" and new["study_format"] == 2
    assert new["ok"] and new["relative_path"] == "caffeine/Example"
    assert new["vocabulary_hash"] == old["vocabulary_hash"]


def validate_json(folder, vocabulary, tmp_path):
    lock = tmp_path / "vocabulary.json"
    vocabulary.save(lock)
    return main(
        [
            "validate",
            str(folder),
            "--offline",
            "--vocabulary",
            str(lock),
            "--format",
            "json",
        ]
    )


def test_merge_conflict_stops_format_and_validate(
    valid_study, sf_vocabulary, tmp_path, capsys
):
    path = valid_study / "outputs_Tab2.tsv"
    header, row = path.read_text(encoding="utf-8").splitlines(keepends=True)
    data = (
        header
        + "<<<<<<< HEAD\n"
        + row
        + "=======\n"
        + row.replace("\t2.5\t", "\t2.6\t")
        + ">>>>>>> feature\n"
    ).encode()
    path.write_bytes(data)
    assert main(["format", str(valid_study), "--format", "json"]) == 1
    entry = lines(capsys)[0]
    assert [issue["code"] for issue in entry["issues"]] == ["merge_conflict"]
    assert "resolve the git conflict" in entry["issues"][0]["message"]
    assert entry["changes"] == []
    assert path.read_bytes() == data
    assert validate_json(valid_study, sf_vocabulary, tmp_path) == 1
    result = lines(capsys)[-1]
    assert not result["ok"]
    assert "merge_conflict" in {issue["code"] for issue in result["report"]["issues"]}


@pytest.mark.parametrize(
    ("text", "code"),
    [
        (
            '{\n<<<<<<< HEAD\n  "format": 2,\n=======\n  "format": 2,\n>>>>>>> b\n}\n',
            "invalid_json",
        ),
        ('{"format": 2, "format": 2}', "duplicate_key"),
    ],
    ids=["conflict", "duplicate_key"],
)
def test_broken_study_json_is_reported_as_format_2(
    valid_study, sf_vocabulary, tmp_path, capsys, text, code
):
    path = valid_study / "study.json"
    path.write_text(text, encoding="utf-8")
    assert main(["format", str(valid_study), "--format", "json"]) == 1
    entry = lines(capsys)[0]
    assert "skipped" not in entry
    assert [issue["code"] for issue in entry["issues"]] == [code]
    assert path.read_text(encoding="utf-8") == text
    assert validate_json(valid_study, sf_vocabulary, tmp_path) == 1
    result = lines(capsys)[-1]
    assert result["study_format"] == 2
    assert code in {issue["code"] for issue in result["report"]["issues"]}


def test_prepared_format_2_folder_reads_its_upload(valid_study, sf_vocabulary):
    # The curation app and Client.upload call prepare() and then read the source.
    from pkdb.preparation import prepare

    before = snapshot(valid_study)
    prepared = prepare(valid_study, vocabulary=sf_vocabulary)
    assert prepared.study.sid == "caffeine/Example"
    assert prepared.study_format == 2
    assert prepared.report.valid
    with prepared.source() as source:
        assert source.study_format == 2
        assert source.prepared.study == prepared.study
        assert source.study == before["study.json"]
        assert source.reference == before["reference.json"]
        assert {name: path.read_bytes() for name, path in source.files.items()} == {
            name: content
            for name, content in before.items()
            if name not in {"study.json", "reference.json"}
        }
        assert source.upload_path == "/api/v2/studies/caffeine/Example"
        assert source.validation_path == "/api/v2/studies/caffeine/Example/validate"
    assert snapshot(valid_study) == before


PENDING = "Workbook changes are not in the tables yet; run pkdb tables sync"


def set_mean(path, sheet, row, mean):
    """Change the mean of a sheet row and save, as in a spreadsheet application."""
    import openpyxl
    from openpyxl.utils import get_column_letter

    from pkdb.studyformat.tables import parse_table_file

    parsed = parse_table_file(f"{sheet}.tsv")
    assert parsed is not None
    workbook = openpyxl.load_workbook(path)
    column = get_column_letter(parsed[0].names.index("mean") + 1)
    workbook[sheet][f"{column}{row}"] = mean
    workbook.save(path)


def file_state(folder):
    return {
        path.name: (path.read_bytes(), path.stat().st_mtime_ns)
        for path in sorted(folder.iterdir())
    }


@pytest.mark.parametrize("command", ["validate", "prepare"])
def test_validate_and_prepare_report_workbook_changes_without_writing(
    valid_study, sf_vocabulary, tmp_path, capsys, command
):
    from pkdb.studyformat import sync_study
    from pkdb.studyformat.workbook.base import workbook_path

    assert sync_study(valid_study, sf_vocabulary).ok
    lock = tmp_path / "vocabulary.json"
    sf_vocabulary.save(lock)
    args = [command, str(valid_study), "--offline", "--vocabulary", str(lock)]
    assert main([*args, "--format", "json"]) == 0
    assert lines(capsys)[-1]["workbook"] == {
        "action": "unchanged",
        "changes": [],
        "conflicts": 0,
        "ok": True,
    }
    assert main([*args, "--format", "human"]) == 0
    assert PENDING not in capsys.readouterr().out
    set_mean(workbook_path(valid_study), "outputs_Tab2", 2, 3.25)
    before = file_state(valid_study)

    assert main([*args, "--format", "json"]) == 0
    result = lines(capsys)[-1]
    assert result["ok"]
    assert result["workbook"] == {
        "action": "unchanged",
        "changes": [{"file": "outputs_Tab2.tsv", "action": "write"}],
        "conflicts": 0,
        "ok": True,
    }
    assert main([*args, "--format", "human"]) == 0
    assert PENDING in capsys.readouterr().out
    assert file_state(valid_study) == before


def test_validate_counts_workbook_conflicts(
    valid_study, sf_vocabulary, tmp_path, capsys
):
    from pkdb.studyformat import sync_study
    from pkdb.studyformat.workbook.base import workbook_path

    assert sync_study(valid_study, sf_vocabulary).ok
    set_mean(workbook_path(valid_study), "outputs_Tab2", 2, 0.25)
    path = valid_study / "outputs_Tab2.tsv"
    path.write_text(path.read_text(encoding="utf-8").replace("\t2.5\t", "\t0.75\t"))
    before = file_state(valid_study)

    assert validate_json(valid_study, sf_vocabulary, tmp_path) == 0
    result = lines(capsys)[-1]
    assert result["workbook"] == {
        "action": "unchanged",
        "changes": [],
        "conflicts": 1,
        "ok": False,
    }
    lock = tmp_path / "vocabulary.json"
    args = ["validate", str(valid_study), "--offline", "--vocabulary", str(lock)]
    assert main([*args, "--format", "human"]) == 0
    assert PENDING in capsys.readouterr().out
    assert file_state(valid_study) == before


def test_validate_names_a_workbook_that_cannot_be_synced(
    valid_study, sf_vocabulary, tmp_path, capsys
):
    from pkdb.studyformat import sync_study
    from pkdb.studyformat.workbook.base import workbook_path

    assert sync_study(valid_study, sf_vocabulary).ok
    set_mean(workbook_path(valid_study), "outputs_Tab2", 2, "#DIV/0!")
    lock = tmp_path / "vocabulary.json"
    sf_vocabulary.save(lock)
    args = ["validate", str(valid_study), "--offline", "--vocabulary", str(lock)]

    assert main([*args, "--format", "json"]) == 0
    assert lines(capsys)[-1]["workbook"]["ok"] is False
    assert main([*args, "--format", "human"]) == 0
    out = capsys.readouterr().out
    assert PENDING not in out
    assert "The workbook cannot be synced with the tables; run pkdb tables sync" in out
