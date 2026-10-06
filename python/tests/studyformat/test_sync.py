"""Tests of the sync engine between the study workbook and the TSV tables.

Performance budget, checked on CI (plan Review Focus 5): a 20,000-row study
syncs a one-cell workbook change together with a tables change, which merges the
table and regenerates the workbook, in under 30 seconds.
"""

import time
import warnings
from pathlib import Path

import openpyxl
import pytest
from openpyxl.utils import get_column_letter

from pkdb.studyformat import SyncResult, add_table, sync, sync_study, workbook_check
from pkdb.studyformat.formatter import FileChange, format_folder
from pkdb.studyformat.sync import SyncConflict
from pkdb.studyformat.tables import parse_table_file
from pkdb.studyformat.workbook.base import (
    read_state,
    state_path,
    workbook_path,
    write_state,
)
from pkdb.studyformat.workbook.read import read_workbook
from pkdb.studyformat.workbook.write import build_workbook

STUDY = "Example"
SUBJECTS = "subjects.tsv"
OUTPUTS = "outputs_Tab2.tsv"
TIMECOURSES = "timecourses_Fig1.tsv"
SCATTERS = "scatters_Fig2.tsv"
LOCK = ".~lock.Example.xlsx#"


def names(file):
    parsed = parse_table_file(file)
    assert parsed is not None
    return parsed[0].names


def letter(file, column):
    return get_column_letter(names(file).index(column) + 1)


def lines(text):
    """Lines of canonical TSV text; str.splitlines would also split at U+2028."""
    return text.removesuffix("\n").split("\n")


def table_lines(folder, file):
    return lines((folder / file).read_text(encoding="utf-8"))


def cell(folder, file, line, column):
    return table_lines(folder, file)[line - 1].split("\t")[names(file).index(column)]


def replaced(line, file, **cells):
    values = line.split("\t")
    for column, value in cells.items():
        values[names(file).index(column)] = value
    return "\t".join(values)


def edit_table(folder, file, line, **cells):
    """Change cells of a TSV line by column name, as in a text editor."""
    rows = table_lines(folder, file)
    rows[line - 1] = replaced(rows[line - 1], file, **cells)
    (folder / file).write_text("\n".join(rows) + "\n", encoding="utf-8", newline="")


def edit(path, change):
    workbook = openpyxl.load_workbook(path)
    change(workbook)
    workbook.save(path)


def set_cells(path, sheet, row, **cells):
    """Change cells of a sheet row by column name, as in a spreadsheet application."""

    def change(workbook):
        for column, value in cells.items():
            workbook[sheet][f"{letter(f'{sheet}.tsv', column)}{row}"] = value

    edit(path, change)


def remove_sheet(name):
    def change(workbook):
        workbook.remove(workbook[name])

    return change


def snapshot(folder):
    """Content and modification time of every file of a folder."""
    return {
        path.name: (path.read_bytes(), path.stat().st_mtime_ns)
        for path in sorted(folder.iterdir())
        if path.is_file()
    }


def tables_of(folder):
    return {
        path.name: path.read_text(encoding="utf-8")
        for path in sorted(folder.glob("*.tsv"))
    }


def content_of(path):
    content = read_workbook(path, STUDY)
    assert content.ok, content.issues
    assert content.base is not None
    return content


def generation_of(path):
    return content_of(path).base.generation


def assert_in_step(folder, path):
    """The workbook holds the tables and was generated from them."""
    content = content_of(path)
    tables = tables_of(folder)
    assert {file: table.text for file, table in content.tables.items()} == tables
    assert dict(content.base.files) == tables
    assert not state_path(path).exists()


def assert_canonical(folder):
    assert format_folder(folder, check=True).changes == []


def codes(result):
    return [issue.code for issue in result.issues]


@pytest.fixture
def study(valid_study, sf_vocabulary):
    """The valid study with the workbook that the first sync creates."""
    result = sync_study(valid_study, sf_vocabulary)
    assert result.ok, result.issues
    assert result.workbook_action == "created"
    return valid_study


@pytest.fixture
def workbook(study):
    return workbook_path(study)


def test_without_a_workbook_the_sync_creates_it(valid_study, sf_vocabulary):
    path = workbook_path(valid_study)
    before = snapshot(valid_study)
    # The state file of an earlier workbook is stale.
    write_state(path, "0" * 32, {OUTPUTS: None})

    result = sync_study(valid_study, sf_vocabulary)

    assert result == SyncResult(valid_study, path, workbook_action="created")
    assert result.ok
    assert_in_step(valid_study, path)
    assert {name: snapshot(valid_study)[name] for name in before} == before


def test_untouched_tables_are_not_reformatted(make_study, valid_files, sf_vocabulary):
    # Before formatting, the owned columns are empty.
    folder = make_study(valid_files)
    before = snapshot(folder)

    created = sync_study(folder, sf_vocabulary)
    in_step = sync_study(folder, sf_vocabulary)

    assert created.workbook_action == "created"
    assert in_step == SyncResult(folder, workbook_path(folder))
    assert {name: snapshot(folder)[name] for name in before} == before


def test_a_workbook_in_step_writes_nothing(study, workbook, sf_vocabulary):
    before = snapshot(study)

    result = sync_study(study, sf_vocabulary)

    assert result == SyncResult(study, workbook)
    assert snapshot(study) == before


def test_a_workbook_change_is_written_in_canonical_form(study, workbook, sf_vocabulary):
    # Time 0 becomes 5, so the row moves to the end of the canonical table.
    set_cells(workbook, "timecourses_Fig1", 2, time=5, mean=0.5)
    edited = workbook.read_bytes()

    result = sync_study(study, sf_vocabulary)

    assert result.ok, result.issues
    assert result.changes == (FileChange(TIMECOURSES, "write"),)
    assert result.workbook_action == "unchanged"
    assert workbook.read_bytes() == edited
    assert [cell(study, TIMECOURSES, line, "time") for line in (2, 3, 4)] == [
        "1",
        "2",
        "5",
    ]
    assert cell(study, TIMECOURSES, 4, "mean") == "0.5"
    assert_canonical(study)
    assert read_state(workbook, generation_of(workbook)) == {
        TIMECOURSES: (study / TIMECOURSES).read_text(encoding="utf-8")
    }
    # The next sync finds the workbook in step.
    before = snapshot(study)
    assert sync_study(study, sf_vocabulary) == SyncResult(study, workbook)
    assert snapshot(study) == before


def test_a_tables_change_regenerates_the_workbook(study, workbook, sf_vocabulary):
    set_cells(workbook, "timecourses_Fig1", 3, mean=2.25)
    assert sync_study(study, sf_vocabulary).ok
    assert state_path(workbook).exists()
    generation = generation_of(workbook)
    edit_table(study, OUTPUTS, 2, mean="3.5")

    result = sync_study(study, sf_vocabulary)

    assert result == SyncResult(study, workbook, workbook_action="regenerated")
    assert generation_of(workbook) != generation
    assert_in_step(study, workbook)
    assert cell(study, OUTPUTS, 2, "mean") == "3.5"


def test_changes_to_different_rows_merge(study, workbook, sf_vocabulary):
    set_cells(workbook, "timecourses_Fig1", 2, mean=0.25)
    edit_table(study, TIMECOURSES, 4, mean="1.5")

    result = sync_study(study, sf_vocabulary)

    assert result.ok, result.issues
    assert result.conflicts == ()
    assert result.changes == (FileChange(TIMECOURSES, "write"),)
    assert result.workbook_action == "regenerated"
    assert [cell(study, TIMECOURSES, line, "mean") for line in (2, 3, 4)] == [
        "0.25",
        "2",
        "1.5",
    ]
    assert_canonical(study)
    assert_in_step(study, workbook)


@pytest.fixture
def conflicting(study, workbook):
    """Both sides change the mean at time 0; the other changes merge.

    Returns the TSV lines of the timecourse before the changes.
    """
    base = table_lines(study, TIMECOURSES)

    def change(book):
        mean = letter(TIMECOURSES, "mean")
        book["timecourses_Fig1"][f"{mean}2"] = 0.25
        book["timecourses_Fig1"][f"{mean}4"] = 1.25
        book["outputs_Tab2"][f"{letter(OUTPUTS, 'sd')}2"] = 0.75

    edit(workbook, change)
    edit_table(study, TIMECOURSES, 2, mean="0.75")
    edit_table(study, SCATTERS, 2, y_mean="2.5")
    return base


def test_changes_to_the_same_cell_conflict(study, workbook, conflicting, sf_vocabulary):
    before = snapshot(study)

    result = sync_study(study, sf_vocabulary)

    assert not result.ok
    assert result.conflicts == (
        SyncConflict(
            TIMECOURSES,
            "timecourses_Fig1",
            workbook_rows=((2, replaced(conflicting[1], TIMECOURSES, mean="0.25")),),
            table_lines=((2, replaced(conflicting[1], TIMECOURSES, mean="0.75")),),
            base_lines=(conflicting[1],),
        ),
    )
    [issue] = result.issues
    assert issue.code == "sync_conflict"
    assert issue.severity == "error"
    assert issue.source is not None
    assert (issue.source.sheet, issue.source.row) == ("timecourses_Fig1", 2)
    assert "line 2 of timecourses_Fig1.tsv" in issue.message
    assert result.changes == ()
    assert result.workbook_action == "unchanged"
    assert snapshot(study) == before


@pytest.mark.parametrize(("keep", "mean"), [("workbook", "0.25"), ("tables", "0.75")])
def test_keep_resolves_a_conflict_and_merges_the_rest(
    study, workbook, conflicting, sf_vocabulary, keep, mean
):
    result = sync_study(study, sf_vocabulary, keep=keep)

    assert result.ok, result.issues
    assert result.issues == ()
    assert [conflict.kept for conflict in result.conflicts] == [keep]
    assert result.changes == (
        FileChange(OUTPUTS, "write"),
        FileChange(TIMECOURSES, "write"),
    )
    assert result.workbook_action == "regenerated"
    assert [cell(study, TIMECOURSES, line, "mean") for line in (2, 3, 4)] == [
        mean,
        "2",
        "1.25",
    ]
    assert cell(study, OUTPUTS, 2, "sd") == "0.75"
    assert cell(study, SCATTERS, 2, "y_mean") == "2.5"
    assert_canonical(study)
    assert_in_step(study, workbook)


def test_a_conflict_names_the_sheet_row_of_a_moved_row(study, workbook, sf_vocabulary):
    base = table_lines(study, TIMECOURSES)

    def change(book):
        sheet = book["timecourses_Fig1"]
        first, last = list(sheet[2]), list(sheet[4])
        for top, bottom in zip(first, last, strict=True):
            top.value, bottom.value = bottom.value, top.value
        # Time 0 is now in row 4.
        sheet[f"{letter(TIMECOURSES, 'mean')}4"] = 0.25

    edit(workbook, change)
    edit_table(study, TIMECOURSES, 2, mean="0.75")

    result = sync_study(study, sf_vocabulary)

    [conflict] = result.conflicts
    assert conflict.workbook_rows == ((4, replaced(base[1], TIMECOURSES, mean="0.25")),)
    assert conflict.table_lines == ((2, replaced(base[1], TIMECOURSES, mean="0.75")),)
    [issue] = result.issues
    assert issue.source is not None
    assert issue.source.row == 4


def test_a_conflict_with_a_removed_row_is_located_at_the_row_before(
    study, workbook, sf_vocabulary
):
    base = table_lines(study, TIMECOURSES)
    edit(workbook, lambda book: book["timecourses_Fig1"].delete_rows(3))
    edit_table(study, TIMECOURSES, 3, mean="2.5")

    result = sync_study(study, sf_vocabulary)

    assert result.conflicts == (
        SyncConflict(
            TIMECOURSES,
            "timecourses_Fig1",
            workbook_rows=(),
            table_lines=((3, replaced(base[2], TIMECOURSES, mean="2.5")),),
            base_lines=(base[2],),
        ),
    )
    [issue] = result.issues
    assert issue.source is not None
    assert issue.source.row == 2
    assert "rows removed after row 2" in issue.message


@pytest.mark.parametrize("lock", ["~$Example.xlsx", ".~lock.Example.xlsx#"])
def test_an_open_workbook_is_regenerated_after_it_is_closed(
    study, workbook, sf_vocabulary, lock
):
    set_cells(workbook, "timecourses_Fig1", 3, mean=2.25)
    edit_table(study, OUTPUTS, 2, mean="3.5")
    (study / lock).write_text("curator", encoding="utf-8")
    opened = workbook.read_bytes()

    result = sync_study(study, sf_vocabulary)

    assert result.ok, result.issues
    # The workbook change still reaches the tables.
    assert result.changes == (FileChange(TIMECOURSES, "write"),)
    assert cell(study, TIMECOURSES, 3, "mean") == "2.25"
    assert result.workbook_action == "close_to_update"
    assert result.lock == study / lock
    assert workbook.read_bytes() == opened
    [issue] = result.issues
    assert issue.code == "workbook_open"
    assert issue.severity == "warning"
    assert issue.message.endswith(f"If the workbook is not open, delete {study / lock}")

    (study / lock).unlink()
    result = sync_study(study, sf_vocabulary)

    assert result == SyncResult(study, workbook, workbook_action="regenerated")
    assert_in_step(study, workbook)


def test_two_saves_while_the_workbook_stays_open(study, workbook, sf_vocabulary):
    """Plan Review Focus 1: the second save does not conflict with the first sync."""
    (study / LOCK).write_text("curator", encoding="utf-8")
    for mean in (2.25, 2.5):
        set_cells(workbook, "timecourses_Fig1", 3, mean=mean)

        result = sync_study(study, sf_vocabulary)

        assert result.ok, result.issues
        assert result.conflicts == ()
        assert result.changes == (FileChange(TIMECOURSES, "write"),)
        assert result.workbook_action == "unchanged"
        assert cell(study, TIMECOURSES, 3, "mean") == str(mean)


def test_a_save_after_a_merge_while_the_workbook_stays_open(
    study, workbook, sf_vocabulary
):
    (study / LOCK).write_text("curator", encoding="utf-8")
    set_cells(workbook, "timecourses_Fig1", 2, mean=0.25)
    edit_table(study, TIMECOURSES, 4, mean="1.5")

    merged = sync_study(study, sf_vocabulary)

    assert merged.ok, merged.issues
    assert merged.changes == (FileChange(TIMECOURSES, "write"),)
    assert merged.workbook_action == "close_to_update"

    # The open workbook still lacks the tables change of line 4.
    set_cells(workbook, "timecourses_Fig1", 2, mean=0.3)

    saved = sync_study(study, sf_vocabulary)

    assert saved.ok, saved.issues
    assert saved.conflicts == ()
    assert saved.changes == (FileChange(TIMECOURSES, "write"),)
    assert saved.workbook_action == "close_to_update"
    assert [cell(study, TIMECOURSES, line, "mean") for line in (2, 3, 4)] == [
        "0.3",
        "2",
        "1.5",
    ]


def test_keep_tables_while_the_workbook_stays_open(study, workbook, sf_vocabulary):
    (study / LOCK).write_text("curator", encoding="utf-8")
    set_cells(workbook, "timecourses_Fig1", 2, mean=0.25)
    edit_table(study, TIMECOURSES, 2, mean="0.75")

    kept = sync_study(study, sf_vocabulary, keep="tables")

    assert kept.ok, kept.issues
    assert [conflict.kept for conflict in kept.conflicts] == ["tables"]
    assert kept.changes == ()
    assert kept.workbook_action == "close_to_update"

    again = sync_study(study, sf_vocabulary)

    assert again.ok, again.issues
    assert again.conflicts == ()
    assert again.changes == ()
    assert again.workbook_action == "close_to_update"

    # Closed without saving: the workbook takes the kept tables.
    (study / LOCK).unlink()

    closed = sync_study(study, sf_vocabulary)

    assert closed == SyncResult(study, workbook, workbook_action="regenerated")
    assert cell(study, TIMECOURSES, 2, "mean") == "0.75"
    assert_in_step(study, workbook)


def test_a_state_file_of_another_generation_is_ignored(study, workbook, sf_vocabulary):
    write_state(workbook, "f" * 32, {TIMECOURSES: "stale\n", OUTPUTS: None})
    set_cells(workbook, "timecourses_Fig1", 3, mean=2.25)

    result = sync_study(study, sf_vocabulary)

    assert result.ok, result.issues
    assert result.conflicts == ()
    assert result.changes == (FileChange(TIMECOURSES, "write"),)
    assert read_state(workbook, generation_of(workbook)) == {
        TIMECOURSES: (study / TIMECOURSES).read_text(encoding="utf-8")
    }


def test_an_added_empty_sheet_survives_a_regeneration(study, workbook, sf_vocabulary):
    # What pkdb tables add does.
    added = build_workbook(
        tables_of(study),
        sf_vocabulary,
        existing=workbook,
        empty_sheets=["outputs_Tab3"],
    )
    assert added.data is not None
    workbook.write_bytes(added.data)
    edit_table(study, OUTPUTS, 2, mean="3.5")

    result = sync_study(study, sf_vocabulary)

    assert result == SyncResult(study, workbook, workbook_action="regenerated")
    content = content_of(workbook)
    assert "outputs_Tab3" in content.sheets
    assert "outputs_Tab3.tsv" not in content.tables
    assert not (study / "outputs_Tab3.tsv").exists()


def test_a_deleted_table_removes_its_sheet(study, workbook, sf_vocabulary):
    (study / SCATTERS).unlink()

    result = sync_study(study, sf_vocabulary)

    assert result == SyncResult(study, workbook, workbook_action="regenerated")
    assert "scatters_Fig2" not in content_of(workbook).sheets
    assert_in_step(study, workbook)


def test_a_copied_and_renamed_sheet_creates_its_table(study, workbook, sf_vocabulary):
    def copy(book):
        sheet = book.copy_worksheet(book["outputs_Tab2"])
        sheet.title = "outputs_Tab3"
        sheet[f"{letter(OUTPUTS, 'mean')}2"] = 7

    edit(workbook, copy)
    original = table_lines(study, OUTPUTS)

    result = sync_study(study, sf_vocabulary)

    assert result.ok, result.issues
    assert result.changes == (FileChange("outputs_Tab3.tsv", "write"),)
    assert result.workbook_action == "unchanged"
    assert table_lines(study, "outputs_Tab3.tsv") == [
        original[0],
        replaced(original[1], OUTPUTS, source="Tab3", mean="7"),
    ]
    assert_canonical(study)


def test_a_removed_sheet_deletes_its_table(study, workbook, sf_vocabulary):
    edit(workbook, remove_sheet("outputs_Tab2"))

    result = sync_study(study, sf_vocabulary)

    assert result.ok, result.issues
    assert result.changes == (FileChange(OUTPUTS, "delete"),)
    assert result.workbook_action == "unchanged"
    assert not (study / OUTPUTS).exists()
    assert read_state(workbook, generation_of(workbook)) == {OUTPUTS: None}
    assert sync_study(study, sf_vocabulary) == SyncResult(study, workbook)


def test_a_removed_sheet_whose_table_changed_conflicts(study, workbook, sf_vocabulary):
    base = table_lines(study, OUTPUTS)
    edit(workbook, remove_sheet("outputs_Tab2"))
    edit_table(study, OUTPUTS, 2, mean="3.5")
    changed = table_lines(study, OUTPUTS)
    before = snapshot(study)

    result = sync_study(study, sf_vocabulary)

    assert not result.ok
    assert result.conflicts == (
        SyncConflict(
            OUTPUTS,
            "outputs_Tab2",
            workbook_rows=(),
            table_lines=tuple(enumerate(changed, start=1)),
            base_lines=tuple(base),
        ),
    )
    [issue] = result.issues
    assert issue.code == "sync_conflict"
    assert issue.source is not None
    assert (issue.source.sheet, issue.source.row) == ("outputs_Tab2", None)
    assert snapshot(study) == before


@pytest.mark.parametrize("keep", ["workbook", "tables"])
def test_keep_resolves_a_removed_sheet_whose_table_changed(
    study, workbook, sf_vocabulary, keep
):
    edit(workbook, remove_sheet("outputs_Tab2"))
    edit_table(study, OUTPUTS, 2, mean="3.5")

    result = sync_study(study, sf_vocabulary, keep=keep)

    assert result.ok, result.issues
    assert [conflict.kept for conflict in result.conflicts] == [keep]
    if keep == "workbook":
        assert result.changes == (FileChange(OUTPUTS, "delete"),)
        assert result.workbook_action == "unchanged"
        assert not (study / OUTPUTS).exists()
    else:
        assert result.changes == ()
        assert result.workbook_action == "regenerated"
        assert cell(study, OUTPUTS, 2, "mean") == "3.5"
        assert_in_step(study, workbook)


@pytest.mark.parametrize("restored", ["same", "changed"])
def test_a_table_removed_on_both_sides_can_be_restored(
    study, workbook, sf_vocabulary, restored
):
    original = (study / OUTPUTS).read_bytes()
    edit(workbook, remove_sheet("outputs_Tab2"))
    (study / OUTPUTS).unlink()

    removed = sync_study(study, sf_vocabulary)

    assert removed == SyncResult(study, workbook)
    # The removal on both sides is the new base of the table.
    assert read_state(workbook, generation_of(workbook)) == {OUTPUTS: None}

    # Restored as git checkout would, or with an edit.
    (study / OUTPUTS).write_bytes(original)
    if restored == "changed":
        edit_table(study, OUTPUTS, 2, mean="3.5")
    expected = (study / OUTPUTS).read_bytes()

    result = sync_study(study, sf_vocabulary)

    assert result == SyncResult(study, workbook, workbook_action="regenerated")
    assert (study / OUTPUTS).read_bytes() == expected
    assert "outputs_Tab2" in content_of(workbook).sheets
    assert_in_step(study, workbook)


@pytest.fixture
def case_insensitive(monkeypatch):
    """The table writes of the sync on a case-insensitive file system, as APFS or NTFS.

    A name refers to the existing file whose name equals it ignoring case, and
    replacing that file keeps its name.
    """
    write, unlink = sync.atomic_text, Path.unlink

    def existing(path):
        for entry in path.parent.iterdir():
            if entry.name.casefold() == path.name.casefold():
                return entry
        return path

    monkeypatch.setattr(
        sync, "atomic_text", lambda path, text: write(existing(path), text)
    )
    monkeypatch.setattr(
        Path,
        "unlink",
        lambda path, missing_ok=False: unlink(existing(path), missing_ok=missing_ok),
    )


def rename_sheet(old, new):
    def change(workbook):
        # openpyxl compares sheet names ignoring case, as Excel does.
        workbook[old].title = "_renaming"
        workbook["_renaming"].title = new

    return change


@pytest.mark.parametrize("file_system", ["case-sensitive", "case-insensitive"])
@pytest.mark.parametrize(("old", "new"), [("Tab2a", "Tab2A"), ("Tab2A", "Tab2a")])
def test_a_sheet_renamed_only_in_case_keeps_its_table(
    valid_study, sf_vocabulary, request, file_system, old, new
):
    (valid_study / OUTPUTS).rename(valid_study / f"outputs_{old}.tsv")
    assert format_folder(valid_study).ok
    assert sync_study(valid_study, sf_vocabulary).workbook_action == "created"
    path = workbook_path(valid_study)
    rows = table_lines(valid_study, f"outputs_{old}.tsv")
    edit(path, rename_sheet(f"outputs_{old}", f"outputs_{new}"))
    if file_system == "case-insensitive":
        request.getfixturevalue("case_insensitive")

    result = sync_study(valid_study, sf_vocabulary)

    assert result.ok, result.issues
    assert table_lines(valid_study, f"outputs_{new}.tsv") == [
        rows[0],
        replaced(rows[1], OUTPUTS, source=new),
    ]
    # The delete comes first, so it never removes the file just written.
    assert result.changes == (
        FileChange(f"outputs_{old}.tsv", "delete"),
        FileChange(f"outputs_{new}.tsv", "write"),
    )
    again = sync_study(valid_study, sf_vocabulary)
    assert again.ok, again.issues
    assert again.changes == ()
    assert [name for name in tables_of(valid_study) if name.startswith("outputs")] == [
        f"outputs_{new}.tsv"
    ]
    content = content_of(path)
    assert f"outputs_{new}" in content.sheets
    assert content.tables[f"outputs_{new}.tsv"].text == (
        valid_study / f"outputs_{new}.tsv"
    ).read_text(encoding="utf-8")


def test_a_table_deleted_while_its_sheet_changed_conflicts(
    study, workbook, sf_vocabulary
):
    base = table_lines(study, OUTPUTS)
    set_cells(workbook, "outputs_Tab2", 2, mean=3.5)
    (study / OUTPUTS).unlink()
    before = snapshot(study)

    result = sync_study(study, sf_vocabulary)

    assert not result.ok
    assert result.conflicts == (
        SyncConflict(
            OUTPUTS,
            "outputs_Tab2",
            workbook_rows=((1, base[0]), (2, replaced(base[1], OUTPUTS, mean="3.5"))),
            table_lines=(),
            base_lines=tuple(base),
        ),
    )
    [issue] = result.issues
    assert issue.code == "sync_conflict"
    assert issue.source is not None
    assert issue.source.row == 1
    assert result.changes == ()
    assert result.workbook_action == "unchanged"
    assert snapshot(study) == before


def test_the_subjects_table_is_never_deleted(study, workbook, sf_vocabulary):
    (study / SUBJECTS).unlink()
    before = snapshot(study)

    result = sync_study(study, sf_vocabulary)

    assert not result.ok
    assert codes(result) == ["missing_file"]
    assert result.workbook_action == "unchanged"
    assert snapshot(study) == before


def formula(study, workbook):
    set_cells(workbook, "scatters_Fig2", 2, y_mean="=1+1")


def conflict_markers(study, workbook):
    with (study / SCATTERS).open("a", encoding="utf-8", newline="") as stream:
        stream.write("<<<<<<< HEAD\n")


def unknown_sheet(study, workbook):
    edit(workbook, lambda book: book.create_sheet("notes"))


@pytest.mark.parametrize(
    ("damage", "code"),
    [
        (formula, "formula_without_value"),
        (conflict_markers, "merge_conflict"),
        (unknown_sheet, "unknown_sheet"),
    ],
)
def test_errors_write_nothing(study, workbook, sf_vocabulary, damage, code):
    set_cells(workbook, "timecourses_Fig1", 3, mean=2.25)
    edit_table(study, OUTPUTS, 2, mean="3.5")
    damage(study, workbook)
    before = snapshot(study)

    result = sync_study(study, sf_vocabulary)

    assert not result.ok
    assert code in codes(result)
    assert result.changes == ()
    assert result.workbook_action == "unchanged"
    assert snapshot(study) == before


def damage_base(workbook):
    workbook["_base"]["A1"] = "something else"


@pytest.mark.parametrize(
    ("damage", "code"),
    [
        (remove_sheet("_base"), "workbook_base_missing"),
        (damage_base, "workbook_base_invalid"),
    ],
)
def test_a_lost_base_is_restored_while_the_workbook_is_in_step(
    study, workbook, sf_vocabulary, damage, code
):
    edit(workbook, damage)
    edit(workbook, lambda book: book.create_sheet("_scratch").cell(1, 1, "kept"))
    tables = tables_of(study)

    result = sync_study(study, sf_vocabulary)

    assert result.ok
    assert codes(result) == [code]
    assert result.changes == ()
    assert result.workbook_action == "regenerated"
    assert tables_of(study) == tables
    assert_in_step(study, workbook)
    assert "_scratch" in openpyxl.load_workbook(workbook).sheetnames

    # The base merges the next changes of both sides.
    set_cells(workbook, "timecourses_Fig1", 2, mean=0.25)
    edit_table(study, TIMECOURSES, 4, mean="7")
    merged = sync_study(study, sf_vocabulary)
    assert merged.ok, merged.issues
    assert merged.issues == ()
    assert [cell(study, TIMECOURSES, line, "mean") for line in (2, 4)] == ["0.25", "7"]


def test_a_lost_base_of_an_open_workbook_is_restored_after_it_is_closed(
    study, workbook, sf_vocabulary
):
    edit(workbook, remove_sheet("_base"))
    (study / LOCK).write_text("lock")
    before = snapshot(study)

    result = sync_study(study, sf_vocabulary)

    assert result.ok
    assert codes(result) == ["workbook_base_missing", "workbook_open"]
    assert result.workbook_action == "close_to_update"
    issue = result.issues[1]
    assert "restore" in issue.message
    assert "restores its base" in issue.suggestions[0].message
    assert snapshot(study) == before

    (study / LOCK).unlink()
    assert sync_study(study, sf_vocabulary).workbook_action == "regenerated"
    assert_in_step(study, workbook)


@pytest.mark.parametrize("code", ["workbook_base_missing", "workbook_base_invalid"])
def test_a_lost_base_names_how_to_restore_it(study, workbook, sf_vocabulary, code):
    edit(workbook, remove_sheet("_base") if code.endswith("missing") else damage_base)
    issue = read_workbook(workbook, STUDY).issues[0]
    assert issue.code == code
    assert issue.suggestions[0].message == (
        "Close the workbook and run pkdb tables sync; if the tables and the "
        "workbook differ, choose a side with --keep"
    )


@pytest.mark.parametrize("keep", [None, "workbook", "tables"])
def test_without_a_base_differences_conflict_unless_kept(
    study, workbook, sf_vocabulary, keep
):
    edit(workbook, remove_sheet("_base"))
    set_cells(workbook, "timecourses_Fig1", 3, mean=2.25)
    sheet = lines(read_workbook(workbook, STUDY).tables[TIMECOURSES].text)
    tables = table_lines(study, TIMECOURSES)
    before = snapshot(study)

    result = sync_study(study, sf_vocabulary, keep=keep)

    assert result.conflicts == (
        SyncConflict(
            TIMECOURSES,
            "timecourses_Fig1",
            workbook_rows=tuple(enumerate(sheet, start=1)),
            table_lines=tuple(enumerate(tables, start=1)),
            base_lines=(),
            kept=keep,
        ),
    )
    if keep is None:
        assert not result.ok
        assert codes(result) == ["workbook_base_missing", "sync_conflict"]
        assert snapshot(study) == before
    elif keep == "workbook":
        assert result.ok
        assert result.changes == (FileChange(TIMECOURSES, "write"),)
        # The workbook gets its base back.
        assert result.workbook_action == "regenerated"
        assert table_lines(study, TIMECOURSES) == sheet
        assert_in_step(study, workbook)
        set_cells(workbook, "timecourses_Fig1", 4, mean=9)
        edit_table(study, OUTPUTS, 2, mean="3")
        merged = sync_study(study, sf_vocabulary)
        assert merged.ok and merged.issues == (), merged.issues
    else:
        assert result.ok
        assert result.changes == ()
        assert result.workbook_action == "regenerated"
        assert_in_step(study, workbook)


def test_without_a_base_a_table_on_one_side_is_kept(study, workbook, sf_vocabulary):
    scatters = (study / SCATTERS).read_bytes()
    edit(workbook, remove_sheet("_base"))
    edit(workbook, remove_sheet("outputs_Tab2"))
    (study / SCATTERS).unlink()

    result = sync_study(study, sf_vocabulary)

    # The union of both sides: nothing tells a removal from an addition.
    assert result.ok, result.issues
    assert codes(result) == ["workbook_base_missing"]
    assert result.changes == (FileChange(SCATTERS, "write"),)
    assert result.workbook_action == "regenerated"
    assert (study / SCATTERS).read_bytes() == scatters
    assert (study / OUTPUTS).exists()
    assert_in_step(study, workbook)


def test_a_failed_write_keeps_every_file_whole(
    study, workbook, sf_vocabulary, monkeypatch
):
    set_cells(workbook, "outputs_Tab2", 2, mean=3.5)
    set_cells(workbook, "timecourses_Fig1", 3, mean=2.25)
    edit_table(study, SCATTERS, 2, y_mean="2.5")
    generation = generation_of(workbook)
    before = snapshot(study)
    replace = Path.replace
    calls = []

    def fail_second(self, target):
        calls.append(target)
        if len(calls) == 2:
            raise OSError("No space left on device")
        return replace(self, target)

    monkeypatch.setattr(Path, "replace", fail_second)

    result = sync_study(study, sf_vocabulary)

    assert not result.ok
    assert codes(result) == ["sync_write_failed"]
    assert "No space left on device" in result.issues[0].message
    assert result.changes == (FileChange(OUTPUTS, "write"),)
    assert result.workbook_action == "unchanged"
    after = snapshot(study)
    assert cell(study, OUTPUTS, 2, "mean") == "3.5"
    assert after[TIMECOURSES] == before[TIMECOURSES]
    assert after[workbook.name] == before[workbook.name]
    assert not [name for name in after if name.startswith(".tmp")]
    assert_canonical(study)
    # The table written from the workbook is recorded, the other one is not.
    assert read_state(workbook, generation) == {
        OUTPUTS: (study / OUTPUTS).read_text(encoding="utf-8")
    }

    monkeypatch.undo()
    result = sync_study(study, sf_vocabulary)

    assert result.ok, result.issues
    assert result.changes == (FileChange(TIMECOURSES, "write"),)
    assert result.workbook_action == "regenerated"
    assert cell(study, TIMECOURSES, 3, "mean") == "2.25"
    assert_in_step(study, workbook)


def test_a_failed_state_write_stops_before_the_workbook(
    study, workbook, sf_vocabulary, monkeypatch
):
    set_cells(workbook, "timecourses_Fig1", 3, mean=2.25)
    edit_table(study, OUTPUTS, 2, mean="3.5")
    opened = workbook.read_bytes()

    def fail(*arguments):
        raise OSError("Read-only file system")

    monkeypatch.setattr("pkdb.studyformat.sync.write_state", fail)

    result = sync_study(study, sf_vocabulary)

    assert not result.ok
    [issue] = result.issues
    assert issue.code == "sync_write_failed"
    assert issue.message.startswith(f"{state_path(workbook).name} cannot be updated")
    assert result.changes == (FileChange(TIMECOURSES, "write"),)
    assert result.workbook_action == "unchanged"
    assert cell(study, TIMECOURSES, 3, "mean") == "2.25"
    assert workbook.read_bytes() == opened


@pytest.mark.parametrize("meanwhile", ["save", "open"])
def test_a_workbook_saved_or_opened_during_the_sync_is_not_replaced(
    study, workbook, sf_vocabulary, monkeypatch, meanwhile
):
    edit_table(study, OUTPUTS, 2, mean="3.5")
    generation = generation_of(workbook)
    build = sync.build_workbook

    def build_meanwhile(*arguments, **options):
        built = build(*arguments, **options)
        if meanwhile == "save":
            # A save that leaves no lock file, such as of a synced folder.
            set_cells(workbook, "timecourses_Fig1", 3, mean=2.25)
        else:
            (study / LOCK).write_text("curator", encoding="utf-8")
        return built

    monkeypatch.setattr(sync, "build_workbook", build_meanwhile)

    result = sync_study(study, sf_vocabulary)
    current = workbook.read_bytes()

    assert result.ok, result.issues
    assert result.changes == ()
    [issue] = result.issues
    assert issue.severity == "warning"
    if meanwhile == "save":
        assert result.workbook_action == "sync_again"
        assert issue.code == "workbook_changed"
        assert result.lock is None
    else:
        assert result.workbook_action == "close_to_update"
        assert issue.code == "workbook_open"
        assert result.lock == study / LOCK
        (study / LOCK).unlink()
    assert generation_of(workbook) == generation
    assert workbook.read_bytes() == current

    monkeypatch.undo()
    result = sync_study(study, sf_vocabulary)

    assert result.ok, result.issues
    assert result.workbook_action == "regenerated"
    assert cell(study, OUTPUTS, 2, "mean") == "3.5"
    if meanwhile == "save":
        assert result.changes == (FileChange(TIMECOURSES, "write"),)
        assert cell(study, TIMECOURSES, 3, "mean") == "2.25"
    assert_in_step(study, workbook)


def test_a_state_file_that_cannot_be_removed_stays(
    valid_study, sf_vocabulary, monkeypatch
):
    path = workbook_path(valid_study)
    write_state(path, "0" * 32, {OUTPUTS: None})

    def fail(workbook):
        raise PermissionError("Permission denied")

    monkeypatch.setattr(sync, "remove_state", fail)

    created = sync_study(valid_study, sf_vocabulary)
    edit_table(valid_study, OUTPUTS, 2, mean="3.5")
    regenerated = sync_study(valid_study, sf_vocabulary)

    assert created == SyncResult(valid_study, path, workbook_action="created")
    assert regenerated == SyncResult(valid_study, path, workbook_action="regenerated")
    # The state file of an old generation is ignored.
    assert state_path(path).exists()
    assert read_state(path, generation_of(path)) == {}


def test_check_mode_reports_the_plan_without_writing(study, workbook, sf_vocabulary):
    set_cells(workbook, "timecourses_Fig1", 3, mean=2.25)
    edit_table(study, OUTPUTS, 2, mean="3.5")
    before = snapshot(study)

    result = sync_study(study, sf_vocabulary, check=True)

    assert result == SyncResult(
        study,
        workbook,
        changes=(FileChange(TIMECOURSES, "write"),),
        workbook_action="regenerated",
        checked=True,
    )
    assert snapshot(study) == before


@pytest.mark.parametrize("check", [False, True])
def test_a_missing_workbook_with_a_lock_file_is_not_created(
    study, workbook, sf_vocabulary, check
):
    # Excel on Windows renames the workbook while it saves it.
    write_state(workbook, "0123456789abcdef0123456789abcdef", {OUTPUTS: None})
    workbook.unlink()
    (study / LOCK).write_text("lock")
    before = snapshot(study)

    result = sync_study(study, sf_vocabulary, check=check)

    assert result.ok
    assert result.workbook_action == "close_to_update"
    assert result.lock == study / LOCK
    [issue] = result.issues
    assert (issue.code, issue.severity) == ("workbook_open", "warning")
    assert "is missing" in issue.message and LOCK in issue.message
    assert not workbook.exists()
    assert snapshot(study) == before


def test_check_mode_without_a_workbook(valid_study, sf_vocabulary):
    before = snapshot(valid_study)

    result = sync_study(valid_study, sf_vocabulary, check=True)

    assert result == SyncResult(
        valid_study,
        workbook_path(valid_study),
        workbook_action="created",
        checked=True,
    )
    assert snapshot(valid_study) == before


def test_a_workbook_saved_by_libreoffice_is_in_step(
    study, workbook, sf_vocabulary, libreoffice_resave
):
    workbook.write_bytes(libreoffice_resave(workbook).read_bytes())
    before = snapshot(study)

    result = sync_study(study, sf_vocabulary)

    assert result == SyncResult(study, workbook)
    assert snapshot(study) == before


@pytest.mark.parametrize("text", ["#N/A", "#DIV/0!", "#REF!"])
def test_text_that_looks_like_an_error_value_syncs(
    study, workbook, sf_vocabulary, libreoffice_resave, text
):
    edit_table(study, TIMECOURSES, 2, comment=text)
    assert sync_study(study, sf_vocabulary).workbook_action == "regenerated"
    workbook.write_bytes(libreoffice_resave(workbook).read_bytes())
    tables = tables_of(study)

    result = sync_study(study, sf_vocabulary)

    assert result.ok, result.issues
    assert result.issues == ()
    assert result.changes == ()
    assert tables_of(study) == tables
    assert cell(study, TIMECOURSES, 2, "comment") == text


def test_a_formula_is_synced_as_its_saved_value(
    study, workbook, sf_vocabulary, libreoffice_resave
):
    set_cells(workbook, "timecourses_Fig1", 3, mean="=2*1.125")
    workbook.write_bytes(libreoffice_resave(workbook).read_bytes())
    edit_table(study, OUTPUTS, 2, mean="3.5")

    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        result = sync_study(study, sf_vocabulary)

    assert [str(warning.message) for warning in caught] == []
    assert result.ok, result.issues
    assert codes(result) == ["formula_value"]
    assert result.changes == (FileChange(TIMECOURSES, "write"),)
    assert result.workbook_action == "regenerated"
    assert cell(study, TIMECOURSES, 3, "mean") == "2.25"
    assert_in_step(study, workbook)


def test_a_scratch_sheet_survives_a_regeneration(study, workbook, sf_vocabulary):
    def notes(book):
        sheet = book.create_sheet("_notes")
        sheet["A1"] = "keep me"
        sheet["B1"] = "=SUM(1,2)"

    edit(workbook, notes)
    edit_table(study, OUTPUTS, 2, mean="3.5")

    result = sync_study(study, sf_vocabulary)

    assert result.workbook_action == "regenerated"
    sheet = openpyxl.load_workbook(workbook)["_notes"]
    assert sheet["A1"].value == "keep me"
    assert sheet["B1"].value == "=SUM(1,2)"


def test_rows_follow_the_subject_order_of_the_merged_subjects(
    study, workbook, sf_vocabulary
):
    # The tables make S1 a child of S2, so S2 now sorts before S1.
    subjects = table_lines(study, SUBJECTS)
    line = next(
        number
        for number, text in enumerate(subjects, start=1)
        if text.split("\t")[names(SUBJECTS).index("name")] == "S1"
    )
    edit_table(study, SUBJECTS, line, parent="S2")
    assert format_folder(study).ok

    # The workbook adds outputs of S1 and S2, sorted by the old subject order.
    def add(book):
        sheet = book["outputs_Tab2"]
        template = [source.value for source in sheet[2]]
        for row, subject in ((3, "S1"), (4, "S2")):
            for column, value in enumerate(template, start=1):
                sheet.cell(row, column).value = value
            sheet[f"{letter(OUTPUTS, 'subjects')}{row}"] = subject

    edit(workbook, add)

    result = sync_study(study, sf_vocabulary)

    assert result.ok, result.issues
    assert result.changes == (FileChange(OUTPUTS, "write"),)
    assert result.workbook_action == "regenerated"
    order = [cell(study, OUTPUTS, number, "subjects") for number in (2, 3, 4)]
    assert order.index("S2") < order.index("S1")
    assert_canonical(study)
    assert_in_step(study, workbook)


def test_changes_to_different_subjects_merge(study, workbook, sf_vocabulary):
    subjects = table_lines(study, SUBJECTS)
    assert [
        line.split("\t")[names(SUBJECTS).index("name")] for line in subjects[1:]
    ] == [
        "all",
        "S1",
        "S2",
    ]
    set_cells(workbook, "subjects", 2, count=3)
    edit_table(study, SUBJECTS, 4, count="2")

    result = sync_study(study, sf_vocabulary)

    assert result.ok, result.issues
    assert result.changes == (FileChange(SUBJECTS, "write"),)
    assert result.workbook_action == "regenerated"
    assert [cell(study, SUBJECTS, line, "count") for line in (2, 3, 4)] == [
        "3",
        "1",
        "2",
    ]
    assert_canonical(study)
    assert_in_step(study, workbook)


def add_table_row(study, workbook):
    with (study / OUTPUTS).open("a", encoding="utf-8", newline="") as stream:
        stream.write(replaced(table_lines(study, OUTPUTS)[1], OUTPUTS, mean="9") + "\n")


def add_sheet_row(study, workbook):
    def change(book):
        sheet = book["outputs_Tab2"]
        for column, source in enumerate(list(sheet[2]), start=1):
            sheet.cell(3, column).value = source.value

    edit(workbook, change)


@pytest.mark.parametrize("grow", [add_table_row, add_sheet_row])
def test_more_rows_than_the_limit_are_an_error(study, workbook, sf_vocabulary, grow):
    rows = sum(len(table_lines(study, file)) - 1 for file in tables_of(study))
    grow(study, workbook)
    before = snapshot(study)

    result = sync_study(study, sf_vocabulary, max_rows=rows)

    assert not result.ok
    assert codes(result) == ["row_limit"]
    assert result.changes == ()
    assert snapshot(study) == before


@pytest.mark.parametrize(
    "damage",
    ["{", '{"format": 2, "reference": 1}'],
    ids=["invalid_json", "invalid_study_json"],
)
def test_a_broken_study_json_does_not_stop_the_sync(
    study, workbook, sf_vocabulary, damage
):
    (study / "study.json").write_text(damage, encoding="utf-8")
    set_cells(workbook, "outputs_Tab2", 2, mean=3.25)

    result = sync_study(study, sf_vocabulary)

    # pkdb validate reports study.json; merging the tables does not need it.
    assert result.ok, result.issues
    assert result.issues == ()
    assert result.changes == (FileChange(OUTPUTS, "write"),)
    assert cell(study, OUTPUTS, 2, "mean") == "3.25"
    # The study name comes from the folder.
    assert cell(study, OUTPUTS, 2, "study") == STUDY
    check = workbook_check(study, sf_vocabulary)
    assert check is not None and check["ok"] is True
    assert add_table(study, sf_vocabulary, "outputs_Tab3").ok


def test_keep_must_name_a_side(study, sf_vocabulary):
    with pytest.raises(ValueError, match="keep"):
        sync_study(study, sf_vocabulary, keep="both")  # ty: ignore[invalid-argument-type]


def test_a_large_study_syncs_quickly(make_study, valid_files, tsv, sf_vocabulary):
    rows = [
        {
            "subjects": "all",
            "interventions": "D1",
            "measurement": "cmax",
            "substance": "drug",
            "tissue": "plasma",
            "time": str(index % 24),
            "time_unit": "h",
            "mean": str(index * 0.25),
            "sd": "0.5",
            "unit": "mg/l",
            "comment": f"row {index}",
        }
        for index in range(20_000)
    ]
    folder = make_study({**valid_files, "outputs_Tab1.tsv": tsv("outputs", *rows)})
    assert format_folder(folder).ok
    assert sync_study(folder, sf_vocabulary).workbook_action == "created"
    path = workbook_path(folder)
    first, last = (
        cell(folder, "outputs_Tab1.tsv", line, "comment") for line in (2, 20_001)
    )
    set_cells(path, "outputs_Tab1", 2, mean=99)
    edit_table(folder, "outputs_Tab1.tsv", 20_001, sd="0.75")

    start = time.perf_counter()
    result = sync_study(folder, sf_vocabulary)
    elapsed = time.perf_counter() - start

    assert result.ok, result.issues[:3]
    assert result.changes == (FileChange("outputs_Tab1.tsv", "write"),)
    assert result.workbook_action == "regenerated"
    # The edits may move the rows, which sort by all their cells in the end.
    rows = {
        row["comment"]: row
        for row in (
            dict(zip(names("outputs_Tab1.tsv"), line.split("\t"), strict=True))
            for line in table_lines(folder, "outputs_Tab1.tsv")[1:]
        )
    }
    assert len(rows) == 20_000
    assert rows[first]["mean"] == "99"
    assert rows[last]["sd"] == "0.75"
    assert elapsed < 30, f"syncing 20,000 rows took {elapsed:.1f} s"


def test_add_table_adds_an_empty_sheet(study, workbook, sf_vocabulary):
    set_cells(workbook, "outputs_Tab2", 2, mean=3.25)

    result = add_table(study, sf_vocabulary, "outputs_Tab3")

    assert result.ok, result.issues
    assert result.sync is not None
    assert result.sync.changes == (FileChange(OUTPUTS, "write"),)
    assert [(issue.code, issue.severity) for issue in result.issues] == [
        ("missing_image", "warning")
    ]
    content = content_of(workbook)
    assert "outputs_Tab3" in content.sheets
    assert "outputs_Tab3.tsv" not in content.tables
    assert_in_step(study, workbook)
    assert cell(study, OUTPUTS, 2, "mean") == "3.25"


@pytest.mark.parametrize(
    ("table", "code"),
    [
        ("results_Tab3", "invalid_table_name"),
        ("interventions", "invalid_table_name"),
        ("outputs_Tab3_with_a_very_long_name_x", "table_name_too_long"),
        ("outputs_Tab2", "table_exists"),
    ],
)
def test_add_table_rejects_a_name_before_the_sync(
    study, workbook, sf_vocabulary, table, code
):
    before = snapshot(study)

    result = add_table(study, sf_vocabulary, table)

    assert not result.ok
    assert result.sync is None
    assert codes(result) == [code]
    assert snapshot(study) == before


def test_add_table_rejects_a_sheet_that_exists_ignoring_case(
    study, workbook, sf_vocabulary
):
    assert add_table(study, sf_vocabulary, "outputs_TabA").ok
    before = snapshot(study)

    result = add_table(study, sf_vocabulary, "outputs_Taba")

    assert not result.ok
    assert result.sync is not None and result.sync.ok
    assert codes(result) == ["table_exists"]
    assert snapshot(study) == before


def test_add_table_needs_a_closed_workbook(study, workbook, sf_vocabulary):
    (study / LOCK).write_text("curator", encoding="utf-8")
    before = snapshot(study)

    result = add_table(study, sf_vocabulary, "outputs_Tab3")

    assert not result.ok
    [issue] = result.issues
    assert (issue.code, issue.severity) == ("workbook_open", "error")
    assert "rename it to outputs_Tab3" in issue.message
    assert snapshot(study) == before


def test_add_table_stops_when_the_sync_fails(study, workbook, sf_vocabulary):
    set_cells(workbook, "timecourses_Fig1", 2, mean=0.25)
    edit_table(study, TIMECOURSES, 2, mean="0.75")
    before = snapshot(study)

    result = add_table(study, sf_vocabulary, "outputs_Tab3")

    assert not result.ok
    assert result.sync is not None and codes(result.sync) == ["sync_conflict"]
    assert result.issues == ()
    assert snapshot(study) == before


def test_add_table_reports_a_workbook_it_cannot_read(
    study, workbook, sf_vocabulary, monkeypatch
):
    from pkdb.schemas.validation import fail

    calls = []
    read = sync.read_workbook

    def limited(*arguments, **options):
        # The sync reads the workbook first; the read before rebuilding fails.
        calls.append(arguments)
        if len(calls) > 1:
            fail("row_limit", "The workbook has too many rows")
        return read(*arguments, **options)

    monkeypatch.setattr(sync, "read_workbook", limited)
    before = snapshot(study)

    result = add_table(study, sf_vocabulary, "outputs_Tab3")

    assert not result.ok
    assert codes(result) == ["row_limit"]
    assert snapshot(study) == before


def test_add_table_never_replaces_a_workbook_saved_meanwhile(
    study, workbook, sf_vocabulary, monkeypatch
):
    build = sync.build_workbook
    saved = []

    def build_meanwhile(*arguments, **options):
        built = build(*arguments, **options)
        set_cells(workbook, "timecourses_Fig1", 3, mean=2.25)
        saved.append(workbook.read_bytes())
        return built

    monkeypatch.setattr(sync, "build_workbook", build_meanwhile)

    result = add_table(study, sf_vocabulary, "outputs_Tab3")

    assert not result.ok
    [issue] = result.issues
    assert (issue.code, issue.severity) == ("workbook_changed", "error")
    # The save during add_table is kept as it was saved.
    assert [workbook.read_bytes()] == saved
    content = content_of(workbook)
    assert "outputs_Tab3" not in content.sheets
    [_, _, row, _] = lines(content.tables[TIMECOURSES].text)
    assert row.split("\t")[names(TIMECOURSES).index("mean")] == "2.25"


def test_workbook_check_plans_without_writing(study, workbook, sf_vocabulary):
    in_step = {"action": "unchanged", "changes": [], "conflicts": 0, "ok": True}
    assert workbook_check(study, sf_vocabulary) == in_step
    set_cells(workbook, "outputs_Tab2", 2, mean=3.25)
    before = snapshot(study)

    assert workbook_check(study, sf_vocabulary) == {
        **in_step,
        "changes": [{"file": OUTPUTS, "action": "write"}],
    }
    assert snapshot(study) == before


def test_workbook_check_without_a_workbook(valid_study, sf_vocabulary):
    assert workbook_check(valid_study, sf_vocabulary) is None
    assert not workbook_path(valid_study).exists()


@pytest.mark.parametrize(("broken", "ok"), [("workbook", False), ("tables", None)])
def test_workbook_check_names_why_it_cannot_sync(
    study, workbook, sf_vocabulary, broken, ok
):
    if broken == "workbook":
        set_cells(workbook, "outputs_Tab2", 2, mean="#DIV/0!")
    else:
        path = study / OUTPUTS
        path.write_text(path.read_text(encoding="utf-8").replace("mean", "x", 1))

    check = workbook_check(study, sf_vocabulary)

    assert check is not None
    assert check["ok"] is ok
