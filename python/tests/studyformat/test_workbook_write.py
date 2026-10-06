"""Tests of the generated study workbook.

The snapshot `data/workbook_template.json` describes the sheets of the template.
After an intended change of the template, regenerate it from `python/` with:

    uv run --locked python -c "import json, pathlib; from pkdb.studyformat.workbook.write import workbook_template; pathlib.Path('tests/studyformat/data/workbook_template.json').write_text(json.dumps(workbook_template(), indent=2, ensure_ascii=False) + chr(10), encoding='utf-8')"
"""

import json
import re
from datetime import UTC, datetime, timedelta
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest
from openpyxl.utils import get_column_letter

from pkdb.studyformat.tables import TABLES
from pkdb.studyformat.text import render_tsv
from pkdb.studyformat.workbook.base import WorkbookBase, parse_base
from pkdb.studyformat.workbook.write import (
    WorkbookError,
    build_workbook,
    workbook_template,
)

GENERATION = "0123456789abcdef0123456789abcdef"
CREATED = datetime(2026, 10, 6, 8, 30, tzinfo=UTC)
SNAPSHOT = Path(__file__).parent / "data" / "workbook_template.json"
REGENERATE = (
    'uv run --locked python -c "import json, pathlib; from pkdb.studyformat.workbook.write '
    "import workbook_template; pathlib.Path('tests/studyformat/data/workbook_template.json')"
    ".write_text(json.dumps(workbook_template(), indent=2, ensure_ascii=False) + chr(10), "
    "encoding='utf-8')\""
)


def canonical_tables(folder):
    return {
        path.name: path.read_text(encoding="utf-8")
        for path in sorted(folder.glob("*.tsv"))
    }


def build(tables, vocabulary, **options):
    return build_workbook(
        tables, vocabulary, generation=GENERATION, created=CREATED, **options
    )


def load(result):
    assert result.issues == []
    assert result.data is not None
    return openpyxl.load_workbook(BytesIO(result.data))


def letter(kind, name):
    return get_column_letter(TABLES[kind].names.index(name) + 1)


def outputs(*rows):
    names = TABLES["outputs"].names
    return render_tsv(
        names,
        [
            tuple(
                {"study": "Example", "source": "Tab1", "subjects": "all", **row}.get(
                    name, ""
                )
                for name in names
            )
            for row in rows
        ],
    )


def validations(sheet):
    return {
        str(validation.sqref): validation
        for validation in sheet.data_validations.dataValidation
    }


def list_columns(workbook):
    sheet = workbook["_lists"]
    return {
        cell.value: get_column_letter(cell.column) for cell in sheet[1] if cell.value
    }


def list_values(workbook, key):
    column = list_columns(workbook)[key]
    sheet = workbook["_lists"]
    return [
        cell.value
        for (cell,) in sheet[f"{column}2:{column}{sheet.max_row}"]
        if cell.value
    ]


def test_sheet_order(valid_study, sf_vocabulary):
    tables = canonical_tables(valid_study)
    del tables["interventions.tsv"]
    workbook = load(
        build(
            tables,
            sf_vocabulary,
            empty_sheets=["outputs_Tab10", "scatters_Fig1", "timecourses_Tab3"],
        )
    )
    assert workbook.sheetnames == [
        "subjects",
        "interventions",
        "characteristica",
        "outputs_Tab2",
        "outputs_Tab10",
        "timecourses_Fig1",
        "timecourses_Tab3",
        "scatters_Fig1",
        "scatters_Fig2",
        "_lists",
        "_base",
    ]
    assert workbook.active.title == "subjects"
    for name in ("interventions", "outputs_Tab10"):
        sheet = workbook[name]
        assert sheet.max_row == 1
        kind = name.partition("_")[0]
        assert tuple(cell.value for cell in sheet[1]) == TABLES[kind].names
    assert workbook["outputs_Tab2"].sheet_state == "visible"
    assert workbook["_lists"].sheet_state == "hidden"
    assert workbook["_base"].sheet_state == "veryHidden"


def test_header_is_bold_frozen_filtered_and_commented(valid_study, sf_vocabulary):
    workbook = load(build(canonical_tables(valid_study), sf_vocabulary))
    sheet = workbook["outputs_Tab2"]
    spec = TABLES["outputs"]
    header = sheet[1]
    assert tuple(cell.value for cell in header) == spec.names
    assert all(cell.font.bold for cell in header)
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == f"A1:{letter('outputs', 'comment')}2"
    measurement = header[spec.names.index("measurement")].comment
    assert measurement.author == "pkdb"
    assert measurement.text == (
        f"{spec.column('measurement').description}\nExample: cmax\nRequired"
    )
    assert (
        header[spec.names.index("comment")].comment.text
        == spec.column("comment").description
    )
    assert workbook["subjects"].auto_filter.ref == "A1:F4"


def test_column_widths_are_clamped(sf_vocabulary):
    tables = {"outputs_Tab1.tsv": outputs({"comment": "x" * 100, "mean": "1"})}
    sheet = load(build(tables, sf_vocabulary))["outputs_Tab1"]
    assert sheet.column_dimensions[letter("outputs", "comment")].width == 60
    assert sheet.column_dimensions[letter("outputs", "sd")].width == 8
    interventions = sheet.column_dimensions[letter("outputs", "interventions")].width
    assert len("interventions") < interventions < 20


def test_cell_types_follow_the_column_type(sf_vocabulary):
    tables = {
        "outputs_Tab1.tsv": outputs(
            {
                "mean": "0.1",
                "count": "12",
                "sd": "0.30000000000000004",
                "se": "123456789012345",
                "cv": "1234567890123456",
                "gmean": "2.5e-07",
                "median": "-3",
                "min": "abc",
                "time": "NR",
                "comment": "1-2",
                "choice": "007",
            },
            {"time": "1.5", "comment": "=x", "unit": "TRUE"},
        )
    }
    sheet = load(build(tables, sf_vocabulary))["outputs_Tab1"]

    def cell(row, name):
        return sheet[f"{letter('outputs', name)}{row}"]

    numbers = {
        (2, "mean"): 0.1,
        (2, "count"): 12,
        (2, "se"): 123456789012345,
        (2, "gmean"): 2.5e-07,
        (2, "median"): -3,
        (3, "time"): 1.5,
    }
    for (row, name), value in numbers.items():
        assert cell(row, name).value == value
        assert type(cell(row, name).value) is type(value)
        assert cell(row, name).data_type == "n"
    texts = {
        (2, "sd"): "0.30000000000000004",
        (2, "cv"): "1234567890123456",
        (2, "min"): "abc",
        (2, "time"): "NR",
        (2, "comment"): "1-2",
        (2, "choice"): "007",
        (3, "comment"): "=x",
        (3, "unit"): "TRUE",
        (2, "measurement"): "",
    }
    for (row, name), value in texts.items():
        assert cell(row, name).value == (value or None)
        if value:
            assert cell(row, name).data_type == "s"
            assert cell(row, name).number_format == "@"
    assert cell(3, "mean").value is None
    columns = sheet.column_dimensions
    assert columns[letter("outputs", "comment")].number_format == "@"
    assert columns[letter("outputs", "interventions")].number_format == "@"
    assert columns[letter("outputs", "mean")].number_format == "General"


def test_dropdowns(valid_study, sf_vocabulary):
    workbook = load(build(canonical_tables(valid_study), sf_vocabulary))
    lists = list_columns(workbook)
    measurements = len(list_values(workbook, "measurements"))

    def target(kind, name):
        column = letter(kind, name)
        return f"{column}2:{column}1048576"

    found = validations(workbook["outputs_Tab2"])
    column = lists["measurements"]
    assert (
        found[target("outputs", "measurement")].formula1
        == f"_lists!${column}$2:${column}${measurements + 1}"
    )
    assert found[target("outputs", "subjects")].formula1 == "subjects!$B$2:$B$1048576"
    assert (
        found[target("outputs", "interventions")].formula1
        == "interventions!$C$2:$C$1048576"
    )
    column = lists["choice:error_type"]
    assert (
        found[target("outputs", "error_type")].formula1
        == f"_lists!${column}$2:${column}$4"
    )
    for name in ("unit", "choice", "time_unit", "mean", "comment", "study"):
        assert target("outputs", name) not in found
    scatter = validations(workbook["scatters_Fig2"])
    for name in ("x_interventions", "y_interventions"):
        assert (
            scatter[target("scatters", name)].formula1
            == "interventions!$C$2:$C$1048576"
        )
    assert target("scatters", "x_unit") not in scatter
    subjects = validations(workbook["subjects"])
    assert list(subjects) == [target("subjects", "parent")]
    for sheet in workbook.worksheets:
        for validation in sheet.data_validations.dataValidation:
            assert validation.type == "list"
            assert validation.allow_blank
            assert not validation.showErrorMessage
            assert not validation.showDropDown


def test_lists_are_hidden_and_leave_out_deprecated_measurements(
    valid_study, sf_vocabulary
):
    workbook = load(build(canonical_tables(valid_study), sf_vocabulary))
    assert workbook["_lists"].sheet_state == "hidden"
    assert list(list_columns(workbook)) == [
        "measurements",
        "calculation_types",
        "substances",
        "tissues",
        "methods",
        "routes",
        "forms",
        "applications",
        "choice:error_type",
    ]
    measurements = list_values(workbook, "measurements")
    assert "old_measure" not in measurements
    assert "cmax" in measurements
    assert measurements == sorted(measurements, key=str.casefold)
    assert list_values(workbook, "calculation_types") == [
        "calculation",
        "sample mean",
        "unspecified summary",
    ]
    assert list_values(workbook, "choice:error_type") == ["gsd", "sd", "se"]


def test_base_holds_the_tables(valid_study, sf_vocabulary):
    tables = canonical_tables(valid_study)
    result = build(tables, sf_vocabulary, empty_sheets=["outputs_Tab3"])
    assert result.base == WorkbookBase(GENERATION, CREATED, tables)
    sheet = load(result)["_base"]
    assert sheet.sheet_state == "veryHidden"
    assert parse_base(sheet.iter_rows(values_only=True)) == result.base


def test_document_properties_use_the_creation_time(valid_study, sf_vocabulary):
    workbook = load(build(canonical_tables(valid_study), sf_vocabulary))
    naive = CREATED.replace(tzinfo=None)
    assert workbook.properties.created == naive
    assert workbook.properties.modified == naive


def test_generation_and_creation_time_default_to_new_values(sf_vocabulary):
    result = build_workbook({}, sf_vocabulary)
    assert re.fullmatch(r"[0-9a-f]{32}", result.base.generation)
    assert result.base.created.tzinfo == UTC
    assert datetime.now(UTC) - result.base.created < timedelta(minutes=1)
    assert build_workbook({}, sf_vocabulary).base.generation != result.base.generation


def test_regenerating_keeps_scratch_sheets(valid_study, sf_vocabulary, tmp_path):
    tables = canonical_tables(valid_study)
    workbook = load(build(tables, sf_vocabulary))
    notes = workbook.create_sheet("_notes", 2)
    notes["A1"] = "note"
    notes["B1"] = "=SUM(1,2)"
    workbook.create_sheet("outputs_Tab9")
    # Excel compares sheet names ignoring case, so this is the list sheet.
    workbook.remove(workbook["_lists"])
    workbook.create_sheet("_Lists")
    workbook.active = notes
    existing = tmp_path / "Example.xlsx"
    workbook.save(existing)

    regenerated = load(build(tables, sf_vocabulary, existing=existing))
    assert regenerated.sheetnames == [
        "subjects",
        "interventions",
        "characteristica",
        "outputs_Tab2",
        "timecourses_Fig1",
        "scatters_Fig2",
        "_notes",
        "_lists",
        "_base",
    ]
    notes = regenerated["_notes"]
    assert notes["A1"].value == "note"
    assert notes["B1"].value == "=SUM(1,2)"
    assert notes["B1"].data_type == "f"
    assert regenerated.active.title == "subjects"
    assert [sheet.sheet_view.tabSelected for sheet in regenerated.worksheets] == [
        True,
        *[False] * 8,
    ]


def test_an_unreadable_existing_workbook_raises(sf_vocabulary, tmp_path):
    existing = tmp_path / "Example.xlsx"
    existing.write_bytes(b"not a workbook")
    with pytest.raises(WorkbookError) as error:
        build({}, sf_vocabulary, existing=existing)
    assert error.value.code == "workbook_unreadable"
    assert existing.read_bytes() == b"not a workbook"


def test_a_long_sheet_name_raises(sf_vocabulary):
    with pytest.raises(WorkbookError) as error:
        build({}, sf_vocabulary, empty_sheets=["outputs_Tab" + "1" * 21])
    assert error.value.code == "table_name_too_long"


def test_sheet_names_equal_ignoring_case_raise(sf_vocabulary):
    tables = {"outputs_Tab1a.tsv": outputs({"mean": "1"})}
    with pytest.raises(WorkbookError) as error:
        build(tables, sf_vocabulary, empty_sheets=["outputs_Tab1A"])
    assert error.value.code == "duplicate_table_name"


@pytest.mark.parametrize(
    "tables",
    [
        {"notes.tsv": "a\n"},
        {"outputs_Tab1.tsv": "subjects\tmeasurement\nall\tcmax\n"},
    ],
)
def test_tables_must_be_canonical_table_files(sf_vocabulary, tables):
    with pytest.raises(ValueError):
        build(tables, sf_vocabulary)


def test_a_cell_longer_than_a_spreadsheet_cell_is_an_issue(sf_vocabulary):
    tables = {"outputs_Tab1.tsv": outputs({"mean": "1"}, {"comment": "x" * 32768})}
    result = build(tables, sf_vocabulary)
    assert result.data is None
    [issue] = result.issues
    assert issue.code == "cell_too_long"
    assert issue.severity == "error"
    assert issue.source.file == "outputs_Tab1.tsv"
    assert issue.source.cell == f"{letter('outputs', 'comment')}3"
    assert result.base.files == tables


def test_excel_counts_a_cell_in_utf16_code_units(sf_vocabulary):
    # 20,000 characters outside the Basic Multilingual Plane are 40,000 units.
    tables = {"outputs_Tab1.tsv": outputs({"comment": chr(0x1F600) * 20000})}
    result = build(tables, sf_vocabulary)
    assert result.data is None
    [issue] = result.issues
    assert issue.code == "cell_too_long"
    assert "40,000" in issue.message


# Characters that XML 1.0, and so a workbook, cannot hold: a control character,
# the noncharacters U+FFFE and U+FFFF, and a lone surrogate.
ILLEGAL = [0x01, 0x1F, 0xFFFE, 0xFFFF, 0xD800]


@pytest.mark.parametrize("code", ILLEGAL)
def test_a_character_xml_cannot_hold_is_an_issue(sf_vocabulary, code):
    tables = {"outputs_Tab1.tsv": outputs({"comment": f"a{chr(code)}b"})}
    result = build(tables, sf_vocabulary)
    assert result.data is None
    [issue] = result.issues
    assert issue.code == "illegal_character"
    assert f"U+{code:04X}" in issue.message
    assert issue.source.sheet == "outputs_Tab1"
    assert issue.source.cell == f"{letter('outputs', 'comment')}2"


@pytest.mark.parametrize("code", ILLEGAL)
def test_a_vocabulary_term_xml_cannot_hold_is_an_issue(sf_vocabulary, code):
    vocabulary = sf_vocabulary.model_copy(
        update={"tissues": ("plasma", f"bad{chr(code)}")}
    )
    result = build({}, vocabulary)
    assert result.data is None
    [issue] = result.issues
    assert issue.code == "illegal_character"
    assert issue.source is None
    assert "tissues" in issue.message
    assert f"U+{code:04X}" in issue.message


def test_tabs_line_breaks_and_other_planes_are_kept(sf_vocabulary):
    # A canonical cell holds no tab or line break, but a vocabulary term may.
    vocabulary = sf_vocabulary.model_copy(update={"tissues": ("a\tb", "c\nd")})
    comment = f"{chr(0x1F600)}{chr(0xFFFD)}{chr(0xE000)}{chr(0x10FFFF)}"
    workbook = load(
        build({"outputs_Tab1.tsv": outputs({"comment": comment})}, vocabulary)
    )
    sheet = workbook["outputs_Tab1"]
    assert sheet[f"{letter('outputs', 'comment')}2"].value == comment
    assert list_values(workbook, "tissues") == ["a\tb", "c\nd"]


def test_repeated_cell_issues_are_capped(sf_vocabulary):
    rows = [{"comment": f"{index}{chr(0x02)}"} for index in range(12)]
    result = build({"outputs_Tab1.tsv": outputs(*rows)}, sf_vocabulary)
    assert [issue.code for issue in result.issues] == ["illegal_character"] * 11
    assert "12" in result.issues[-1].message


def test_template_matches_the_snapshot():
    expected = json.loads(SNAPSHOT.read_text(encoding="utf-8"))
    assert workbook_template() == expected, (
        f"The workbook template changed; if intended, regenerate the snapshot from python/ with: {REGENERATE}"
    )
