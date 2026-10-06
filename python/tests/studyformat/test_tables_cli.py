"""The pkdb tables commands: open, sync and add a study workbook from the command line."""

import json
import os
import shlex
import shutil
import subprocess
from datetime import datetime

import openpyxl
import pytest
from openpyxl.utils import get_column_letter

from pkdb.cli import main
from pkdb.studyformat import sync_study
from pkdb.studyformat.tables import parse_table_file
from pkdb.studyformat.workbook.base import state_path, workbook_path
from pkdb.studyformat.workbook.read import read_workbook
from pkdb.studyformat.workbook.write import build_workbook

OUTPUTS = "outputs_Tab2.tsv"
TIMECOURSES = "timecourses_Fig1.tsv"
LOCK = ".~lock.Example.xlsx#"
CLOSE_FIRST = (
    "Close the workbook first, or copy a sheet in the spreadsheet application "
    "and rename it to outputs_Tab3"
)


def names(file):
    parsed = parse_table_file(file)
    assert parsed is not None
    return parsed[0].names


def letter(file, column):
    return get_column_letter(names(file).index(column) + 1)


def set_cells(path, sheet, row, **cells):
    """Change cells of a sheet row by column name, as in a spreadsheet application."""
    workbook = openpyxl.load_workbook(path)
    for column, value in cells.items():
        workbook[sheet][f"{letter(f'{sheet}.tsv', column)}{row}"] = value
    workbook.save(path)


def table_lines(folder, file):
    return (folder / file).read_text(encoding="utf-8").removesuffix("\n").split("\n")


def cell(folder, file, line, column):
    return table_lines(folder, file)[line - 1].split("\t")[names(file).index(column)]


def edit_table(folder, file, line, **cells):
    rows = table_lines(folder, file)
    values = rows[line - 1].split("\t")
    for column, value in cells.items():
        values[names(file).index(column)] = value
    rows[line - 1] = "\t".join(values)
    (folder / file).write_text("\n".join(rows) + "\n", encoding="utf-8", newline="")


def snapshot(folder):
    return {
        path.name: (path.read_bytes(), path.stat().st_mtime_ns)
        for path in sorted(folder.iterdir())
        if path.is_file()
    }


def entries(capsys):
    return [
        json.loads(line)
        for line in capsys.readouterr().out.splitlines()
        if line.startswith("{")
    ]


@pytest.fixture(autouse=True)
def isolated_git(monkeypatch, tmp_path):
    """git without the configuration and ignore files of the user and the system.

    git finds no repository above tmp_path, wherever pytest keeps it, so every
    test of the module sees the same git, also when pytest keeps tmp_path in a
    git work tree.
    """
    for name in ("GIT_DIR", "GIT_WORK_TREE", "GIT_INDEX_FILE", "GIT_COMMON_DIR"):
        monkeypatch.delenv(name, raising=False)
    monkeypatch.setenv("GIT_CONFIG_GLOBAL", os.devnull)
    monkeypatch.setenv("GIT_CONFIG_NOSYSTEM", "1")
    # The default global ignore file is $XDG_CONFIG_HOME/git/ignore.
    monkeypatch.setenv("XDG_CONFIG_HOME", str(tmp_path / "config"))
    monkeypatch.setenv("GIT_CEILING_DIRECTORIES", str(tmp_path.parent))


@pytest.fixture
def vocabulary(sf_vocabulary, tmp_path):
    """The vocabulary options of a command, pinned to the test vocabulary."""
    lock = tmp_path / "vocabulary.json"
    sf_vocabulary.save(lock)
    return ["--vocabulary", str(lock)]


def tables(*arguments):
    return main(["tables", *map(str, arguments)])


@pytest.fixture
def study(valid_study, sf_vocabulary):
    """The valid study with its workbook."""
    assert sync_study(valid_study, sf_vocabulary).workbook_action == "created"
    return valid_study


@pytest.fixture
def conflicting(study):
    """The workbook and the tables change the mean at time 0 differently."""
    set_cells(workbook_path(study), "timecourses_Fig1", 2, mean=0.25)
    edit_table(study, TIMECOURSES, 2, mean="0.75")
    return study


def test_sync_reports_every_study_as_a_json_line(valid_study, vocabulary, capsys):
    old = valid_study.parent / "Old"
    old.mkdir()
    (old / "study.json").write_text('{"sid": "X"}')

    code = tables("sync", valid_study.parent, "--format", "json", *vocabulary)

    assert code == 0
    found = {entry["path"]: entry for entry in entries(capsys)}
    assert found[str(old)] == {
        "path": str(old),
        "skipped": "study format 1, left unchanged",
    }
    assert found[str(valid_study)] == {
        "path": str(valid_study),
        "ok": True,
        "workbook": str(workbook_path(valid_study)),
        "workbook_action": "created",
        "lock": None,
        "changes": [],
        "conflicts": [],
        "issues": [],
    }
    assert workbook_path(valid_study).is_file()
    assert not (old / "Old.xlsx").exists()


def test_sync_writes_workbook_changes_and_names_them(study, vocabulary, capsys):
    set_cells(workbook_path(study), "outputs_Tab2", 2, mean=3.25)

    assert tables("sync", study, "--format", "human", *vocabulary) == 0

    out = capsys.readouterr().out.splitlines()
    assert out[:3] == [
        "caffeine/Example: synced",
        f"  wrote {OUTPUTS}",
        "  Example.xlsx: unchanged",
    ]
    assert cell(study, OUTPUTS, 2, "mean") == "3.25"
    assert tables("sync", study, "--format", "human", *vocabulary) == 0
    assert capsys.readouterr().out.splitlines() == [
        "caffeine/Example: in sync",
        "  Example.xlsx: unchanged",
    ]


def test_sync_names_the_workbook_action_in_plain_words(study, vocabulary, capsys):
    edit_table(study, OUTPUTS, 2, mean="3.5")
    (study / LOCK).write_text("lock")

    assert tables("sync", study, "--format", "human", *vocabulary) == 0
    out = capsys.readouterr().out
    assert out.startswith(
        "caffeine/Example: tables synced, workbook not updated (close it and sync "
        "again)\n"
    )
    assert "  Example.xlsx: not updated because it is open" in out
    assert "close it and sync again" in out and "[workbook_open]" in out

    (study / LOCK).unlink()
    assert tables("sync", study, "--format", "human", *vocabulary) == 0
    assert "  Example.xlsx: updated with the tables" in capsys.readouterr().out


def test_sync_names_a_save_during_the_sync(study, vocabulary, capsys, monkeypatch):
    path = workbook_path(study)
    edit_table(study, OUTPUTS, 2, mean="3.5")

    def save_meanwhile(*args, **kwargs):
        built = build_workbook(*args, **kwargs)
        set_cells(path, "timecourses_Fig1", 3, mean=2.25)
        return built

    monkeypatch.setattr("pkdb.studyformat.sync.build_workbook", save_meanwhile)

    assert tables("sync", study, "--format", "human", *vocabulary) == 0

    out = capsys.readouterr().out.splitlines()
    assert out[0] == (
        "caffeine/Example: tables synced, the workbook was saved during the sync; "
        "sync again"
    )
    assert "  Example.xlsx: not updated because it was saved during the sync" in out


def test_check_writes_nothing_and_fails_on_planned_changes(study, vocabulary, capsys):
    set_cells(workbook_path(study), "timecourses_Fig1", 3, mean=2.25)
    before = snapshot(study)

    code = tables("sync", study, "--check", "--format", "json", *vocabulary)

    assert code == 1
    [entry] = entries(capsys)
    assert entry["ok"]
    assert entry["changes"] == [{"file": TIMECOURSES, "action": "write"}]
    assert snapshot(study) == before
    assert tables("sync", study, "--check", "--format", "human", *vocabulary) == 1
    assert capsys.readouterr().out.splitlines()[:2] == [
        "caffeine/Example: out of sync",
        f"  would write {TIMECOURSES}",
    ]
    assert snapshot(study) == before
    assert tables("sync", study, "--format", "json", *vocabulary) == 0
    capsys.readouterr()
    assert tables("sync", study, "--check", "--format", "json", *vocabulary) == 0


def test_check_without_a_workbook_plans_to_create_it(valid_study, vocabulary, capsys):
    code = tables("sync", valid_study, "--check", "--format", "human", *vocabulary)

    assert code == 1
    assert capsys.readouterr().out.splitlines() == [
        "caffeine/Example: out of sync",
        "  Example.xlsx: would be created",
    ]
    assert not workbook_path(valid_study).exists()


def test_a_conflict_fails_and_lists_both_versions(conflicting, vocabulary, capsys):
    before = snapshot(conflicting)

    code = tables("sync", conflicting, "--format", "json", *vocabulary)

    assert code == 1
    [entry] = entries(capsys)
    assert not entry["ok"]
    [conflict] = entry["conflicts"]
    assert conflict["file"] == TIMECOURSES
    assert conflict["sheet"] == "timecourses_Fig1"
    assert conflict["kept"] is None
    [workbook_row] = conflict["workbook_rows"]
    [table_line] = conflict["table_lines"]
    [base_line] = conflict["base_lines"]
    assert workbook_row["row"] == 2 and table_line["line"] == 2
    mean = names(TIMECOURSES).index("mean")
    assert workbook_row["text"].split("\t")[mean] == "0.25"
    assert table_line["text"].split("\t")[mean] == "0.75"
    assert base_line.split("\t")[mean] == "0"
    assert [issue["code"] for issue in entry["issues"]] == ["sync_conflict"]
    assert snapshot(conflicting) == before

    assert tables("sync", conflicting, "--format", "human", *vocabulary) == 1
    out = capsys.readouterr().out.splitlines()
    assert out[0] == "caffeine/Example: cannot sync, resolve the conflicts below"
    start = out.index("  conflict in sheet timecourses_Fig1")
    workbook, tsv = out[start + 1 : start + 3]
    assert workbook.startswith("    workbook row 2: ") and "mean=0.25" in workbook
    assert tsv.startswith(f"    {TIMECOURSES} line 2: ") and "mean=0.75" in tsv
    assert "label=drug_plasma" in workbook and "\t" not in workbook
    assert any("--keep workbook or --keep tables" in line for line in out)


def test_a_conflict_is_located_at_the_sheet_row(study, vocabulary, capsys):
    path = workbook_path(study)
    workbook = openpyxl.load_workbook(path)
    # Two empty rows above the data, so that time 0 is in row 4.
    workbook["timecourses_Fig1"].insert_rows(2, 2)
    workbook.save(path)
    set_cells(path, "timecourses_Fig1", 4, mean=0.25)
    edit_table(study, TIMECOURSES, 2, mean="0.75")

    assert tables("sync", study, "--format", "json", *vocabulary) == 1
    [issue] = entries(capsys)[0]["issues"]
    assert issue["source"]["file"] == "Example.xlsx"
    assert issue["source"]["sheet"] == "timecourses_Fig1"
    assert issue["source"]["row"] == 4
    assert issue["source"].get("cell") is None
    assert issue["message"].endswith(
        "row 4 of the sheet, line 2 of timecourses_Fig1.tsv"
    )

    assert tables("sync", study, "--format", "human", *vocabulary) == 1
    out = capsys.readouterr().out.splitlines()
    assert "    workbook row 4: " in "\n".join(out)
    [line] = [line for line in out if line.endswith("[sync_conflict]")]
    assert line.startswith(
        "  Example.xlsx, sheet timecourses_Fig1, row 4: The workbook and the tables "
    )


def test_issues_name_the_workbook_or_the_table_file(study, vocabulary, capsys):
    path = workbook_path(study)
    comment = f"{letter(TIMECOURSES, 'comment')}3"
    workbook = openpyxl.load_workbook(path)
    workbook["timecourses_Fig1"][comment] = datetime(2020, 1, 2)
    workbook.save(path)
    original = (study / OUTPUTS).read_bytes()
    edit_table(study, OUTPUTS, 1, mean="average")

    # A table that does not load stops the sync before the workbook is read.
    assert tables("sync", study, "--format", "json", *vocabulary) == 1
    [issue] = entries(capsys)[0]["issues"]
    assert issue["code"] == "unknown_column"
    assert issue["source"]["file"] == OUTPUTS
    assert issue["source"]["sheet"] == "outputs_Tab2"
    (study / OUTPUTS).write_bytes(original)

    assert tables("sync", study, "--format", "json", *vocabulary) == 1
    [issue] = entries(capsys)[0]["issues"]
    assert issue["code"] == "cell_date"
    assert issue["source"]["file"] == "Example.xlsx"
    assert issue["source"]["sheet"] == "timecourses_Fig1"
    assert issue["source"]["cell"] == comment
    assert issue["source"]["header"] == "comment"

    assert tables("sync", study, "--format", "human", *vocabulary) == 1
    assert (
        f"  Example.xlsx, sheet timecourses_Fig1, cell {comment}: The cell holds"
        in capsys.readouterr().out
    )


def test_keep_tables_resolves_a_conflict(conflicting, vocabulary, capsys):
    code = tables(
        "sync", conflicting, "--keep", "tables", "--format", "json", *vocabulary
    )

    assert code == 0
    [entry] = entries(capsys)
    assert entry["ok"] and entry["issues"] == []
    assert [conflict["kept"] for conflict in entry["conflicts"]] == ["tables"]
    assert entry["workbook_action"] == "regenerated"
    assert cell(conflicting, TIMECOURSES, 2, "mean") == "0.75"
    assert tables("sync", conflicting, "--format", "human", *vocabulary) == 0
    assert "conflict" not in capsys.readouterr().out


def test_keep_names_the_kept_side(conflicting, vocabulary, capsys):
    code = tables(
        "sync", conflicting, "--keep", "workbook", "--format", "human", *vocabulary
    )

    assert code == 0
    out = capsys.readouterr().out
    assert "  conflict in sheet timecourses_Fig1, kept workbook" in out
    assert cell(conflicting, TIMECOURSES, 2, "mean") == "0.25"


def test_sync_without_study_folder(tmp_path, vocabulary, capsys):
    assert tables("sync", tmp_path / "missing", *vocabulary) == 1
    assert "Study path must be a directory" in capsys.readouterr().err


def test_sync_reports_an_unreadable_vocabulary(valid_study, tmp_path, capsys):
    broken = tmp_path / "broken.json"
    broken.write_text("{")

    assert tables("sync", valid_study, "--vocabulary", broken) == 1
    assert "vocabulary" in capsys.readouterr().err
    assert not workbook_path(valid_study).exists()


def test_open_without_opening_creates_the_workbook(valid_study, vocabulary, capsys):
    assert tables("open", valid_study, "--no-open", *vocabulary) == 0

    assert "  Example.xlsx: created" in capsys.readouterr().out
    assert read_workbook(workbook_path(valid_study), "Example").ok


def test_open_opens_the_workbook(valid_study, vocabulary, capsys, monkeypatch):
    opened = []
    monkeypatch.setattr(
        "pkdb.curation.launch.open_path", lambda path, **kwargs: opened.append(path)
    )

    assert tables("open", valid_study, *vocabulary) == 0

    assert opened == [workbook_path(valid_study)]
    assert f"Opened {workbook_path(valid_study)}" in capsys.readouterr().out


def test_open_reports_a_failed_sync_and_still_opens(
    conflicting, vocabulary, capsys, monkeypatch
):
    opened = []
    monkeypatch.setattr(
        "pkdb.curation.launch.open_path", lambda path, **kwargs: opened.append(path)
    )

    assert tables("open", conflicting, *vocabulary) == 1

    assert opened == [workbook_path(conflicting)]
    assert "[sync_conflict]" in capsys.readouterr().out


def test_open_reports_an_application_that_cannot_open(
    study, vocabulary, capsys, monkeypatch
):
    def fail(path, **kwargs):
        raise subprocess.CalledProcessError(4, ["xdg-open", str(path)])

    monkeypatch.setattr("pkdb.curation.launch.open_path", fail)

    assert tables("open", study, *vocabulary) == 1
    assert "Cannot open" in capsys.readouterr().err


@pytest.mark.parametrize(
    ("make", "message"),
    [
        (lambda tmp_path: tmp_path, "has no study.json"),
        (
            lambda tmp_path: (
                (tmp_path / "Old").mkdir(),
                (tmp_path / "Old" / "study.json").write_text('{"sid": "X"}'),
                tmp_path / "Old",
            )[-1],
            "study format 1",
        ),
    ],
    ids=["no_study", "format_1"],
)
@pytest.mark.parametrize("command", ["open", "add"])
def test_open_and_add_need_a_format_2_study(
    tmp_path, vocabulary, capsys, make, message, command
):
    folder = make(tmp_path)
    extra = ["outputs_Tab3"] if command == "add" else ["--no-open"]

    assert tables(command, folder, *extra, *vocabulary) == 1
    assert message in capsys.readouterr().err


def sheet_header(path, sheet):
    workbook = openpyxl.load_workbook(path, read_only=True)
    try:
        return [cell.value for cell in next(workbook[sheet].iter_rows(max_row=1))]
    finally:
        workbook.close()


def test_add_creates_an_empty_sheet_with_the_header(study, vocabulary, capsys):
    path = workbook_path(study)
    set_cells(path, "outputs_Tab2", 2, mean=3.25)

    assert tables("add", study, "outputs_Tab3", *vocabulary) == 0

    out = capsys.readouterr().out
    assert f"  wrote {OUTPUTS}" in out
    assert "Added the sheet outputs_Tab3 to Example.xlsx" in out
    assert "Example_Tab3.png is missing" in out and "[missing_image]" in out
    content = read_workbook(path, "Example")
    assert content.ok, content.issues
    assert "outputs_Tab3" in content.sheets
    assert "outputs_Tab3.tsv" not in content.tables
    assert sheet_header(path, "outputs_Tab3") == list(names("outputs_Tab3.tsv"))
    assert not (study / "outputs_Tab3.tsv").exists()
    assert not state_path(path).exists()
    # The workbook change was synced before the workbook was rebuilt.
    assert cell(study, OUTPUTS, 2, "mean") == "3.25"
    assert content.tables[OUTPUTS].text == (study / OUTPUTS).read_text("utf-8")


def test_a_row_in_an_added_sheet_becomes_its_table(study, vocabulary, capsys):
    path = workbook_path(study)
    assert tables("add", study, "outputs_Tab3", *vocabulary) == 0
    assert tables("add", study, "timecourses_Text", *vocabulary) == 0
    out = capsys.readouterr().out
    assert out.count("[missing_image]") == 1
    # The second sheet kept the first.
    assert {"outputs_Tab3", "timecourses_Text"} <= set(
        read_workbook(path, "Example").sheets
    )
    row = dict(
        subjects="all",
        interventions="D1",
        measurement="cmax",
        substance="drug",
        tissue="plasma",
        mean=4,
        unit="mg/l",
    )
    set_cells(path, "outputs_Tab3", 2, **row)

    assert tables("sync", study, "--format", "json", *vocabulary) == 0

    [entry] = entries(capsys)
    assert entry["changes"] == [{"file": "outputs_Tab3.tsv", "action": "write"}]
    assert cell(study, "outputs_Tab3.tsv", 2, "mean") == "4"
    assert cell(study, "outputs_Tab3.tsv", 2, "source") == "Tab3"
    assert "timecourses_Text" in read_workbook(path, "Example").sheets


@pytest.mark.parametrize(
    ("table", "message"),
    [
        ("results_Tab3", "is not a table name"),
        ("subjects", "is not a table name"),
        ("outputs_Tab3_with_a_very_long_name_x", "Excel limits sheet names to 31"),
        ("outputs_Tab2", "already exists"),
        ("outputs_tab2", "is not a table name"),
    ],
)
def test_add_rejects_a_bad_or_existing_table(study, vocabulary, capsys, table, message):
    before = snapshot(study)

    assert tables("add", study, table, *vocabulary) == 1

    assert message in capsys.readouterr().err
    assert snapshot(study) == before


def test_add_rejects_a_table_twice_and_names_differing_in_case(
    study, vocabulary, capsys
):
    assert tables("add", study, "outputs_TabA", *vocabulary) == 0
    capsys.readouterr()
    before = snapshot(study)

    assert tables("add", study, "outputs_TabA", *vocabulary) == 1
    err = capsys.readouterr().err.splitlines()
    assert err == [
        "Cannot add outputs_TabA: fix the problems below",
        "  Example.xlsx, sheet outputs_TabA: The sheet outputs_TabA already exists "
        "in Example.xlsx [table_exists]",
    ]
    assert tables("add", study, "outputs_Taba", *vocabulary) == 1
    assert "already exists" in capsys.readouterr().err
    assert snapshot(study) == before


def test_add_needs_a_closed_workbook(study, vocabulary, capsys):
    (study / LOCK).write_text("lock")
    before = snapshot(study)

    assert tables("add", study, "outputs_Tab3", *vocabulary) == 1

    err = capsys.readouterr().err.splitlines()
    assert err[0] == "Cannot add outputs_Tab3: fix the problems below"
    assert err[1].startswith("  Example.xlsx: Example.xlsx is open")
    assert CLOSE_FIRST in err[1] and err[1].endswith("[workbook_open]")
    assert err[2] == f"    If the workbook is not open, delete {study / LOCK}"
    assert snapshot(study) == before


def test_add_stops_when_the_sync_fails(conflicting, vocabulary, capsys):
    before = snapshot(conflicting)

    assert tables("add", conflicting, "outputs_Tab3", *vocabulary) == 1

    assert "[sync_conflict]" in capsys.readouterr().out
    assert snapshot(conflicting) == before


def test_add_creates_a_missing_workbook(valid_study, vocabulary, capsys):
    assert tables("add", valid_study, "scatters_Fig3", *vocabulary) == 0

    assert "  Example.xlsx: created" in capsys.readouterr().out
    assert (
        "scatters_Fig3" in read_workbook(workbook_path(valid_study), "Example").sheets
    )


def test_add_never_replaces_a_workbook_saved_meanwhile(
    study, vocabulary, capsys, monkeypatch
):
    path = workbook_path(study)

    def save_meanwhile(*args, **kwargs):
        built = build_workbook(*args, **kwargs)
        set_cells(path, "outputs_Tab2", 2, mean=9.5)
        return built

    monkeypatch.setattr("pkdb.studyformat.sync.build_workbook", save_meanwhile)

    assert tables("add", study, "outputs_Tab3", *vocabulary) == 1

    assert "saved" in capsys.readouterr().err
    assert "outputs_Tab3" not in read_workbook(path, "Example").sheets


git = pytest.mark.skipif(shutil.which("git") is None, reason="git is not installed")


def git_repository(folder):
    subprocess.run(["git", "init", "-q", str(folder)], check=True)


@git
def test_git_ignore_warning(study, vocabulary, capsys):
    from pkdb.studyformat.workbook.git import ignored_by_git

    root = study.parent.parent
    path = workbook_path(study)
    assert ignored_by_git(path) is None
    git_repository(root)
    assert ignored_by_git(path) is False

    assert tables("sync", study, "--format", "json", *vocabulary) == 0
    [entry] = entries(capsys)
    assert entry["ok"]
    [issue] = entry["issues"]
    assert issue["code"] == "workbook_not_ignored"
    assert issue["severity"] == "warning"
    assert issue["suggestions"][0]["candidates"] == ["*.xlsx", ".*.pkdb-base"]

    assert tables("open", study, "--no-open", *vocabulary) == 0
    out = capsys.readouterr().out.splitlines()
    assert "      *.xlsx" in out and "      .*.pkdb-base" in out

    (root / ".gitignore").write_text("*.xlsx\n")
    assert tables("sync", study, "--format", "json", *vocabulary) == 0
    [issue] = entries(capsys)[0]["issues"]
    assert issue["suggestions"][0]["candidates"] == [".*.pkdb-base"]

    (root / ".gitignore").write_text("*.xlsx\n.*.pkdb-base\n")
    assert ignored_by_git(path) is True
    assert ignored_by_git(state_path(path)) is True
    assert tables("sync", study, "--format", "json", *vocabulary) == 0
    assert entries(capsys)[0]["issues"] == []
    assert tables("open", study, "--no-open", *vocabulary) == 0
    assert "workbook_not_ignored" not in capsys.readouterr().out


@git
def test_a_workbook_that_git_tracks_is_a_warning(
    study, vocabulary, capsys, monkeypatch
):
    from pkdb.studyformat.workbook.git import git_issues, tracked_by_git

    root = study.parent.parent
    path = workbook_path(study)
    assert tracked_by_git(path) is None
    git_repository(root)
    (root / ".gitignore").write_text("*.xlsx\n.*.pkdb-base\n")
    assert tracked_by_git(path) is False
    subprocess.run(["git", "add", "-f", str(path)], cwd=root, check=True)
    assert tracked_by_git(path) is True
    # Outside the current directory, the command names the absolute folder.
    [issue] = git_issues(path)
    assert issue.suggestions[0].candidates == [
        f"git -C {shlex.quote(str(study))} rm --cached Example.xlsx"
    ]
    monkeypatch.chdir(root)

    assert tables("sync", "caffeine/Example", "--format", "json", *vocabulary) == 0
    [issue] = entries(capsys)[0]["issues"]
    assert issue["code"] == "workbook_tracked"
    assert issue["severity"] == "warning"
    assert issue["source"]["file"] == "Example.xlsx"
    [command] = issue["suggestions"][0]["candidates"]
    assert command == "git -C caffeine/Example rm --cached Example.xlsx"

    # The state file is not ignored and the workbook is tracked.
    (root / ".gitignore").write_text("")
    assert tables("sync", study, "--format", "json", *vocabulary) == 0
    issues = {issue["code"]: issue for issue in entries(capsys)[0]["issues"]}
    assert set(issues) == {"workbook_tracked", "workbook_not_ignored"}
    assert issues["workbook_not_ignored"]["suggestions"][0]["candidates"] == [
        ".*.pkdb-base"
    ]

    # The command works from the current directory.
    subprocess.run([*shlex.split(command), "-q"], check=True)
    assert tracked_by_git(path) is False
    (root / ".gitignore").write_text("*.xlsx\n.*.pkdb-base\n")
    assert tables("sync", study, "--format", "json", *vocabulary) == 0
    assert entries(capsys)[0]["issues"] == []


def test_without_git_the_ignore_check_is_skipped(
    study, vocabulary, capsys, monkeypatch
):
    from pkdb.studyformat.workbook.git import ignored_by_git

    monkeypatch.setattr("shutil.which", lambda name: None)
    assert ignored_by_git(workbook_path(study)) is None
    assert tables("sync", study, "--format", "json", *vocabulary) == 0
    assert entries(capsys)[0]["issues"] == []


def test_a_removed_sheet_whose_table_changed_lists_the_table_lines(
    study, vocabulary, capsys
):
    path = workbook_path(study)
    workbook = openpyxl.load_workbook(path)
    workbook.remove(workbook["outputs_Tab2"])
    workbook.save(path)
    edit_table(study, OUTPUTS, 2, mean="3.5")

    assert tables("sync", study, "--format", "human", *vocabulary) == 1

    out = capsys.readouterr().out.splitlines()
    start = out.index("  conflict in sheet outputs_Tab2")
    assert out[start + 1] == "    workbook: rows removed"
    assert out[start + 2].startswith(f"    {OUTPUTS} line 2: ")
    assert "mean=3.5" in out[start + 2]


def test_a_long_conflict_lists_the_first_rows(capsys):
    from pkdb.tables_cli import LISTED_ROWS, _print_rows

    rows = [(number, "x\tTab2\tall") for number in range(1, LISTED_ROWS + 4)]

    _print_rows(OUTPUTS, "workbook row", rows, "workbook: rows removed")

    out = capsys.readouterr().out.splitlines()
    # The header row is left out, and the owned columns are not shown.
    assert out[0] == "    workbook row 2: subjects=all"
    assert len(out) == LISTED_ROWS + 1
    assert out[-1] == "    and 2 more; --format json lists all"


@pytest.mark.parametrize(
    ("check", "words"),
    [([], "not created"), (["--check"], "cannot be created")],
)
def test_a_workbook_that_was_not_created_says_why(
    valid_study, vocabulary, capsys, check, words
):
    path = valid_study / OUTPUTS
    path.write_text(path.read_text(encoding="utf-8").replace("mean", "average", 1))

    assert tables("sync", valid_study, "--format", "human", *check, *vocabulary) == 1

    out = capsys.readouterr().out.splitlines()
    assert out[1] == f"  Example.xlsx: {words} because of the problems below"
    assert "unchanged" not in "\n".join(out)
    assert any("[unknown_column]" in line for line in out)
    assert not workbook_path(valid_study).exists()


def test_a_missing_workbook_that_is_open_is_not_created(study, vocabulary, capsys):
    workbook_path(study).unlink()
    (study / LOCK).write_text("lock")

    assert tables("sync", study, "--format", "human", *vocabulary) == 0

    out = capsys.readouterr().out.splitlines()
    assert out[0] == (
        "caffeine/Example: tables synced, workbook not created (it is open; close "
        "it and sync again)"
    )
    assert out[1] == "  Example.xlsx: not created because it is open"
    assert any("[workbook_open]" in line for line in out)
    assert not workbook_path(study).exists()


def test_open_does_not_open_a_workbook_that_was_not_created(
    valid_study, vocabulary, capsys, monkeypatch
):
    opened = []
    monkeypatch.setattr(
        "pkdb.curation.launch.open_path", lambda path, **kwargs: opened.append(path)
    )
    path = valid_study / OUTPUTS
    path.write_text(path.read_text(encoding="utf-8").replace("mean", "average", 1))

    assert tables("open", valid_study, *vocabulary) == 1

    captured = capsys.readouterr()
    assert "  Example.xlsx: not created because of the problems below" in captured.out
    assert captured.err == ""
    assert opened == []
