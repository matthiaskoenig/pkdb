"""Optimized parsing must preserve values, provenance, and resource ownership."""

from copy import deepcopy
from datetime import datetime

import openpyxl
import pytest
from openpyxl.styles import Font

from pkdb.importers import workbook as module
from pkdb.schemas.source import SourceLocation
from pkdb.schemas.validation import StudyValidationError


def test_header_location_preserves_serialized_contract_and_isolates_maps():
    source = SourceLocation(file="data.xlsx", sheet="Results", row=7, path=("x", 1))
    source._columns["mean"] = "C"
    source._fields["value"] = SourceLocation(file="study.json", path=("value",))
    for header in ("mean", "absent"):
        column = source._columns.get(header)
        expected = SourceLocation.model_validate(
            {
                **source.model_dump(),
                "header": header,
                "column": column,
                "cell": f"{column}7" if column else None,
            }
        )
        actual = source.for_header(header)
        assert actual.model_dump() == expected.model_dump()
        assert actual.model_dump(exclude_unset=True) == expected.model_dump(
            exclude_unset=True
        )
        assert actual._columns == actual._fields == {}
        actual._columns["new"] = "D"
        actual._fields["new"] = actual
        assert "new" not in source._columns
        assert "new" not in source._fields
    other = SourceLocation(file="other.xlsx")
    assert other._columns == other._fields == {}


def test_sparse_iterator_matches_openpyxl_values_and_physical_rows(tmp_path):
    path = tmp_path / "values.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["notes"])
    sheet.append(["date", "number", "boolean", "text", "formula", "error"])
    sheet.append([datetime(2024, 1, 2), 1.5, True, "inline text", "=1+1", "#DIV/0!"])
    sheet.cell(7, 2, -3)
    sheet.cell(10, 6).font = Font(bold=True)
    book.save(path)
    book.close()
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book.active
        sheet.reset_dimensions()
        expected = [
            (i, row)
            for i, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2)
            if i == 2 or any(v is not None for v in row)
        ]
        actual = [
            (i, row)
            for i, row in module._xlsx_rows(sheet)
            if i == 2 or any(v is not None for v in row)
        ]
        assert actual == expected
    finally:
        book.close()


def test_sparse_millionth_row_preserves_cells_comments_and_limits(tmp_path):
    path = tmp_path / "sparse.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Results"
    sheet.cell(2, 1, "time")
    sheet.cell(2, 2, "mean")
    sheet.cell(3, 1, "# comment")
    sheet.cell(3, 2, 99)
    sheet.cell(1000000, 1, 4)
    sheet.cell(1000000, 2, 5)
    sheet.cell(1048576, 1).font = Font(bold=True)
    book.save(path)
    book.close()
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        # Bounded by stored XML rows, not the million-row gap.
        assert len(list(module._xlsx_rows(book.active))) == 4
    finally:
        book.close()
    rows = module.read_table(path, "Results", 1)
    assert len(rows) == 1
    values, source = rows[0]
    assert values == {"time": 4, "mean": 5}
    assert source.row == 1000000
    assert source.for_header("mean").cell == "B1000000"
    book = openpyxl.load_workbook(path)
    book.active.cell(4, 1, 1)
    book.save(path)
    book.close()
    with pytest.raises(StudyValidationError) as error:
        module.read_table(path, "Results", 1)
    assert error.value.report.issues[0].code == "row_limit"


def test_missing_header_is_not_taken_from_first_data_row(tmp_path):
    path = tmp_path / "no-header.xlsx"
    book = openpyxl.Workbook()
    book.active.cell(5, 1, "not a header")
    book.save(path)
    book.close()
    rows = module.read_table(path, "Sheet", 10)
    assert rows[0][0] == {}
    assert rows[0][1].row == 5


@pytest.mark.parametrize("fail", [False, True])
def test_workbook_reuse_closes_on_success_and_failure_without_stale_cache(
    tmp_path, monkeypatch, fail
):
    path = tmp_path / "sheets.xlsx"
    book = openpyxl.Workbook()
    book.active.title = "First"
    book.create_sheet("Second")
    for sheet in book:
        sheet.append(["notes"])
        sheet.append(["value"])
        sheet.append([1])
    book.save(path)
    book.close()
    opened = []
    original = openpyxl.load_workbook

    def load(*args, **kwargs):
        result = original(*args, **kwargs)
        opened.append(result)
        return result

    monkeypatch.setattr(module.openpyxl, "load_workbook", load)
    try:
        with module.table_reader() as read:
            assert read(path, "First", 10)[0][0] == {"value": 1}
            assert read(path, "Second", 10)[0][0] == {"value": 1}
            assert len(opened) == 1
            if fail:
                read(path, "Absent", 10)
    except StudyValidationError as error:
        assert fail
        assert error.report.issues[0].code == "missing_sheet"
    assert opened[0]._archive.fp is None
    book = original(path)
    book["First"]["A3"] = 2
    book.save(path)
    book.close()
    with module.table_reader() as read:
        assert read(path, "First", 10)[0][0] == {"value": 2}
    assert len(opened) == 2
    assert opened[1]._archive.fp is None


def test_location_deepcopy_preserves_aliases_and_isolates_mutable_maps():
    row = SourceLocation(file="data.xlsx", sheet="Results", row=3, path=("x", 1))
    row._columns["value"] = "B"
    cell = row.for_header("value")
    row._fields.update(mean=cell, value=cell)
    row._fields["self"] = row
    copied = deepcopy(row)
    assert copied.model_dump() == row.model_dump()
    assert copied is not row
    assert copied._fields["self"] is copied
    assert copied._fields["mean"] is copied._fields["value"]
    assert copied._fields["mean"] is not cell
    copied._columns["value"] = "C"
    copied._fields["mean"]._fields["new"] = copied
    assert row._columns["value"] == "B"
    assert not cell._fields
    copied.model_fields_set.clear()
    assert row.model_fields_set


def test_location_subclass_deepcopy_keeps_mutable_fields_independent():
    class ExtendedLocation(SourceLocation):
        notes: list[str]

    source = ExtendedLocation(file="data.xlsx", notes=["original"])
    copied = deepcopy(source)
    copied.notes.append("copy")
    assert source.notes == ["original"]
