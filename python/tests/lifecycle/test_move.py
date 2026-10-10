import json
from pathlib import Path

import openpyxl
import pytest
from digitize_fixtures import GOOD, png, project
from openpyxl.utils import get_column_letter

from pkdb.cli import main
from pkdb.identity import Author
from pkdb.lifecycle.move import MoveIncomplete, MoveRefused, move_study
from pkdb.schemas.provenance import AutomaticCuration
from pkdb.schemas.review import ReviewTarget
from pkdb.schemas.validation import ValidationIssue
from pkdb.studyformat.formatter import FormatResult, format_folder
from pkdb.studyformat.load import load_study
from pkdb.studyformat.metadata import patch_metadata, read_metadata
from pkdb.studyformat.review_edit import add_item, read_review
from pkdb.studyformat.sync import sync_study
from pkdb.studyformat.tables import parse_table_file
from pkdb.studyformat.validation import validate_folder
from pkdb.studyformat.workbook.base import state_path, workbook_path

RAW = "Group\tAge (yr)\nmen\t30\nwomen\t40\n"
OLD = "caffeine/Example"
NEW = "codeine/Renamed"


@pytest.fixture
def moved_checkout(tmp_path, valid_files):
    """A checkout with the valid study caffeine/Example, a raw table, a digitization and two review targets."""
    folder = tmp_path / "studies" / "caffeine" / "Example"
    folder.mkdir(parents=True)
    files = {
        **valid_files,
        "Example_Tab9.tsv": RAW,
        "Example_Tab9.png": b"png",
        "Example_Fig1.png": png(100, 100),
        "Example_Fig1.wpd.json": json.dumps(project(GOOD)),
    }
    for name, content in files.items():
        path = folder / name
        if isinstance(content, bytes):
            path.write_bytes(content)
        else:
            path.write_text(content, encoding="utf-8", newline="")
    assert format_folder(folder).ok
    author = Author("curator")
    add_item(
        folder,
        author,
        kind="question",
        text="Is the second point right?",
        target=ReviewTarget(file="Example_Fig1.wpd.json", key="drug_plasma"),
    )
    add_item(
        folder,
        author,
        kind="issue",
        text="The raw table is fine.",
        target=ReviewTarget(file="Example_Tab9.tsv"),
        acknowledges="raw_table_unused",
    )
    add_item(
        folder,
        author,
        kind="question",
        text="Who are the subjects?",
        target=ReviewTarget(file="subjects.tsv", column="count"),
    )
    return tmp_path


def study_folder(root, location=OLD):
    return root / "studies" / location


def snapshot(folder):
    """Every file of a folder by name with its bytes."""
    return {path.name: path.read_bytes() for path in sorted(folder.iterdir())}


def study_column(folder):
    """The values of the study column in every table of a folder, as the files hold them."""
    study = load_study(folder)
    return {row.cells["study"] for table in study.tables for row in table.rows}


def set_cell(workbook, sheet, row, column, value):
    """Change a cell of a data sheet by column name, as in a spreadsheet application."""
    parsed = parse_table_file(f"{sheet}.tsv")
    assert parsed is not None
    names = parsed[0].names
    book = openpyxl.load_workbook(workbook)
    book[sheet][f"{get_column_letter(names.index(column) + 1)}{row}"] = value
    book.save(workbook)


def test_a_move_renames_the_folder_and_every_file_of_the_study(
    moved_checkout, sf_vocabulary
):
    root = moved_checkout
    before = snapshot(study_folder(root))
    result = move_study(root, OLD, NEW, sf_vocabulary)
    folder = study_folder(root, NEW)
    assert result.folder == folder and not (root / "studies" / "caffeine").exists()
    names = sorted(p.name for p in folder.iterdir())
    assert "Renamed.pdf" in names and "Renamed_Tab1.png" in names
    assert "Renamed_Tab9.tsv" in names and "Renamed_Fig1.wpd.json" in names
    assert not [n for n in names if n.startswith("Example")]
    assert sorted(result.renamed) == sorted(
        (name, "Renamed" + name.removeprefix("Example"))
        for name in before
        if name.startswith("Example")
    )
    # A rename keeps the content of the file.
    for old, new in result.renamed:
        if new != "Renamed_Tab9.tsv":
            assert (folder / new).read_bytes() == before[old]
    targets = [item.target for item in read_review(folder).review.items]
    assert [target.file for target in targets if target] == [
        "Renamed_Fig1.wpd.json",
        "Renamed_Tab9.tsv",
        "subjects.tsv",
    ]
    assert targets[0] == ReviewTarget(file="Renamed_Fig1.wpd.json", key="drug_plasma")
    assert result.targets == 2
    assert result.issue is None and result.workbook is None
    assert (result.assets, result.reference) == (0, True)
    assert study_column(folder) == {"Renamed"}
    assert format_folder(folder, check=True).changes == []
    assert not [
        i
        for i in validate_folder(folder, sf_vocabulary).issues
        if i.severity == "error"
    ]


def test_a_move_within_the_substance_keeps_its_other_studies(
    moved_checkout, sf_vocabulary
):
    other = study_folder(moved_checkout, "caffeine/Other")
    other.mkdir()
    (moved_checkout / "studies" / "codeine").mkdir()
    result = move_study(moved_checkout, OLD, "caffeine/Renamed", sf_vocabulary)
    assert result.folder == study_folder(moved_checkout, "caffeine/Renamed")
    assert other.is_dir() and not study_folder(moved_checkout).exists()
    assert (moved_checkout / "studies" / "codeine").is_dir()


def test_a_move_names_the_issue_of_the_study(moved_checkout, sf_vocabulary):
    folder = study_folder(moved_checkout)
    patch_metadata(folder, {"issue": 7}, read_metadata(folder).revision)
    assert move_study(moved_checkout, OLD, NEW, sf_vocabulary).issue == 7


def automatic_curation(folder):
    """Give the study the provenance of an agent that read the PDF and a file outside the study."""
    provenance = {
        "kind": "automatic_curation",
        "source_key": "pkdb.ai",
        "method": "claude",
        "version": "5.5",
        "run_id": "run-1",
        "assets": [
            {"url": "Example.pdf", "sha256": "a" * 64},
            {"url": "supplement.pdf", "sha256": "b" * 64},
        ],
    }
    patch_metadata(folder, {"provenance": provenance}, read_metadata(folder).revision)


def test_provenance_assets_and_the_reference_name_follow_the_move(
    moved_checkout, sf_vocabulary
):
    automatic_curation(study_folder(moved_checkout))
    folder = move_study(moved_checkout, OLD, NEW, sf_vocabulary).folder
    provenance = read_metadata(folder).metadata.provenance
    assert isinstance(provenance, AutomaticCuration)
    assert [asset.url for asset in provenance.assets] == [
        "Renamed.pdf",
        "supplement.pdf",
    ]
    assert [asset.sha256 for asset in provenance.assets] == ["a" * 64, "b" * 64]
    reference = json.loads((folder / "reference.json").read_text(encoding="utf-8"))
    assert reference == {
        "sid": "123",
        "name": "Renamed",
        "pmid": "123",
        "title": "Example study",
    }
    assert format_folder(folder, check=True).changes == []
    assert not [
        i
        for i in validate_folder(folder, sf_vocabulary).issues
        if i.severity == "error"
    ]


def test_moved_counts_the_assets_and_the_reference_name(moved_checkout, sf_vocabulary):
    automatic_curation(study_folder(moved_checkout))
    moved = move_study(moved_checkout, OLD, NEW, sf_vocabulary)
    assert (moved.assets, moved.reference) == (1, True)
    assert moved.pkdb_id is None


# Attachments named after the study that are no PDF, image, raw table or digitization.
ATTACHMENTS = [
    "Example_Supp.pdf",
    "Example_data_p145.png",
    "Example_Fig1.2.png",
    "Example.docx",
]


def test_other_attachments_named_after_the_study_follow_the_move(
    moved_checkout, sf_vocabulary
):
    folder = study_folder(moved_checkout)
    for name in [*ATTACHMENTS, "Examples.pdf", "Example"]:
        (folder / name).write_bytes(name.encode())
    add_item(
        folder,
        Author("curator"),
        kind="question",
        text="Is the supplement complete?",
        target=ReviewTarget(file="Example_Supp.pdf"),
    )
    provenance = {
        "kind": "automatic_curation",
        "source_key": "pkdb.ai",
        "method": "claude",
        "version": "5.5",
        "run_id": "run-1",
        "assets": [{"url": "Example_Supp.pdf", "sha256": "a" * 64}],
    }
    patch_metadata(folder, {"provenance": provenance}, read_metadata(folder).revision)
    moved = move_study(moved_checkout, OLD, NEW, sf_vocabulary)
    for name in ATTACHMENTS:
        renamed = "Renamed" + name.removeprefix("Example")
        assert (name, renamed) in moved.renamed
        assert (moved.folder / renamed).read_bytes() == name.encode()
    # Only <name>.<extension> and <name>_<anything> are named after the study.
    assert (moved.folder / "Examples.pdf").exists() and (
        moved.folder / "Example"
    ).exists()
    targets = [
        item.target.file
        for item in read_review(moved.folder).review.items
        if item.target
    ]
    assert "Renamed_Supp.pdf" in targets and "Example_Supp.pdf" not in targets
    metadata = read_metadata(moved.folder).metadata
    assert isinstance(metadata.provenance, AutomaticCuration)
    assert [asset.url for asset in metadata.provenance.assets] == ["Renamed_Supp.pdf"]
    assert (moved.targets, moved.assets) == (3, 1)


@pytest.mark.parametrize("name", ["study", "review", "reference", "subjects"])
def test_the_fixed_files_of_a_study_named_like_them_stay(
    moved_checkout, sf_vocabulary, name
):
    # Built by hand: pkdb new and pkdb move refuse such names now.
    (moved_checkout / "studies" / "caffeine" / "Example").rename(
        moved_checkout / "studies" / "caffeine" / name
    )
    folder = study_folder(moved_checkout, f"caffeine/{name}")
    for path in folder.glob("Example*"):
        path.rename(folder / path.name.replace("Example", name, 1))
    target = folder / "review.json"
    target.write_text(
        target.read_text(encoding="utf-8").replace('"Example', f'"{name}'),
        encoding="utf-8",
    )
    assert format_folder(folder).ok
    fixed = {"study.json", "review.json", "reference.json"} | {
        path.name for path in folder.iterdir() if parse_table_file(path.name)
    }
    assert f"{name}.json" in fixed or f"{name}.tsv" in fixed
    moved = move_study(
        moved_checkout, f"caffeine/{name}", "caffeine/Other", sf_vocabulary
    )
    assert fixed <= {path.name for path in moved.folder.iterdir()}
    assert not fixed & {before for before, _ in moved.renamed}
    assert (moved.folder / "Other.pdf").exists()
    assert (moved.folder / "Other_Tab9.tsv").exists()
    assert not [
        issue
        for issue in validate_folder(moved.folder, sf_vocabulary).issues
        if issue.severity == "error"
    ]


def test_a_workbook_in_another_case_is_not_an_attachment(moved_checkout, sf_vocabulary):
    folder = study_folder(moved_checkout)
    (folder / "Example.XLSX").write_bytes(b"workbook")
    try:
        moved = move_study(moved_checkout, OLD, NEW, sf_vocabulary)
    except MoveRefused as error:
        # A case-insensitive file system finds it as the workbook, which is damaged.
        assert "Example.xlsx cannot be read" in str(error)
        return
    assert "Example.XLSX" not in dict(moved.renamed)
    assert (moved.folder / "Example.XLSX").read_bytes() == b"workbook"


def test_an_attachment_whose_new_name_is_taken_is_refused(
    moved_checkout, sf_vocabulary
):
    folder = study_folder(moved_checkout)
    (folder / "Example_Supp.pdf").write_bytes(b"supplement")
    (folder / "renamed_supp.pdf").write_bytes(b"other")
    before = snapshot(folder)
    with pytest.raises(MoveRefused) as error:
        move_study(moved_checkout, OLD, NEW, sf_vocabulary)
    assert str(error.value) == (
        "Example_Supp.pdf would become Renamed_Supp.pdf, which collides with "
        "renamed_supp.pdf of studies/caffeine/Example; rename or remove "
        "renamed_supp.pdf first"
    )
    assert snapshot(folder) == before


def test_a_reference_name_other_than_the_study_name_stays(
    moved_checkout, sf_vocabulary
):
    path = study_folder(moved_checkout) / "reference.json"
    text = path.read_text(encoding="utf-8").replace('"Example"', '"Jönsson2015"')
    path.write_text(text, encoding="utf-8")
    before = path.read_bytes()
    folder = move_study(moved_checkout, OLD, NEW, sf_vocabulary).folder
    assert (folder / "reference.json").read_bytes() == before
    assert read_metadata(folder).metadata.provenance.kind == "manual_curation"


REGISTRY = {
    "PKDB00010": ["other/Study", "2020-01-01"],
    "PKDB00009": ["caffeine/Example", "2026-09-28"],
}


def released(root, text):
    """Release the study as PKDB00009 and write the registry file `text`."""
    folder = study_folder(root)
    release = {"release": {"pkdb_id": "PKDB00009", "date": "2026-09-28"}}
    patch_metadata(folder, release, read_metadata(folder).revision)
    path = root / "studies" / "study_identifiers.json"
    path.write_text(text, encoding="utf-8", newline="")
    return path


@pytest.mark.parametrize("end", ["", "\n"])
def test_a_move_updates_the_registry_file_entry(moved_checkout, sf_vocabulary, end):
    from pkdb.lifecycle.registry import duplicates, registry_problems, scan

    path = released(moved_checkout, json.dumps(REGISTRY, indent=2) + end)
    moved = move_study(moved_checkout, OLD, NEW, sf_vocabulary)
    assert moved.registry and moved.pkdb_id == "PKDB00009"
    entry = {"PKDB00009": [NEW, "2026-09-28"]}
    # The file keeps its order and layout.
    assert (
        path.read_text(encoding="utf-8") == json.dumps(REGISTRY | entry, indent=2) + end
    )
    found = scan(moved_checkout)
    assert duplicates(found) == registry_problems(found, moved_checkout) == []


@pytest.mark.parametrize(
    "registry",
    [
        {"PKDB00010": ["other/Study", "2020-01-01"]},
        {"PKDB00009": ["caffeine/Another", "2026-09-28"]},
    ],
)
def test_a_registry_file_without_the_old_location_stays(
    moved_checkout, sf_vocabulary, registry
):
    path = released(moved_checkout, json.dumps(registry, indent=2))
    before = path.read_bytes()
    assert not move_study(moved_checkout, OLD, NEW, sf_vocabulary).registry
    assert path.read_bytes() == before


def test_a_malformed_registry_file_refuses_the_move_of_a_released_study(
    moved_checkout, sf_vocabulary
):
    path = released(moved_checkout, "[]")
    before = snapshot(study_folder(moved_checkout))
    with pytest.raises(MoveRefused) as error:
        move_study(moved_checkout, OLD, NEW, sf_vocabulary)
    assert str(error.value) == (
        "studies/study_identifiers.json: must be a JSON object; fix it before "
        "moving a released study"
    )
    assert snapshot(study_folder(moved_checkout)) == before
    assert path.read_text(encoding="utf-8") == "[]"


def test_an_unreadable_registry_file_refuses_the_move_of_a_released_study(
    moved_checkout, sf_vocabulary
):
    path = released(moved_checkout, "{}")
    path.unlink()
    path.mkdir()
    before = snapshot(study_folder(moved_checkout))
    with pytest.raises(MoveRefused) as error:
        move_study(moved_checkout, OLD, NEW, sf_vocabulary)
    assert str(error.value).startswith(
        "studies/study_identifiers.json: cannot be read: "
    )
    assert str(error.value).endswith("; fix it before moving a released study")
    assert snapshot(study_folder(moved_checkout)) == before


def test_a_failed_registry_write_undoes_the_move(
    moved_checkout, sf_vocabulary, monkeypatch
):
    from pkdb.lifecycle import move

    path = released(moved_checkout, json.dumps(REGISTRY, indent=2))
    folder = study_folder(moved_checkout)
    before, registry = snapshot(folder), path.read_bytes()
    write = move.write_checked

    def flaky(target, text, expected):
        if target.name == "study_identifiers.json":
            raise PermissionError(13, "Permission denied", str(target))
        return write(target, text, expected)

    monkeypatch.setattr(move, "write_checked", flaky)
    with pytest.raises(MoveRefused, match="Nothing was changed"):
        move_study(moved_checkout, OLD, NEW, sf_vocabulary)
    assert snapshot(folder) == before and path.read_bytes() == registry


@pytest.mark.parametrize("failing", ["study.json", "reference.json"])
def test_a_failed_json_write_undoes_the_move(
    moved_checkout, sf_vocabulary, monkeypatch, failing
):
    from pkdb.lifecycle import move

    folder = study_folder(moved_checkout)
    automatic_curation(folder)
    before = snapshot(folder)
    write = move.write_checked

    def flaky(path, text, expected):
        if path.name == failing:
            raise PermissionError(13, "Permission denied", str(path))
        return write(path, text, expected)

    monkeypatch.setattr(move, "write_checked", flaky)
    with pytest.raises(MoveRefused, match="Nothing was changed"):
        move_study(moved_checkout, OLD, NEW, sf_vocabulary)
    assert snapshot(folder) == before
    assert not (moved_checkout / "studies" / "codeine").exists()


def test_a_case_only_rename_suggests_a_temporary_name(moved_checkout, sf_vocabulary):
    with pytest.raises(MoveRefused) as error:
        move_study(moved_checkout, OLD, "caffeine/example", sf_vocabulary)
    assert str(error.value) == (
        "studies/caffeine/example differs from studies/caffeine/Example only in "
        "case; to change only the case, move the study to a temporary name first "
        "and then to caffeine/example"
    )


def test_an_existing_target_is_refused(moved_checkout, sf_vocabulary):
    (moved_checkout / "studies" / "codeine" / "Renamed").mkdir(parents=True)
    with pytest.raises(MoveRefused, match="exists"):
        move_study(moved_checkout, OLD, NEW, sf_vocabulary)
    assert (moved_checkout / "studies" / "caffeine" / "Example").is_dir()


@pytest.mark.parametrize(
    "old, new, message",
    [
        (OLD, OLD, "The study is at caffeine/Example already"),
        (OLD, "codeine/Renamed 2", "letters, digits"),
        (OLD, "codeine/outputs", "reserved"),
        (OLD, "Renamed", "<substance>/<name>"),
        ("caffeine/Missing", NEW, "studies/caffeine/Missing is not a study"),
        ("caffeine", NEW, "'caffeine' is not <substance>/<name>"),
        ("../caffeine/Example", NEW, "is not <substance>/<name>"),
        ("caffeine/..", NEW, "is not <substance>/<name>"),
        (OLD, "caffeine/example", "differs from studies/caffeine/Example only in case"),
        (OLD, "Caffeine/Renamed", "differs from studies/caffeine only in case"),
    ],
)
def test_wrong_locations_are_refused(moved_checkout, sf_vocabulary, old, new, message):
    before = snapshot(study_folder(moved_checkout))
    with pytest.raises(MoveRefused, match=message):
        move_study(moved_checkout, old, new, sf_vocabulary)
    assert snapshot(study_folder(moved_checkout)) == before
    assert sorted(p.name for p in (moved_checkout / "studies").iterdir()) == [
        "caffeine"
    ]


@pytest.mark.parametrize(
    "old", ["caffeine/example", "Caffeine/Example", "CAFFEINE/EXAMPLE"]
)
def test_an_old_location_in_another_case_is_refused(moved_checkout, sf_vocabulary, old):
    # On macOS and Windows the folder is found in any case, but the files are
    # named after the exact study name, so a move under another case would
    # leave them behind.
    before = snapshot(study_folder(moved_checkout))
    with pytest.raises(MoveRefused) as error:
        move_study(moved_checkout, old, NEW, sf_vocabulary)
    assert str(error.value) == (
        f"studies/{old} is studies/caffeine/Example; give it in that spelling"
    )
    assert snapshot(study_folder(moved_checkout)) == before
    assert not (moved_checkout / "studies" / "codeine").exists()


def test_a_format_1_folder_is_refused(moved_checkout, sf_vocabulary):
    legacy = study_folder(moved_checkout, "caffeine/Legacy")
    legacy.mkdir()
    (legacy / "study.json").write_text('{"name": "Legacy"}', encoding="utf-8")
    with pytest.raises(MoveRefused, match="is not a study format 2 folder"):
        move_study(moved_checkout, "caffeine/Legacy", NEW, sf_vocabulary)
    assert legacy.is_dir()


def test_a_trailing_slash_of_the_old_location_is_accepted(
    moved_checkout, sf_vocabulary
):
    result = move_study(moved_checkout, OLD + "/", NEW, sf_vocabulary)
    assert result.folder == study_folder(moved_checkout, NEW)


def test_a_file_of_the_new_name_in_the_folder_is_refused(moved_checkout, sf_vocabulary):
    folder = study_folder(moved_checkout)
    (folder / "renamed_Tab1.png").write_bytes(b"other")
    before = snapshot(folder)
    with pytest.raises(MoveRefused, match="renamed_Tab1.png"):
        move_study(moved_checkout, OLD, NEW, sf_vocabulary)
    assert snapshot(folder) == before


@pytest.mark.parametrize("stray", ["Renamed.xlsx", ".renamed.xlsx.pkdb-base"])
def test_a_workbook_file_of_the_new_name_in_the_folder_is_refused(
    moved_checkout, sf_vocabulary, stray
):
    folder = study_folder(moved_checkout)
    (folder / stray).write_bytes(b"stray")
    before = snapshot(folder)
    with pytest.raises(MoveRefused) as error:
        move_study(moved_checkout, OLD, NEW, sf_vocabulary)
    assert str(error.value) == (
        f"{stray} in studies/caffeine/Example would belong to the workbook of "
        "the moved study; remove it first"
    )
    assert snapshot(folder) == before


def test_a_raw_table_name_too_long_for_a_sheet_is_refused(
    moved_checkout, sf_vocabulary
):
    folder = study_folder(moved_checkout)
    (folder / "Example_Tab3_part2.tsv").write_text(RAW, encoding="utf-8")
    with pytest.raises(MoveRefused, match="31 characters"):
        move_study(moved_checkout, OLD, "codeine/" + "R" * 24, sf_vocabulary)
    assert (folder / "Example_Tab3_part2.tsv").exists()


def test_a_study_that_pkdb_format_cannot_read_is_refused(moved_checkout, sf_vocabulary):
    folder = study_folder(moved_checkout)
    (folder / "subjects.tsv").write_bytes(b"name\tparent\xff\n")
    before = snapshot(folder)
    with pytest.raises(MoveRefused, match="subjects.tsv"):
        move_study(moved_checkout, OLD, NEW, sf_vocabulary)
    assert snapshot(folder) == before
    assert not (moved_checkout / "studies" / "codeine").exists()


def test_an_open_workbook_refuses_the_move(moved_checkout, sf_vocabulary):
    folder = study_folder(moved_checkout)
    assert sync_study(folder, sf_vocabulary).workbook_action == "created"
    (folder / ".~lock.Example.xlsx#").write_text("lock", encoding="utf-8")
    before = snapshot(folder)
    with pytest.raises(MoveRefused, match="open"):
        move_study(moved_checkout, OLD, NEW, sf_vocabulary)
    assert snapshot(folder) == before


def test_a_workbook_with_unsynced_edits_refuses_the_move(moved_checkout, sf_vocabulary):
    folder = study_folder(moved_checkout)
    assert sync_study(folder, sf_vocabulary).workbook_action == "created"
    set_cell(workbook_path(folder), "outputs_Tab2", 2, "mean", "3.5")
    before = snapshot(folder)
    with pytest.raises(MoveRefused, match="pkdb tables sync"):
        move_study(moved_checkout, OLD, NEW, sf_vocabulary)
    assert (folder / "Example.xlsx").exists()
    assert snapshot(folder) == before
    assert not (moved_checkout / "studies" / "codeine").exists()


def test_a_workbook_that_conflicts_with_the_tables_refuses_the_move(
    moved_checkout, sf_vocabulary
):
    folder = study_folder(moved_checkout)
    assert sync_study(folder, sf_vocabulary).workbook_action == "created"
    set_cell(workbook_path(folder), "outputs_Tab2", 2, "mean", "3.5")
    table = folder / "outputs_Tab2.tsv"
    text = table.read_text(encoding="utf-8").replace("\t2.5\t", "\t4.5\t")
    table.write_text(text, encoding="utf-8")
    with pytest.raises(MoveRefused, match="pkdb tables sync"):
        move_study(moved_checkout, OLD, NEW, sf_vocabulary)
    assert (folder / "Example.xlsx").exists()


def test_a_workbook_with_a_scratch_sheet_refuses_the_move(
    moved_checkout, sf_vocabulary
):
    folder = study_folder(moved_checkout)
    assert sync_study(folder, sf_vocabulary).workbook_action == "created"
    book = openpyxl.load_workbook(workbook_path(folder))
    book.create_sheet("_notes")["A1"] = "keep me"
    book.save(workbook_path(folder))
    before = snapshot(folder)
    with pytest.raises(MoveRefused, match="_notes"):
        move_study(moved_checkout, OLD, NEW, sf_vocabulary)
    assert snapshot(folder) == before


def test_a_damaged_workbook_refuses_the_move(moved_checkout, sf_vocabulary):
    folder = study_folder(moved_checkout)
    (folder / "Example.xlsx").write_bytes(b"not a workbook")
    before = snapshot(folder)
    with pytest.raises(MoveRefused) as error:
        move_study(moved_checkout, OLD, NEW, sf_vocabulary)
    # A sync cannot read the workbook either.
    assert str(error.value) == (
        "Example.xlsx cannot be read: File is not a zip file. Repair it in a "
        "spreadsheet program, or remove it when it holds no edits that are not in "
        "the tables, and move again"
    )
    assert snapshot(folder) == before


def test_a_synced_workbook_is_removed(moved_checkout, sf_vocabulary):
    folder = study_folder(moved_checkout)
    assert sync_study(folder, sf_vocabulary).workbook_action == "created"
    # A workbook that only lacks a change of the tables holds no edits.
    table = folder / "outputs_Tab2.tsv"
    text = table.read_text(encoding="utf-8").replace("\t2.5\t", "\t4.5\t")
    table.write_text(text, encoding="utf-8")
    state_path(workbook_path(folder)).write_text("{}", encoding="utf-8")
    result = move_study(moved_checkout, OLD, NEW, sf_vocabulary)
    assert result.workbook == "Example.xlsx"
    names = [path.name for path in result.folder.iterdir()]
    assert not [name for name in names if name.endswith((".xlsx", ".pkdb-base"))]
    table = result.folder / "outputs_Tab2.tsv"
    assert "\t4.5\t" in table.read_text(encoding="utf-8")


def test_a_substance_only_move_removes_the_synced_workbook(
    moved_checkout, sf_vocabulary
):
    folder = study_folder(moved_checkout)
    assert sync_study(folder, sf_vocabulary).workbook_action == "created"
    state_path(workbook_path(folder)).write_text("{}", encoding="utf-8")
    before = {
        name: content
        for name, content in snapshot(folder).items()
        if not name.endswith((".xlsx", ".pkdb-base"))
    }
    result = move_study(moved_checkout, OLD, "codeine/Example", sf_vocabulary)
    assert result.folder == study_folder(moved_checkout, "codeine/Example")
    assert result.workbook == "Example.xlsx"
    # Nothing is named after another study, so nothing is renamed or rewritten.
    assert (result.renamed, result.targets) == ([], 0)
    assert (result.assets, result.reference) == (0, False)
    assert snapshot(result.folder) == before
    assert not (moved_checkout / "studies" / "caffeine").exists()


def test_a_move_to_the_name_in_another_case_removes_the_synced_workbook(
    moved_checkout, sf_vocabulary
):
    # The study's own workbook equals the workbook of the new name ignoring
    # case; the move checks and removes it, so it is no stray file.
    folder = study_folder(moved_checkout)
    assert sync_study(folder, sf_vocabulary).workbook_action == "created"
    result = move_study(moved_checkout, OLD, "codeine/example", sf_vocabulary)
    assert result.workbook == "Example.xlsx"
    assert ("Example.pdf", "example.pdf") in result.renamed
    assert not [p for p in result.folder.iterdir() if p.suffix == ".xlsx"]


def test_a_substance_only_move_of_a_workbook_with_unsynced_edits_is_refused(
    moved_checkout, sf_vocabulary
):
    folder = study_folder(moved_checkout)
    assert sync_study(folder, sf_vocabulary).workbook_action == "created"
    set_cell(workbook_path(folder), "outputs_Tab2", 2, "mean", "3.5")
    before = snapshot(folder)
    with pytest.raises(MoveRefused, match="pkdb tables sync"):
        move_study(moved_checkout, OLD, "codeine/Example", sf_vocabulary)
    assert snapshot(folder) == before
    assert not (moved_checkout / "studies" / "codeine").exists()


@pytest.mark.parametrize("change", ["save", "open"])
def test_a_workbook_changed_during_the_move_is_kept(
    moved_checkout, sf_vocabulary, monkeypatch, change
):
    from pkdb.lifecycle import move

    folder = study_folder(moved_checkout)
    assert sync_study(folder, sf_vocabulary).workbook_action == "created"
    state = state_path(workbook_path(folder))
    state.write_text("{}", encoding="utf-8")
    rename = Path.rename

    def edit_workbook(path, target):
        # The curator saves or opens the workbook while the files are renamed.
        moved = rename(path, target)
        workbook = Path(target) / "Example.xlsx"
        if workbook.exists():
            if change == "save":
                set_cell(workbook, "outputs_Tab2", 2, "mean", "3.5")
            else:
                lock = Path(target) / ".~lock.Example.xlsx#"
                lock.write_text("lock", encoding="utf-8")
        return moved

    monkeypatch.setattr(move.Path, "rename", edit_workbook)
    with pytest.raises(MoveIncomplete) as error:
        move_study(moved_checkout, OLD, NEW, sf_vocabulary)
    assert str(error.value) == (
        "Moved studies/caffeine/Example to studies/codeine/Renamed, but "
        "Example.xlsx changed or was opened during the move, so it stays with its "
        "sync state; carry its edits over to the tables by hand and remove it. "
        "The study stays at its new place."
    )
    moved = error.value.moved
    assert moved is not None and moved.workbook is None
    assert (moved.folder / "Example.xlsx").exists()
    assert (moved.folder / state.name).exists()
    assert study_column(moved.folder) == {"Renamed"}


@pytest.mark.parametrize("failing", [1, 4])
def test_a_failed_rename_undoes_the_move(
    moved_checkout, sf_vocabulary, monkeypatch, failing
):
    folder = study_folder(moved_checkout)
    before = snapshot(folder)
    rename = Path.rename
    calls = []

    def flaky(path, target):
        calls.append(path)
        if len(calls) == failing:
            raise PermissionError(13, "Permission denied")
        return rename(path, target)

    monkeypatch.setattr(Path, "rename", flaky)
    with pytest.raises(MoveRefused, match="Permission denied. Nothing was changed"):
        move_study(moved_checkout, OLD, NEW, sf_vocabulary)
    assert snapshot(folder) == before
    assert not (moved_checkout / "studies" / "codeine").exists()


def test_a_failed_undo_is_reported(moved_checkout, sf_vocabulary, monkeypatch):
    rename = Path.rename
    calls = []

    def flaky(path, target):
        calls.append(path)
        # The third file rename fails, and so does undoing the second.
        if len(calls) in (4, 5):
            raise PermissionError(13, "Permission denied")
        return rename(path, target)

    monkeypatch.setattr(Path, "rename", flaky)
    with pytest.raises(MoveIncomplete) as error:
        move_study(moved_checkout, OLD, NEW, sf_vocabulary)
    assert error.value.moved is None
    assert str(error.value) == (
        "Cannot move studies/caffeine/Example to studies/codeine/Renamed: Permission "
        "denied. Undoing the move failed (Renamed_Fig1.png to Example_Fig1.png: "
        "Permission denied); finish or undo it by hand."
    )
    assert (study_folder(moved_checkout) / "Renamed_Fig1.png").exists()


def test_a_review_change_during_the_move_is_kept(
    moved_checkout, sf_vocabulary, monkeypatch
):
    from pkdb.lifecycle import move

    folder = study_folder(moved_checkout)
    rename = Path.rename

    def edit_review(path, target):
        # Another writer changes review.json while the files are renamed.
        moved = rename(path, target)
        review = Path(target) / "review.json"
        if review.exists():
            text = review.read_text(encoding="utf-8")
            review.write_text(text.replace("Who are", "Which are"), encoding="utf-8")
        return moved

    monkeypatch.setattr(move.Path, "rename", edit_review)
    with pytest.raises(MoveRefused, match="review.json changed"):
        move_study(moved_checkout, OLD, NEW, sf_vocabulary)
    assert "Which are" in (folder / "review.json").read_text(encoding="utf-8")
    assert (folder / "Example.pdf").exists()
    assert not (moved_checkout / "studies" / "codeine").exists()


def test_a_format_failure_leaves_the_study_at_its_new_place(
    moved_checkout, sf_vocabulary, monkeypatch
):
    calls = []

    def broken(folder, *, check=False):
        calls.append(folder)
        if check:
            return format_folder(folder, check=True)
        issue = ValidationIssue(code="invalid_json", message="subjects.tsv is broken")
        return FormatResult(folder, issues=[issue])

    monkeypatch.setattr("pkdb.lifecycle.move.format_folder", broken)
    with pytest.raises(MoveIncomplete, match="subjects.tsv is broken") as error:
        move_study(moved_checkout, OLD, NEW, sf_vocabulary)
    assert "Moved studies/caffeine/Example to studies/codeine/Renamed" in str(
        error.value
    )
    moved = error.value.moved
    assert moved is not None
    assert moved.folder == study_folder(moved_checkout, NEW) and moved.folder.is_dir()
    assert (moved.folder / "Renamed.pdf").exists()
    assert not (moved_checkout / "studies" / "caffeine").exists()


@pytest.fixture
def github(monkeypatch):
    from fake_github import FakeGitHub

    fake = FakeGitHub(issues=[{"number": 7, "title": OLD}])
    monkeypatch.setattr(
        "pkdb.lifecycle_cli.github_client",
        lambda repository, token, **options: fake.client(**options),
    )
    monkeypatch.delenv("GH_TOKEN", raising=False)
    monkeypatch.setenv("GITHUB_TOKEN", "token")
    return fake


@pytest.fixture
def with_issue(moved_checkout):
    folder = study_folder(moved_checkout)
    patch_metadata(folder, {"issue": 7}, read_metadata(folder).revision)
    return moved_checkout


def test_move_command_renames_the_issue(with_issue, github, monkeypatch, capsys):
    monkeypatch.chdir(with_issue / "studies")
    assert main(["move", OLD, NEW, "--format", "human"]) == 0
    output = capsys.readouterr()
    lines = output.out.splitlines()
    assert lines[0] == "Moved studies/caffeine/Example to studies/codeine/Renamed"
    assert "Renamed Example.pdf to Renamed.pdf" in lines
    assert "Renamed Example_Fig1.wpd.json to Renamed_Fig1.wpd.json" in lines
    assert lines[-3:] == [
        "Updated 2 review targets",
        "Updated the name in reference.json",
        "Issue #7 renamed to codeine/Renamed",
    ]
    assert output.err == ""
    assert github.issues[7]["title"] == NEW
    assert github.writes == [("PATCH", "/issues/7", {"title": NEW})]


def test_move_command_json(with_issue, github, capsys):
    arguments = ["move", OLD, NEW, "--root", str(with_issue), "--format", "json"]
    assert main(arguments) == 0
    output = json.loads(capsys.readouterr().out)
    assert output["registry"] is False
    assert output["from"] == OLD and output["location"] == NEW
    assert output["path"] == str(study_folder(with_issue, NEW))
    assert {"from": "Example.pdf", "to": "Renamed.pdf"} in output["renamed"]
    assert output["targets"] == 2 and output["workbook"] is None
    assert output["assets"] == 0 and output["reference"] is True
    assert output["issue"] == 7 and output["issue_renamed"] is True
    assert output["warnings"] == []


def test_move_command_names_the_registry_update(
    with_issue, github, monkeypatch, capsys
):
    released(with_issue, json.dumps(REGISTRY, indent=2))
    monkeypatch.chdir(with_issue)
    assert main(["move", OLD, NEW, "--format", "human"]) == 0
    lines = capsys.readouterr().out.splitlines()
    assert "Updated PKDB00009 in studies/study_identifiers.json" in lines


def test_move_command_without_a_token_warns(with_issue, monkeypatch, capsys):
    monkeypatch.delenv("GH_TOKEN", raising=False)
    monkeypatch.delenv("GITHUB_TOKEN", raising=False)
    monkeypatch.chdir(with_issue)
    assert main(["move", OLD, NEW, "--format", "human"]) == 0
    output = capsys.readouterr()
    assert output.out.splitlines()[0] == (
        "Moved studies/caffeine/Example to studies/codeine/Renamed"
    )
    assert output.err == (
        "Warning: Set GH_TOKEN or GITHUB_TOKEN to rename issue #7 now; "
        "pkdb issues sync renames it later\n"
    )
    assert study_folder(with_issue, NEW).is_dir()


def test_move_command_warns_when_github_fails(with_issue, github, monkeypatch, capsys):
    github.fail[("PATCH", 7)] = 500
    monkeypatch.chdir(with_issue)
    assert main(["move", OLD, NEW, "--format", "json"]) == 0
    output = json.loads(capsys.readouterr().out)
    assert output["issue_renamed"] is False
    [warning] = output["warnings"]
    assert warning.startswith("Issue #7 was not renamed: ")
    assert warning.endswith("; pkdb issues sync renames it later")
    assert github.issues[7]["title"] == OLD
    assert study_folder(with_issue, NEW).is_dir()


def test_move_command_announces_waits_for_github(
    with_issue, github, monkeypatch, capsys
):
    github.fail[("PATCH", 7)] = 429
    monkeypatch.chdir(with_issue)
    assert main(["move", OLD, NEW, "--format", "human"]) == 0
    err = capsys.readouterr().err.splitlines()
    assert err[0] == "Waiting 60 seconds: GitHub limits the requests."
    assert err[-1].startswith("Warning: Issue #7 was not renamed: ")


@pytest.mark.parametrize("output", ["human", "json"])
def test_an_interrupted_issue_rename_keeps_the_move(
    with_issue, github, monkeypatch, capsys, output
):
    answer = github.handler

    def handler(request):
        if request.method == "PATCH":
            raise KeyboardInterrupt
        return answer(request)

    monkeypatch.setattr(github, "handler", handler)
    monkeypatch.chdir(with_issue)
    assert main(["move", OLD, NEW, "--format", output]) == 130
    captured = capsys.readouterr()
    warning = "Renaming issue #7 was interrupted; pkdb issues sync renames it later"
    if output == "json":
        data = json.loads(captured.out)
        assert data["location"] == NEW and data["issue_renamed"] is False
        assert data["warnings"] == [warning]
    else:
        lines = captured.out.splitlines()
        assert lines[0] == "Moved studies/caffeine/Example to studies/codeine/Renamed"
        assert captured.err == f"Warning: {warning}\n"
    assert study_folder(with_issue, NEW).is_dir()
    assert github.issues[7]["title"] == OLD


NOT_FOLLOWED = (
    "The study has no issue and no release, so PK-DB cannot follow the move: if "
    "PK-DB stores it as caffeine/Example, uploading codeine/Renamed is refused. "
    "Then move it back, run pkdb issues sync --adopt, upload it, and move it again."
)


def test_move_command_without_an_issue_or_release_warns(
    moved_checkout, github, monkeypatch, capsys
):
    automatic_curation(study_folder(moved_checkout))
    monkeypatch.chdir(moved_checkout)
    assert main(["move", OLD, NEW, "--format", "human"]) == 0
    output = capsys.readouterr()
    lines = output.out.splitlines()
    assert not [line for line in lines if "Issue" in line]
    assert lines[-3:] == [
        "Updated 2 review targets",
        "Updated 1 provenance asset of study.json",
        "Updated the name in reference.json",
    ]
    assert output.err == f"Warning: {NOT_FOLLOWED}\n"
    assert github.writes == []
    arguments = ["move", NEW, OLD, "--format", "json"]
    assert main(arguments) == 0
    output = json.loads(capsys.readouterr().out)
    assert output["assets"] == 1 and output["issue"] is None
    assert output["warnings"] == [
        NOT_FOLLOWED.replace("caffeine/Example", "X")
        .replace("codeine/Renamed", "caffeine/Example")
        .replace("X", "codeine/Renamed")
    ]


def test_move_command_of_a_released_study_without_an_issue_does_not_warn(
    moved_checkout, monkeypatch, capsys
):
    folder = study_folder(moved_checkout)
    release = {"release": {"pkdb_id": "PKDB00001", "date": "2026-10-01"}}
    patch_metadata(folder, release, read_metadata(folder).revision)
    monkeypatch.delenv("GH_TOKEN", raising=False)
    monkeypatch.delenv("GITHUB_TOKEN", raising=False)
    monkeypatch.chdir(moved_checkout)
    assert main(["move", OLD, NEW, "--format", "json"]) == 0
    assert json.loads(capsys.readouterr().out)["warnings"] == []


def test_move_command_refusal_exits_1(with_issue, github, monkeypatch, capsys):
    (with_issue / "studies" / "codeine" / "Renamed").mkdir(parents=True)
    monkeypatch.chdir(with_issue)
    assert main(["move", OLD, NEW, "--format", "json"]) == 1
    assert json.loads(capsys.readouterr().out) == {
        "from": OLD,
        "location": NEW,
        "error": "studies/codeine/Renamed exists already",
    }
    assert main(["move", OLD, NEW, "--format", "human"]) == 1
    assert capsys.readouterr().err == "studies/codeine/Renamed exists already\n"
    assert github.writes == []


def test_move_command_reports_an_incomplete_move(
    with_issue, github, monkeypatch, capsys
):
    def broken(folder, *, check=False):
        if check:
            return format_folder(folder, check=True)
        issue = ValidationIssue(code="invalid_json", message="subjects.tsv is broken")
        return FormatResult(folder, issues=[issue])

    monkeypatch.setattr("pkdb.lifecycle.move.format_folder", broken)
    monkeypatch.chdir(with_issue)
    assert main(["move", OLD, NEW, "--format", "human"]) == 1
    output = capsys.readouterr()
    assert "subjects.tsv is broken" in output.err
    assert "Issue #7 renamed to codeine/Renamed" in output.out
    assert github.issues[7]["title"] == NEW


def test_move_command_bad_token_moves_nothing(with_issue, monkeypatch, capsys):
    monkeypatch.delenv("GH_TOKEN", raising=False)
    monkeypatch.setenv("GITHUB_TOKEN", "bad token")
    monkeypatch.chdir(with_issue)
    assert main(["move", OLD, NEW]) == 2
    assert "whitespace" in capsys.readouterr().err
    assert study_folder(with_issue).is_dir()
    assert not study_folder(with_issue, NEW).exists()
