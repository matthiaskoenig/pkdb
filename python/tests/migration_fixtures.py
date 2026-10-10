"""Synthetic format 1 studies for the migration tests."""

import io
import json
import re
import zipfile
from pathlib import Path
from typing import Any, NamedTuple

import openpyxl
from openpyxl.utils import get_column_letter
from PIL import Image

# study.json of the twin; Any lets tests derive variants from its nested parts.
STUDY: dict[str, Any] = {
    "sid": "Example",
    "name": "Example",
    "reference": "123",
    "creator": "curator",
    "curators": [["curator", 3]],
    "licence": "open",
    "access": "private",
    "groupset": {
        "groups": [
            {
                "name": "all",
                "count": 2,
                "image": "Tab1",
                "characteristica": [
                    {
                        "measurement_type": "species",
                        "choice": "Homo sapiens",
                        "image": "Tab1",
                    },
                    {"measurement_type": "healthy", "choice": "Y", "image": "Tab1"},
                    {"measurement_type": "sex", "choice": "M", "image": "Tab1"},
                ],
            }
        ]
    },
    "individualset": {
        "individuals": [
            {
                "name": name,
                "group": "all",
                "image": "TabA",
                "characteristica": [
                    {
                        "measurement_type": "age",
                        "value": age,
                        "unit": "yr",
                        "image": "TabA",
                    }
                ],
            }
            for name, age in (("S1", 30), ("S2", 40))
        ]
    },
    "interventionset": {
        "interventions": [
            {
                "name": "D1",
                "measurement_type": "dosing",
                "substance": "drug",
                "route": "oral",
                "form": "tablet",
                "application": "single dose",
                "time": 0,
                "time_unit": "h",
                "value": 100,
                "unit": "mg",
            }
        ]
    },
    "outputset": {
        "outputs": [
            {
                "source": "Tab2",
                "image": "Tab2",
                "output_type": "output",
                "group": "all",
                "interventions": ["D1"],
                "measurement_type": "cmax",
                "substance": "drug",
                "tissue": "plasma",
                "mean": "col==mean",
                "sd": "col==sd",
                "unit": "mg/l",
            },
            {
                "source": "Fig1",
                "image": "Fig1",
                "output_type": "timecourse",
                "label": "drug_plasma",
                "group": "all",
                "interventions": ["D1"],
                "measurement_type": "concentration",
                "substance": "drug",
                "tissue": "plasma",
                "time": "col==time",
                "time_unit": "h",
                "mean": "col==mean",
                "unit": "mg/l",
            },
        ]
    },
}
SHEETS = {
    "Tab2": [["mean", "sd"], [2.5, 0.5]],
    "Fig1": [["time", "mean"], [0, 0], [1, 2], [2, 1]],
}
IMAGES = ("Tab1", "TabA", "Tab2", "Fig1")
# The scatter of the twin: dataset age_vs_cmax, x age of S1 and S2, y cmax after D1.
SCATTER_OUTPUTS = [
    {
        "source": "Fig2",
        "image": "Fig2",
        "output_type": "output",
        "label": "age_vs_cmax_x",
        "individual": "col==subject",
        "measurement_type": "age",
        "mean": "col==age",
        "unit": "yr",
    },
    {
        "source": "Fig2",
        "image": "Fig2",
        "output_type": "output",
        "label": "age_vs_cmax_y",
        "individual": "col==subject",
        "interventions": ["D1"],
        "measurement_type": "cmax",
        "substance": "drug",
        "tissue": "plasma",
        "mean": "col==cmax",
        "unit": "mg/l",
    },
]
DATASET: dict[str, Any] = {
    "data": [
        {
            "name": "age_vs_cmax",
            "data_type": "scatter",
            "image": "Fig2",
            "subsets": [
                {
                    "name": "age_vs_cmax",
                    "dimensions": ["age_vs_cmax_x", "age_vs_cmax_y"],
                    "shared": ["individual"],
                }
            ],
        }
    ]
}
SCATTER_SHEET = {"Fig2": [["subject", "age", "cmax"], ["S1", 30, 2], ["S2", 40, 3]]}


class Formula(NamedTuple):
    """A formula cell and the value that a spreadsheet application saved with it."""

    text: str
    value: float


def _save_formula_values(path: Path, values: dict[int, dict[str, float]]) -> None:
    """Write the saved values of formula cells, which openpyxl leaves empty.

    `values` maps the index of a sheet (1 for the first) to cell values by
    coordinate. openpyxl writes a formula cell as `<c r="B3"><f>...</f><v /></c>`.
    """
    with zipfile.ZipFile(path) as archive:
        files = {info: archive.read(info) for info in archive.infolist()}
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for info, data in files.items():
            index = re.fullmatch(r"xl/worksheets/sheet([0-9]+)\.xml", info.filename)
            for cell, value in values.get(int(index[1]) if index else 0, {}).items():
                data, count = re.subn(
                    rf'<c r="{cell}"><f>(.*?)</f><v ?/></c>'.encode(),
                    rf'<c r="{cell}" t="n"><f>\1</f><v>{value!r}</v></c>'.encode(),
                    data,
                )
                assert count == 1, f"no formula cell {cell} in {info.filename}"
            archive.writestr(info, data)


def write_sheets(folder: Path, name: str, sheets: dict, *, workbook: bool) -> None:
    """Format 1 tables: a workbook with a notes row above the header, or hidden TSVs.

    Both give the same rows: the importer reads the header of a sheet from its
    second row and the header of a TSV from its first line. A `Formula` cell is
    a formula with its saved value in a workbook, and that value in a TSV.
    """
    if not sheets:
        return
    if workbook:
        book = openpyxl.Workbook()
        book.remove(book.active)
        values: dict[int, dict[str, float]] = {}
        for index, (title, rows) in enumerate(sheets.items(), 1):
            sheet = book.create_sheet(title)
            sheet.append(["Curator notes"])
            for number, row in enumerate(rows, 2):
                for column, cell in enumerate(row, 1):
                    if isinstance(cell, Formula):
                        coordinate = f"{get_column_letter(column)}{number}"
                        values.setdefault(index, {})[coordinate] = cell.value
                sheet.append([c.text if isinstance(c, Formula) else c for c in row])
        book.save(folder / f"{name}.xlsx")
        book.close()
        if values:
            _save_formula_values(folder / f"{name}.xlsx", values)
    else:
        for title, rows in sheets.items():
            text = "".join(
                "\t".join(str(c.value if isinstance(c, Formula) else c) for c in row)
                + "\n"
                for row in rows
            )
            (folder / f".{name}_{title}.tsv").write_text(text, encoding="utf-8")


def tiny_png() -> bytes:
    """A decodable PNG image."""
    buffer = io.BytesIO()
    Image.new("RGB", (2, 2), "white").save(buffer, format="PNG")
    return buffer.getvalue()


def v1_study(
    root: Path,
    study: dict,
    sheets: dict,
    images: tuple[str, ...],
    *,
    substance: str = "caffeine",
    workbook: bool = True,
    reference: dict | None = None,
) -> Path:
    """A format 1 study folder at `root/studies/<substance>/<name>`."""
    name = study["name"]
    folder = root / "studies" / substance / name
    folder.mkdir(parents=True)
    (folder / "study.json").write_text(json.dumps(study), encoding="utf-8")
    (folder / "reference.json").write_text(
        json.dumps(
            reference
            or {"sid": "123", "name": name, "pmid": "123", "title": "Example study"}
        ),
        encoding="utf-8",
    )
    write_sheets(folder, name, sheets, workbook=workbook)
    (folder / f"{name}.pdf").write_bytes(b"%PDF")
    for source in images:
        (folder / f"{name}_{source}.png").write_bytes(tiny_png())
    return folder


def v1_example(root: Path, *, workbook: bool = True) -> Path:
    """The format 1 twin of the `valid_files` fixture, without its scatters."""
    return v1_study(root, STUDY, SHEETS, IMAGES, workbook=workbook)


def v1_full_example(root: Path, *, workbook: bool = True) -> Path:
    """The format 1 twin of the whole `valid_files` study, scatters included."""
    study = {
        **STUDY,
        "outputset": {"outputs": [*STUDY["outputset"]["outputs"], *SCATTER_OUTPUTS]},
        "dataset": DATASET,
    }
    return v1_study(
        root, study, {**SHEETS, **SCATTER_SHEET}, (*IMAGES, "Fig2"), workbook=workbook
    )
