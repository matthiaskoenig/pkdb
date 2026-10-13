"""The read and write routes of a study page on a real loopback server over a real engine."""

import errno
import json
import os
import re
import sys
import threading
import time
from collections import Counter
from types import SimpleNamespace
from unittest.mock import ANY
from urllib.parse import quote

import openpyxl
import pytest
from curation_http import authenticate, error_body, request
from digitize_fixtures import GOOD, png, project
from open_commands import (
    FAILING_OPENER,
    command_line,
    holding_opener,
    late_opener,
    stop_program,
    wait_for_open_commands,
)

from pkdb.curation import studies
from pkdb.curation.engine import CurationEngine
from pkdb.curation.server import create_server
from pkdb.identity import Author, IdentityError
from pkdb.schemas.validation import fail
from pkdb.studyformat.formatter import format_folder
from pkdb.studyformat.issues import make_issue
from pkdb.studyformat.jsonio import dump_json
from pkdb.studyformat.review_edit import read_review
from pkdb.studyformat.sync import sync_study
from pkdb.studyformat.ulid import new_ulid
from pkdb.studyformat.workbook.base import workbook_path

DETAIL = "/local/studies/caffeine/Example"
METADATA = {"format": 2, "creator": "curator", "licence": "open", "access": "private"}
# The kinds of tables that the study page describes to the app, in the order of the sheets.
TABLE_KINDS_OF_THE_APP = [
    {
        "kind": "subjects",
        "label": "Subjects",
        "info": "subjects",
        "per_source": False,
        "addable": False,
    },
    {
        "kind": "characteristica",
        "label": "Characteristica",
        "info": "subjects",
        "per_source": False,
        "addable": False,
    },
    {
        "kind": "interventions",
        "label": "Interventions",
        "info": "interventions",
        "per_source": False,
        "addable": False,
    },
    {
        "kind": "outputs",
        "label": "Outputs",
        "info": "data",
        "per_source": True,
        "addable": True,
    },
    {
        "kind": "timecourses",
        "label": "Timecourses",
        "info": "data",
        "per_source": True,
        "addable": True,
    },
    {
        "kind": "scatters",
        "label": "Scatters",
        "info": "data",
        "per_source": True,
        "addable": True,
    },
    {
        "kind": "raw",
        "label": "Raw table",
        "info": "raw",
        "per_source": True,
        "addable": True,
    },
]


@pytest.fixture
def api(tmp_path_factory, make_study, valid_files):
    folder = make_study(
        {
            **valid_files,
            "Example_Tab2.tsv": "cmax\t2.5 ± 0.5\n",
            "Example_Fig1.png": png(100, 100),
            "Example_Fig1.wpd.json": json.dumps(project(GOOD)),
        }
    )
    assert format_folder(folder).ok
    engine = CurationEngine(
        folder.parent.parent,
        state_dir=tmp_path_factory.mktemp("state"),
        offline=True,
        start=False,
    )
    engine.user = "curator"
    server = create_server(engine)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    yield server, engine, folder
    server.shutdown()
    server.server_close()
    thread.join(timeout=3)


def test_study_detail(api):
    server, engine, folder = api
    headers = authenticate(server)
    status, response_headers, data = request(server, "GET", DETAIL, headers=headers)
    assert status == 200
    detail = json.loads(data)
    assert detail["id"] == "caffeine/Example"
    assert detail["metadata"]["value"]["licence"] == "open"
    assert detail["metadata"]["revision"] and detail["metadata"]["issues"] == []
    assert detail["review"]["value"]["status"] == "draft"
    assert {s["source"] for s in detail["sources"]} >= {"Fig1", "Tab2"}
    tab2 = next(s for s in detail["sources"] if s["source"] == "Tab2")
    assert (tab2["kind"], tab2["missing_image"]) == ("table", None)
    assert "missing_raw" in tab2
    assert "Example_Fig1.wpd.json" in detail["files"]
    assert detail["tables"] == [
        {
            "file": "subjects.tsv",
            "kind": "subjects",
            "info": "subjects",
            "source": None,
            "rows": 3,
        },
        {
            "file": "characteristica.tsv",
            "kind": "characteristica",
            "info": "subjects",
            "source": None,
            "rows": 5,
        },
        {
            "file": "interventions.tsv",
            "kind": "interventions",
            "info": "interventions",
            "source": None,
            "rows": 1,
        },
        {
            "file": "outputs_Tab2.tsv",
            "kind": "outputs",
            "info": "data",
            "source": "Tab2",
            "rows": 1,
        },
        {
            "file": "timecourses_Fig1.tsv",
            "kind": "timecourses",
            "info": "data",
            "source": "Fig1",
            "rows": 3,
        },
        {
            "file": "scatters_Fig2.tsv",
            "kind": "scatters",
            "info": "data",
            "source": "Fig2",
            "rows": 2,
        },
        # A raw table belongs to its paper table; its rows are its non-empty lines.
        {
            "file": "Example_Tab2.tsv",
            "kind": "raw",
            "info": "raw",
            "source": "Tab2",
            "rows": 1,
        },
    ]
    assert detail["table_kinds"] == TABLE_KINDS_OF_THE_APP
    etag = response_headers["ETag"]
    again = request(server, "GET", DETAIL, headers={**headers, "If-None-Match": etag})
    assert again[0] == 304 and again[2] == b"" and again[1]["ETag"] == etag
    (folder / "study.json").write_text(
        (folder / "study.json")
        .read_text(encoding="utf-8")
        .replace('"open"', '"closed"'),
        encoding="utf-8",
        newline="",
    )
    engine.scan()
    changed = request(server, "GET", DETAIL, headers={**headers, "If-None-Match": etag})
    assert changed[0] == 200 and changed[1]["ETag"] != etag
    assert json.loads(changed[2])["metadata"]["value"]["licence"] == "closed"


def test_study_detail_names_the_image_url_of_each_source(api):
    server, engine, folder = api
    headers = authenticate(server)
    (folder / "Example_Tab2.png").unlink()
    detail = json.loads(request(server, "GET", DETAIL, headers=headers)[2])
    sources = {source["source"]: source for source in detail["sources"]}
    # The same URL as the source view, so that the browser shows one image for both.
    assert sources["Fig1"]["image_url"] == f"{DETAIL}/files/Example_Fig1.png"
    view = json.loads(
        request(server, "GET", f"{DETAIL}/sources/Fig1", headers=headers)[2]
    )
    assert view["image_url"] == sources["Fig1"]["image_url"]
    # A source without an image has no URL.
    assert sources["Tab2"]["image"] is None and sources["Tab2"]["image_url"] is None


def test_study_detail_gives_no_image_url_for_an_image_that_the_file_route_refuses(
    api, tmp_path
):
    server, engine, folder = api
    headers = authenticate(server)
    outside = tmp_path / "outside.png"
    outside.write_bytes((folder / "Example_Fig1.png").read_bytes())
    (folder / "Example_Fig1.png").unlink()
    (folder / "Example_Fig1.png").symlink_to(outside)
    engine.scan()
    detail = json.loads(request(server, "GET", DETAIL, headers=headers)[2])
    fig1 = next(source for source in detail["sources"] if source["source"] == "Fig1")
    url = f"{DETAIL}/files/Example_Fig1.png"
    assert request(server, "GET", url, headers=headers)[0] == 404
    # A symlink is no file of the study: the image counts as missing and has no URL, which
    # would show a broken thumbnail in the navigator.
    assert (fig1["image"], fig1["missing_image"]) == (None, "Example_Fig1.png")
    assert fig1["image_url"] is None
    view = json.loads(
        request(server, "GET", f"{DETAIL}/sources/Fig1", headers=headers)[2]
    )
    assert view["image_url"] is None


def test_study_detail_lists_a_table_that_cannot_be_read_without_a_row_count(api):
    server, engine, folder = api
    # A git conflict marker makes the table unreadable; it stays a table of the study.
    (folder / "timecourses_Fig1.tsv").write_text("<<<<<<< HEAD\n", encoding="utf-8")
    tables = {
        entry["file"]: entry
        for entry in engine.study_detail("caffeine/Example")["tables"]
    }
    assert tables["timecourses_Fig1.tsv"] == {
        "file": "timecourses_Fig1.tsv",
        "kind": "timecourses",
        "info": "data",
        "source": "Fig1",
        "rows": None,
    }
    assert tables["outputs_Tab2.tsv"]["rows"] == 1


def test_detail_etag_follows_a_file_that_changes_and_changes_back_between_two_scans(
    api,
):
    # The engine runs no watcher: the scan of the files is older than both changes.
    server, engine, folder = api
    headers = authenticate(server)
    study_json = folder / "study.json"
    original = study_json.read_text(encoding="utf-8")
    study_json.write_text(
        original.replace('"open"', '"closed"'), encoding="utf-8", newline=""
    )
    status, response_headers, data = request(server, "GET", DETAIL, headers=headers)
    assert status == 200
    assert json.loads(data)["metadata"]["value"]["licence"] == "closed"
    study_json.write_text(original, encoding="utf-8", newline="")
    status, _, data = request(
        server,
        "GET",
        DETAIL,
        headers={**headers, "If-None-Match": response_headers["ETag"]},
    )
    assert status == 200
    assert json.loads(data)["metadata"]["value"]["licence"] == "open"


def test_table_etag_follows_a_table_that_changes_and_changes_back_between_two_scans(
    api,
):
    server, engine, folder = api
    headers = authenticate(server)
    path = f"{DETAIL}/tables/Example_Tab2.tsv"
    table = folder / "Example_Tab2.tsv"
    original = table.read_text(encoding="utf-8")
    table.write_text("cmax\t3.5 ± 0.5\n", encoding="utf-8", newline="")
    status, response_headers, data = request(server, "GET", path, headers=headers)
    assert status == 200
    assert json.loads(data)["rows"][0]["cells"] == ["cmax", "3.5 ± 0.5"]
    table.write_text(original, encoding="utf-8", newline="")
    status, _, data = request(
        server,
        "GET",
        path,
        headers={**headers, "If-None-Match": response_headers["ETag"]},
    )
    assert status == 200
    assert json.loads(data)["rows"][0]["cells"] == ["cmax", "2.5 ± 0.5"]


def test_detail_etag_changes_with_a_job(api):
    server, engine, folder = api
    headers = authenticate(server)
    status, response_headers, _ = request(server, "GET", DETAIL, headers=headers)
    assert status == 200
    engine.enqueue(["caffeine/Example"], "validate")
    status, changed_headers, data = request(
        server,
        "GET",
        DETAIL,
        headers={**headers, "If-None-Match": response_headers["ETag"]},
    )
    assert status == 200 and changed_headers["ETag"] != response_headers["ETag"]
    detail = json.loads(data)
    assert detail["status"] == "queued"
    assert [job["status"] for job in detail["jobs"]] == ["queued"]


def test_invalid_documents_keep_their_revision(api):
    server, engine, folder = api
    (folder / "review.json").write_text(
        '{"status": "finished"}', encoding="utf-8", newline=""
    )
    detail = json.loads(request(server, "GET", DETAIL, headers=authenticate(server))[2])
    assert detail["review"]["value"] is None and detail["review"]["revision"]
    assert detail["review"]["issues"][0]["code"] == "invalid_review_json"
    assert detail["metadata"]["value"]["licence"] == "open"
    # The targets come from the same read of review.json as the items.
    assert detail["targets"] == {}


def test_symlinked_review_json_is_not_read(api, tmp_path_factory):
    server, engine, folder = api
    outside = tmp_path_factory.mktemp("outside") / "review.json"
    outside.write_text('{"status": "in_review"}', encoding="utf-8", newline="")
    (folder / "review.json").unlink()
    (folder / "review.json").symlink_to(outside)
    status, _, data = request(server, "GET", DETAIL, headers=authenticate(server))
    assert status == 200
    review = json.loads(data)["review"]
    assert review["value"] is None and review["revision"] is None
    assert "symlink" in {issue["code"] for issue in review["issues"]}


def test_directory_named_review_json_is_reported(api):
    server, engine, folder = api
    (folder / "review.json").unlink()
    (folder / "review.json").mkdir()
    status, _, data = request(server, "GET", DETAIL, headers=authenticate(server))
    assert status == 200
    detail = json.loads(data)
    assert detail["review"]["value"] is None
    assert {"unknown_directory", "missing_file"} <= {
        issue["code"] for issue in detail["review"]["issues"]
    }
    assert detail["metadata"]["value"]["licence"] == "open"


def test_missing_study_json_has_the_absent_revision(api):
    server, engine, folder = api
    (folder / "study.json").unlink()
    detail = json.loads(request(server, "GET", DETAIL, headers=authenticate(server))[2])
    assert detail["metadata"]["value"] is None
    assert detail["metadata"]["revision"] == "absent"
    assert detail["metadata"]["issues"][0]["code"] == "missing_file"


@pytest.mark.parametrize(
    ("limit", "value", "code"),
    [("MAX_ROWS", 3, "row_limit"), ("MAX_FILES", 2, "file_limit")],
)
def test_upload_limits_bound_the_read_routes(api, monkeypatch, limit, value, code):
    server, engine, folder = api
    monkeypatch.setattr(studies, limit, value)
    headers = authenticate(server)
    status, _, data = request(server, "GET", DETAIL, headers=headers)
    assert status == 200
    detail = json.loads(data)
    assert detail["metadata"]["value"]["licence"] == "open"
    assert detail["review"]["value"]["status"] == "draft"
    assert detail["problems"][0]["code"] == code
    assert detail["sources"] == [] and detail["files"] == [] and detail["tables"] == []
    # Add table and the notes of study.json still need the kinds of tables.
    assert detail["table_kinds"] == TABLE_KINDS_OF_THE_APP
    assert detail["targets"] == {}
    for path in (f"{DETAIL}/tables/timecourses_Fig1.tsv", f"{DETAIL}/sources/Fig1"):
        status, _, data = request(server, "GET", path, headers=headers)
        assert status == 413 and "more than" in json.loads(data)["error"]
    body = {
        "study": "caffeine/Example",
        "target": {"file": "timecourses_Fig1.tsv", "rows": {"time": "1"}},
    }
    status, _, data = request(
        server, "POST", "/local/studies/review/preview", body, headers
    )
    assert status == 413 and "more than" in json.loads(data)["error"]


def test_a_study_that_cannot_be_read_is_not_beyond_the_limits(api, monkeypatch):
    server, engine, folder = api

    def unreadable(folder, **limits):
        fail("invalid_tsv", "timecourses_Fig1.tsv cannot be read")

    monkeypatch.setattr(studies, "load_study", unreadable)
    headers = authenticate(server)
    body = {
        "study": "caffeine/Example",
        "target": {"file": "timecourses_Fig1.tsv", "rows": {"time": "1"}},
    }
    for method, path, payload in [
        ("POST", "/local/studies/review/preview", body),
        ("GET", f"{DETAIL}/tables/timecourses_Fig1.tsv", None),
    ]:
        status, _, data = request(server, method, path, payload, headers)
        refused = json.loads(data)
        assert status == 422, path
        assert [issue["code"] for issue in refused["issues"]] == ["invalid_tsv"]
    detail = _detail(server, headers)
    assert detail["problems"][0]["code"] == "invalid_tsv"
    assert detail["tables"] == [] and detail["targets"] == {}


def test_detail_lists_sync_conflicts(api, sf_vocabulary, monkeypatch):
    server, engine, folder = api
    # The bundled vocabulary lacks the substance of the test study.
    monkeypatch.setattr(engine, "_local_vocabulary", lambda folder=None: sf_vocabulary)
    headers = authenticate(server)
    assert (
        json.loads(request(server, "GET", DETAIL, headers=headers)[2])["conflicts"]
        == []
    )
    assert sync_study(folder, sf_vocabulary).ok
    book = openpyxl.load_workbook(workbook_path(folder))
    sheet = book["timecourses_Fig1"]
    header = [cell.value for cell in sheet[1]]
    sheet.cell(3, header.index("mean") + 1).value = 5
    book.save(workbook_path(folder))
    table = folder / "timecourses_Fig1.tsv"
    lines = table.read_text(encoding="utf-8").splitlines()
    cells = lines[2].split("\t")
    cells[lines[0].split("\t").index("mean")] = "7"
    lines[2] = "\t".join(cells)
    table.write_text("\n".join(lines) + "\n", encoding="utf-8", newline="")
    engine.scan()
    detail = json.loads(request(server, "GET", DETAIL, headers=headers)[2])
    assert detail["sync"]["status"] == "conflict"
    [conflict] = detail["conflicts"]
    assert conflict["file"] == "timecourses_Fig1.tsv"
    assert conflict["workbook_rows"][0]["row"] == 3
    assert conflict["table_lines"][0]["line"] == 2


def test_acknowledged_warnings_are_listed_until_dismissed(api):
    server, engine, folder = api
    created = "2026-10-01T10:00:00Z"
    acknowledged, dismissed = new_ulid(), new_ulid()
    target = {"file": "timecourses_Fig1.tsv", "rows": {"time": "1"}, "column": "mean"}
    review = {
        "status": "draft",
        "items": [
            {
                "id": acknowledged,
                "kind": "issue",
                "state": "resolved",
                "target": target,
                "acknowledges": "digitized_mismatch",
                "text": "The figure is blurred here",
                "author": "curator",
                "created": created,
                "resolved_by": "curator",
                "resolved": created,
            },
            {
                "id": dismissed,
                "kind": "issue",
                "state": "dismissed",
                "acknowledges": "digitized_mismatch",
                "text": "Withdrawn",
                "author": "curator",
                "created": created,
                "resolved_by": "curator",
                "resolved": created,
            },
            {
                "id": new_ulid(),
                "kind": "question",
                "text": "Which dose was given?",
                "author": "curator",
                "created": created,
            },
        ],
    }
    (folder / "review.json").write_text(
        json.dumps(review), encoding="utf-8", newline=""
    )
    detail = json.loads(request(server, "GET", DETAIL, headers=authenticate(server))[2])
    assert detail["acknowledged"] == [
        {
            "id": acknowledged,
            "code": "digitized_mismatch",
            "target": target,
            "text": "The figure is blurred here",
            "author": "curator",
            "resolved_by": "curator",
            "resolved": created,
            "scope": "rows",
        }
    ]


def test_tables_sources_and_images(api):
    server, engine, folder = api
    headers = authenticate(server)
    table = json.loads(
        request(
            server, "GET", f"{DETAIL}/tables/timecourses_Fig1.tsv", headers=headers
        )[2]
    )
    assert table["kind"] == "table" and table["rows"][0]["line"] == 1
    assert table["rows"][0]["cells"][table["header"].index("label")] == "drug_plasma"
    raw = json.loads(
        request(server, "GET", f"{DETAIL}/tables/Example_Tab2.tsv", headers=headers)[2]
    )
    assert raw["kind"] == "raw"
    assert raw["rows"] == [{"line": 1, "cells": ["cmax", "2.5 ± 0.5"]}]
    source = json.loads(
        request(server, "GET", f"{DETAIL}/sources/Fig1", headers=headers)[2]
    )
    assert source["image_url"] == f"{DETAIL}/files/Example_Fig1.png"
    assert source["overlay"] and source["layout"] == "overlay"
    # The points, the series with their colors, and the texts of the hover pass through.
    assert source["points"][0]["series"] == "drug_plasma"
    assert {"x_text", "y_text", "error_bar"} <= set(source["points"][0])
    assert source["series"][0]["name"] == "drug_plasma"
    assert {"color", "dark_color", "x_label", "y_label"} <= set(source["series"][0])
    assert {"x_text", "y_text"} <= set(source["overlay"][0])
    status, image_headers, data = request(
        server, "GET", source["image_url"], headers=headers
    )
    assert status == 200 and image_headers["Content-Type"] == "image/png"
    assert data.startswith(b"\x89PNG") and "ETag" not in image_headers


@pytest.mark.parametrize(
    "path", [f"{DETAIL}/tables/timecourses_Fig1.tsv", f"{DETAIL}/sources/Fig1"]
)
def test_tables_and_sources_answer_304(api, path):
    server, engine, folder = api
    headers = authenticate(server)
    status, response_headers, _ = request(server, "GET", path, headers=headers)
    assert status == 200
    again = request(
        server,
        "GET",
        path,
        headers={**headers, "If-None-Match": response_headers["ETag"]},
    )
    assert again[0] == 304 and again[2] == b""


def test_a_table_304_does_not_load_the_study(api, monkeypatch):
    server, engine, folder = api
    headers = authenticate(server)
    path = f"{DETAIL}/tables/timecourses_Fig1.tsv"
    etag = request(server, "GET", path, headers=headers)[1]["ETag"]
    loads = []
    monkeypatch.setattr(
        studies, "load_study", lambda *args, **kwargs: loads.append(args)
    )
    assert (
        request(server, "GET", path, headers={**headers, "If-None-Match": etag})[0]
        == 304
    )
    assert loads == []


def test_table_etag_changes_with_the_table(api):
    server, engine, folder = api
    headers = authenticate(server)
    path = f"{DETAIL}/tables/Example_Tab2.tsv"
    etag = request(server, "GET", path, headers=headers)[1]["ETag"]
    (folder / "Example_Tab2.tsv").write_text(
        "cmax\t3.5 ± 0.5\n", encoding="utf-8", newline=""
    )
    engine.scan()
    status, response_headers, data = request(
        server, "GET", path, headers={**headers, "If-None-Match": etag}
    )
    assert status == 200 and response_headers["ETag"] != etag
    assert json.loads(data)["rows"][0]["cells"] == ["cmax", "3.5 ± 0.5"]


@pytest.mark.parametrize(
    "path",
    [
        f"{DETAIL}/files/..%2Fstudy.json",
        f"{DETAIL}/files/study.json",
        f"{DETAIL}/files/%2Fetc%2Fpasswd",
        f"{DETAIL}/files/%2E%2E",
        f"{DETAIL}/files/..%5Cstudy.json",
        f"{DETAIL}/files/Example_Fig1.png%00",
        f"{DETAIL}/files/",
        "/local/studies/caffeine/../Example",
        f"{DETAIL}/tables/Example.pdf",
        f"{DETAIL}/tables/..%2F..%2FExample%2Fsubjects.tsv",
        "/local/studies/caffeine/Missing",
        "/local/studies/caffeine",
        f"{DETAIL}/sources/Fig9",
        f"{DETAIL}/notes/Fig1",
    ],
)
def test_crafted_paths_are_not_found(api, path):
    server, engine, folder = api
    assert request(server, "GET", path, headers=authenticate(server))[0] == 404


def test_symlinked_image_is_refused(api, tmp_path):
    server, engine, folder = api
    outside = tmp_path / "outside.png"
    outside.write_bytes((folder / "Example_Fig1.png").read_bytes())
    (folder / "Example_Fig3.png").symlink_to(outside)
    engine.scan()
    status = request(
        server, "GET", f"{DETAIL}/files/Example_Fig3.png", headers=authenticate(server)
    )[0]
    assert status == 404


@pytest.mark.parametrize(
    "path",
    [
        DETAIL,
        f"{DETAIL}/files/Example_Fig1.png",
        f"{DETAIL}/tables/Example_Tab2.tsv",
        f"{DETAIL}/sources/Fig1",
    ],
)
def test_study_folder_replaced_by_a_symlink_is_refused(api, tmp_path_factory, path):
    server, engine, folder = api
    headers = authenticate(server)
    etag = request(server, "GET", path, headers=headers)[1].get("ETag")
    moved = tmp_path_factory.mktemp("outside") / "Example"
    folder.rename(moved)
    folder.symlink_to(moved, target_is_directory=True)
    # The scan skips a symlinked folder and keeps its row.
    engine.scan()
    assert request(server, "GET", path, headers=headers)[0] == 404
    if etag:
        cached = {**headers, "If-None-Match": etag}
        assert request(server, "GET", path, headers=cached)[0] == 404


def test_duplicate_identity_answers_409(api, valid_files):
    server, engine, folder = api
    copy = engine.root / "copies" / "caffeine" / "Example"
    copy.mkdir(parents=True)
    for file, content in valid_files.items():
        (copy / file).write_bytes(
            content if isinstance(content, bytes) else content.encode()
        )
    engine.scan()
    status, _, data = request(server, "GET", DETAIL, headers=authenticate(server))
    body = json.loads(data)
    assert status == 409 and "identity of two folders" in body["error"]
    # The app lists the folders without parsing the message.
    assert sorted(body["paths"]) == ["caffeine/Example", "copies/caffeine/Example"]


def test_ambiguous_study_counts_and_lists_its_folders():
    error = studies.AmbiguousStudy(
        "caffeine/Example", ["a, b/caffeine/Example", "c", "d"]
    )
    assert str(error) == (
        "caffeine/Example is the identity of 3 folders: "
        "a, b/caffeine/Example, c, d; rename all but one"
    )
    assert error.paths == ["a, b/caffeine/Example", "c", "d"]


def test_study_routes_require_the_session(api):
    server, engine, folder = api
    assert request(server, "GET", DETAIL)[0] == 401
    assert request(server, "GET", f"{DETAIL}/files/Example_Fig1.png")[0] == 401


@pytest.mark.parametrize(
    ("path", "body"),
    [
        (
            "/local/studies/tables/preview",
            {"study": "caffeine/Example", "kind": "outputs", "source": "Tab3"},
        ),
        (
            "/local/studies/review/preview",
            {
                "study": "caffeine/Example",
                "target": {"file": "timecourses_Fig1.tsv", "rows": {"time": "1"}},
            },
        ),
    ],
)
def test_preview_routes_require_the_session_and_the_action_token(api, path, body):
    server, engine, folder = api
    headers = authenticate(server)
    without_session = {"X-CSRF-Token": headers["X-CSRF-Token"]}
    assert request(server, "POST", path, body, without_session)[0] == 401
    status, _, data = request(server, "POST", path, body, {"Cookie": headers["Cookie"]})
    assert status == 403
    assert json.loads(data)["code"] == "action_token"
    assert request(server, "POST", path, body, headers)[0] == 200


def test_state_etag(api):
    server, engine, folder = api
    headers = authenticate(server)
    status, response_headers, _ = request(
        server, "GET", "/local/state", headers=headers
    )
    assert status == 200
    again = request(
        server,
        "GET",
        "/local/state",
        headers={**headers, "If-None-Match": response_headers["ETag"]},
    )
    assert again[0] == 304 and again[2] == b""
    engine.set_paused(True)
    changed = request(
        server,
        "GET",
        "/local/state",
        headers={**headers, "If-None-Match": response_headers["ETag"]},
    )
    assert changed[0] == 200 and changed[1]["ETag"] != response_headers["ETag"]


STATE = "/local/state"


def test_the_state_answers_304_until_the_engine_state_changes(api):
    server, engine, _ = api
    headers = authenticate(server)
    status, response_headers, _ = request(server, "GET", STATE, headers=headers)
    assert status == 200
    etag = response_headers["ETag"]
    again = request(server, "GET", STATE, headers={**headers, "If-None-Match": etag})
    assert again[0] == 304 and again[2] == b""
    with engine.lock:
        engine.theme = "dark"
    changed = request(server, "GET", STATE, headers={**headers, "If-None-Match": etag})
    assert changed[0] == 200 and changed[1]["ETag"] != etag
    assert json.loads(changed[2])["theme"] == "dark"


def test_the_state_answers_304_across_watcher_ticks(api):
    server, engine, _ = api
    headers = authenticate(server)
    for _ in range(2):
        engine._watch_tick()
    first = request(server, "GET", STATE, headers=headers)
    assert first[0] == 200
    conditional = {**headers, "If-None-Match": first[1]["ETag"]}
    for _ in range(2):
        engine._watch_tick()
        assert request(server, "GET", STATE, headers=conditional)[0] == 304


def test_the_state_is_built_once_per_revision_and_never_for_a_poll(api, monkeypatch):
    server, engine, _ = api
    headers = authenticate(server)
    built = []
    snapshot = engine.snapshot

    def counted(*args, **kwargs):
        built.append(1)
        return snapshot(*args, **kwargs)

    monkeypatch.setattr(engine, "snapshot", counted)
    first = request(server, "GET", STATE, headers=headers)
    second = request(server, "GET", STATE, headers=headers)
    poll = request(
        server, "GET", STATE, headers={**headers, "If-None-Match": first[1]["ETag"]}
    )
    assert (first[0], second[0], poll[0]) == (200, 200, 304)
    assert second[2] == first[2] and len(built) == 1
    with engine.lock:
        engine.paused = True
    request(server, "GET", STATE, headers=headers)
    assert len(built) == 2


def test_the_state_etag_changes_when_a_recent_workspace_disappears(api, tmp_path):
    server, engine, _ = api
    headers = authenticate(server)
    gone = tmp_path / "gone"
    gone.mkdir()
    with engine.lock:
        engine.recent_workspaces.append(str(gone))
    etag = request(server, "GET", STATE, headers=headers)[1]["ETag"]
    gone.rmdir()
    status, _, data = request(
        server, "GET", STATE, headers={**headers, "If-None-Match": etag}
    )
    assert status == 200
    assert {"path": str(gone), "exists": False} in json.loads(data)["recent_workspaces"]


def test_the_state_keeps_whether_a_recent_workspace_exists_as_its_version_has_it(
    api, tmp_path, monkeypatch
):
    server, engine, _ = api
    headers = authenticate(server)
    mount = tmp_path / "mount"
    mount.mkdir()
    with engine.lock:
        engine.recent_workspaces.append(str(mount))
    snapshot = engine.snapshot

    def unmounted_while_built(*args, **kwargs):
        # The mount goes away while the state is built, and is back before the next poll.
        mount.rmdir()
        try:
            return snapshot(*args, **kwargs)
        finally:
            mount.mkdir()

    monkeypatch.setattr(engine, "snapshot", unmounted_while_built)
    status, response_headers, data = request(server, "GET", STATE, headers=headers)
    assert status == 200
    monkeypatch.setattr(engine, "snapshot", snapshot)
    poll = request(
        server,
        "GET",
        STATE,
        headers={**headers, "If-None-Match": response_headers["ETag"]},
    )
    kept = json.loads(data if poll[0] == 304 else poll[2])
    assert {"path": str(mount), "exists": True} in kept["recent_workspaces"]


def test_a_change_while_the_state_is_built_is_served_by_the_next_poll(api, monkeypatch):
    server, engine, _ = api
    headers = authenticate(server)
    snapshot = engine.snapshot

    def changed_after_the_read(*args, **kwargs):
        state = snapshot(*args, **kwargs)
        with engine.lock:
            engine.theme = "dark"
        return state

    monkeypatch.setattr(engine, "snapshot", changed_after_the_read)
    status, response_headers, data = request(server, "GET", STATE, headers=headers)
    assert status == 200 and json.loads(data)["theme"] != "dark"
    monkeypatch.setattr(engine, "snapshot", snapshot)
    status, _, data = request(
        server,
        "GET",
        STATE,
        headers={**headers, "If-None-Match": response_headers["ETag"]},
    )
    assert status == 200 and json.loads(data)["theme"] == "dark"


def test_every_answer_of_the_state_holds_the_state_of_its_version(api):
    server, engine, _ = api
    headers = authenticate(server)
    stop = threading.Event()

    def change():
        # Writing sections that change a field of the state, as the watcher and jobs do.
        dark = False
        while not stop.is_set():
            dark = not dark
            with engine.lock:
                engine.theme = "dark" if dark else "light"
            time.sleep(0.002)

    writer = threading.Thread(target=change)
    writer.start()
    kept, etag, answers = None, None, Counter()
    try:
        for _ in range(300):
            condition = {"If-None-Match": etag} if etag else {}
            status, response_headers, data = request(
                server, "GET", STATE, headers={**headers, **condition}
            )
            if status == 200:
                kept, etag = json.loads(data), response_headers["ETag"]
            assert status in {200, 304} and response_headers["ETag"] == etag
            # The state as it is now, if the version is still the one of the answer.
            with engine.lock.reading():
                if f'"{engine.state_version().tag}"' == etag:
                    current = {**engine.snapshot(), "csrf_token": server.csrf_token}
                    assert kept == current
                    answers[status] += 1
    finally:
        stop.set()
        writer.join()
    assert answers[200] and answers[304]


def _detail(server, headers):
    return json.loads(request(server, "GET", DETAIL, headers=headers)[2])


def test_detail_and_preview_match_review_targets(api):
    server, engine, folder = api
    created = "2026-10-01T10:00:00Z"
    item = new_ulid()
    review = {
        "status": "draft",
        "items": [
            {
                "id": item,
                "kind": "question",
                "target": {
                    "file": "timecourses_Fig1.tsv",
                    "rows": {"label": "drug_plasma"},
                },
                "text": "Which dose?",
                "author": "curator",
                "created": created,
            },
            {
                "id": new_ulid(),
                "kind": "question",
                "text": "Is the study complete?",
                "author": "curator",
                "created": created,
            },
        ],
    }
    (folder / "review.json").write_text(
        json.dumps(review), encoding="utf-8", newline=""
    )
    # The TSV line of each time point of the formatted table.
    lines = (folder / "timecourses_Fig1.tsv").read_text(encoding="utf-8").splitlines()
    time = lines[0].split("\t").index("time")
    at = {
        row.split("\t")[time]: number for number, row in enumerate(lines[1:], start=1)
    }
    headers = authenticate(server)
    detail = _detail(server, headers)
    # Only items with a target file have a match.
    assert detail["targets"] == {
        item: {
            "lines": sorted(at.values()),
            "series": {"source": "Fig1", "series": "drug_plasma"},
            "total": len(at),
        }
    }
    preview = "/local/studies/review/preview"
    body = {
        "study": "caffeine/Example",
        "target": {"file": "timecourses_Fig1.tsv", "rows": {"time": "1"}},
    }
    status, _, data = request(server, "POST", preview, body, headers)
    assert status == 200
    assert json.loads(data) == {"lines": [at["1"]], "series": None, "total": len(at)}
    status, _, data = request(
        server, "POST", preview, {**body, "target": {"rows": {"time": "1"}}}, headers
    )
    assert status == 422
    assert json.loads(data)["issues"][0]["code"] == "invalid_review_json"
    # A preview writes nothing, needs no user and records no activity.
    engine.user = ""
    assert request(server, "POST", preview, body, headers)[0] == 200
    assert _detail(server, headers)["jobs"] == []
    assert json.loads((folder / "review.json").read_text(encoding="utf-8")) == review


def test_files_version_changes_with_the_files_only(api):
    server, engine, folder = api
    headers = authenticate(server)
    _, response_headers, data = request(server, "GET", DETAIL, headers=headers)
    version = json.loads(data)["files_version"]
    assert isinstance(version, str) and version
    # A write of the app lists an entry in the activity: the page changes, its files do not.
    engine._record_write("caffeine/Example", "Saved study.json")
    _, changed, data = request(server, "GET", DETAIL, headers=headers)
    assert changed["ETag"] != response_headers["ETag"]
    assert json.loads(data)["files_version"] == version
    timecourses = folder / "timecourses_Fig1.tsv"
    timecourses.write_text(
        timecourses.read_text(encoding="utf-8").replace("drug_plasma", "drug_urine"),
        encoding="utf-8",
        newline="",
    )
    engine.scan()
    assert _detail(server, headers)["files_version"] != version


def test_metadata_write_and_conflict(api):
    server, engine, folder = api
    headers = authenticate(server)
    detail = _detail(server, headers)
    metadata = {**detail["metadata"]["value"], "licence": "closed"}
    body = {
        "study": "caffeine/Example",
        "revision": detail["metadata"]["revision"],
        "metadata": metadata,
    }
    status, _, data = request(server, "POST", "/local/studies/metadata", body, headers)
    assert status == 200
    assert (
        json.loads((folder / "study.json").read_text(encoding="utf-8"))["licence"]
        == "closed"
    )
    written = json.loads(data)
    assert written == {
        "revision": written["revision"],
        "reference": None,
        "reference_error": None,
    }
    stale = request(server, "POST", "/local/studies/metadata", body, headers)
    assert stale[0] == 409
    conflict = json.loads(stale[2])
    assert conflict["file"] == "study.json"
    assert conflict["revision"] == written["revision"]
    assert json.loads(conflict["content"])["licence"] == "closed"
    invalid = {
        **body,
        "revision": written["revision"],
        "metadata": {**metadata, "licence": "maybe"},
    }
    rejected = request(server, "POST", "/local/studies/metadata", invalid, headers)
    assert rejected[0] == 422
    issues = json.loads(rejected[2])["issues"]
    assert issues and issues[0]["code"] == "invalid_study_json"
    assert (
        json.loads((folder / "study.json").read_text(encoding="utf-8"))["licence"]
        == "closed"
    )


def test_review_actions_and_approval_refusal(api):
    server, engine, folder = api
    headers = authenticate(server)
    revision = _detail(server, headers)["review"]["revision"]
    status, _, data = request(
        server,
        "POST",
        "/local/studies/review",
        {
            "study": "caffeine/Example",
            "revision": revision,
            "action": "add",
            "kind": "question",
            "text": "Why?",
            "target": {"file": "timecourses_Fig1.tsv", "column": "mean"},
        },
        headers,
    )
    assert status == 200
    item, revision = json.loads(data)["item"], json.loads(data)["revision"]
    assert item["author"] == "curator" and item["state"] == "open"
    assert item["target"] == {
        "file": "timecourses_Fig1.tsv",
        "rows": {},
        "column": "mean",
    }
    refused = request(
        server,
        "POST",
        "/local/studies/review",
        {
            "study": "caffeine/Example",
            "revision": revision,
            "action": "status",
            "status": "approved",
        },
        headers,
    )
    assert refused[0] == 422 and json.loads(refused[2])["code"] == "approval_refused"
    for action, extra in [
        ("reply", {"text": "Because."}),
        ("resolve", {}),
        ("reopen", {"text": "Not yet."}),
        ("dismiss", {"text": "Withdrawn."}),
    ]:
        done = request(
            server,
            "POST",
            "/local/studies/review",
            {
                "study": "caffeine/Example",
                "revision": revision,
                "action": action,
                "item": item["id"],
                **extra,
            },
            headers,
        )
        assert done[0] == 200, action
        assert set(json.loads(done[2])) == {"revision"}
        revision = json.loads(done[2])["revision"]
    [stored] = read_review(folder).review.items
    assert stored.state == "dismissed" and len(stored.thread) == 3
    assert stored.resolved_by == "curator"
    status, _, data = request(
        server,
        "POST",
        "/local/studies/review",
        {
            "study": "caffeine/Example",
            "revision": revision,
            "action": "status",
            "status": "in_review",
        },
        headers,
    )
    assert status == 200 and read_review(folder).review.status == "in_review"


def test_review_write_errors(api):
    server, engine, folder = api
    headers = authenticate(server)
    revision = _detail(server, headers)["review"]["revision"]
    review = {"study": "caffeine/Example", "revision": revision}

    def post(body):
        status, _, data = request(
            server, "POST", "/local/studies/review", {**review, **body}, headers
        )
        return status, json.loads(data)

    status, data = post({"action": "add", "kind": "praise", "text": "Nice."})
    assert status == 422 and data["issues"][0]["code"] == "invalid_review_json"
    status, data = post(
        {"action": "add", "kind": "issue", "text": "?", "target": {"column": "mean"}}
    )
    assert status == 422 and data["issues"]
    status, data = post({"action": "resolve", "item": new_ulid()})
    assert status == 422 and "does not exist" in data["error"]
    assert post({"action": "close", "item": new_ulid()})[0] == 400
    status, data = post({"action": "add", "kind": "issue", "text": ""})
    assert status == 422 and data["issues"][0]["field"] == "text"
    assert post({"action": "add", "kind": "issue"})[0] == 400
    assert (
        post({"action": "add", "kind": "issue", "text": "?", "revision": None})[0]
        == 400
    )
    stale = post({"action": "add", "kind": "issue", "text": "?", "revision": "0" * 64})
    assert stale[0] == 409 and stale[1]["file"] == "review.json"
    assert json.loads(stale[1]["content"]) == {"status": "draft"}
    assert read_review(folder).review.items == []


@pytest.mark.parametrize(
    ("key", "change", "action"),
    [
        (
            "review",
            {"status": "in_review"},
            {"action": "add", "kind": "issue", "text": "Check the dose."},
        ),
        ("metadata", {"licence": "closed"}, {}),
    ],
)
def test_a_stale_write_shows_the_file_on_disk_at_once(api, key, change, action):
    """The study page loads the study after a 409 and gets the file on disk, not a 304."""
    server, engine, folder = api
    headers = authenticate(server)
    status, response_headers, data = request(server, "GET", DETAIL, headers=headers)
    document = json.loads(data)[key]
    on_disk = {**document["value"], **change}
    name = "review.json" if key == "review" else "study.json"
    (folder / name).write_text(
        json.dumps(on_disk, indent=2) + "\n", encoding="utf-8", newline=""
    )
    body = action or {"metadata": document["value"]}
    route = f"/local/studies/{key}"
    write = {"study": "caffeine/Example", "revision": document["revision"], **body}
    stale = request(server, "POST", route, write, headers)
    assert stale[0] == 409
    etag = {**headers, "If-None-Match": response_headers["ETag"]}
    status, _, data = request(server, "GET", DETAIL, headers=etag)
    assert status == 200
    reloaded = json.loads(data)[key]
    assert reloaded["revision"] == json.loads(stale[2])["revision"]
    assert reloaded["value"] == on_disk


def test_writes_need_a_user(api):
    server, engine, folder = api
    engine.user = ""
    headers = authenticate(server)
    response = request(
        server,
        "POST",
        "/local/studies/review",
        {
            "study": "caffeine/Example",
            "revision": None,
            "action": "add",
            "kind": "question",
            "text": "?",
        },
        headers,
    )
    assert response[0] == 403 and json.loads(response[2])["error"] == "no_user"
    assert "Connection settings" in json.loads(response[2])["message"]
    detail = _detail(server, headers)
    metadata = request(
        server,
        "POST",
        "/local/studies/metadata",
        {
            "study": "caffeine/Example",
            "revision": detail["metadata"]["revision"],
            "metadata": detail["metadata"]["value"],
        },
        headers,
    )
    assert metadata[0] == 403 and json.loads(metadata[2])["error"] == "no_user"
    for body in [
        {"action": "open"},
        {"action": "sync"},
        {"action": "resolve", "keep": "workbook"},
        {"action": "add", "kind": "raw", "source": "Tab3"},
    ]:
        tables = request(
            server,
            "POST",
            "/local/studies/tables",
            {"study": "caffeine/Example", **body},
            headers,
        )
        assert tables[0] == 403, body["action"]
        assert json.loads(tables[2])["error"] == "no_user"
    assert not workbook_path(folder).exists()


def test_the_author_is_the_account_of_a_checked_key(api):
    server, engine, folder = api
    assert engine.author() == Author("curator")
    assert engine.author("claude-opus-5-5") == Author("curator", "claude-opus-5-5")
    assert engine.snapshot()["author"] == {"user": "curator", "reason": None}
    engine.offline, engine.endpoint = False, "https://pk-db.test"
    engine.api_key, engine.account, engine.checked_at = "key", "account", "now"
    assert engine.author() == Author("account")
    engine.connection_error = "The server rejected the API key."
    assert engine.author() == Author("curator")
    # The last check found that the key belongs to another account than the user.
    engine.connection_error = engine.user_mismatch = (
        "The API key belongs to PK-DB user 'other', not the expected user 'curator'"
    )
    with pytest.raises(IdentityError, match="belongs to PK-DB user 'other'"):
        engine.author()
    assert engine.snapshot()["author"] == {
        "user": None,
        "reason": engine.user_mismatch,
    }
    engine.connection_error = engine.user_mismatch = None
    engine.api_key = None
    assert engine.author() == Author("curator")
    engine.user = ""
    with pytest.raises(IdentityError, match="PK-DB user name is required"):
        engine.author()
    assert engine.snapshot()["author"]["user"] is None
    assert "Connection settings" in engine.snapshot()["author"]["reason"]
    engine.user = "two words"
    with pytest.raises(IdentityError, match="not a PK-DB user name"):
        engine.author()


def test_writes_with_a_key_of_another_account_are_refused(api):
    server, engine, folder = api
    engine.offline, engine.endpoint, engine.api_key = (
        False,
        "https://pk-db.test",
        "pk-secret-7",
    )
    engine.user_mismatch = "The API key belongs to PK-DB user 'other'"
    headers = authenticate(server)
    review = _detail(server, headers)["review"]
    response = request(
        server,
        "POST",
        "/local/studies/review",
        {
            "study": "caffeine/Example",
            "revision": review["revision"],
            "action": "add",
            "kind": "question",
            "text": "?",
        },
        headers,
    )
    assert response[0] == 403
    assert error_body(response[2]) == {
        "error": "user_mismatch",
        "message": "The API key belongs to PK-DB user 'other'",
        "code": "user_mismatch",
        "detail": None,
        "action": "Add a review item (POST /local/studies/review)",
    }
    assert read_review(folder).review.items == []


def test_write_after_the_job_formatted_the_file_conflicts(api):
    server, engine, folder = api
    headers = authenticate(server)
    revision = _detail(server, headers)["metadata"]["revision"]
    study = json.loads((folder / "study.json").read_text(encoding="utf-8"))
    # An external edit, which the watcher job then formats.
    (folder / "study.json").write_text(
        dump_json({**study, "descriptions": ["Edited outside the app."]}),
        encoding="utf-8",
        newline="",
    )
    format_folder(folder)
    response = request(
        server,
        "POST",
        "/local/studies/metadata",
        {
            "study": "caffeine/Example",
            "revision": revision,
            "metadata": METADATA,
        },
        headers,
    )
    assert response[0] == 409
    assert json.loads((folder / "study.json").read_text(encoding="utf-8"))[
        "descriptions"
    ] == ["Edited outside the app."]


def test_a_write_changes_the_detail_etag(api):
    server, engine, folder = api
    headers = authenticate(server)
    status, response_headers, data = request(server, "GET", DETAIL, headers=headers)
    etag = response_headers["ETag"]
    response = request(
        server,
        "POST",
        "/local/studies/review",
        {
            "study": "caffeine/Example",
            "revision": json.loads(data)["review"]["revision"],
            "action": "add",
            "kind": "question",
            "text": "Which dose?",
        },
        headers,
    )
    assert response[0] == 200
    status, changed, data = request(
        server, "GET", DETAIL, headers={**headers, "If-None-Match": etag}
    )
    assert status == 200 and changed["ETag"] != etag
    assert json.loads(data)["summary"]["open_items"] == 1


def test_acknowledge_one_warning(api, sf_vocabulary, monkeypatch):
    server, engine, folder = api
    monkeypatch.setattr(engine, "_local_vocabulary", lambda folder=None: sf_vocabulary)
    timecourses = folder / "timecourses_Fig1.tsv"
    lines = timecourses.read_text(encoding="utf-8").splitlines()
    header = lines[0].split("\t")
    for index in (2, 3):  # the means at times 1 and 2 lie outside their range
        row = lines[index].split("\t")
        row[header.index("min")], row[header.index("max")] = "3", "4"
        lines[index] = "\t".join(row)
    timecourses.write_text("\n".join(lines) + "\n", encoding="utf-8", newline="")
    headers = authenticate(server)
    body = {
        "study": "caffeine/Example",
        "revision": _detail(server, headers)["review"]["revision"],
        "action": "acknowledge",
        "code": "outside_range",
        "file": "timecourses_Fig1.tsv",
        "text": "As printed.",
    }
    status, _, data = request(server, "POST", "/local/studies/review", body, headers)
    assert status == 422
    # A line and a column that are left out match every line and column.
    assert error_body(data) == {
        "error": "2 warnings [outside_range] match in timecourses_Fig1.tsv at line 2 "
        "column mean, line 3 column mean; give the line, column and key of one",
        "issues": [],
        "code": "ambiguous_warning",
        "detail": None,
        "action": "Acknowledge a warning (POST /local/studies/review)",
    }
    missing = {**body, "code": "missing_image"}
    status, _, data = request(server, "POST", "/local/studies/review", missing, headers)
    assert status == 422
    assert error_body(data) == {
        "error": "No warning [missing_image] in timecourses_Fig1.tsv matches",
        "issues": [],
        "code": "no_such_warning",
        "detail": None,
        "action": "Acknowledge a warning (POST /local/studies/review)",
    }
    # A null line or column matches only warnings without one.
    for exact in ({"line": None}, {"column": None}, {"line": 3, "column": None}):
        status, _, data = request(
            server, "POST", "/local/studies/review", {**body, **exact}, headers
        )
        assert status == 422 and json.loads(data)["code"] == "no_such_warning"
    assert read_review(folder).review.items == []
    status, _, data = request(
        server,
        "POST",
        "/local/studies/review",
        {**body, "line": 3, "column": "mean"},
        headers,
    )
    assert status == 200
    item = json.loads(data)["item"]
    assert item["acknowledges"] == "outside_range" and item["state"] == "resolved"
    assert item["target"]["rows"] == {"label": "drug_plasma", "time": "2"}
    assert json.loads(data)["revision"] == read_review(folder).revision


def test_acknowledge_one_dataset_by_its_key(api, sf_vocabulary, monkeypatch):
    server, engine, folder = api
    monkeypatch.setattr(engine, "_local_vocabulary", lambda folder=None: sf_vocabulary)
    (folder / "Example_Fig1.wpd.json").write_text(
        json.dumps(project(GOOD, extra=("legend", "axis labels"))),
        encoding="utf-8",
        newline="",
    )
    assert format_folder(folder).ok
    headers = authenticate(server)
    body = {
        "study": "caffeine/Example",
        "revision": _detail(server, headers)["review"]["revision"],
        "action": "acknowledge",
        "code": "unknown_dataset",
        "file": "Example_Fig1.wpd.json",
        "line": None,
        "column": None,
        "text": "Not data.",
    }
    status, _, data = request(server, "POST", "/local/studies/review", body, headers)
    assert status == 422
    assert error_body(data) == {
        "error": "2 warnings [unknown_dataset] match in Example_Fig1.wpd.json at key "
        "axis labels, key legend; give the line, column and key of one",
        "issues": [],
        "code": "ambiguous_warning",
        "detail": None,
        "action": "Acknowledge a warning (POST /local/studies/review)",
    }
    # A null key matches only warnings without one.
    status, _, data = request(
        server, "POST", "/local/studies/review", {**body, "key": None}, headers
    )
    assert status == 422 and json.loads(data)["code"] == "no_such_warning"
    status, _, data = request(
        server, "POST", "/local/studies/review", {**body, "key": "legend"}, headers
    )
    assert status == 200
    target = json.loads(data)["item"]["target"]
    assert (target["file"], target["key"], "column" in target) == (
        "Example_Fig1.wpd.json",
        "legend",
        False,
    )
    [entry] = _detail(server, headers)["acknowledged"]
    assert entry["scope"] == "key"
    assert entry["target"] == {
        "file": "Example_Fig1.wpd.json",
        "rows": {},
        "key": "legend",
    }


def test_a_warning_without_a_row_or_a_key_is_refused_plainly(
    api, sf_vocabulary, monkeypatch
):
    server, engine, folder = api
    monkeypatch.setattr(engine, "_local_vocabulary", lambda folder=None: sf_vocabulary)
    unkeyed = make_issue("unknown_dataset", "Old.", file="Example_Fig1.wpd.json")
    monkeypatch.setattr(
        studies,
        "validate_folder",
        lambda folder, vocabulary: SimpleNamespace(issues=[unkeyed]),
    )
    headers = authenticate(server)
    body = {
        "study": "caffeine/Example",
        "revision": _detail(server, headers)["review"]["revision"],
        "action": "acknowledge",
        "code": "unknown_dataset",
        "file": "Example_Fig1.wpd.json",
        "text": "Not data.",
    }
    status, _, data = request(server, "POST", "/local/studies/review", body, headers)
    assert status == 422
    answer = json.loads(data)
    assert answer["code"] == "no_exact_target"
    assert answer["error"].startswith(
        "The warning [unknown_dataset] in Example_Fig1.wpd.json has no row of a data "
        "table and no key, so it cannot be acknowledged alone."
    )
    assert read_review(folder).review.items == []


def test_tables_open_uses_the_opener(api, monkeypatch):
    server, engine, folder = api
    opened = []
    monkeypatch.setattr(
        "pkdb.curation.studies.open_path", lambda path, **kw: opened.append(path)
    )
    headers = authenticate(server)
    status, _, data = request(
        server,
        "POST",
        "/local/studies/tables",
        {"study": "caffeine/Example", "action": "open"},
        headers,
    )
    assert status == 200 and opened == [folder / "Example.xlsx"]
    result = json.loads(data)
    assert result["ok"] and result["workbook_action"] == "created"
    assert result["changes"] == [] and result["conflicts"] == []


def _open(server, headers, **place):
    status, _, data = request(
        server,
        "POST",
        "/local/studies/tables",
        {"study": "caffeine/Example", "action": "open", **place},
        headers,
    )
    return status, json.loads(data)


@pytest.fixture
def opened(api, monkeypatch):
    """The api with a recording opener, and the workbook that a first Open tables created."""
    server, engine, folder = api
    paths = []
    monkeypatch.setattr(
        "pkdb.curation.studies.open_path", lambda path, **kw: paths.append(path)
    )
    headers = authenticate(server)
    status, result = _open(server, headers)
    assert status == 200 and result["opened"] and result["notice"] is None
    return server, engine, folder, headers, paths


def test_open_tables_at_a_cell_opens_the_workbook_at_its_sheet(opened):
    server, engine, folder, headers, paths = opened

    status, result = _open(server, headers, file="timecourses_Fig1.tsv", cell="O6")

    assert status == 200 and result["opened"] and result["notice"] is None
    assert paths == [workbook_path(folder)] * 2
    sheet = openpyxl.load_workbook(workbook_path(folder)).active
    assert sheet.title == "timecourses_Fig1"
    assert sheet.sheet_view.selection[0].activeCell == "O6"


def test_open_tables_at_a_raw_table_opens_its_sheet(opened):
    server, engine, folder, headers, paths = opened

    status, result = _open(server, headers, file="Example_Tab2.tsv")

    assert status == 200 and result["notice"] is None
    assert openpyxl.load_workbook(workbook_path(folder)).active.title == "Example_Tab2"


def test_open_tables_without_a_place_keeps_the_last_sheet(opened):
    server, engine, folder, headers, paths = opened
    assert _open(server, headers, file="outputs_Tab2.tsv")[0] == 200
    before = workbook_path(folder).read_bytes()

    status, result = _open(server, headers)

    assert status == 200 and result["notice"] is None
    assert workbook_path(folder).read_bytes() == before


def test_open_tables_at_a_cell_of_an_open_workbook_names_the_place(opened):
    server, engine, folder, headers, paths = opened
    workbook = workbook_path(folder)
    (folder / f".~lock.{workbook.name}#").write_text("curator", encoding="utf-8")
    before = workbook.read_bytes()

    status, result = _open(server, headers, file="timecourses_Fig1.tsv", cell="O6")

    assert status == 200 and result["opened"]
    assert result["notice"] == (
        "The workbook is already open. Go to sheet timecourses_Fig1, cell O6."
    )
    assert workbook.read_bytes() == before
    # It opens as it is: the program brings its window to the front.
    assert paths == [workbook] * 2
    status, result = _open(server, headers, file="outputs_Tab2.tsv")
    assert result["notice"] == "The workbook is already open. Go to sheet outputs_Tab2."


def test_open_tables_at_a_table_without_a_sheet_says_so(opened):
    server, engine, folder, headers, paths = opened
    before = workbook_path(folder).read_bytes()

    status, result = _open(server, headers, file="outputs_Tab9.tsv", cell="A2")

    assert status == 200 and result["opened"]
    assert result["notice"] == "The workbook has no sheet outputs_Tab9."
    assert workbook_path(folder).read_bytes() == before


def test_open_tables_at_a_hidden_sheet_says_to_show_it(opened):
    server, engine, folder, headers, paths = opened
    book = openpyxl.load_workbook(workbook_path(folder))
    book["timecourses_Fig1"].sheet_state = "hidden"
    book.save(workbook_path(folder))

    status, result = _open(server, headers, file="timecourses_Fig1.tsv", cell="O6")

    assert status == 200 and result["opened"]
    assert result["notice"] == (
        "Sheet timecourses_Fig1 is hidden in the workbook. Show it, then go to cell O6."
    )
    status, result = _open(server, headers, file="timecourses_Fig1.tsv")
    assert (
        result["notice"] == "Sheet timecourses_Fig1 is hidden in the workbook. Show it."
    )


@pytest.mark.parametrize(
    ("error", "reason"),
    [
        (
            lambda path: OSError(errno.ENOSPC, os.strerror(errno.ENOSPC), str(path)),
            os.strerror(errno.ENOSPC),
        ),
        (lambda path: OSError(f"Could not write {path}"), "OSError"),
    ],
)
def test_open_tables_names_a_failed_write_without_its_path(
    opened, monkeypatch, error, reason
):
    server, engine, folder, headers, paths = opened

    def failing(path, data):
        raise error(path)

    monkeypatch.setattr("pkdb.studyformat.workbook.view.atomic_bytes", failing)

    status, result = _open(server, headers, file="timecourses_Fig1.tsv", cell="O6")

    assert status == 200 and result["opened"]
    assert result["notice"] == (
        f"Could not select sheet timecourses_Fig1: {reason}. "
        "Go to sheet timecourses_Fig1, cell O6."
    )
    assert str(folder) not in result["notice"]


def test_open_tables_at_a_cell_while_a_job_of_the_study_runs(opened):
    server, engine, folder, headers, paths = opened
    before = workbook_path(folder).read_bytes()
    engine.active = "caffeine/Example"

    status, result = _open(server, headers, file="timecourses_Fig1.tsv", cell="O6")

    assert status == 200 and result["opened"]
    assert result["notice"] == (
        "A job of the study is running. Go to sheet timecourses_Fig1, cell O6."
    )
    assert workbook_path(folder).read_bytes() == before
    assert paths == [workbook_path(folder)] * 2


@pytest.mark.parametrize(
    ("place", "detail"),
    [
        ({"file": "study.json"}, "study.json is no table file of Example"),
        (
            {"file": "../Other/outputs_Tab2.tsv"},
            "../Other/outputs_Tab2.tsv is no table file of Example",
        ),
        (
            {"file": "outputs_Tab2.tsv", "cell": "y97"},
            "'y97' is no workbook cell such as Y97",
        ),
        ({"file": "outputs_Tab2.tsv", "cell": 97}, "Expected text in cell"),
        ({"cell": "Y97"}, "A cell needs the file of its table"),
    ],
)
def test_open_tables_refuses_a_place_that_is_no_table_or_cell(opened, place, detail):
    server, engine, folder, headers, paths = opened
    before = workbook_path(folder).read_bytes()

    status, result = _open(server, headers, **place)

    assert status == 400 and result["code"] == "invalid_request"
    assert result["detail"] == f"ValueError: {detail}"
    assert workbook_path(folder).read_bytes() == before and len(paths) == 1


def test_opening_at_a_place_starts_no_job(opened):
    server, engine, folder, headers, paths = opened
    engine.scan()
    row = engine.studies["caffeine/Example"]
    # The workbook that the first Open tables created is a change; it counts as seen now.
    row["_pending"], row["stale"] = False, False
    seen = row["_fingerprint"]

    status, result = _open(server, headers, file="timecourses_Fig1.tsv", cell="O6")
    engine.scan()

    assert status == 200 and result["notice"] is None
    assert openpyxl.load_workbook(workbook_path(folder)).active.title == (
        "timecourses_Fig1"
    )
    # The workbook changed, but no cell: the watcher queues no validation or upload.
    assert row["_fingerprint"] != seen
    assert row["_pending"] is False and row["stale"] is False


def test_tables_sync_resolve_and_add(api, sf_vocabulary, monkeypatch):
    server, engine, folder = api
    monkeypatch.setattr(engine, "_local_vocabulary", lambda folder=None: sf_vocabulary)
    headers = authenticate(server)

    def tables(**body):
        status, _, data = request(
            server,
            "POST",
            "/local/studies/tables",
            {"study": "caffeine/Example", **body},
            headers,
        )
        return status, json.loads(data)

    status, result = tables(action="sync")
    assert status == 200 and result["workbook_action"] == "created"
    book = openpyxl.load_workbook(workbook_path(folder))
    sheet = book["timecourses_Fig1"]
    column = [cell.value for cell in sheet[1]].index("mean") + 1
    sheet.cell(3, column).value = 5
    book.save(workbook_path(folder))
    table = folder / "timecourses_Fig1.tsv"
    lines = table.read_text(encoding="utf-8").splitlines()
    cells = lines[2].split("\t")
    cells[lines[0].split("\t").index("mean")] = "7"
    lines[2] = "\t".join(cells)
    table.write_text("\n".join(lines) + "\n", encoding="utf-8", newline="")
    status, result = tables(action="sync")
    assert status == 200 and not result["ok"]
    assert [c["file"] for c in result["conflicts"]] == ["timecourses_Fig1.tsv"]
    assert tables(action="resolve")[0] == 400
    assert tables(action="resolve", keep="neither")[0] == 400
    status, result = tables(action="resolve", keep="workbook")
    assert status == 200 and result["ok"]
    assert result["conflicts"][0]["kept"] == "workbook"
    assert (
        table.read_text(encoding="utf-8")
        .splitlines()[2]
        .split("\t")[lines[0].split("\t").index("mean")]
        == "5"
    )
    status, result = tables(action="add", kind="raw", source="Tab3")
    assert status == 200 and result["ok"] and result["table"] == "Example_Tab3"
    assert "Example_Tab3" in openpyxl.load_workbook(workbook_path(folder)).sheetnames
    status, result = tables(action="add", kind="outputs", source="Tab3")
    assert status == 200 and result["ok"]
    status, result = tables(action="add", kind="outputs", source="nonsense")
    assert status == 200 and not result["ok"]
    assert result["issues"][0]["code"] == "invalid_table_name"
    assert tables(action="add")[0] == 400
    assert tables(action="add", kind="outputs")[0] == 400
    assert tables(action="close")[0] == 400


def test_table_preview_and_add_by_kind_and_source(api, sf_vocabulary, monkeypatch):
    server, engine, folder = api
    monkeypatch.setattr(engine, "_local_vocabulary", lambda folder=None: sf_vocabulary)
    headers = authenticate(server)
    preview = "/local/studies/tables/preview"
    body = {"study": "caffeine/Example", "kind": "outputs", "source": "Tab3"}
    status, _, data = request(server, "POST", preview, body, headers)
    assert status == 200
    assert json.loads(data) == {
        "table": "outputs_Tab3",
        "file": "outputs_Tab3.tsv",
        "image": "Example_Tab3.png",
        "image_found": False,
        "issues": [],
    }
    status, _, data = request(
        server, "POST", preview, {**body, "kind": "raw", "source": "Fig1"}, headers
    )
    assert [issue["message"] for issue in json.loads(data)["issues"]] == [
        "A raw table needs a paper table source such as Tab3"
    ]
    # A preview writes nothing and needs no user.
    engine.user = ""
    assert request(server, "POST", preview, body, headers)[0] == 200
    assert not (folder / "outputs_Tab3.tsv").exists()


@pytest.mark.parametrize(
    ("path", "body"),
    [
        ("/local/studies/metadata", {"revision": "absent", "metadata": {}}),
        ("/local/studies/review", {"revision": "absent", "action": "status"}),
        ("/local/studies/tables", {"action": "sync"}),
    ],
)
def test_write_routes_refuse_an_ambiguous_identity(api, valid_files, path, body):
    server, engine, folder = api
    copy = engine.root / "copies" / "caffeine" / "Example"
    copy.mkdir(parents=True)
    for file, content in valid_files.items():
        (copy / file).write_bytes(
            content if isinstance(content, bytes) else content.encode()
        )
    engine.scan()
    headers = authenticate(server)
    status, _, data = request(
        server, "POST", path, {"study": "caffeine/Example", **body}, headers
    )
    refused = json.loads(data)
    assert status == 409 and "identity of two folders" in refused["error"]
    assert len(refused["paths"]) == 2
    for route, existing in [
        ("/local/jobs", {"ids": ["caffeine/Example"], "action": "validate"}),
        ("/local/mode", {"ids": ["caffeine/Example"], "mode": "off"}),
        ("/local/files/open", {"study_id": "caffeine/Example"}),
    ]:
        assert request(server, "POST", route, existing, headers)[0] == 409, route


def test_write_routes_answer_404_for_an_unknown_study(api):
    server, engine, folder = api
    headers = authenticate(server)
    for path, body in [
        ("/local/studies/metadata", {"revision": "absent", "metadata": {}}),
        ("/local/studies/review", {"revision": "absent", "action": "status"}),
        ("/local/studies/tables", {"action": "sync"}),
    ]:
        response = request(
            server, "POST", path, {"study": "caffeine/Missing", **body}, headers
        )
        assert response[0] == 404, path


@pytest.mark.parametrize(
    ("file", "path", "body"),
    [
        ("study.json", "/local/studies/metadata", {"metadata": METADATA}),
        (
            "review.json",
            "/local/studies/review",
            {"action": "add", "kind": "question", "text": "?"},
        ),
    ],
)
def test_writes_refuse_a_symlinked_document(api, tmp_path_factory, file, path, body):
    server, engine, folder = api
    outside = tmp_path_factory.mktemp("outside") / file
    original = (folder / file).read_bytes()
    outside.write_bytes(original)
    (folder / file).unlink()
    (folder / file).symlink_to(outside)
    status, _, data = request(
        server,
        "POST",
        path,
        {"study": "caffeine/Example", "revision": "0" * 64, **body},
        authenticate(server),
    )
    assert status == 422
    refused = json.loads(data)
    assert "symlink" in {issue["code"] for issue in refused["issues"]}
    assert "content" not in refused
    assert (folder / file).is_symlink() and outside.read_bytes() == original


def test_curator_roster(api):
    server, engine, folder = api
    headers = authenticate(server)
    status, response_headers, data = request(
        server, "GET", "/local/curators", headers=headers
    )
    assert status == 200
    curators = {
        curator["username"]: curator for curator in json.loads(data)["curators"]
    }
    assert curators["mkoenig"]["display_name"] == "Matthias König"
    assert curators["mkoenig"]["avatar_url"] == "/avatars/matthias_koenig.webp"
    assert curators["Ahmed-fub"]["avatar_url"] is None
    assert all(
        {"username", "display_name", "avatar_url"} <= set(curator)
        for curator in curators.values()
    )
    again = request(
        server,
        "GET",
        "/local/curators",
        headers={**headers, "If-None-Match": response_headers["ETag"]},
    )
    assert again[0] == 304
    assert request(server, "GET", "/local/curators")[0] == 401


def test_detail_has_the_people_with_their_profiles(api):
    server, engine, folder = api
    study = json.loads((folder / "study.json").read_text(encoding="utf-8"))
    (folder / "study.json").write_text(
        dump_json(
            {
                **study,
                "creator": "mkoenig",
                "curators": [{"user": "mkoenig", "rating": 4.5}, {"user": "curator"}],
                "collaborators": ["Jane Doe"],
            }
        ),
        encoding="utf-8",
        newline="",
    )
    engine.scan()
    people = _detail(server, authenticate(server))["people"]
    assert people["creator"]["display_name"] == "Matthias König"
    assert people["creator"]["avatar_url"] == "/avatars/matthias_koenig.webp"
    assert [(c["user"], c["rating"]) for c in people["curators"]] == [
        ("mkoenig", 4.5),
        ("curator", 0),
    ]
    assert people["curators"][0]["profile"]["username"] == "mkoenig"
    assert people["curators"][1]["profile"]["display_name"] == "curator"
    assert [c["display_name"] for c in people["collaborators"]] == ["Jane Doe"]


def test_app_writes_are_listed_in_the_activity(api, sf_vocabulary, monkeypatch):
    server, engine, folder = api
    monkeypatch.setattr(engine, "_local_vocabulary", lambda folder=None: sf_vocabulary)
    monkeypatch.setattr("pkdb.curation.studies.open_path", lambda path, **kw: None)
    headers = authenticate(server)
    detail = _detail(server, headers)
    request(
        server,
        "POST",
        "/local/studies/metadata",
        {
            "study": "caffeine/Example",
            "revision": detail["metadata"]["revision"],
            "metadata": {**detail["metadata"]["value"], "licence": "closed"},
        },
        headers,
    )
    added = json.loads(
        request(
            server,
            "POST",
            "/local/studies/review",
            {
                "study": "caffeine/Example",
                "revision": detail["review"]["revision"],
                "action": "add",
                "kind": "question",
                "text": "Why?",
            },
            headers,
        )[2]
    )
    item = added["item"]["id"]
    request(
        server,
        "POST",
        "/local/studies/review",
        {
            "study": "caffeine/Example",
            "revision": added["revision"],
            "action": "resolve",
            "item": item,
        },
        headers,
    )
    for body in [
        {"action": "open"},
        {"action": "sync"},
        {"action": "add", "kind": "raw", "source": "Tab3"},
    ]:
        request(
            server,
            "POST",
            "/local/studies/tables",
            {"study": "caffeine/Example", **body},
            headers,
        )
    # Recording a write starts nothing; the watcher validates the changed files.
    assert engine.queue == {}
    writes = [
        job for job in _detail(server, headers)["jobs"] if job["action"] == "write"
    ]
    assert [job["message"] for job in writes] == [
        "Added raw table Example_Tab3",
        "Synced the workbook and the tables",
        "Opened the workbook",
        f"Resolved review item {item}",
        f"Added review item {item}",
        "Saved study.json",
    ]
    assert {job["status"] for job in writes} == {"succeeded"}
    assert all(
        job["study_id"] == "caffeine/Example"
        and job["study_name"] == "Example"
        and job["created_at"]
        for job in writes
    )
    saved = json.loads((engine.state_dir / "state.json").read_text(encoding="utf-8"))[
        "jobs"
    ]
    assert [job["message"] for job in saved if job["action"] == "write"][0] == (
        "Saved study.json"
    )
    by_message = {job["message"]: job for job in writes}
    assert by_message[f"Added review item {item}"]["item"] == item
    assert by_message[f"Added review item {item}"]["parts"] == [
        {"text": "Added "},
        {"item": item},
    ]
    assert by_message[f"Resolved review item {item}"]["parts"] == [
        {"text": "Resolved "},
        {"item": item},
    ]
    assert "item" not in by_message["Saved study.json"]
    assert "parts" not in by_message["Saved study.json"]


def test_detail_has_the_reference_and_the_state_of_the_row(api):
    server, engine, folder = api
    headers = authenticate(server)
    detail = _detail(server, headers)
    assert detail["reference"]["title"] == "Example study"
    assert detail["reference"]["pmid"] == "123"
    assert detail["reference_match"] is True
    assert detail["message"] is None and detail["last_upload"] is None
    reference = json.loads((folder / "reference.json").read_text(encoding="utf-8"))
    (folder / "reference.json").write_text(
        dump_json({**reference, "pmid": "456", "doi": "10.1000/ABC"}),
        encoding="utf-8",
        newline="",
    )
    engine.scan()
    assert _detail(server, headers)["reference_match"] is False
    study = json.loads((folder / "study.json").read_text(encoding="utf-8"))
    (folder / "study.json").write_text(
        dump_json({**study, "reference": {"doi": "10.1000/abc"}}),
        encoding="utf-8",
        newline="",
    )
    engine.scan()
    assert _detail(server, headers)["reference_match"] is True
    del study["reference"]
    (folder / "study.json").write_text(dump_json(study), encoding="utf-8", newline="")
    engine.scan()
    assert _detail(server, headers)["reference_match"] is None
    etag = request(server, "GET", DETAIL, headers=headers)[1]["ETag"]
    upload = {"persistence": "created", "at": "2026-10-07T10:00:00+00:00"}
    with engine.lock:
        row = engine.studies["caffeine/Example"]
        row["message"] = "Upload on save suspended: connect an authorized account"
        row["last_upload"] = upload
    status, _, data = request(
        server, "GET", DETAIL, headers={**headers, "If-None-Match": etag}
    )
    assert status == 200
    detail = json.loads(data)
    assert detail["message"].startswith("Upload on save suspended")
    assert detail["last_upload"] == upload


def test_tables_refuse_a_symlinked_workbook(api, tmp_path_factory, monkeypatch):
    server, engine, folder = api
    opened = []
    monkeypatch.setattr(
        "pkdb.curation.studies.open_path", lambda path, **kw: opened.append(path)
    )
    outside = tmp_path_factory.mktemp("outside") / "Example.xlsx"
    outside.write_bytes(b"not for the app")
    workbook_path(folder).symlink_to(outside)
    headers = authenticate(server)
    for body in [
        {"action": "open"},
        {"action": "sync"},
        {"action": "resolve", "keep": "tables"},
        {"action": "add", "kind": "raw", "source": "Tab3"},
    ]:
        status, _, data = request(
            server,
            "POST",
            "/local/studies/tables",
            {"study": "caffeine/Example", **body},
            headers,
        )
        assert status == 400, body
        assert "Example.xlsx is a symlink" in json.loads(data)["error"]
    assert opened == [] and outside.read_bytes() == b"not for the app"


def test_a_write_succeeds_when_the_rescan_fails(api, monkeypatch):
    server, engine, folder = api

    def broken():
        raise RuntimeError("unexpected")

    monkeypatch.setattr(engine, "scan", broken)
    headers = authenticate(server)
    detail = _detail(server, headers)
    status, _, _ = request(
        server,
        "POST",
        "/local/studies/metadata",
        {
            "study": "caffeine/Example",
            "revision": detail["metadata"]["revision"],
            "metadata": {**detail["metadata"]["value"], "licence": "closed"},
        },
        headers,
    )
    assert status == 200
    assert (
        json.loads((folder / "study.json").read_text(encoding="utf-8"))["licence"]
        == "closed"
    )


def test_an_unexpected_error_is_logged_with_its_traceback_and_named_in_the_terminal(
    api, monkeypatch, capsys
):
    server, engine, folder = api
    engine.api_key = "private-api-key"
    headers = authenticate(server)

    def broken(identity, payload):
        raise RuntimeError(f"broken with private-api-key and {server.csrf_token}")

    monkeypatch.setattr(engine, "table_preview", broken)
    status, _, data = request(
        server,
        "POST",
        "/local/studies/tables/preview",
        {"study": "caffeine/Example", "kind": "outputs", "source": "Tab3"},
        headers,
    )
    text = data.decode()
    body = json.loads(data)
    assert status == 500
    assert body["error"] == "Could not preview the new table of caffeine/Example."
    assert body["code"] == "unexpected"
    assert body["detail"] == "RuntimeError: broken with [redacted] and [redacted]"
    assert body["action"] == "Preview a table (POST /local/studies/tables/preview)"
    assert "Traceback (most recent call last)" in body["traceback"]
    assert "private-api-key" not in text and server.csrf_token not in text
    # A preview writes nothing: no entry in the activity.
    assert "activity_id" not in body
    log = engine.log.path.read_text(encoding="utf-8")
    assert f"Error {body['request_id']} in Preview a table" in log
    assert "Traceback (most recent call last)" in log
    assert "private-api-key" not in log and server.csrf_token not in log
    # Request lines stay unlogged.
    assert not re.search(r"HTTP/1\.[01]", log)
    assert capsys.readouterr().err == (
        f"Error {body['request_id']} in Preview a table: RuntimeError "
        f"(see {engine.log.path})\n"
    )


def test_a_failed_write_is_listed_in_the_activity_without_its_traceback(
    api, monkeypatch
):
    server, engine, folder = api
    headers = authenticate(server)

    def broken(identity, payload):
        raise RuntimeError("the workbook library broke")

    monkeypatch.setattr(engine, "tables_action", broken)
    status, _, data = request(
        server,
        "POST",
        "/local/studies/tables",
        {"study": "caffeine/Example", "action": "sync"},
        headers,
    )
    body = json.loads(data)
    assert status == 500
    assert body["error"] == "Could not sync the workbook of caffeine/Example."
    entry = next(
        job
        for job in _detail(server, headers)["jobs"]
        if job["id"] == body["activity_id"]
    )
    assert (entry["action"], entry["status"], entry["message"]) == (
        "write",
        "failed",
        "Could not sync the workbook of caffeine/Example.",
    )
    assert entry["error"] == {
        key: value
        for key, value in body.items()
        if key not in {"traceback", "activity_id"}
    }


def test_only_a_source_that_the_study_lacks_is_not_found(api, monkeypatch):
    server, engine, folder = api
    headers = authenticate(server)
    missing = request(server, "GET", f"{DETAIL}/sources/Fig9", headers=headers)
    assert (missing[0], error_body(missing[2])["code"]) == (404, "not_found")

    def broken(study, source):
        raise KeyError("series")

    monkeypatch.setattr(studies, "source_view", broken)
    status, _, data = request(server, "GET", f"{DETAIL}/sources/Fig1", headers=headers)
    body = json.loads(data)
    assert (status, body["code"]) == (500, "unexpected")
    assert "Traceback (most recent call last)" in body["traceback"]


def test_a_failing_open_command_is_listed_in_the_activity_with_its_exit_status_and_stderr(
    api, monkeypatch
):
    server, engine, folder = api
    monkeypatch.setattr("pkdb.curation.launch.EARLY_FAILURE_SECONDS", 30)
    monkeypatch.setenv("PKDB_OPEN_COMMAND", command_line(FAILING_OPENER))
    headers = authenticate(server)
    status, _, data = request(
        server,
        "POST",
        "/local/studies/tables",
        {"study": "caffeine/Example", "action": "open"},
        headers,
    )
    body = json.loads(data)
    assert status == 500
    assert (body["error"], body["code"]) == (
        "Could not open the tables of caffeine/Example.",
        "open_failed",
    )
    program = sys.executable.replace("\\", "/").rsplit("/", 1)[-1]
    assert body["detail"] == (
        f"CalledProcessError: {program} exited with status 4\n"
        "no method available for opening"
    )
    assert "traceback" not in body
    entry = next(
        job
        for job in _detail(server, headers)["jobs"]
        if job["id"] == body["activity_id"]
    )
    assert (entry["status"], entry["message"]) == (
        "failed",
        "Could not open the tables of caffeine/Example.",
    )
    assert entry["error"]["detail"] == body["detail"]
    assert "exited with status 4" in engine.log.path.read_text(encoding="utf-8")


def test_open_tables_answers_while_the_program_that_it_started_runs(
    api, monkeypatch, tmp_path
):
    server, engine, folder = api
    pid_file = tmp_path / "program.pid"
    monkeypatch.setenv("PKDB_OPEN_COMMAND", command_line(holding_opener(pid_file)))
    monkeypatch.setattr("pkdb.curation.launch.EARLY_FAILURE_SECONDS", 0.5)
    headers = authenticate(server)
    try:
        # request() gives up after 3 seconds; the program runs for 30.
        status, _, data = request(
            server,
            "POST",
            "/local/studies/tables",
            {"study": "caffeine/Example", "action": "open"},
            headers,
        )
    finally:
        wait_for_open_commands()
        stop_program(pid_file)
    assert status == 200
    assert json.loads(data)["opened"] is True


#: A PDF larger than one piece of the file route (64 kB), with an end to ask for by a range.
PDF = b"%PDF-1.7\n" + bytes(range(256)) * 400 + b"%%EOF\n"


def test_a_pdf_of_the_study_is_served_inline_whole_and_in_ranges(api):
    server, engine, folder = api
    (folder / "Example.pdf").write_bytes(PDF)
    engine.scan()
    headers = authenticate(server)
    url = f"{DETAIL}/files/Example.pdf"
    status, answer, body = request(server, "GET", url, headers=headers)
    assert (status, body) == (200, PDF)
    assert (
        answer["Content-Type"],
        answer["Content-Disposition"],
        answer["Accept-Ranges"],
    ) == ("application/pdf", 'inline; filename="Example.pdf"', "bytes")
    status, answer, body = request(
        server, "GET", url, headers={**headers, "Range": "bytes=-6"}
    )
    assert (status, body) == (206, b"%%EOF\n")
    assert answer["Content-Range"] == f"bytes {len(PDF) - 6}-{len(PDF) - 1}/{len(PDF)}"
    over = {**headers, "Range": f"bytes={len(PDF)}-"}
    assert request(server, "GET", url, headers=over)[0] == 416
    several = {**headers, "Range": "bytes=0-1,4-5"}
    status, answer, body = request(server, "GET", url, headers=several)
    assert (status, body) == (200, PDF)
    assert "Content-Range" not in answer


def test_a_pdf_named_beyond_ascii_is_served_with_both_names(api):
    server, engine, folder = api
    (folder / "Müller Supp.pdf").write_bytes(PDF)
    engine.scan()
    status, answer, body = request(
        server,
        "GET",
        f"{DETAIL}/files/{quote('Müller Supp.pdf')}",
        headers=authenticate(server),
    )
    assert (status, body) == (200, PDF)
    assert answer["Content-Disposition"] == (
        "inline; filename=\"M_ller Supp.pdf\"; filename*=UTF-8''M%C3%BCller%20Supp.pdf"
    )


def test_a_symlinked_pdf_and_files_that_are_no_image_or_pdf_are_refused(api, tmp_path):
    server, engine, folder = api
    outside = tmp_path / "outside.pdf"
    outside.write_bytes(PDF)
    (folder / "Linked.pdf").symlink_to(outside)
    (folder / "notes.txt").write_text("A note", encoding="utf-8")
    engine.scan()
    headers = authenticate(server)
    for file in ("Linked.pdf", "notes.txt", "subjects.tsv", "Missing.pdf"):
        assert (
            request(server, "GET", f"{DETAIL}/files/{file}", headers=headers)[0] == 404
        )


def test_a_named_pipe_named_like_a_pdf_is_refused_without_waiting_for_a_writer(api):
    if sys.platform == "win32":
        pytest.skip("Named pipes are POSIX only")
    server, engine, folder = api
    os.mkfifo(folder / "Pipe.pdf")
    engine.scan()
    # The request times out after 3 s if opening the pipe waits for a writer.
    status = request(
        server, "GET", f"{DETAIL}/files/Pipe.pdf", headers=authenticate(server)
    )[0]
    assert status == 404


def test_the_study_page_lists_its_pdfs_with_the_paper_first_in_any_case(api, tmp_path):
    server, engine, folder = api
    sizes = {
        "Example.pdf": 30,
        "Example_Supp.pdf": 20,
        "Zeta.PDF": 7,
        "appendix10.pdf": 10,
        "appendix9.Pdf": 5,
        ".hidden.pdf": 1,
        "notes.txt": 3,
    }
    for name, size in sizes.items():
        (folder / name).write_bytes(b"%" * size)
    outside = tmp_path / "outside.pdf"
    outside.write_bytes(b"%PDF")
    (folder / "Linked.pdf").symlink_to(outside)
    engine.scan()
    detail = json.loads(request(server, "GET", DETAIL, headers=authenticate(server))[2])
    assert detail["papers"] == [
        {
            "file": name,
            "url": f"{DETAIL}/files/{name}",
            "bytes": sizes[name],
            "version": ANY,
        }
        for name in (
            "Example.pdf",
            "appendix9.Pdf",
            "appendix10.pdf",
            "Example_Supp.pdf",
            "Zeta.PDF",
        )
    ]


def test_the_version_of_a_pdf_changes_when_it_is_written_also_at_the_same_size(api):
    server, engine, folder = api
    headers = authenticate(server)
    pdf = folder / "Example.pdf"
    pdf.write_bytes(b"%PDF-1.7\n%first\n")
    os.utime(pdf, ns=(1_700_000_000_000_000_000, 1_700_000_000_000_000_000))
    engine.scan()
    version = _detail(server, headers)["papers"][0]["version"]
    assert isinstance(version, str) and version
    engine.scan()
    assert _detail(server, headers)["papers"][0]["version"] == version
    # A PDF program saves another PDF of the same size: PDF.js must not mix their ranges.
    pdf.write_bytes(b"%PDF-1.7\n%other\n")
    os.utime(pdf, ns=(1_700_000_001_000_000_000, 1_700_000_001_000_000_000))
    engine.scan()
    paper = _detail(server, headers)["papers"][0]
    assert paper["bytes"] == len(b"%PDF-1.7\n%first\n")
    assert paper["version"] != version


def test_the_study_page_lists_its_pdfs_beyond_the_upload_limits(api, monkeypatch):
    server, engine, folder = api
    (folder / "Example.pdf").write_bytes(b"%PDF-1.7\n")
    engine.scan()
    monkeypatch.setattr(studies, "MAX_FILES", 2)
    detail = json.loads(request(server, "GET", DETAIL, headers=authenticate(server))[2])
    assert detail["problems"][0]["code"] == "file_limit"
    assert [paper["file"] for paper in detail["papers"]] == ["Example.pdf"]


def test_a_hidden_pdf_and_a_hidden_image_are_refused_by_the_file_route(api):
    server, engine, folder = api
    (folder / ".hidden.pdf").write_bytes(PDF)
    (folder / "._Example.pdf").write_bytes(b"\x00\x05\x16\x07")
    (folder / ".hidden.png").write_bytes(png(2, 2))
    engine.scan()
    headers = authenticate(server)
    for file in (".hidden.pdf", "._Example.pdf", ".hidden.png"):
        assert (
            request(server, "GET", f"{DETAIL}/files/{file}", headers=headers)[0] == 404
        )


@pytest.mark.skipif(
    sys.platform == "win32" or os.geteuid() == 0,
    reason="Windows ignores the file mode, and root reads any file",
)
def test_a_pdf_that_cannot_be_read_is_a_file_error_and_not_a_missing_file(api):
    server, engine, folder = api
    pdf = folder / "Example.pdf"
    pdf.write_bytes(PDF)
    engine.scan()
    pdf.chmod(0o000)
    try:
        status, _, data = request(
            server, "GET", f"{DETAIL}/files/Example.pdf", headers=authenticate(server)
        )
    finally:
        pdf.chmod(0o644)
    # The viewer says that the PDF cannot be read, with the reason of the server; a 404 would
    # say that it is gone.
    assert status == 400
    body = error_body(data)
    assert (body["error"], body["code"]) == (
        "Could not read Example.pdf of caffeine/Example.",
        "file_error",
    )
    assert body["detail"].startswith("PermissionError: [Errno 13] Permission denied")
    assert body["action"] == f"Read a file (GET {DETAIL}/files/Example.pdf)"


def test_a_pdf_that_vanished_since_the_check_is_no_longer_in_the_study_folder(
    api, monkeypatch
):
    server, engine, folder = api
    (folder / "Example.pdf").write_bytes(PDF)
    engine.scan()
    real_open = os.open

    def vanished(path, flags, *args, **kwargs):
        # Deleted after the scan listed it, before the file route opened it.
        if os.fspath(path).endswith("Example.pdf"):
            raise FileNotFoundError(2, "No such file or directory", os.fspath(path))
        return real_open(path, flags, *args, **kwargs)

    monkeypatch.setattr(studies.os, "open", vanished)
    status, _, data = request(
        server, "GET", f"{DETAIL}/files/Example.pdf", headers=authenticate(server)
    )
    assert status == 404
    body = error_body(data)
    assert (body["error"], body["code"], body["detail"]) == (
        "Example.pdf is no longer in the study folder",
        "not_found",
        None,
    )


def test_every_pdf_that_the_study_page_lists_is_served_by_the_file_route(api):
    server, engine, folder = api
    # four names that stay four files on the case-insensitive file systems of
    # macOS and Windows
    for name in ("Example.pdf", "Appendix.PDF", "Example_Supp.pdf", "Zeta.Pdf"):
        (folder / name).write_bytes(PDF)
    engine.scan()
    headers = authenticate(server)
    detail = json.loads(request(server, "GET", DETAIL, headers=headers)[2])
    assert len(detail["papers"]) == 4
    for paper in detail["papers"]:
        status, response_headers, _ = request(
            server, "GET", paper["url"], headers=headers
        )
        assert status == 200
        assert response_headers["Content-Type"] == "application/pdf"


def test_a_failure_of_open_tables_after_it_answered_goes_to_the_log(api, monkeypatch):
    server, engine, folder = api
    monkeypatch.setattr("pkdb.curation.launch.EARLY_FAILURE_SECONDS", 0.2)
    monkeypatch.setenv(
        "PKDB_OPEN_COMMAND", command_line(late_opener(3, 3, "lost the display"))
    )
    headers = authenticate(server)
    status, _, data = request(
        server,
        "POST",
        "/local/studies/tables",
        {"study": "caffeine/Example", "action": "open"},
        headers,
    )
    # The command still ran after 0.2 seconds: it counted as started.
    assert (status, json.loads(data)["opened"]) == (200, True)
    wait_for_open_commands()
    text = engine.log.path.read_text(encoding="utf-8")
    program = sys.executable.replace("\\", "/").rsplit("/", 1)[-1]
    assert (
        "in Open tables (POST /local/studies/tables): "
        f"CalledProcessError: {program} exited with status 3\nlost the display"
    ) in text
    assert "Could not open the tables of caffeine/Example." in text
