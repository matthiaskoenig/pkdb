"""Read tables with one-based physical coordinates, including blank/comment rows."""

import csv
from contextlib import contextmanager
from functools import partial
from pathlib import Path
from zipfile import BadZipFile

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet._reader import WorkSheetParser

from pkdb.importers.expressions import clean
from pkdb.schemas.source import SourceLocation
from pkdb.schemas.validation import StudyValidationError, fail


@contextmanager
def table_reader():
    """Reuse workbooks only within one immutable source bundle, then close them."""
    workbooks = {}
    try:
        yield partial(read_table, workbooks=workbooks)
    finally:
        for workbook in workbooks.values():
            workbook.close()


def _xlsx_rows(worksheet):
    """Read physical rows without materializing gaps up to Excel's row limit.

    This adapter uses openpyxl's read-only parser and value conversion, with
    the same date, formula, shared-string, and cell-padding settings as its
    ReadOnlyWorksheet iterator. Keep compatibility tests when upgrading it.
    """
    parent = worksheet.parent
    with worksheet._get_source() as stream:
        parser = WorkSheetParser(
            stream,
            worksheet._shared_strings,
            data_only=parent.data_only,
            epoch=parent.epoch,
            date_formats=parent._date_formats,
            timedelta_formats=parent._timedelta_formats,
        )
        header_pending = True
        for number, cells in parser.parse():
            if number < 2:
                continue
            if header_pending:
                header_pending = False
                if number != 2:
                    yield 2, ()
            yield number, worksheet._get_row(cells, values_only=True)
        if header_pending:
            yield 2, ()


def read_table(
    path: Path, sheet: str | None, max_rows: int, *, workbooks: dict | None = None
) -> list[tuple[dict, SourceLocation]]:
    if max_rows < 1:
        raise ValueError("max_rows must be positive")
    location = SourceLocation(file=path.name, sheet=sheet)
    workbook = None
    handle = None
    iterator = None
    try:
        if path.suffix.lower() == ".xlsx":
            if sheet is None:
                fail("missing_sheet", "Workbook sheet is required", location)
            workbook = workbooks.get(path) if workbooks is not None else None
            if workbook is None:
                workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
                if workbooks is not None:
                    workbooks[path] = workbook
            if sheet not in workbook.sheetnames:
                fail(
                    "missing_sheet",
                    "Referenced sheet is absent from the workbook",
                    location,
                    category="reference",
                    stage="read",
                    field="sheet",
                    actual=sheet,
                    expected={"available_sheets": workbook.sheetnames},
                )
            worksheet = workbook[sheet]
            iterator = _xlsx_rows(worksheet)
        else:
            handle = path.open(newline="", encoding="utf-8-sig")
            iterator = enumerate(csv.reader(handle, delimiter="\t"), 1)
        _, headers = next(iterator, (None, ()))
        columns = {}
        for index, header in enumerate(headers):
            if header is None or header == "":
                continue
            if not isinstance(header, str):
                fail("invalid_header", "Column headers must be strings", location)
            header = header.strip()
            if not header:
                continue
            if header in columns:
                fail("duplicate_column", "Column headers must be unique", location)
            columns[header] = index
        rows = []
        column_letters = {
            header: get_column_letter(index + 1) for header, index in columns.items()
        }
        for physical_row, raw in iterator:
            values = []
            commented = False
            for value in raw:
                if (
                    path.suffix.lower() == ".xlsx"
                    and isinstance(value, str)
                    and "#" in value
                ):
                    value = value.split("#", 1)[0]
                    commented = True
                    values.append(clean(value) or None)
                elif commented:
                    values.append(None)
                else:
                    values.append(clean(value))
            if not any(value is not None and value != "" for value in values):
                continue
            if len(rows) >= max_rows:
                fail(
                    "row_limit",
                    "Source table exceeds configured row limit",
                    location,
                    category="limit",
                    stage="read",
                    expected={"maximum_rows": max_rows},
                )
            source = SourceLocation(file=path.name, sheet=sheet, row=physical_row)
            source._columns.update(column_letters)
            rows.append(
                (
                    {
                        header: values[index] if index < len(values) else None
                        for header, index in columns.items()
                    },
                    source,
                )
            )
        return rows
    except StudyValidationError:
        raise
    except ValueError, KeyError, OSError, BadZipFile:
        fail(
            "invalid_table",
            "Cannot read the source table. Check its format and encoding and that the workbook is not corrupted.",
            location,
            category="parsing",
            stage="read",
        )
    finally:
        if iterator is not None and hasattr(iterator, "close"):
            iterator.close()
        if workbook is not None and workbooks is None:
            workbook.close()
        if handle is not None:
            handle.close()
