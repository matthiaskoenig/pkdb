"""Map OSP assessment records to source-qualified PK-DB study folders."""

import hashlib
import math
import re
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from urllib.parse import unquote

import openpyxl

from pkdb.cache import atomic_json, bundled_vocabulary
from pkdb.domain.units import ureg
from pkdb.domain.vocabulary import MeasurementRule, SubstanceDefinition
from pkdb.importers.osp.release import RELEASE, REVISION, SHA256, URL

SOURCE = "osp.observed-data"
VERSION = "1"


def number(value):
    if isinstance(value, bool):
        return None
    try:
        value = float(value)
        return value if math.isfinite(value) else None
    except TypeError, ValueError:
        return None


def text(value):
    return str(value).strip() if value is not None else ""


def unit(value):
    return (
        text(value)
        .replace("day(s)", "day")
        .replace("years", "yr")
        .replace("m2", "m^2")
        .replace("²", "^2")
        .replace("m3", "m^3")
        .replace("CM", "cm")
        .replace("%", "percent")
        .replace("1.73m^2", "(1.73*m^2)")
        .replace("µ", "u")
        or None
    )


def identity(reference, name):
    value = unquote(text(reference)).strip('"')
    pmid = re.search(
        r"(?:pubmed\.ncbi\.nlm\.nih\.gov/|ncbi\.nlm\.nih\.gov/pubmed/)(\d+)", value
    )
    if pmid:
        return "pmid:" + str(int(pmid[1])), {"pmid": str(int(pmid[1]))}
    doi = re.search(r"10\.\d{4,9}/[^\s?#]+", value, re.I)
    if doi:
        return "doi:" + doi[0].lower(), {"doi": doi[0].lower()}
    if value.startswith(("http://", "https://")):
        return "url:" + value, {"url": value}
    return "osp-reference:" + hashlib.sha256(
        (value or text(name)).encode()
    ).hexdigest(), {}


def json_value(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def rows(workbook, sheet, start):
    for index, row in enumerate(
        workbook[sheet].iter_rows(min_row=start, values_only=True), start
    ):
        if any(v is not None for v in row):
            yield index, [json_value(v) for v in row]


def statistics(
    avg,
    avg_type,
    var,
    var_unit,
    var_type,
    units,
    individual,
    warnings,
    location,
):
    value = number(avg)
    if value is None:
        warnings.append(
            {**location, "code": "non_numeric_or_missing_value", "value": avg}
        )
        return None
    kind = re.sub(r"[ .]", "", text(avg_type).lower())
    if kind == "indvidual":
        kind = "individual"
    if individual and kind != "individual":
        warnings.append(
            {
                **location,
                "code": "individual_summary_retained_in_source",
                "value": avg_type,
            }
        )
        return None
    geometric = kind in {"geomean", "geommean", "gmean"}
    if kind in {"individual"}:
        field = "value" if individual else None
    elif kind == "median":
        field = "median"
    elif geometric or kind in {
        "arithmean",
        "arithmmean",
        "arithmeticmean",
        "mean",
        "average",
        "avg",
    }:
        field = "mean"
    else:
        field = None
    if field is None:
        warnings.append(
            {**location, "code": "unsupported_statistic", "value": avg_type}
        )
        return None
    result = {field: value}
    if geometric:
        result["calculation_type"] = "geometric mean"
    variation = number(var)
    if var is not None and not individual:
        error_kind = re.sub(r"[ .]", "", text(var_type).lower())
        key = {
            "arithsd": "sd",
            "sd": "sd",
            "arithsem": "se",
            "arithse": "se",
            "sem": "se",
            "se": "se",
        }.get(error_kind)
        if variation is not None and variation >= 0 and key and not geometric:
            try:
                result[key] = float(
                    (variation * ureg(unit(var_unit) or units)).to(units).magnitude
                )
            except Exception:
                warnings.append(
                    {**location, "code": "uncertainty_unit_mismatch", "value": var_unit}
                )
        elif (
            variation is not None
            and variation >= 0
            and not geometric
            and "cv" in error_kind
            and "geo" not in error_kind
            and ("%" in error_kind or text(var_unit) == "%")
        ):
            result["cv"] = variation / 100
        else:
            warnings.append(
                {
                    **location,
                    "code": "uncertainty_retained_in_source",
                    "value": var_type,
                }
            )
    return result


def schedule(value):
    scalar = number(value)
    if scalar is not None:
        return [scalar]
    value = text(value).replace(" ", "").strip("()")
    match = re.fullmatch(r"S(\d+(?:\.\d+)?)-T(\d+(?:\.\d+)?)-R(\d+)", value)
    if match:
        start, tau, repeat = float(match[1]), float(match[2]), int(match[3])
        if 0 < repeat <= 1000 and tau > 0:
            return [start + tau * i for i in range(repeat)]
    if re.fullmatch(r"\d+(?:\.\d+)?(?:-\d+(?:\.\d+)?)+", value):
        return [float(item) for item in value.split("-")]
    return None


def import_workbook(path, output, *, creator, vocabulary=None, verify=True):
    path, output = Path(path), Path(output)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    if verify and digest != SHA256:
        raise ValueError("Workbook checksum does not match pinned OSP v1.9 release")
    if output.exists() and any(output.iterdir()):
        raise ValueError(
            "Import output directory must be empty; use a new directory to preserve previous runs"
        )
    output.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        source = {
            sheet: list(rows(workbook, sheet, 3 if sheet == "Studies" else 2))
            for sheet in (
                "Studies",
                "PK-Profiles",
                "PK-Parameter",
                "DDI",
                "Analyte",
                "Projects",
            )
        }
        headers = {
            sheet: [
                json_value(v)
                for v in next(
                    workbook[sheet].iter_rows(
                        min_row=2 if sheet == "Studies" else 1,
                        max_row=2 if sheet == "Studies" else 1,
                        values_only=True,
                    )
                )
            ]
            for sheet in source
        }
    finally:
        workbook.close()
    vocabulary = vocabulary or bundled_vocabulary()
    substances = {item.name.casefold(): item.name for item in vocabulary.substances}
    masses = {text(row[0]).casefold(): number(row[1]) for _, row in source["Analyte"]}
    synonyms = {
        "hydroxy-itraconazole": "hydroxyitraconazole",
        "keto-itraconazole": "ketoitraconazole",
        "r-metoprolol": "(R)-metoprolol",
        "s-metoprolol": "(S)-metoprolol",
        "s-mephenytoine": "s-mephenytoin",
    }
    additional = {}

    def substance(value):
        name = text(value)
        if not name or name == "#TODO#":
            return None
        key = name.casefold()
        key = synonyms.get(key, key).casefold()
        if key in substances:
            return substances[key]
        if key not in additional:
            sid = "osp-" + hashlib.sha256(key.encode()).hexdigest()[:16]
            additional[key] = SubstanceDefinition(
                name="OSP: " + key,
                sid=sid,
                mass=(masses.get(text(value).casefold()) or None),
            )
        return additional[key].name

    for _, row in source["Studies"]:
        substance(row[4])
    for _, row in source["DDI"]:
        substance(row[4])
        substance(row[5])
    measurements = list(vocabulary.measurements)
    known_measurements = vocabulary.measurement_map()
    if "cmax_relative" not in known_measurements:
        measurements.append(
            MeasurementRule(
                name="cmax_relative", sid="cmax_relative", units=("dimensionless",)
            )
        )
    vocabulary = vocabulary.model_copy(
        update={
            "substances": (*vocabulary.substances, *additional.values()),
            "measurements": tuple(measurements),
            "version": vocabulary.version + "+osp-v1.9",
        }
    )
    vocabulary.save(output / "vocabulary.json")
    rules = vocabulary.measurement_map()
    tissue_map = {t.casefold(): t for t in vocabulary.tissues}
    tissue_map.update(
        {
            "whole blood": "blood",
            "arterial whole blood": "arterial blood",
            "liver intracellular unbound": "liver",
        }
    )
    by_id, grouped = {}, {}
    totals, warnings = Counter(), []

    def get_study(reference, label):
        key, identifiers = identity(reference, label)
        if key not in grouped:
            sid = "OSP_OBS_" + hashlib.sha256(key.encode()).hexdigest()[:20]
            year = re.search(r"\b((?:19|20)\d{2})\b", text(label))
            ref = {
                "sid": sid + ":reference",
                "name": text(label) or sid,
                **identifiers,
                "publication_date": year[1] if year else None,
                "provenance": {"source_reference": reference, "source_label": label},
            }
            if text(reference).startswith(("http://", "https://")):
                ref["url"] = text(reference)
            study = {
                "sid": sid,
                "name": sid,
                "reference": ref["sid"],
                "creator": creator,
                "access": "private",
                "licence": "closed",
                "provenance": {
                    "kind": "data_import",
                    "source_key": SOURCE,
                    "release": RELEASE,
                    "revision": REVISION,
                    "importer": "pkdb.osp",
                    "importer_version": VERSION,
                    "assets": [{"url": URL, "sha256": digest}],
                    "dataset_ids": [],
                    "report_file": "osp-source.json",
                },
                "descriptions": [
                    f"Automatic import of OSP observed data {RELEASE}: {text(label)}"
                ],
                "groupset": {"groups": []},
                "individualset": {"individuals": []},
                "interventionset": {"interventions": []},
                "outputset": {"outputs": []},
            }
            grouped[key] = {
                "study": study,
                "reference": ref,
                "records": defaultdict(list),
                "warnings": [],
                "mapped": Counter(),
            }
        return grouped[key]

    def warn(target, location, code, value=None):
        target["warnings"].append({**location, "code": code, "value": value})

    for rownum, row in source["Studies"]:
        target = get_study(row[2], row[1])
        target["records"]["Studies"].append({"row": rownum, "values": row})
        if row[0] is None:
            warn(target, {"sheet": "Studies", "row": rownum}, "missing_assessment_id")
            continue
        identifier = text(row[0])
        if identifier in by_id:
            raise ValueError(f"Duplicate OSP assessment ID {identifier}")
        individual = text(row[6]).lower() in {"individual", "typical"}
        n = number(row[26])
        count = int(n) if n is not None and n >= 0 and n.is_integer() else None
        subject = {
            "name": "osp_" + identifier,
            "descriptions": [
                f"OSP assessment {identifier}: {text(row[3])}; type={text(row[6])}"
            ],
            "characteristica": [],
        }
        if not individual:
            subject["count"] = count
        if text(row[25]).casefold() == "human":
            subject["characteristica"].append(
                {"measurement_type": "species", "choice": "homo sapiens"}
            )
        # Demographics are mapped only when their units/statistic are explicit.
        for offset, measurement in [
            (28, "age"),
            (36, "weight"),
            (44, "height"),
            (52, "bmi"),
        ]:
            if row[offset] is None:
                continue
            units = unit(row[offset + 1])
            if units is None:
                warn(
                    target,
                    {"sheet": "Studies", "row": rownum},
                    "demographic_unit_missing",
                    measurement,
                )
                continue
            try:
                if not any(
                    ureg(units).dimensionality == ureg(expected).dimensionality
                    for expected in rules[measurement].units
                ):
                    raise ValueError("incompatible demographic unit")
            except Exception:
                warn(
                    target,
                    {"sheet": "Studies", "row": rownum},
                    "incompatible_demographic_unit",
                    [measurement, units],
                )
                continue
            stats = statistics(
                row[offset],
                row[offset + 2],
                row[offset + 3],
                row[offset + 4],
                row[offset + 5],
                units,
                individual,
                target["warnings"],
                {"sheet": "Studies", "row": rownum, "field": measurement},
            )
            if stats:
                subject["characteristica"].append(
                    {"measurement_type": measurement, "unit": units, **stats}
                )
        section = "individualset" if individual else "groupset"
        target["study"][section]["individuals" if individual else "groups"].append(
            subject
        )
        target["study"]["provenance"]["dataset_ids"].append(identifier)
        info = {
            "target": target,
            "row": row,
            "subject": "osp_" + identifier,
            "individual": individual,
            "interventions": [],
        }
        by_id[identifier] = info
        location = {"sheet": "Studies", "row": rownum, "dataset_id": identifier}
        dose, times, route = (
            number(row[10]) if number(row[10]) is not None else number(row[9]),
            schedule(row[13]),
            {"po": "oral", "iv": "iv", "intraduodenal": "intraduodenal"}.get(
                text(row[12]).lower()
            ),
        )
        ambiguous_analyte = any(
            x in text(row[4]).lower()
            for x in (
                "hydroxy",
                "norverapamil",
                "epoxide",
                "gluc",
                "metabolite",
                "desethyl",
                "keto-",
                "desalkyl",
                "total",
                "(e)",
                "r-",
                "s-",
                "coproporph",
                "cyp",
                "p-gp",
            )
        )
        if (
            dose is not None
            and dose >= 0
            and times
            and route
            and unit(row[11])
            and unit(row[14])
            and not ambiguous_analyte
            and substance(row[4])
        ):
            for i, at in enumerate(times):
                name = f"dose_{identifier}_{i}"
                intervention = {
                    "name": name,
                    "measurement_type": "dosing",
                    "substance": substance(row[4]),
                    "value": dose,
                    "unit": unit(row[11]),
                    "time": at,
                    "time_unit": unit(row[14]),
                    "route": route,
                    "form": "NR",
                    "application": "single dose",
                    "descriptions": [
                        f"OSP Studies row {rownum}; original dose (free API when reported) and schedule"
                    ],
                }
                duration = number(row[18])
                if duration and duration > 0:
                    intervention.update(
                        application="constant infusion",
                        time_end=at
                        + float((duration * ureg.min).to(unit(row[14])).magnitude),
                    )
                target["study"]["interventionset"]["interventions"].append(intervention)
                info["interventions"].append(name)
        else:
            warn(target, location, "regimen_retained_in_source")
    for sheet in ("PK-Profiles", "PK-Parameter", "DDI", "Projects"):
        for rownum, row in source[sheet]:
            identifier = text(row[0])
            info = by_id.get(identifier)
            target = info["target"] if info else get_study(row[2], row[1])
            target["records"][sheet].append({"row": rownum, "values": row})
            location = {"sheet": sheet, "row": rownum, "dataset_id": identifier}
            if not info:
                warn(target, location, "unknown_assessment_id")
                continue
            if sheet == "Projects":
                continue

            def output_record(
                avg,
                avg_type,
                variance,
                var_unit,
                var_type,
                units,
                measurement,
                analyte,
                *,
                time_value=None,
                time_unit=None,
                tissue=None,
                suffix="",
            ):
                if avg is None:
                    return
                substance_name = substance(analyte)
                if not units or not substance_name or measurement not in rules:
                    warn(
                        target,
                        location,
                        "unmapped_measurement_context",
                        [measurement, units, analyte],
                    )
                    return
                stats = statistics(
                    avg,
                    avg_type,
                    variance,
                    var_unit,
                    var_type,
                    units,
                    info["individual"],
                    target["warnings"],
                    location,
                )
                if stats is None:
                    return
                if (
                    any(
                        stats.get(field, 0) < 0 for field in ("mean", "value", "median")
                    )
                    and not rules[measurement].can_negative
                ):
                    warn(target, location, "negative_value_retained")
                    return
                # Reject incompatible unit mappings before writing a source study.
                try:
                    if not any(
                        ureg(units).dimensionality == ureg(expected).dimensionality
                        for expected in rules[measurement].units
                    ):
                        raise ValueError("incompatible units")
                except Exception:
                    warn(
                        target,
                        location,
                        "incompatible_measurement_unit",
                        [measurement, units],
                    )
                    return
                record = {
                    "measurement_type": measurement,
                    "substance": substance_name,
                    "unit": units,
                    **stats,
                    "individual" if info["individual"] else "group": info["subject"],
                    "interventions": info["interventions"],
                    "descriptions": [
                        f"OSP {sheet} row {rownum}, assessment {identifier}; original data in osp-source.json"
                    ],
                }
                if tissue:
                    record["tissue"] = tissue
                if time_value is not None:
                    record.update(time=time_value, time_unit=time_unit)
                elif rules[measurement].time_required:
                    record.update(time="NR", time_unit="NR")
                if (
                    sheet == "PK-Profiles"
                    and time_value is not None
                    and measurement in {"concentration", "concentration_unbound"}
                ):
                    record.update(
                        output_type="timecourse",
                        label=f"osp_{identifier}_{measurement}_{text(avg_type)}_{units}",
                    )
                target["study"]["outputset"]["outputs"].append(record)
                target["mapped"][sheet] += 1

            if sheet == "PK-Profiles":
                compartment = text(row[5]).casefold()
                units = unit(row[9])
                measurement = "concentration"
                tissue = tissue_map.get(compartment)
                if not tissue:
                    warn(target, location, "unmapped_compartment", row[5])
                    continue
                if units == "percent" and compartment in {"urine", "feces", "bile"}:
                    measurement = "recovery"
                elif units in {"mg", "ug", "g"}:
                    measurement = "amount"
                elif compartment == "liver intracellular unbound":
                    measurement = "concentration_unbound"
                if number(row[6]) is None or not unit(row[7]):
                    warn(target, location, "missing_profile_time")
                    continue
                output_record(
                    row[8],
                    row[10],
                    row[11],
                    row[12],
                    row[13],
                    units,
                    measurement,
                    row[4],
                    time_value=number(row[6]),
                    time_unit=unit(row[7]),
                    tissue=tissue,
                )
            elif sheet == "PK-Parameter":
                for offset, types in [
                    (
                        8,
                        {
                            "auc_inf": "auc_inf",
                            "auc_tend": "auc_end",
                            "auc_tau": "auc_end",
                        },
                    ),
                    (15, {"cmax": "cmax", "cmax_total": "cmax"}),
                    (
                        22,
                        {
                            "cl/f": "oral clearance",
                            "cliv": "clearance",
                            "cl total": "clearance",
                            "cliv_unbound": "clearance_unbound",
                        },
                    ),
                ]:
                    measurement = types.get(text(row[offset + 6]).lower())
                    if row[offset] is not None:
                        output_record(
                            row[offset],
                            row[offset + 2],
                            row[offset + 3],
                            row[offset + 4],
                            row[offset + 5],
                            unit(row[offset + 1]),
                            measurement,
                            row[4],
                            time_value=number(row[6])
                            if measurement == "auc_end"
                            else None,
                            time_unit=unit(row[7])
                            if measurement == "auc_end"
                            else None,
                            tissue=tissue_map.get(text(info["row"][5]).casefold()),
                        )
            else:
                for offset, measurement in [
                    (10, "auc_relative"),
                    (14, "cmax_relative"),
                ]:
                    avg_type = text(row[offset + 1])
                    norm = re.sub(r"[ .]", "", avg_type.lower())
                    if norm in {
                        "gmr",
                        "geommeanratio",
                        "geommean(gmr)",
                        "geomavg",
                        "geoavg",
                        "geomavg.",
                    }:
                        avg_type = "geom. mean"
                    output_record(
                        row[offset],
                        avg_type,
                        row[offset + 2],
                        None,
                        row[offset + 3],
                        "dimensionless",
                        measurement,
                        row[4],
                        tissue=tissue_map.get(text(row[9]).casefold()),
                    )
    # Substance discoveries in output sheets are included too.
    vocabulary = vocabulary.model_copy(
        update={
            "substances": tuple(
                s
                for s in vocabulary.substances
                if s.sid not in {a.sid for a in additional.values()}
            )
            + tuple(additional.values())
        }
    )
    vocabulary.save(output / "vocabulary.json")
    atomic_json(
        output / "source-metadata.json",
        {
            "release": RELEASE,
            "revision": REVISION,
            "sha256": digest,
            "headers": headers,
            "analytes": [{"row": n, "values": row} for n, row in source["Analyte"]],
        },
    )
    entries = []
    for key, target in sorted(grouped.items()):
        study = target["study"]
        series = defaultdict(list)
        for record in study["outputset"]["outputs"]:
            if record.get("label"):
                series[record["label"]].append(record["time"])
        for record in study["outputset"]["outputs"]:
            if record.get("label") and (
                len(set(series[record["label"]])) < 2
                or len(set(series[record["label"]])) != len(series[record["label"]])
            ):
                record.pop("label")
                record.pop("output_type")
        folder = output / "studies" / study["name"]
        folder.mkdir(parents=True)
        atomic_json(folder / "study.json", study)
        atomic_json(folder / "reference.json", target["reference"])
        atomic_json(
            folder / "osp-source.json",
            {
                "release": RELEASE,
                "revision": REVISION,
                "sha256": digest,
                "headers": headers,
                "records": target["records"],
                "warnings": target["warnings"],
                "mapped_measurements": dict(target["mapped"]),
            },
        )
        entries.append(
            {
                "sid": study["sid"],
                "publication_key": key,
                "datasets": len(study["provenance"]["dataset_ids"]),
                "mapped_measurements": dict(target["mapped"]),
                "source_rows": {k: len(v) for k, v in target["records"].items()},
                "warnings": len(target["warnings"]),
            }
        )
        totals.update(target["mapped"])
        warnings.extend(target["warnings"])
    report = {
        "source": SOURCE,
        "release": RELEASE,
        "revision": REVISION,
        "sha256": digest,
        "studies": entries,
        "study_count": len(entries),
        "source_rows": {k: len(v) for k, v in source.items()},
        "mapped_measurements": dict(totals),
        "warning_counts": dict(Counter(w["code"] for w in warnings)),
        "warnings": warnings,
        "additional_substances": [
            v.model_dump(mode="json") for v in additional.values()
        ],
    }
    atomic_json(output / "import-report.json", report)
    return report
