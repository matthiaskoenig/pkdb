"""Read the study workbook back into canonical TSV tables (spec 10.2, 10.3 and 10.6)."""

import re
import warnings
from collections.abc import Callable, Iterable, Iterator
from contextlib import ExitStack
from dataclasses import dataclass
from datetime import date, time, timedelta
from io import BytesIO
from pathlib import Path

from openpyxl.cell.text import Text
from openpyxl.reader.excel import ExcelReader
from openpyxl.styles.numbers import is_date_format
from openpyxl.workbook import Workbook
from openpyxl.xml.constants import SHARED_STRINGS, SHEET_MAIN_NS
from openpyxl.xml.functions import iterparse

from pkdb.schemas.validation import ValidationIssue
from pkdb.studyformat.formatter import subject_order, table_rows
from pkdb.studyformat.issues import LISTED, IssueCap, column_letter, make_issue
from pkdb.studyformat.load import STRUCTURAL, LoadedTable, RowLimit, load_table
from pkdb.studyformat.raw import load_raw, parse_raw_file, raw_lines, render_raw
from pkdb.studyformat.tables import TableSpec, parse_table_file
from pkdb.studyformat.text import format_number, render_tsv
from pkdb.studyformat.workbook.base import (
    BASE_SHEET,
    BaseError,
    WorkbookBase,
    parse_base,
)
from pkdb.studyformat.workbook.write import not_xml

SUBJECTS = "subjects"
# The error value of a date-formatted number that openpyxl cannot convert.
OUTSIDE_CALENDAR = "#VALUE!"
# The OOXML escape of a character, such as `_x0041_` for A and `_x005F_` for _.
ESCAPE = re.compile(r"_x([0-9A-Fa-f]{4})_")
SHARED_STRING = f"{{{SHEET_MAIN_NS}}}si"
LINE_BREAK = re.compile(r"[\t\n\r]")
LINE_BREAKS = {"\t": "tab", "\n": "line break", "\r": "carriage return"}
DATE_HINT = (
    "The spreadsheet converted this cell to a date; format the column as text "
    "and enter the value again."
)
CELL_ISSUES = {
    "cell_line_break": "contain a tab or a line break",
    "illegal_character": "contain characters a table cannot hold",
    "cell_date": "hold a date or a time",
    "cell_percent": "are formatted as percentages",
    "cell_error": "hold an error value",
    "formula_without_value": "hold a formula without a saved value",
    "formula_value": "hold formulas that are stored as their values",
    "value_outside_table": "have values outside the header columns",
}
BASE_HINT = (
    "Close the workbook and run pkdb tables sync; if the tables and the workbook "
    "differ, choose a side with --keep"
)
SHEET_HINT = (
    "Rename the sheet to subjects, interventions, characteristica or "
    "<kind>_<source>, such as outputs_Tab2, or <study>_<source> for the raw table "
    "of a paper table, such as Example_Tab2, or start its name with _ to keep it "
    "as a scratch sheet."
)

# The code, message and hint of an issue at a cell.
type Problem = tuple[str, str, str | None]
# Reports a problem at a sheet row and a 0-based column, with the column header.
type Report = Callable[[Problem, int, int, str | None], None]


@dataclass(frozen=True)
class SheetTable:
    """A data or raw sheet as the canonical TSV text of its table file.

    `rows[i]` is the sheet row of canonical line `i + 1`; for a data sheet,
    `rows[0]` is 1, the header.
    """

    file: str
    text: str
    rows: tuple[int, ...]


@dataclass(frozen=True)
class WorkbookContent:
    """The tables of a workbook, its sheets in workbook order, its base and issues.

    The sheets are the data and raw sheets. An optional table or a raw table
    without rows is left out of `tables`, as if absent, but its sheet is listed in
    `sheets`. `base` is None when `_base` is missing or invalid.
    """

    tables: dict[str, SheetTable]
    sheets: tuple[str, ...]
    base: WorkbookBase | None
    issues: list[ValidationIssue]

    @property
    def ok(self) -> bool:
        return not any(issue.severity == "error" for issue in self.issues)


class _Unreadable(Exception):
    """openpyxl cannot read the workbook or one of its sheets."""


class _Reader(ExcelReader):
    """The reader of openpyxl, keeping the shared strings as the file stores them.

    openpyxl removes every `x005F_` from a shared string, which turns the stored
    `a_x005F_x005F_b` of the text `a_x005F_b` into `a_b`, and it keeps the
    escapes of inline strings. `_decoded` decodes both kinds alike.
    """

    def read_strings(self) -> None:
        part = self.package.find(SHARED_STRINGS)
        if part is None:
            return
        strings = []
        with self.archive.open(part.PartName[1:]) as source:
            for _, node in iterparse(source):
                if node.tag == SHARED_STRING:
                    strings.append(Text.from_tree(node).content)
                    node.clear()
        self.shared_strings = strings


def _load(data: bytes, *, data_only: bool) -> Workbook:
    try:
        reader = _Reader(BytesIO(data), read_only=True, data_only=data_only)
        reader.read()
    except Exception as error:
        # openpyxl raises many kinds of errors for a damaged or foreign file.
        raise _Unreadable(str(error)) from error
    return reader.wb


def _character(match: re.Match[str]) -> str:
    code = int(match.group(1), 16)
    # A lone surrogate is no character; its escape stays text.
    return match.group() if 0xD800 <= code <= 0xDFFF else chr(code)


def _decoded(text: str) -> str:
    """A stored text with its OOXML escapes decoded, as spreadsheet applications read it."""
    return ESCAPE.sub(_character, text) if "_x" in text else text


def _rows(sheet, *, values_only: bool = False) -> Iterator[tuple]:
    """The rows of a read-only sheet, ignoring the stored dimension.

    A row missing in the file is an empty tuple, so the n-th row is sheet row n.
    """
    sheet.reset_dimensions()
    try:
        yield from sheet.iter_rows(values_only=values_only)
    except Exception as error:
        raise _Unreadable(
            f"the sheet {sheet.title!r} cannot be read: {error}"
        ) from error


def _value_text(cell, *, formula: bool = False) -> tuple[str, Problem | None]:
    """Text of a cell value, and the problem of a value the spreadsheet converted.

    Only an error cell holds an error value, also as the saved value of a
    formula; a text such as `#N/A` in a text cell is text. `formula` tells that
    the cell holds the saved value of a formula.
    """
    value = cell.value
    if value is None:
        return "", None
    if cell.data_type == "e":
        if (
            value == OUTSIDE_CALENDAR
            and not formula
            and is_date_format(cell.number_format)
        ):
            # openpyxl reads a date serial beyond the calendar as this error;
            # a formula whose value is an error is reported as one.
            return "", (
                "cell_date",
                "The cell is formatted as a date, and its value is outside the "
                "calendar",
                DATE_HINT,
            )
        return str(value), _error_value(str(value))
    if isinstance(value, str):
        value = _decoded(value)
        if (match := LINE_BREAK.search(value)) is not None:
            name = LINE_BREAKS[match.group()]
            return LINE_BREAK.sub(" ", value), (
                "cell_line_break",
                f"The cell contains a {name}; a table cell holds a single line",
                f"Remove the {name} from the cell.",
            )
        if (character := not_xml(value)) is not None:
            # Excel stores a control character as its escape, such as _x0001_.
            return value, (
                "illegal_character",
                f"The cell contains {character}, which a table cannot hold",
                "Remove the character.",
            )
        return value, None
    if isinstance(value, bool):
        return ("TRUE" if value else "FALSE"), None
    if isinstance(value, int | float):
        text = str(value) if isinstance(value, int) else format_number(value)
        if "%" in (cell.number_format or ""):
            return text, (
                "cell_percent",
                f"The cell is formatted as a percentage and holds {text}",
                "20% is stored as 0.2; enter 20.",
            )
        return text, None
    if isinstance(value, date | time | timedelta):
        return str(value), (
            "cell_date",
            f"The cell holds the date or time {value}",
            DATE_HINT,
        )
    return str(value), None


def _error_value(value: str) -> Problem:
    return (
        "cell_error",
        f"The cell holds the error value {value}",
        "Correct the formula or enter the value instead.",
    )


def _formula_source(value: object) -> str:
    # A formula is text such as `=B2*2`; an array formula keeps it in `text`.
    text = value if isinstance(value, str) else getattr(value, "text", None)
    return f" {text}" if isinstance(text, str) else ""


def _cell_text(formula, value) -> tuple[str, Problem | None]:
    """Text of a cell, from the cells of the formula and the value pass."""
    if formula.data_type != "f":
        return _value_text(value)
    source = _formula_source(formula.value)
    # A formula whose value is empty text keeps an empty cached value of type
    # str; a formula that no spreadsheet application computed has none.
    if value.value is None and value.data_type != "str":
        return "", (
            "formula_without_value",
            f"The formula{source} has no saved value",
            "Save the workbook in Excel or LibreOffice.",
        )
    text, problem = _value_text(value, formula=True)
    if problem is not None:
        return text, problem
    stored = f"its value {text}" if text else "an empty cell"
    return text, ("formula_value", f"The formula{source} is stored as {stored}", None)


def _outside(index: int) -> Problem:
    return (
        "value_outside_table",
        f"The value is outside the table: column {column_letter(index)} has no header",
        "Move the value into a column of the table, or into a scratch sheet whose "
        "name starts with _.",
    )


def _lines(rows: Iterable[tuple[tuple, tuple]], report: Report) -> Iterator[bytes]:
    """TSV lines of a sheet, one per sheet row, so that line numbers are sheet rows.

    The header is row 1 without trailing empty cells. A value in a column past
    the header is reported. Empty rows become empty lines, which `load_table`
    skips; trailing empty rows are left out.
    """
    header: list[str] | None = None
    empty = 0
    for number, (formulas, values) in enumerate(rows, start=1):
        cells: list[str] = []
        problems: list[tuple[int, Problem]] = []
        for index, (formula, value) in enumerate(zip(formulas, values, strict=True)):
            if value.value is None and formula.data_type != "f":
                cells.append("")
                continue
            text, problem = _cell_text(formula, value)
            cells.append(text)
            if problem is not None:
                problems.append((index, problem))
        if header is None:
            while cells and not cells[-1].strip():
                cells.pop()
            header = cells
        # Without a header, `load_table` reports the missing header instead.
        width = len(header) or len(cells)
        for index, problem in problems:
            if index < width:
                report(problem, number, index, header[index] if header else None)
        for index in range(width, len(cells)):
            if cells[index].strip():
                report(_outside(index), number, index, None)
        del cells[width:]
        if number > 1 and not any(cell.strip() for cell in cells):
            empty += 1
            continue
        for _ in range(empty):
            yield b"\n"
        empty = 0
        yield ("\t".join(cells) + "\n").encode("utf-8")


def _raw_lines(rows: Iterable[tuple[tuple, tuple]], report: Report) -> Iterator[bytes]:
    """TSV lines of a raw sheet, one per sheet row, so that line numbers are sheet rows.

    Every cell is read as its text. A raw table is as wide as its widest row, so
    no value is outside the table.
    """
    for number, (formulas, values) in enumerate(rows, start=1):
        cells: list[str] = []
        for index, (formula, value) in enumerate(zip(formulas, values, strict=True)):
            if value.value is None and formula.data_type != "f":
                cells.append("")
                continue
            text, problem = _cell_text(formula, value)
            cells.append(text)
            if problem is not None:
                report(problem, number, index, None)
        yield ("\t".join(cells) + "\n").encode("utf-8")


def _in_sheet(issue: ValidationIssue, workbook: str) -> ValidationIssue:
    """An issue of the table text of a sheet, located in the workbook.

    The lines of the text are the sheet rows, so the row and the cell stay.
    """
    if issue.source is None:
        return issue
    return issue.model_copy(
        update={"source": issue.source.model_copy(update={"file": workbook})}
    )


def _read_lines[T](
    workbooks: tuple[Workbook, Workbook],
    workbook: str,
    name: str,
    lines: Callable[[Iterable[tuple[tuple, tuple]], Report], Iterator[bytes]],
    load: Callable[[Iterator[bytes]], tuple[T | None, list[ValidationIssue]]],
) -> tuple[T | None, list[ValidationIssue]]:
    """Load the TSV lines of a sheet; its cell issues and the structural issues.

    The issues are located in the sheet of the workbook.
    """
    formulas, values = (book[name] for book in workbooks)
    issues: list[ValidationIssue] = []
    cap = IssueCap()

    def report(problem: Problem, number: int, index: int, header: str | None) -> None:
        code, message, hint = problem
        if cap.admit(code):
            issues.append(
                make_issue(
                    code,
                    message,
                    file=workbook,
                    sheet=name,
                    line=number,
                    column=index,
                    header=header or None,
                    hint=hint,
                )
            )

    rows = zip(_rows(formulas), _rows(values), strict=True)
    loaded, found = load(lines(rows, report))
    issues.extend(
        make_issue(
            code,
            f"{total:,} cells {CELL_ISSUES[code]}; {LISTED}",
            file=workbook,
            sheet=name,
        )
        for code, total in cap.beyond()
    )
    # The content is judged by validation, as when formatting.
    issues.extend(
        _in_sheet(issue, workbook) for issue in found if issue.code in STRUCTURAL
    )
    return loaded, issues


def _read_sheet(
    workbooks: tuple[Workbook, Workbook],
    workbook: str,
    name: str,
    parsed: tuple[TableSpec, str | None],
    study_name: str,
    limit: RowLimit,
) -> tuple[LoadedTable | None, list[ValidationIssue]]:
    """Load a data sheet as a table; its cell issues and the structural issues."""
    file = f"{name}.tsv"
    spec, source = parsed
    return _read_lines(
        workbooks,
        workbook,
        name,
        _lines,
        lambda lines: load_table(
            file, lines, spec, source, study=study_name, limit=limit
        ),
    )


def _read_raw_sheet(
    workbooks: tuple[Workbook, Workbook],
    workbook: str,
    name: str,
    source: str,
    limit: RowLimit,
) -> tuple[SheetTable | None, list[ValidationIssue]]:
    """A raw sheet as the canonical text of its raw table, None without cells.

    Its lines are read as those of the TSV file, so a cell reads as in the
    file. The issues are those of the cells and of lines a TSV file cannot hold.
    """
    file = f"{name}.tsv"
    raw, issues = _read_lines(
        workbooks,
        workbook,
        name,
        _raw_lines,
        lambda lines: load_raw(file, lines, source, limit=limit),
    )
    if raw is None or (text := render_raw(raw)) is None:
        return None, issues
    return SheetTable(file, text, tuple(line for line, _ in raw_lines(raw))), issues


def _read_base(
    values: Workbook, workbook: str
) -> tuple[WorkbookBase | None, list[ValidationIssue]]:
    if BASE_SHEET not in values.sheetnames:
        return None, [
            make_issue(
                "workbook_base_missing",
                f"The workbook has no {BASE_SHEET} sheet, so changes in both the "
                "workbook and the tables cannot be merged",
                file=workbook,
                hint=BASE_HINT,
            )
        ]
    rows = (
        tuple(_decoded(cell) if isinstance(cell, str) else cell for cell in row)
        for row in _rows(values[BASE_SHEET], values_only=True)
    )
    try:
        return parse_base(rows), []
    except BaseError as error:
        # A newer workbook format is an error that a sync cannot repair.
        hint = BASE_HINT if error.code == "workbook_base_invalid" else None
        return None, [make_issue(error.code, error.message, file=workbook, hint=hint)]


def _sheet_table(
    table: LoadedTable | None, study_name: str, order: dict[str, int]
) -> SheetTable | None:
    """The canonical text of a loaded table, or None for no table or no rows."""
    rows = None if table is None else table_rows(table, study_name, order)
    if table is None or rows is None:
        return None
    return SheetTable(
        table.file,
        render_tsv(table.spec.names, [cells for _, cells in rows]),
        (1, *(line for line, _ in rows)),
    )


def _read(
    workbooks: tuple[Workbook, Workbook],
    workbook: str,
    study_name: str,
    limit: RowLimit,
) -> WorkbookContent:
    values = workbooks[1]
    issues: list[ValidationIssue] = []
    worksheets = {sheet.title for sheet in values.worksheets}
    data: dict[str, tuple[TableSpec, str | None]] = {}
    raw: dict[str, str] = {}
    for name in values.sheetnames:
        if name.startswith("_"):
            # `_lists`, `_base` and scratch sheets.
            continue
        parsed = parse_table_file(f"{name}.tsv")
        source = parse_raw_file(f"{name}.tsv", study_name)
        if (parsed is None and source is None) or name not in worksheets:
            issues.append(
                make_issue(
                    "unknown_sheet",
                    f"The sheet {name!r} is not a table of the study",
                    file=workbook,
                    sheet=name,
                    hint=SHEET_HINT,
                )
            )
        elif parsed is not None:
            data[name] = parsed
        elif source is not None:
            raw[name] = source
    if SUBJECTS not in data:
        issues.append(
            make_issue(
                "missing_sheet",
                f"The workbook has no {SUBJECTS} sheet",
                file=workbook,
                sheet=SUBJECTS,
                hint=f"Add the {SUBJECTS} sheet again; every study needs it.",
            )
        )
    base, found = _read_base(values, workbook)
    issues.extend(found)
    read: dict[str, tuple[SheetTable | None, list[ValidationIssue]]] = {}
    order: dict[str, int] = {}
    # The subjects come first: their order sorts the rows of every table.
    for name in sorted(data, key=lambda name: name != SUBJECTS):
        table, found = _read_sheet(
            workbooks, workbook, name, data[name], study_name, limit
        )
        if name == SUBJECTS:
            order = subject_order(table)
        read[name] = _sheet_table(table, study_name, order), found
    for name, source in raw.items():
        read[name] = _read_raw_sheet(workbooks, workbook, name, source, limit)
    tables: dict[str, SheetTable] = {}
    # In workbook order.
    sheets = tuple(name for name in values.sheetnames if name in read)
    for name in sheets:
        sheet, found = read[name]
        if sheet is not None:
            tables[sheet.file] = sheet
        issues.extend(found)
    return WorkbookContent(tables, sheets, base, issues)


def sheet_names(path: Path) -> tuple[str, ...] | None:
    """The names of the sheets of a workbook in workbook order, without reading a cell.

    None when the workbook cannot be read.
    """
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", module="openpyxl")
        try:
            reader = ExcelReader(Path(path), read_only=True, keep_links=False)
        except Exception:
            # openpyxl raises many kinds of errors for a damaged or foreign file.
            return None
        try:
            reader.read_manifest()
            reader.read_workbook()
            return tuple(sheet.name for sheet in reader.parser.sheets)
        except Exception:
            return None
        finally:
            reader.archive.close()


def read_workbook(
    path: Path, study_name: str, *, max_rows: int | None = None
) -> WorkbookContent:
    """Read the sheets of a workbook into the canonical text of their table files.

    Formula cells are read as the value the spreadsheet application saved. The
    file is read once, and both passes of openpyxl, with formulas and with
    values, read these bytes, so a save in between cannot mix two versions.
    `max_rows` limits the data rows of all sheets together; more raise
    BeyondLimits `row_limit`, as `load_study` does.
    """
    path = Path(path)
    with warnings.catch_warnings(), ExitStack() as stack:
        # openpyxl warns about parts of Excel files it does not support and
        # about dates beyond the calendar; the reader reports what matters.
        warnings.filterwarnings("ignore", module="openpyxl")
        try:
            data = path.read_bytes()
            formulas = _load(data, data_only=False)
            stack.callback(formulas.close)
            values = _load(data, data_only=True)
            stack.callback(values.close)
            return _read((formulas, values), path.name, study_name, RowLimit(max_rows))
        except (_Unreadable, OSError) as error:
            return WorkbookContent(
                {},
                (),
                None,
                [
                    make_issue(
                        "workbook_unreadable",
                        f"{path.name} cannot be read: {error}",
                        file=path.name,
                    )
                ],
            )
