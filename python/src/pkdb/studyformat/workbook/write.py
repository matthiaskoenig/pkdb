"""Generate the study workbook from the canonical TSV tables (spec 10.2 and 10.6)."""

import re
import uuid
import warnings
from collections.abc import Iterable, Mapping
from copy import copy
from dataclasses import dataclass
from datetime import UTC, datetime
from io import BytesIO
from math import ceil
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl
from openpyxl.cell.cell import Cell
from openpyxl.comments import Comment
from openpyxl.styles.fonts import DEFAULT_FONT
from openpyxl.workbook import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.writer.excel import ExcelWriter

from pkdb.domain.vocabulary import Vocabulary
from pkdb.schemas.validation import ValidationIssue
from pkdb.studyformat.columns import Column, ColumnType
from pkdb.studyformat.issues import LISTED, IssueCap, column_letter, make_issue
from pkdb.studyformat.tables import (
    KIND_ORDER,
    TABLES,
    TableSpec,
    parse_table_file,
    table_file,
)
from pkdb.studyformat.terms import vocabulary_terms
from pkdb.studyformat.text import format_number, natural_key, parse_number
from pkdb.studyformat.workbook.base import (
    BASE_SHEET,
    LISTS_SHEET,
    SHEET_NAME_LIMIT,
    WORKBOOK_FORMAT,
    WorkbookBase,
    base_rows,
)

AUTHOR = "pkdb"
TEXT_FORMAT = "@"
GENERAL_FORMAT = "General"
# Rows of a sheet in Excel and LibreOffice; dropdowns cover every data row.
MAX_ROW = 1_048_576
# Characters a spreadsheet cell holds, counted in UTF-16 code units like Excel.
CELL_LIMIT = 32_767
# Characters XML 1.0, and so a workbook, cannot hold: control characters other
# than tab, line feed and carriage return, surrogates and U+FFFE and U+FFFF.
NOT_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
# OOXML stores a character as the escape `_xHHHH_`, and `_x005F_` is the
# underscore; this matches every underscore that starts such an escape.
ESCAPE_START = re.compile(r"_(?=x[0-9A-Fa-f]{4}_)")
ESCAPED_UNDERSCORE = "_x005F_"
# Two such escapes that share an underscore, as in `_x005F_x0041_`, which
# LibreOffice 24.2 saves unescaped, so that it reads back as `_x0041_`.
OVERLAPPING_ESCAPES = re.compile(r"_x[0-9A-Fa-f]{4}_x[0-9A-Fa-f]{4}_")
# LibreOffice saves numbers with at most 15 significant digits.
NUMBER_DIGITS = 15
NUMBER_TYPES = frozenset({ColumnType.NUMBER, ColumnType.INTEGER, ColumnType.TIME})
# Column widths in characters, from the header and the first WIDTH_ROWS cells.
MIN_WIDTH, MAX_WIDTH, WIDTH_ROWS = 8, 60, 200
# Room around a cell text, and around the bold header text with its filter button.
CELL_PADDING, HEADER_PADDING = 2, 4
# Size of the comment boxes in pixels: characters per line and line height.
COMMENT_WIDTH, COMMENT_CHARACTERS, COMMENT_LINE, COMMENT_MARGIN = 320, 48, 16, 12
# The font of the cells, in bold; a font without a name looks different in LibreOffice.
BOLD = copy(DEFAULT_FONT)
BOLD.bold = True
CELL_ISSUES = {
    "cell_too_long": "cells are longer than a spreadsheet cell",
    "illegal_character": "cells contain characters a workbook cannot hold",
    "cell_escape_text": "cells contain text that LibreOffice changes when it saves",
}


@dataclass(frozen=True)
class WorkbookBuild:
    """A generated workbook; `data` is None when `issues` has an error."""

    data: bytes | None
    base: WorkbookBase
    issues: list[ValidationIssue]


class WorkbookError(Exception):
    """The workbook cannot be generated, and an existing one must be kept as it is."""

    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass(frozen=True)
class _Sheet:
    # A data sheet: its table file, specification, source and canonical text,
    # which is None for a sheet with the header only.
    name: str
    file: str
    spec: TableSpec
    source: str | None
    text: str | None


def _is_text(column: Column) -> bool:
    return column.type not in NUMBER_TYPES


def _comment_text(spec: TableSpec, column: Column) -> str:
    lines = [column.description]
    if column.example:
        lines.append(f"Example: {column.example}")
    if column.name in spec.required_columns:
        lines.append("Required")
    return "\n".join(lines)


def _dropdown(column: Column) -> tuple[str, str] | None:
    """Sheet and header of the column that lists the values of a column, or None.

    Units and choices depend on the measurement and have no dropdown.
    """
    if column.references:
        return column.references, "name"
    if column.vocabulary:
        return LISTS_SHEET, column.vocabulary
    if column.type is ColumnType.ENUM:
        return LISTS_SHEET, f"choice:{column.name}"
    return None


def _list_columns(specs: Iterable[TableSpec]) -> dict[str, Column]:
    """The `_lists` columns the sheets use, in order of first use, with a column using each."""
    found: dict[str, Column] = {}
    for spec in specs:
        for column in spec.columns:
            target = _dropdown(column)
            if target is not None and target[0] == LISTS_SHEET:
                found.setdefault(target[1], column)
    return found


def _number(text: str) -> int | float | None:
    """Value of the number cell for a text of a number column; None keeps the text.

    Only numbers that come back as the same canonical text are number cells, also
    after LibreOffice saved them with at most 15 significant digits.
    """
    value = parse_number(text)
    if value is None or format_number(value) != text:
        return None
    mantissa = text.lstrip("+-").lower().partition("e")[0]
    if len(mantissa.replace(".", "").strip("0")) > NUMBER_DIGITS:
        return None
    return value if any(mark in text for mark in ".eE") else int(text)


def not_xml(text: str) -> str | None:
    """The first character of a text that a workbook cannot hold, described, or None."""
    match = NOT_XML.search(text)
    if match is None:
        return None
    code = ord(match.group())
    kind = (
        "control character"
        if code < 0x20
        else "surrogate"
        if 0xD800 <= code <= 0xDFFF
        else "noncharacter"
    )
    return f"the {kind} U+{code:04X}"


def escape_text(text: str) -> str:
    """A text as a workbook stores it, so that spreadsheet applications read it back.

    Excel and LibreOffice read `_x0041_` as the escape of the character A, as
    OOXML defines, so every underscore that starts such an escape is stored as
    the escape `_x005F_` of the underscore: `_x0041_` becomes `_x005F_x0041_`.
    """
    return ESCAPE_START.sub(ESCAPED_UNDERSCORE, text)


def _put_text(cell: Cell, text: str) -> None:
    """Write a text cell; a text starting with `=` stays text instead of a formula."""
    # openpyxl checks the characters and cuts a text after 32,767 characters,
    # which the escaped text may exceed while the text itself fits.
    cell.value = text
    cell._value = escape_text(text)
    cell.data_type = "s"
    cell.number_format = TEXT_FORMAT


def _formula(target: tuple[str, str], lists: Mapping[str, list[str]]) -> str:
    """The cell range of a dropdown: a `_lists` column, or a whole `name` column."""
    sheet, header = target
    if sheet == LISTS_SHEET:
        index, end = list(lists).index(header), max(2, len(lists[header]) + 1)
    else:
        index, end = TABLES[sheet].names.index(header), MAX_ROW
    letter = column_letter(index)
    return f"{sheet}!${letter}$2:${letter}${end}"


def _comment(text: str) -> Comment:
    lines = sum(
        max(1, ceil(len(line) / COMMENT_CHARACTERS)) for line in text.split("\n")
    )
    return Comment(
        text,
        AUTHOR,
        height=lines * COMMENT_LINE + COMMENT_MARGIN,
        width=COMMENT_WIDTH,
    )


def _plan(tables: Mapping[str, str], empty_sheets: Iterable[str]) -> list[_Sheet]:
    """The data sheets in workbook order: the tables not split by source come first."""
    files = [
        *tables,
        *(f"{name}.tsv" for name in empty_sheets),
        *(table_file(spec.kind) for spec in TABLES.values() if not spec.per_source),
    ]
    # Excel compares sheet names ignoring case.
    sheets: dict[str, _Sheet] = {}
    for file in files:
        parsed = parse_table_file(file)
        if parsed is None:
            raise ValueError(f"{file} is not a table file")
        name = file.removesuffix(".tsv")
        if len(name) > SHEET_NAME_LIMIT:
            raise WorkbookError(
                "table_name_too_long",
                f"The sheet name {name!r} has {len(name)} characters; "
                f"Excel limits sheet names to {SHEET_NAME_LIMIT} characters",
            )
        sheet = sheets.setdefault(
            name.casefold(), _Sheet(name, file, *parsed, tables.get(file))
        )
        if sheet.name != name:
            raise WorkbookError(
                "duplicate_table_name",
                f"The sheet name {name!r} equals {sheet.name!r} ignoring case; "
                "Excel sheet names are case-insensitive",
            )
    return sorted(
        sheets.values(),
        key=lambda sheet: (
            KIND_ORDER[sheet.spec.kind],
            natural_key(sheet.source or ""),
        ),
    )


def _open(existing: Path | None) -> Workbook:
    """A new workbook, or the scratch sheets of the existing one."""
    if existing is None:
        workbook = Workbook()
        workbook.remove(workbook.active)
        workbook.properties.creator = AUTHOR
        return workbook
    try:
        with warnings.catch_warnings():
            # openpyxl warns about parts of Excel files it does not support,
            # which only the scratch sheets keep from the existing workbook.
            warnings.filterwarnings("ignore", module="openpyxl")
            # Full mode with formulas, so that the scratch sheets keep them.
            workbook = openpyxl.load_workbook(existing)
    except Exception as error:
        # openpyxl raises many kinds of errors for a damaged or foreign file.
        raise WorkbookError(
            "workbook_unreadable", f"{existing.name} cannot be read: {error}"
        ) from error
    generated = {LISTS_SHEET, BASE_SHEET}
    for name in workbook.sheetnames:
        if not name.startswith("_") or name.casefold() in generated:
            workbook.remove(workbook[name])
    return workbook


def _write_table(
    sheet: Worksheet, plan: _Sheet, lists: Mapping[str, list[str]]
) -> list[ValidationIssue]:
    spec = plan.spec
    widths = [len(name) + HEADER_PADDING for name in spec.names]
    for index, column in enumerate(spec.columns):
        cell = sheet.cell(1, index + 1)
        _put_text(cell, column.name)
        cell.font = BOLD
        cell.comment = _comment(_comment_text(spec, column))
    issues: list[ValidationIssue] = []
    cap = IssueCap()
    last = 1
    if plan.text is not None:
        # Canonical text: LF line ends, a final newline and no blank lines. It is
        # split as text, so that a text UTF-8 cannot encode becomes an issue.
        lines = plan.text.removesuffix("\n").split("\n")
        if tuple(lines[0].split("\t")) != spec.names:
            raise ValueError(f"{plan.file} is not in canonical form")
        for number, line in enumerate(lines[1:], start=2):
            last = number
            for index, (column, text) in enumerate(
                zip(spec.columns, line.split("\t"), strict=True)
            ):
                if not text:
                    continue
                if number <= WIDTH_ROWS + 1:
                    widths[index] = max(widths[index], len(text) + CELL_PADDING)
                problem = _write_cell(sheet.cell(number, index + 1), column, text)
                if problem is not None and cap.admit(problem[0]):
                    code, message, hint = problem
                    issues.append(
                        make_issue(
                            code,
                            message,
                            file=plan.file,
                            line=number,
                            column=index,
                            header=column.name,
                            hint=hint,
                        )
                    )
    issues.extend(
        make_issue(code, f"{total:,} {CELL_ISSUES[code]}; {LISTED}", file=plan.file)
        for code, total in cap.beyond()
    )
    for index, column in enumerate(spec.columns):
        letter = column_letter(index)
        dimension = sheet.column_dimensions[letter]
        dimension.width = min(MAX_WIDTH, max(MIN_WIDTH, widths[index]))
        if _is_text(column):
            # New cells typed into a text column are text, not dates.
            dimension.number_format = TEXT_FORMAT
        if (target := _dropdown(column)) is not None:
            # A suggestion only: other values stay allowed and pkdb validate checks them.
            validation = DataValidation(
                type="list",
                formula1=_formula(target, lists),
                allow_blank=True,
                showErrorMessage=False,
            )
            validation.add(f"{letter}2:{letter}{MAX_ROW}")
            sheet.add_data_validation(validation)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{column_letter(len(spec.columns) - 1)}{last}"
    return issues


def _write_cell(cell: Cell, column: Column, text: str) -> tuple[str, str, str] | None:
    """Write one cell; the code, message and hint of an issue, or None.

    A cell with an error is not written; a cell with a warning is.
    """
    # A text has at most two UTF-16 code units per character.
    if len(text) * 2 > CELL_LIMIT:
        units = len(text.encode("utf-16-le", "surrogatepass")) // 2
        if units > CELL_LIMIT:
            return (
                "cell_too_long",
                f"The cell has {units:,} characters as Excel counts them; a "
                f"spreadsheet cell holds at most {CELL_LIMIT:,}",
                "Shorten the text.",
            )
    if (character := not_xml(text)) is not None:
        return (
            "illegal_character",
            f"The cell contains {character}, which a workbook cannot hold",
            "Remove the character.",
        )
    if not _is_text(column) and (number := _number(text)) is not None:
        cell.value = number
        return None
    _put_text(cell, text)
    if (match := OVERLAPPING_ESCAPES.search(text)) is not None:
        return (
            "cell_escape_text",
            f"The cell contains {match.group()}, two escapes of the form _xHHHH_ "
            "that share an underscore; LibreOffice changes such text when it "
            "saves the workbook",
            "Edit this cell in the TSV file or in Excel, or change the text.",
        )
    return None


def _write_lists(
    sheet: Worksheet, lists: Mapping[str, list[str]]
) -> list[ValidationIssue]:
    issues = []
    for index, (key, values) in enumerate(lists.items(), start=1):
        for row, value in enumerate((key, *values), start=1):
            if (character := not_xml(value)) is not None:
                issues.append(
                    make_issue(
                        "illegal_character",
                        f"The vocabulary term {value!r} of {key} contains "
                        f"{character}, which a workbook cannot hold",
                    )
                )
            else:
                _put_text(sheet.cell(row, index), value)
    sheet.sheet_state = "hidden"
    return issues


def _write_base(sheet: Worksheet, base: WorkbookBase) -> None:
    for row, values in enumerate(base_rows(base), start=1):
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row, column)
            if isinstance(value, str):
                # A base64 chunk may start with `=`.
                _put_text(cell, value)
            else:
                cell.value = value
    sheet.sheet_state = "veryHidden"


def _save(workbook: Workbook, created: datetime) -> bytes:
    # openpyxl writes the document times as UTC without a zone.
    workbook.properties.created = created.replace(tzinfo=None)
    workbook.properties.modified = created.replace(tzinfo=None)
    buffer = BytesIO()
    # Workbook.save would set the modification time to now; the writer closes the archive.
    ExcelWriter(workbook, ZipFile(buffer, "w", ZIP_DEFLATED, allowZip64=True)).save()
    return buffer.getvalue()


def build_workbook(
    tables: Mapping[str, str],
    vocabulary: Vocabulary,
    *,
    existing: Path | None = None,
    empty_sheets: Iterable[str] = (),
    generation: str | None = None,
    created: datetime | None = None,
) -> WorkbookBuild:
    """Generate the workbook of the canonical TSV text of each table file.

    `empty_sheets` names further sheets with the header only, such as
    `outputs_Tab3`. The scratch sheets of an `existing` workbook are kept. A new
    generation and the current time are used unless given. Cells a workbook
    cannot hold are issues, and then no workbook is written. A sheet name that
    Excel does not accept or an unreadable existing workbook raise WorkbookError.
    """
    base = WorkbookBase(
        generation or uuid.uuid4().hex,
        (created or datetime.now(UTC)).astimezone(UTC),
        dict(tables),
    )
    plan = _plan(tables, empty_sheets)
    terms = vocabulary_terms(vocabulary)
    deprecated = {rule.name for rule in vocabulary.measurements if rule.deprecated}
    terms["measurements"] -= deprecated
    lists = {
        key: sorted(
            terms[column.vocabulary] if column.vocabulary else column.choices,
            key=natural_key,
        )
        for key, column in _list_columns(sheet.spec for sheet in plan).items()
    }
    workbook = _open(existing)
    issues: list[ValidationIssue] = []
    for index, sheet in enumerate(plan):
        issues.extend(
            _write_table(workbook.create_sheet(sheet.name, index), sheet, lists)
        )
    issues.extend(_write_lists(workbook.create_sheet(LISTS_SHEET), lists))
    if any(issue.severity == "error" for issue in issues):
        return WorkbookBuild(None, base, issues)
    # Only now: the base encodes the texts as UTF-8.
    _write_base(workbook.create_sheet(BASE_SHEET), base)
    workbook.active = 0
    for index, sheet in enumerate(workbook.worksheets):
        sheet.sheet_view.tabSelected = index == 0
    return WorkbookBuild(_save(workbook, base.created), base, issues)


def workbook_template() -> dict:
    """Header, column formats, comments and dropdowns of every sheet of a full template.

    A sheet split by source is named `<kind>_<source>`. A dropdown names the
    sheet and the header of the column that lists its values, such as
    `subjects!name` or `_lists!measurements`. `lists` are the `_lists` columns.
    """
    specs = list(TABLES.values())
    return {
        "format": WORKBOOK_FORMAT,
        "sheets": {
            (f"{spec.kind}_<source>" if spec.per_source else spec.kind): {
                "header": list(spec.names),
                "columns": {
                    column.name: {
                        "format": TEXT_FORMAT if _is_text(column) else GENERAL_FORMAT,
                        "comment": _comment_text(spec, column),
                        "dropdown": "!".join(target)
                        if (target := _dropdown(column))
                        else None,
                    }
                    for column in spec.columns
                },
            }
            for spec in specs
        },
        "lists": list(_list_columns(specs)),
    }
