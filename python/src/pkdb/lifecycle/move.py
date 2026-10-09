"""Rename or move a study format 2 folder: `pkdb move`."""

import warnings
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path

from pkdb.domain.vocabulary import Vocabulary
from pkdb.lifecycle.names import case_twin, parse_location
from pkdb.lifecycle.new import error_messages
from pkdb.repository import STUDIES, location
from pkdb.schemas.review import Review
from pkdb.studyformat.digitize import digitization_file, parse_digitization_file
from pkdb.studyformat.formatter import format_folder
from pkdb.studyformat.metadata import read_metadata
from pkdb.studyformat.models import canonical_review_json
from pkdb.studyformat.raw import parse_raw_file, raw_file
from pkdb.studyformat.review_edit import read_review
from pkdb.studyformat.revision import RevisionConflict, write_checked
from pkdb.studyformat.sync import sync_study
from pkdb.studyformat.tables import REVIEW_JSON, SOURCE_PATTERN, image_file
from pkdb.studyformat.text import natural_key
from pkdb.studyformat.validation import is_v2_folder
from pkdb.studyformat.workbook.base import (
    SHEET_NAME_LIMIT,
    is_scratch_sheet,
    open_lock,
    state_path,
    workbook_path,
)

EDITS_IN_WORKBOOK = (
    "The workbook holds edits that are not in the tables; run pkdb tables sync first"
)


@dataclass(frozen=True)
class Moved:
    """A moved study: its new folder and what changed in it.

    `renamed` lists the old and new name of every renamed file, `targets`
    counts the review targets that now name a renamed file, `issue` is the
    issue number of study.json, and `workbook` names the removed workbook.
    """

    folder: Path
    renamed: list[tuple[str, str]]
    targets: int
    issue: int | None
    workbook: str | None = None


class MoveRefused(ValueError):
    """The study cannot be moved; nothing was changed."""


class MoveIncomplete(RuntimeError):
    """The move failed part way.

    `moved` is the study at its new place when only a step after the renames
    failed, and None when the renames could not be undone.
    """

    def __init__(self, message: str, moved: Moved | None):
        super().__init__(message)
        self.moved = moved


@dataclass(frozen=True)
class _Plan:
    """The checked move: folders, file renames and the new review.json text, if it changes."""

    old: Path
    new: Path
    renames: list[tuple[str, str]]
    review: str | None
    revision: str
    targets: int
    issue: int | None
    workbook: Path | None


def move_study(root: Path, old: str, new: str, vocabulary: Vocabulary) -> Moved:
    """Move the study `studies/<old>` of a checkout to `studies/<new>`.

    Every check runs before the first change. The folder and file renames and
    the review targets are undone together when one of them fails, so the move
    either happens or leaves the study as it was. Steps after them, removing
    the workbook and formatting, raise MoveIncomplete with the study at its new
    place.
    """
    plan = _plan(Path(root), old, new, vocabulary)
    _apply(plan)
    _remove_if_empty(plan.old.parent)
    moved = Moved(
        plan.new,
        plan.renames,
        plan.targets,
        plan.issue,
        None if plan.workbook is None else plan.workbook.name,
    )
    problems = []
    # Only a checked workbook goes, which holds nothing beyond the tables. Its
    # state file is named after the old study and goes in any case.
    workbook = plan.new / workbook_path(plan.old).name
    removed = [] if plan.workbook is None else [workbook]
    for path in [*removed, state_path(workbook)]:
        try:
            path.unlink(missing_ok=True)
        except OSError as error:
            problems.append(
                f"{path.name} could not be removed ({_reason(error)}); remove it by hand"
            )
    try:
        formatted = format_folder(plan.new)
    except OSError as error:
        problems.append(f"pkdb format failed ({_reason(error)}); run it again")
    else:
        if not formatted.ok:
            problems.append(
                f"pkdb format found errors ({error_messages(formatted.issues)}); "
                "fix them and run pkdb format"
            )
    if problems:
        raise MoveIncomplete(
            f"Moved {_shown(plan.old)} to {_shown(plan.new)}, but "
            f"{'; and '.join(problems)}. The study stays at its new place.",
            moved,
        )
    return moved


def _plan(root: Path, old: str, new: str, vocabulary: Vocabulary) -> _Plan:
    """Check the move without changing anything."""
    old_folder = root / STUDIES / _old_location(old)
    if not old_folder.is_dir(follow_symlinks=False) or not is_v2_folder(old_folder):
        raise MoveRefused(f"{_shown(old_folder)} is not a study format 2 folder")
    try:
        substance, name = parse_location(new)
    except ValueError as error:
        raise MoveRefused(str(error)) from None
    studies = root / STUDIES
    if (substance, name) == (old_folder.parent.name, old_folder.name):
        raise MoveRefused(f"The study is at {new} already")
    _refuse_case_twin(studies, substance, STUDIES)
    substances = studies / substance
    if substances.exists(follow_symlinks=False) and not substances.is_dir(
        follow_symlinks=False
    ):
        raise MoveRefused(f"{STUDIES}/{substance} is not a folder")
    _refuse_case_twin(substances, name, f"{STUDIES}/{substance}")
    target = substances / name
    if target.exists(follow_symlinks=False):
        raise MoveRefused(f"{_shown(target)} exists already")

    renames = _renames(old_folder, name)
    _refuse_collisions(old_folder, renames)
    _refuse_long_sheets(renames)
    # A study that pkdb format cannot read would stop the move half way.
    formatted = format_folder(old_folder, check=True)
    if not formatted.ok:
        raise MoveRefused(
            f"{_shown(old_folder)} has errors ({error_messages(formatted.issues)}); "
            "fix them before the move"
        )
    workbook = _removable_workbook(old_folder, vocabulary)
    try:
        issue = read_metadata(old_folder).metadata.issue
        document = read_review(old_folder)
    except ValueError as error:
        raise MoveRefused(f"{_shown(old_folder)}: {error}") from None
    review, targets = _retarget(document.review, dict(renames))
    return _Plan(
        old_folder,
        target,
        renames,
        canonical_review_json(review) if targets else None,
        document.revision,
        targets,
        issue,
        workbook,
    )


def _old_location(value: str) -> str:
    """The `<substance>/<name>` of an existing study; a trailing slash is ignored."""
    parts = value.removesuffix("/").split("/")
    if len(parts) != 2 or not all(parts) or any(p.startswith(".") for p in parts):
        raise MoveRefused(f"{value!r} is not <substance>/<name>")
    return "/".join(parts)


def _shown(folder: Path) -> str:
    """`studies/<substance>/<name>` of a study folder."""
    return f"{STUDIES}/{location(folder)}"


def _refuse_case_twin(folder: Path, name: str, shown: str) -> None:
    if (twin := case_twin(folder, name)) is not None:
        raise MoveRefused(f"{shown}/{name} differs from {shown}/{twin} only in case")


def _renamed(file: str, old: str, new: str) -> str | None:
    """The name of a file of the study `old` in the study `new`, or None for a file not named after it.

    These are the PDF, the images, the raw tables and the digitizations.
    """
    if file == f"{old}.pdf":
        return f"{new}.pdf"
    if (source := parse_raw_file(file, old)) is not None:
        return raw_file(new, source)
    if (source := parse_digitization_file(file, old)) is not None:
        return digitization_file(new, source)
    prefix = f"{old}_"
    if file.startswith(prefix) and file.endswith(".png"):
        source = file[len(prefix) : -len(".png")]
        if SOURCE_PATTERN.fullmatch(source):
            return image_file(new, source)
    return None


def _renames(folder: Path, new: str) -> list[tuple[str, str]]:
    """The old and new names of the files named after the study, in natural order."""
    renames = []
    for path in sorted(folder.iterdir(), key=lambda item: natural_key(item.name)):
        if path.is_dir(follow_symlinks=False):
            continue
        if (renamed := _renamed(path.name, folder.name, new)) is not None:
            renames.append((path.name, renamed))
    return renames


def _refuse_collisions(folder: Path, renames: list[tuple[str, str]]) -> None:
    """Refuse a new file name that another file of the folder has, also ignoring case.

    A rename would replace that file, at least on the file systems of macOS and Windows.
    """
    names: dict[str, set[str]] = defaultdict(set)
    for path in folder.iterdir():
        names[path.name.casefold()].add(path.name)
    for old, new in renames:
        if others := sorted(names[new.casefold()] - {old}):
            raise MoveRefused(
                f"{old} would become {new}, which collides with {others[0]} of "
                f"{_shown(folder)}; rename or remove {others[0]} first"
            )


def _refuse_long_sheets(renames: list[tuple[str, str]]) -> None:
    """Refuse a raw table whose workbook sheet name the move makes too long for Excel."""
    for old, new in renames:
        if not new.endswith(".tsv"):
            continue
        sheet = new.removesuffix(".tsv")
        if len(sheet) > SHEET_NAME_LIMIT >= len(old.removesuffix(".tsv")):
            raise MoveRefused(
                f"The raw table {old} would become {new}, whose sheet name has "
                f"{len(sheet)} characters; Excel limits sheet names to "
                f"{SHEET_NAME_LIMIT} characters. Choose a shorter study name."
            )


def _removable_workbook(folder: Path, vocabulary: Vocabulary) -> Path | None:
    """The workbook of the study when it holds nothing beyond the tables, None without one.

    It must be closed, hold no edits that are not in the tables, and have no
    scratch sheets, which only a regeneration of the workbook keeps.
    """
    workbook = workbook_path(folder)
    if not workbook.exists():
        return None
    if (lock := open_lock(workbook)) is not None:
        raise MoveRefused(
            f"{workbook.name} is open in a spreadsheet program ({lock.name}); "
            "close it and run the move again"
        )
    result = sync_study(folder, vocabulary, check=True)
    if result.changes or result.conflicts:
        raise MoveRefused(EDITS_IN_WORKBOOK)
    if errors := [issue for issue in result.issues if issue.severity == "error"]:
        raise MoveRefused(
            f"{workbook.name} cannot be compared with the tables "
            f"({error_messages(errors)}); run pkdb tables sync first"
        )
    if scratch := _scratch_sheets(workbook):
        raise MoveRefused(
            f"{workbook.name} has the scratch sheets {', '.join(scratch)}, which the "
            "move would remove with the workbook; copy what you need and delete them first"
        )
    return workbook


def _scratch_sheets(workbook: Path) -> list[str]:
    import openpyxl

    try:
        with warnings.catch_warnings():
            # openpyxl warns about parts of Excel files it does not support.
            warnings.filterwarnings("ignore", module="openpyxl")
            book = openpyxl.load_workbook(workbook, read_only=True)
    except Exception as error:
        # openpyxl raises many kinds of errors for a damaged or foreign file.
        raise MoveRefused(f"{workbook.name} cannot be read: {error}") from None
    try:
        return [name for name in book.sheetnames if is_scratch_sheet(name)]
    finally:
        book.close()


def _retarget(review: Review, renamed: dict[str, str]) -> tuple[Review, int]:
    """The review with the targets of renamed files renamed, and how many changed.

    An acknowledgement of a warning is a review item with a target, so it follows too.
    """
    items, count = [], 0
    for item in review.items:
        target = item.target
        if target is not None and target.file in renamed:
            target = target.model_copy(update={"file": renamed[target.file]})
            item = item.model_copy(update={"target": target})
            count += 1
        items.append(item)
    return review.model_copy(update={"items": items}), count


def _apply(plan: _Plan) -> None:
    """Rename the folder and its files and write review.json, or undo every step."""
    substances = plan.new.parent
    created = False
    done: list[tuple[Path, Path]] = []
    try:
        if not substances.exists():
            substances.mkdir()
            created = True
        if plan.new.exists(follow_symlinks=False):
            raise MoveRefused(f"{_shown(plan.new)} exists already")
        _rename(plan.old, plan.new, done)
        for old, new in plan.renames:
            _rename(plan.new / old, plan.new / new, done)
        if plan.review is not None:
            write_checked(plan.new / REVIEW_JSON, plan.review, plan.revision)
    except BaseException as error:
        failures = _undo(done)
        if created:
            _remove_if_empty(substances)
        what = f"move {_shown(plan.old)} to {_shown(plan.new)}"
        if failures:
            raise MoveIncomplete(
                f"Cannot {what}: {_reason(error)}. Undoing the move failed "
                f"({'; '.join(failures)}); finish or undo it by hand.",
                None,
            ) from error
        if isinstance(error, (OSError, RevisionConflict)):
            raise MoveRefused(
                f"Cannot {what}: {_reason(error)}. Nothing was changed."
            ) from error
        raise


def _rename(source: Path, target: Path, done: list[tuple[Path, Path]]) -> None:
    source.rename(target)
    done.append((source, target))


def _undo(done: list[tuple[Path, Path]]) -> list[str]:
    """Rename back in reverse order; the renames that failed."""
    failures = []
    for source, target in reversed(done):
        try:
            target.rename(source)
        except OSError as error:
            failures.append(f"{target.name} to {source.name}: {_reason(error)}")
    return failures


def _reason(error: BaseException) -> str:
    if isinstance(error, OSError):
        reason = str(error.strerror or error)
        return f"{reason} ({error.filename})" if error.filename else reason
    return str(error)


def _remove_if_empty(folder: Path) -> None:
    try:
        folder.rmdir()
    except OSError:
        pass
