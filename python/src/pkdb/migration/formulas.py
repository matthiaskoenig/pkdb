"""Formulas of format 1 workbooks, such as error bars computed as ABS(X - mean)."""

import re
from collections import defaultdict
from collections.abc import Collection
from dataclasses import dataclass
from pathlib import Path

import openpyxl
from openpyxl.formula.tokenizer import Token, Tokenizer, TokenizerError
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import coordinate_to_tuple

from pkdb.schemas.study import CanonicalStudy, Observation

CELL = r"\$?([A-Z]{1,3})\$?([0-9]+)"
ABS = re.compile(rf"=\s*ABS\(\s*{CELL}\s*-\s*{CELL}\s*\)\s*", re.IGNORECASE)
SPREADS = ("sd", "se")


def _records(study: CanonicalStudy) -> list[Observation]:
    records: list[Observation] = [*study.interventions, *study.measurements]
    for subject in [*study.groups, *study.individuals]:
        records.extend(subject.characteristica)
    return records


def _cells(
    workbook: Path, wanted: dict[str, set[str]], *, data_only: bool
) -> dict[tuple[str, str], object]:
    """The wanted cells by sheet and coordinate: formulas, or the values saved with them.

    Each sheet is read once up to its last wanted row, since a read-only sheet
    reads from its start for every single cell that is looked up.
    """
    book = openpyxl.load_workbook(workbook, read_only=True, data_only=data_only)
    found: dict[tuple[str, str], object] = {}
    try:
        for name, coordinates in wanted.items():
            if name not in book.sheetnames:
                continue
            sheet = book[name]
            # Some writers save wrong sheet dimensions, which would cut rows short.
            sheet.reset_dimensions()
            last = max(coordinate_to_tuple(coordinate)[0] for coordinate in coordinates)
            rows = sheet.iter_rows(min_row=1, max_row=last, values_only=True)
            for row, values in enumerate(rows, 1):
                for column, value in enumerate(values, 1):
                    coordinate = f"{get_column_letter(column)}{row}"
                    if coordinate in coordinates:
                        found[name, coordinate] = value
    finally:
        book.close()
    return found


def error_bars(workbook: Path, study: CanonicalStudy) -> dict[str, tuple[float, str]]:
    """Record key to (error bar, `sd` or `se`) for spreads computed as ABS(X - mean)."""
    # (record key, sheet, sd or se, its cell, the mean cell of the same row)
    spreads = []
    for record in _records(study):
        source = record.source
        if source is None or source.sheet is None:
            continue
        mean = source.for_field("mean").cell
        for kind in SPREADS:
            # A field without a cell of its own is located at its row, without a cell.
            cell = source.for_field(kind).cell
            if mean is not None and cell is not None:
                spreads.append((record.key, source.sheet, kind, cell, mean))
    wanted = defaultdict(set)
    for _, sheet, _, cell, _ in spreads:
        wanted[sheet].add(cell)
    formulas = _cells(workbook, wanted, data_only=False)
    bars = []
    for key, sheet, kind, cell, mean in spreads:
        formula = formulas.get((sheet, cell))
        match = ABS.fullmatch(formula) if isinstance(formula, str) else None
        if match is None:
            continue
        first = f"{match[1]}{match[2]}".upper()
        second = f"{match[3]}{match[4]}".upper()
        other = second if first == mean else first if second == mean else None
        if other is not None and other != mean:
            bars.append((key, sheet, kind, other))
    wanted = defaultdict(set)
    for _, sheet, _, other in bars:
        wanted[sheet].add(other)
    values = _cells(workbook, wanted, data_only=True)
    found = {}
    for key, sheet, kind, other in bars:
        bar = values.get((sheet, other))
        # Format 2 refuses negative error bars of concentrations; the saved
        # spread says the same.
        if isinstance(bar, (int, float)) and not isinstance(bar, bool) and bar >= 0:
            found[key] = (float(bar), kind)
    return found


# A range of more cells than this, such as a whole column, reads no cells that
# a record could hold one by one.
MAX_RANGE = 10_000


@dataclass(frozen=True)
class References:
    """The cells of its sheet that a formula reads.

    `cells` are the cells of the same sheet, as (row, column letter);
    `outside` says whether it also reads cells of other sheets, whole columns
    or rows, ranges of more than `MAX_RANGE` cells, or names, or cannot be
    parsed.
    """

    cells: frozenset[tuple[int, str]]
    outside: bool


def sheet_formulas(
    workbook: Path, sheets: Collection[str]
) -> dict[str, dict[tuple[int, str], str]]:
    """The formulas of each wanted sheet, by row and column letter."""
    book = openpyxl.load_workbook(workbook, read_only=True, data_only=False)
    found: dict[str, dict[tuple[int, str], str]] = {}
    try:
        for name in sheets:
            if name not in book.sheetnames:
                continue
            sheet = book[name]
            # Some writers save wrong sheet dimensions, which would cut rows short.
            sheet.reset_dimensions()
            found[name] = {
                (number, get_column_letter(column)): value
                for number, values in enumerate(sheet.iter_rows(values_only=True), 1)
                for column, value in enumerate(values, 1)
                if isinstance(value, str) and value.startswith("=")
            }
    finally:
        book.close()
    return found


def references(formula: str, sheet: str) -> References:
    """The cells of `sheet` that a formula of that sheet reads."""
    try:
        tokens = Tokenizer(formula).items
    except TokenizerError:
        return References(frozenset(), outside=True)
    cells: set[tuple[int, str]] = set()
    outside = False
    for token in tokens:
        if token.type != Token.OPERAND or token.subtype != Token.RANGE:
            continue
        place, _, value = token.value.rpartition("!")
        if place and place.strip("'").replace("''", "'") != sheet:
            outside = True
            continue
        corners = [re.fullmatch(CELL, corner) for corner in value.split(":")]
        if not corners or len(corners) > 2 or None in corners:
            outside = True
            continue
        columns = [column_index_from_string(corner[1]) for corner in corners if corner]
        rows = [int(corner[2]) for corner in corners if corner]
        width = max(columns) - min(columns) + 1
        height = max(rows) - min(rows) + 1
        if width * height > MAX_RANGE:
            outside = True
            continue
        cells.update(
            (row, get_column_letter(column))
            for row in range(min(rows), max(rows) + 1)
            for column in range(min(columns), max(columns) + 1)
        )
    return References(frozenset(cells), outside)
