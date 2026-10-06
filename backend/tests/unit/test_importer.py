"""Study files must expand without losing zeros, rows, or input diagnostics."""

import json

import openpyxl
import pytest

from pkdb.importers.folder import load_folder, parse_bundle
from pkdb.schemas.validation import StudyValidationError


@pytest.fixture
def study_folder(tmp_path):
    folder = tmp_path / "Example"
    folder.mkdir()
    data = {
        "sid": "PK1",
        "name": "Example",
        "date": "2026-01-01",
        "reference": 123,
        "creator": "curator",
        "curators": [["curator", 3]],
        "access": "public",
        "licence": "open",
        "groupset": {
            "groups": [
                {"name": "all", "count": 2, "characteristica": []},
            ]
        },
        "outputset": {
            "outputs": [
                {
                    "source": "Results",
                    "output_type": "output",
                    "group": "all",
                    "measurement_type": "concentration",
                    "substance": "apixaban",
                    "mean": "col==mean",
                    "unit": "ng/ml",
                }
            ]
        },
    }
    (folder / "study.json").write_text(json.dumps(data))
    (folder / "reference.json").write_text(
        json.dumps(
            {
                "sid": 123,
                "name": "Example",
                "date": "2026-01-01",
                "authors": [],
            }
        )
    )
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Results"
    sheet.append(["Comment row", "ignored"])
    sheet.append(["id", "mean"])
    sheet.append([1, 0])
    sheet.append([2, None])
    book.save(folder / "Example.xlsx")
    return folder


def test_missing_and_zero_are_distinct(study_folder):
    before = (study_folder / "Example.xlsx").read_bytes()
    study = parse_bundle(load_folder(study_folder))
    assert [m.statistics.mean for m in study.measurements] == [0.0, None]
    assert [m.source.row for m in study.measurements if m.source is not None] == [3, 4]
    assert study.reference.sid == "123"
    assert study.metadata.curators[0].rating == 3
    assert (study_folder / "Example.xlsx").read_bytes() == before
    assert not list(study_folder.glob("*.tsv"))


def test_unknown_column_has_source_location(study_folder):
    p = study_folder / "study.json"
    data = json.loads(p.read_text())
    data["outputset"]["outputs"][0]["mean"] = "col==missing"
    p.write_text(json.dumps(data))
    with pytest.raises(StudyValidationError) as error:
        parse_bundle(load_folder(study_folder))
    assert error.value.report.issues[0].code == "unknown_column"
    assert error.value.report.issues[0].source is not None
    assert error.value.report.issues[0].source.sheet == "Results"


def test_unknown_field_is_rejected(study_folder):
    p = study_folder / "study.json"
    data = json.loads(p.read_text())
    data["unexpected"] = 123
    p.write_text(json.dumps(data))
    with pytest.raises(StudyValidationError):
        parse_bundle(load_folder(study_folder))


def test_symlink_escape_is_rejected(study_folder, tmp_path):
    secret = tmp_path / "outside.txt"
    secret.write_text("private")
    (study_folder / "Example.txt").symlink_to(secret)
    with pytest.raises(StudyValidationError):
        load_folder(study_folder)


def test_expansion_limit_is_enforced(study_folder):
    with pytest.raises(StudyValidationError) as error:
        parse_bundle(load_folder(study_folder), max_rows=1)
    assert error.value.report.issues[0].code == "row_limit"


def test_workbook_is_authoritative_over_generated_tsv(study_folder):
    (study_folder / ".Example_Results.tsv").write_text("id\tstale_column\n1\t99\n")
    study = parse_bundle(load_folder(study_folder))
    assert study.measurements[0].statistics.mean == 0


def test_unreported_time_preserves_marker(study_folder):
    path = study_folder / "study.json"
    data = json.loads(path.read_text())
    data["outputset"]["outputs"][0].update(time="NR", time_unit="NR")
    path.write_text(json.dumps(data))
    record = parse_bundle(load_folder(study_folder)).measurements[0]
    assert record.time is None
    assert record.time_not_reported is True
    assert record.time_unit is None
    assert record.time_unit_not_reported is True


def test_malformed_section_is_a_validation_error(study_folder):
    path = study_folder / "study.json"
    data = json.loads(path.read_text())
    data["groupset"] = 123
    path.write_text(json.dumps(data))
    with pytest.raises(StudyValidationError):
        parse_bundle(load_folder(study_folder))


def test_upload_cannot_claim_calculated_provenance(study_folder):
    path = study_folder / "study.json"
    data = json.loads(path.read_text())
    data["outputset"]["outputs"][0]["origin"] = "calculated"
    path.write_text(json.dumps(data))
    with pytest.raises(StudyValidationError) as error:
        parse_bundle(load_folder(study_folder))
    assert error.value.report.issues[0].code == "reserved_field"


@pytest.mark.parametrize("kind", ["xlsx", "tsv"])
def test_duplicate_source_headers_are_rejected(study_folder, kind):
    from pkdb.importers.workbook import read_table

    path = study_folder / f"duplicate.{kind}"
    if kind == "xlsx":
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Results"
        sheet.append(["comments"])
        sheet.append(["mean", "mean"])
        sheet.append([1, 2])
        book.save(path)
    else:
        path.write_text("mean\tmean\n1\t2\n")
    with pytest.raises(StudyValidationError) as error:
        read_table(path, "Results" if kind == "xlsx" else None, 100)
    assert error.value.report.issues[0].code == "duplicate_column"


def test_canonical_shape_errors_report_full_count(valid_bundle):
    from copy import deepcopy

    output = valid_bundle.study["outputset"]["outputs"][0]
    valid_bundle.study["outputset"]["outputs"] = [
        dict(deepcopy(output), unknown_field=True) for _ in range(150)
    ]
    with pytest.raises(StudyValidationError) as error:
        parse_bundle(valid_bundle)
    assert error.value.report.error_count == 150
    assert len(error.value.report.issues) == 100
    assert error.value.report.truncated


def test_nonfinite_in_memory_bundle_reports_source(valid_bundle):
    valid_bundle.study["outputset"]["outputs"][0]["mean"] = float("inf")
    with pytest.raises(StudyValidationError) as error:
        parse_bundle(valid_bundle)
    issue = error.value.report.issues[0]
    assert issue.code == "invalid_number"
    assert issue.source is not None
    assert issue.source.path == ("outputset", "outputs", 0, "mean")
    assert error.value.report.error_count == 1


def test_external_characteristic_alias_preserves_values_and_source(valid_bundle):
    from copy import deepcopy

    from pkdb.importers.folder import parse_bundle

    original = parse_bundle(valid_bundle)
    changed = valid_bundle.model_copy(deep=True)
    group = changed.study["groupset"]["groups"][0]
    group["characteristica_ex"] = group.pop("characteristica")
    before = deepcopy(changed.study)
    actual = parse_bundle(changed)
    assert actual.groups == original.groups
    assert actual.source_digest != original.source_digest
    assert changed.study == before


@pytest.mark.parametrize("name", [1, 1.5])
def test_numeric_subject_names_and_references_use_legacy_text_conversion(
    study_folder, name
):
    path = study_folder / "study.json"
    data = json.loads(path.read_text())
    data["individualset"] = {"individuals": [{"name": name, "group": "all"}]}
    output = data["outputset"]["outputs"][0]
    output.pop("group")
    output["individual"] = name
    path.write_text(json.dumps(data))
    study = parse_bundle(load_folder(study_folder))
    assert study.individuals[0].name == str(name)
    assert study.individuals[0].key == str(name)
    assert all(record.individual == str(name) for record in study.measurements)


def test_boolean_subject_name_is_not_coerced_to_text(study_folder):
    path = study_folder / "study.json"
    data = json.loads(path.read_text())
    data["individualset"] = {"individuals": [{"name": True, "group": "all"}]}
    path.write_text(json.dumps(data))
    with pytest.raises(StudyValidationError):
        parse_bundle(load_folder(study_folder))


@pytest.mark.parametrize("label", [1, 1.5])
def test_numeric_output_labels_use_legacy_text_conversion(study_folder, label):
    path = study_folder / "study.json"
    data = json.loads(path.read_text())
    data["outputset"]["outputs"][0]["label"] = label
    path.write_text(json.dumps(data))
    study = parse_bundle(load_folder(study_folder))
    assert all(record.label == str(label) for record in study.measurements)


@pytest.mark.parametrize("choice", [0, 1, 1.5])
def test_numeric_characteristic_choices_use_legacy_text_conversion(
    study_folder, choice
):
    path = study_folder / "study.json"
    data = json.loads(path.read_text())
    data["individualset"] = {
        "individuals": [
            {
                "name": "subject",
                "group": "all",
                "characteristica": [
                    {"measurement_type": "nat2 activity", "choice": choice}
                ],
            }
        ]
    }
    path.write_text(json.dumps(data))
    study = parse_bundle(load_folder(study_folder))
    assert study.individuals[0].characteristica[0].choice == str(choice)


def write_value_output(study_folder, **fields):
    book = openpyxl.load_workbook(study_folder / "Example.xlsx")
    book["Results"]["B2"] = "concentration"
    book["Results"]["B3"] = -1
    book.save(study_folder / "Example.xlsx")
    path = study_folder / "study.json"
    data = json.loads(path.read_text())
    data["individualset"] = {"individuals": [{"name": "person", "group": "all"}]}
    output = data["outputset"]["outputs"][0]
    del output["mean"]
    output["value"] = "col==concentration"
    output.update(fields)
    path.write_text(json.dumps(data))


def test_format_1_value_becomes_mean_located_at_its_cell(study_folder, vocabulary):
    from pkdb.domain.validation import prepare_study

    write_value_output(study_folder, group=None, individual="person")
    study = parse_bundle(load_folder(study_folder))
    assert [m.statistics.mean for m in study.measurements] == [-1.0, None]
    source = study.measurements[0].source
    assert source is not None
    located = source.for_field("mean")
    assert (located.sheet, located.cell, located.header) == (
        "Results",
        "B3",
        "concentration",
    )
    with pytest.raises(StudyValidationError) as error:
        prepare_study(study, vocabulary)
    negative = next(
        issue for issue in error.value.report.issues if issue.code == "negative_value"
    )
    assert negative.field == "mean"
    assert negative.source is not None and negative.source.cell == "B3"


def test_format_1_group_value_is_rejected_at_its_cell(study_folder):
    write_value_output(study_folder)
    with pytest.raises(StudyValidationError) as error:
        parse_bundle(load_folder(study_folder))
    issue = error.value.report.issues[0]
    assert (issue.code, issue.field) == ("group_value", "value")
    assert issue.source is not None
    assert (issue.source.sheet, issue.source.cell) == ("Results", "B3")


def test_format_1_group_value_of_an_unspecified_summary_becomes_mean(study_folder):
    write_value_output(study_folder, calculation_type="unspecified summary")
    study = parse_bundle(load_folder(study_folder))
    assert [(m.group, m.statistics.mean) for m in study.measurements] == [
        ("all", -1.0),
        ("all", None),
    ]
    assert all(m.calculation_type == "unspecified summary" for m in study.measurements)


def write_characteristic(study_folder, section, **fields):
    path = study_folder / "study.json"
    data = json.loads(path.read_text())
    characteristic = {"measurement_type": "weight", "unit": "kg", **fields}
    if section == "groupset":
        data["groupset"]["groups"][0]["characteristica"] = [characteristic]
    else:
        data["individualset"] = {
            "individuals": [
                {"name": "person", "group": "all", "characteristica": [characteristic]}
            ]
        }
    path.write_text(json.dumps(data))


def test_format_1_group_characteristic_value_is_rejected(study_folder):
    write_characteristic(study_folder, "groupset", value=70)
    with pytest.raises(StudyValidationError) as error:
        parse_bundle(load_folder(study_folder))
    assert error.value.report.issues[0].code == "group_value"
    write_characteristic(
        study_folder, "groupset", value=70, calculation_type="unspecified summary"
    )
    characteristic = (
        parse_bundle(load_folder(study_folder)).groups[0].characteristica[0]
    )
    assert characteristic.statistics.mean == 70
    assert characteristic.calculation_type == "unspecified summary"


def test_format_1_individual_characteristic_value_becomes_mean(study_folder):
    write_characteristic(study_folder, "individualset", value=70)
    study = parse_bundle(load_folder(study_folder))
    characteristic = study.individuals[0].characteristica[0]
    assert characteristic.statistics.mean == 70
    assert characteristic.calculation_type is None


def test_format_1_value_and_mean_together_are_rejected(study_folder):
    path = study_folder / "study.json"
    data = json.loads(path.read_text())
    data["outputset"]["outputs"][0].update(
        value=1, calculation_type="unspecified summary"
    )
    path.write_text(json.dumps(data))
    with pytest.raises(StudyValidationError) as error:
        parse_bundle(load_folder(study_folder))
    assert error.value.report.issues[0].code == "conflicting_statistics"


def write_intervention(study_folder, **fields):
    path = study_folder / "study.json"
    data = json.loads(path.read_text())
    data["interventionset"] = {
        "interventions": [
            {
                "name": "dose",
                "measurement_type": "dosing",
                "substance": "apixaban",
                "value": 2.5,
                "unit": "mg",
                "time_unit": "h",
                "route": "oral",
                "form": "tablet",
                "application": "multiple dose",
                **fields,
            }
        ]
    }
    path.write_text(json.dumps(data))


@pytest.mark.parametrize(
    "text,expected",
    [
        ("0|12|40", {"time": [0.0, 12.0, 40.0]}),
        (" -6 | -5.5 | 1e1 ", {"time": [-6.0, -5.5, 10.0]}),
        ("S0T12R3", {"time": 0.0, "interval": 12.0, "doses": 3}),
        ("S-6T0.5R12", {"time": -6.0, "interval": 0.5, "doses": 12}),
        ("S0T6R2 | S24T6R2 | 144", {"time": [0.0, 6.0, 24.0, 30.0, 144.0]}),
        ("12", {"time": 12.0}),
        (12, {"time": 12.0}),
    ],
)
def test_format_1_schedule_strings_become_structured(study_folder, text, expected):
    write_intervention(study_folder, time=text)
    intervention = parse_bundle(load_folder(study_folder)).interventions[0]
    assert intervention.statistics.mean == 2.5
    assert {
        "time": intervention.time,
        "interval": intervention.interval,
        "doses": intervention.doses,
    } == {"time": None, "interval": None, "doses": None} | expected


@pytest.mark.parametrize(
    "fields",
    [
        {"time": "S0T12"},
        {"time": "0|"},
        {"time": "a|1"},
        {"time": "S0T12R0"},
        {"time": "S0T-12R3"},
        {"time": "0;12"},
        {"time": "s0t12r3"},
        {"time": "S0T12R3x"},
        {"time": "1e400|1"},
        {"time": "S0T1R10001|1"},
        {"time": "S0T12R3", "doses": 3},
        {"time": "S0T12R3", "interval": 12},
    ],
)
def test_invalid_format_1_schedule_is_located(study_folder, fields):
    write_intervention(study_folder, **fields)
    with pytest.raises(StudyValidationError) as error:
        parse_bundle(load_folder(study_folder))
    issue = error.value.report.issues[0]
    assert issue.code == "invalid_schedule"
    assert issue.source is not None
    assert issue.source.path == ("interventionset", "interventions", 0, "time")


def test_format_1_schedule_cell_is_parsed_with_its_location(study_folder):
    book = openpyxl.load_workbook(study_folder / "Example.xlsx")
    sheet = book.create_sheet("Doses")
    sheet.append(["comment"])
    sheet.append(["name", "schedule"])
    sheet.append(["first", "S0T24R7"])
    sheet.append(["second", "0 | 24 | x"])
    book.save(study_folder / "Example.xlsx")
    write_intervention(study_folder, source="Doses", name="col==name")
    path = study_folder / "study.json"
    data = json.loads(path.read_text())
    data["interventionset"]["interventions"][0]["time"] = "col==schedule"
    path.write_text(json.dumps(data))
    with pytest.raises(StudyValidationError) as error:
        parse_bundle(load_folder(study_folder))
    issue = error.value.report.issues[0]
    assert issue.code == "invalid_schedule"
    assert issue.source is not None
    assert (issue.source.sheet, issue.source.cell) == ("Doses", "B4")
